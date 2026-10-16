"""Copy verified design observations into workbooks with checkable row bindings.

Selection, additional calculations and prose remain the writer's responsibility.
Native identities and measurements are never reconstructed by the writer.
"""
from copy import deepcopy
import hashlib
import json
import math
from pathlib import Path

from .splice_prediction_records import METRICS
from .splice_recommendations import finite_number


CORE_COLUMNS = (
    "binding_id", "target_5to3", "antisense_5to3", "program_ids", "transcript_ids",
    "genomic_blocks_bed0_half_open", "strand", "recorded_chemistry", "length_nt",
    "gc_bases", "gc_fraction", "longest_homopolymer", "off_target_exact_count",
    "off_target_1mm_count", "off_target_search_status", "off_target_mismatch_budget",
    "primary_delta_deletion", "primary_delta_replacement", "physical_splice_elements",
    "stage_roles", "source_observation_ids", "evidence_conflicts",
    "genomic_blocks_1based_inclusive", "annotated_splice_contexts",
)
PREDICTION_SHEET = "Recorded splice evidence"
PREDICTION_COLUMNS = ("binding_id", "observation_id", "source", "source_pointer", "site_identity",
                      *METRICS, "recorded_values", "unresolved_fields")
_RECORDED = {
    "off_target_exact_count": ("Off_target_exact_count", "off_target_exact_count"),
    "off_target_1mm_count": ("Off_target_1mm_count", "off_target_1mm_count"),
    "off_target_search_status": ("off_target_search_status",),
    "off_target_mismatch_budget": ("Off_target_mismatch_budget", "off_target_mismatch_budget"),
    "primary_delta_deletion": ("primary_delta_deletion",),
    "primary_delta_replacement": ("primary_delta_replacement",),
}


def _json(value):
    return json.dumps(value, sort_keys=True, ensure_ascii=False, allow_nan=False)


def _unique(values):
    return list({_json(value): value for value in values}.values())


def _fingerprint(fields):
    return hashlib.sha256(_json(fields).encode()).hexdigest()


def _cell(value):
    if isinstance(value, (dict, list, tuple)):
        value = _json(value)
    if isinstance(value, str) and len(value) > 32767:
        raise ValueError("Design evidence exceeds the Excel cell limit; use a separate linked evidence file")
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError("Nonfinite design value cannot be exported as a measured number")
    return value


def _source_rows(declared):
    evidence = (declared.get("completed_run") or {}).get("design_evidence") or []
    observations, sources = [], []
    for entry in evidence:
        meta = entry.get("observations")
        if not meta:
            continue
        content = Path(meta["path"]).read_bytes()
        if hashlib.sha256(content).hexdigest() != meta["sha256"]:
            raise ValueError("Design observation source hash changed")
        data = json.loads(content)
        if data.get("schema") != "modal-aso-design-observations/1" or data.get("task_id") != entry["task_id"]:
            raise ValueError("Design observation schema or producer does not match its declaration")
        sources.append({"task_id": entry["task_id"], "path": meta["path"], "sha256": meta["sha256"]})
        observations.extend(data["rows"])
    return observations, sources


def _catalog(declared):
    observations, sources = _source_rows(declared)
    groups = {}
    for observation in observations:
        # Unknown identities are retained separately, never joined by sequence.
        key = observation.get("binding_match_key") or observation["observation_id"]
        groups.setdefault(key, []).append(observation)
    rows = []
    for key, group in groups.items():
        first = group[0]
        identity = first["binding_identity"]
        if any(item["binding_identity"] != identity for item in group):
            raise ValueError("Conflicting identities share a design binding key")
        metrics = first.get("derived_sequence_metrics") or {}
        row = {"binding_id": key, "task_id": identity["task_id"],
               **{name: deepcopy(identity.get(name)) for name in (
                   "target_5to3", "antisense_5to3", "program_ids", "transcript_ids", "strand", "recorded_chemistry")},
               "genomic_blocks_bed0_half_open": deepcopy(identity["binding_intervals_bed0_half_open"]),
               "genomic_blocks_1based_inclusive": [[chrom, start + 1, end]
                   for chrom, start, end in identity["binding_intervals_bed0_half_open"]],
               "length_nt": metrics.get("length"), "gc_bases": metrics.get("gc_bases"),
               "gc_fraction": metrics["gc_bases"] / metrics["length"] if metrics.get("length") else None,
               "longest_homopolymer": metrics.get("longest_homopolymer_run"),
               "physical_splice_elements": _unique([element for item in group for element in item["physical_splice_elements"]]),
               "annotated_splice_contexts": _unique([element for item in group
                   for element in item.get("annotated_splice_contexts", [])]),
               "stage_roles": sorted({role for item in group for role in item["stage_roles"]}),
               "source_observation_ids": [item["observation_id"] for item in group],
               "source_row_fingerprints": sorted({_fingerprint(item["recorded_fields"]) for item in group}),
               "evidence_conflicts": {}, "recorded_fields": {}, "splice_predictions": []}
        # Preserve null/absent observations in the source; a later measurement
        # can supply a value, but disagreeing measured values remain unresolved.
        for field in {field for item in group for field in item["recorded_fields"]}:
            values = _unique([item["recorded_fields"][field] for item in group
                              if field in item["recorded_fields"] and item["recorded_fields"][field] is not None])
            row["recorded_fields"][field] = values[0] if len(values) == 1 else None
            if len(values) > 1:
                row["evidence_conflicts"][field] = values
        for name, aliases in _RECORDED.items():
            values = []
            for item in group:
                for prefix in ("/", "/source_annotations/"):
                    for alias in aliases:
                        value = item["recorded_fields"].get(prefix + alias)
                        if value is not None and value != "":
                            if name != "off_target_search_status":
                                number = finite_number(value)
                                if number is None:
                                    continue
                                value = number
                            values.append(value)
            values = _unique(values)
            row[name] = values[0] if len(values) == 1 else None
            if len(values) > 1:
                row["evidence_conflicts"][name] = values
        for item in group:
            for prediction in item["splice_predictions"]:
                row["splice_predictions"].append({"observation_id": item["observation_id"],
                    "source": item["source"], **deepcopy(prediction)})
        rows.append(row)
    return rows, sources


def _source_value(row, field):
    if isinstance(field, dict):
        # The writer selects an attested reference frame, never supplies an
        # anchor or recalculates offsets. Exact IDs preserve transcript versions.
        required = {"kind", "transcript_id", "exon_ordinal", "element", "field"}
        values = {"relative_interval_inclusive", "target_interval_1based_inclusive",
                  "boundary_1based", "exonic_overlap_bp", "intronic_overlap_bp",
                  "exon_interval_1based_inclusive", "intron_interval_1based_inclusive"}
        if (set(field) != required or field["kind"] != "splice_context"
                or not isinstance(field["transcript_id"], str)
                or type(field["exon_ordinal"]) is not int or field["exon_ordinal"] < 1
                or field["element"] not in {"donor", "acceptor"} or field["field"] not in values):
            raise ValueError("Invalid annotated splice context source column")
        matches = [element for element in row["annotated_splice_contexts"]
                   if all(element.get(key) == field[key] for key in ("transcript_id", "exon_ordinal", "element"))]
        if len(matches) > 1:
            raise ValueError("Annotated splice context is ambiguous across target blocks or source observations; retain the complete context array")
        return matches[0].get(field["field"]) if matches else None
    if not isinstance(field, str):
        raise ValueError("Design source column must be a catalog field or an annotated context selector")
    if field.startswith("recorded:"):
        pointer = field.removeprefix("recorded:")
        if pointer not in row["recorded_fields"]:
            return None
        return row["recorded_fields"][pointer]
    if field not in row:
        raise ValueError("Unknown design source column: " + field)
    return row[field]


def write_design_catalog(declared, path):
    """Expose the renderer's exact row schema as a protected analysis input."""
    rows, sources = _catalog(declared)
    path = Path(path)
    if path.is_symlink() or path.resolve() in {Path(source["path"]).resolve() for source in sources}:
        raise ValueError("Design catalog destination overlaps a source")
    path.parent.mkdir(parents=True, exist_ok=True)
    content = (_json({"schema": "modal-aso-design-catalog/1", "rows": rows,
                      "observation_sources": sources}) + "\n").encode()
    path.write_bytes(content)
    return {"path": str(path), "sha256": hashlib.sha256(content).hexdigest(),
            "schema": "modal-aso-design-catalog/1", "row_collection": "rows", "row_count": len(rows),
            "columns": list(rows[0]) if rows else [],
            "task_ids": sorted({row["task_id"] for row in rows}),
            "scope": "One row per verified binding; unresolved source identities remain separate. Stage roles distinguish complete raw libraries from selected subsets."}


def _prediction_rows(indexed, tables):
    seen = set()
    for table in tables:
        for key in table["row_ids"]:
            if key in seen:
                continue
            seen.add(key)
            for record in indexed[key]["splice_predictions"]:
                yield [key, record["observation_id"], record["source"], record["source_pointer"],
                       record["identity"], *[record["numeric_values"].get(name) for name in METRICS],
                       record["recorded_values"], record["unresolved_fields"]]


class DesignWorkbook:
    """A source-bound workbook; add custom columns as parallel row dictionaries.

    wb = DesignWorkbook(json.load(open('declared_inputs.json')))
    rows = wb.rows(stage='raw_generated', task_id='recorded-task-id')
    wb.add_table('Library', rows, coverage='all_raw', task_id='recorded-task-id',
                 extra_columns=['ID'], extra_values=[{'ID': 'candidate-1'}, ...])
    result = wb.save('/output/library.xlsx')
    # Write result into sandbox_result.json; add prose on wb.workbook separately.
    """
    def __init__(self, declared):
        from openpyxl import Workbook
        rows, self.sources = _catalog(declared)
        self._rows = {row["binding_id"]: row for row in rows}
        self.workbook = Workbook()
        self.workbook.remove(self.workbook.active)
        self._tables = []

    def rows(self, stage=None, task_id=None):
        return deepcopy([row for row in self._rows.values()
                         if (stage is None or stage in row["stage_roles"])
                         and (task_id is None or row["task_id"] == task_id)])

    def _source_row(self, row, task_id):
        if not isinstance(row, dict):
            raise ValueError("Design table rows must be source record objects")
        key = row.get("binding_id")
        expected = self._rows.get(key) if isinstance(key, str) else None
        if expected is not None and row == expected:
            matches = [expected]
        else:
            # Exact, complete native records are another valid representation
            # of a catalog row. Never use partial sequence/coordinate matching.
            from .design_observation_tables import _fields
            fingerprint = _fingerprint(_fields(row))
            matches = [item for item in self._rows.values()
                       if fingerprint in item["source_row_fingerprints"]]
        matches = [item for item in matches if task_id is None or item["task_id"] == task_id]
        if len(matches) != 1:
            if expected is not None and not matches:
                raise ValueError("Design row was changed; use extra_values for derived calculations")
            raise ValueError("Native design row does not identify exactly one unchanged source binding in the requested task; "
                             "use the unmodified rows returned by DesignWorkbook.rows()")
        return matches[0]

    def add_table(self, title, rows, *, coverage="subset", task_id=None,
                  source_columns=None, extra_columns=(), extra_values=None):
        columns = dict(source_columns or {})  # Display header -> source field.
        columns = {**{name: name for name in CORE_COLUMNS}, **columns}
        if any(columns[name] != name for name in CORE_COLUMNS):
            raise ValueError("Core design columns cannot be rebound")
        if len(set(extra_columns)) != len(extra_columns) or set(columns).intersection(extra_columns):
            raise ValueError("Extra columns must be unique and separate from source facts")
        if title in self.workbook.sheetnames:
            raise ValueError("Duplicate design table sheet")
        if coverage not in {"subset", "all_raw", "all_final"}:
            raise ValueError("Unknown design table coverage")
        extra_values = [{} for _ in rows] if extra_values is None else extra_values
        if len(extra_values) != len(rows):
            raise ValueError("Extra values must correspond one-to-one with source rows")
        source_rows = [self._source_row(row, task_id) for row in rows]
        ids = [row["binding_id"] for row in source_rows]
        if len(ids) != len(set(ids)):
            raise ValueError("Duplicate binding in a design table")
        sheet = self.workbook.create_sheet(title)
        sheet.append([*columns, *extra_columns])
        for expected, extras in zip(source_rows, extra_values):
            sheet.append([_cell(_source_value(expected, field)) for field in columns.values()]
                         + [_cell(extras.get(name)) for name in extra_columns])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        self._tables.append({"sheet": title, "columns": columns, "row_ids": ids,
                             "coverage": coverage, "task_id": task_id})
        return sheet

    def save(self, path):
        path = Path(path)
        if not self._tables:
            raise ValueError("A design workbook needs a source-bound candidate table")
        # Every included row's site-specific predictions survive in a separate
        # table, including nulls, zeros, replacement/deletion and source identity.
        title = PREDICTION_SHEET
        if title in self.workbook.sheetnames:
            raise ValueError("Reserved design evidence sheet already exists")
        sheet = self.workbook.create_sheet(title)
        sheet.append(PREDICTION_COLUMNS)
        for row in _prediction_rows(self._rows, self._tables):
            sheet.append([_cell(value) for value in row])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        path.parent.mkdir(parents=True, exist_ok=True)
        self.workbook.save(path)
        return {"design_tables": [{"path": str(path.resolve()), **table} for table in self._tables],
                "design_observation_sources": deepcopy(self.sources), "design_prediction_sheet": title}


def render_design_workbook_plans(payload, result, required_outputs):
    """Render a model's analysis table plan with platform-owned source cells.

    The program chooses rows and additional calculations, never reconstructs
    native source columns or controls the workbook's provenance machinery.
    """
    required = {Path(p).resolve() for p in payload.get("source_bound_design_workbooks") or []}
    required &= set(required_outputs)
    if not required:
        return
    plans = result.get("design_workbooks")
    if plans is None and result.get("design_tables"):
        return  # Previously supported source-bound writer output; still verified below.
    if not isinstance(plans, list) or not plans:
        raise ValueError("Return design_workbooks in sandbox_result.json: each workbook has path, tables, and notes; "
                         "each table has title, task_id, coverage, extra_columns, and rows of {binding_id, values}. "
                         "Read binding_id from completed_run.design_catalog. The platform writes the Excel files.")
    declared = payload["declared_inputs"]
    scope = {Path(entry["path"]).resolve(): entry["task_id"]
             for entry in declared.get("requested_design_workbooks") or []}
    seen, tables, sources = set(), [], []
    try:
        for plan in plans:
            path = Path(plan["path"]).resolve()
            if path not in required or path in seen or path.suffix.lower() != ".xlsx":
                raise ValueError("Workbook plan path is duplicate or outside its requested design scope")
            seen.add(path)
            writer = DesignWorkbook(declared)
            catalog = {row["binding_id"]: row for row in writer.rows()}
            for table in plan["tables"]:
                task_id = table.get("task_id", scope.get(path))
                if not task_id or scope.get(path) != task_id:
                    raise ValueError("Workbook table must use its requested producer task")
                columns = table.get("extra_columns", [])
                if not isinstance(columns, list) or not all(isinstance(name, str) for name in columns):
                    raise ValueError("Workbook extra_columns must be an array of column names")
                rows, values = [], []
                for entry in table["rows"]:
                    if not isinstance(entry, dict) or set(entry) - {"binding_id", "values"}:
                        raise ValueError("Workbook rows contain only binding_id and derived values; source facts are supplied by the renderer")
                    key, extra = entry["binding_id"], entry.get("values", {})
                    if key not in catalog:
                        raise ValueError("Workbook row binding_id is absent from the canonical design catalog")
                    if not isinstance(extra, dict) or set(extra) - set(columns):
                        raise ValueError("Workbook values must use declared extra_columns")
                    rows.append(catalog[key])
                    values.append(extra)
                writer.add_table(table["title"], rows, task_id=task_id,
                                 coverage=table["coverage"], extra_columns=columns, extra_values=values,
                                 source_columns=table.get("source_columns"))
            for note in plan.get("notes", []):
                title = note["title"]
                if title in writer.workbook.sheetnames or title == PREDICTION_SHEET:
                    raise ValueError("Workbook note sheet collides with a source table")
                sheet = writer.workbook.create_sheet(title)
                for cells in note["rows"]:
                    if not isinstance(cells, list):
                        raise ValueError("Workbook note rows must be arrays of cells")
                    sheet.append([_cell(value) for value in cells])
            rendered = writer.save(path)
            tables.extend(rendered["design_tables"])
            sources.extend(rendered["design_observation_sources"])
        if seen != required:
            raise ValueError("A rendering plan is required for every requested design workbook")
    except (KeyError, TypeError, AttributeError) as exc:
        raise ValueError("Malformed design workbook rendering plan: " + str(exc)) from exc
    result["design_tables"] = tables
    result["design_observation_sources"] = _unique(sources)
    result["design_prediction_sheet"] = PREDICTION_SHEET
    result["design_workbook_rendering"] = {"owner": "platform", "workbooks": len(seen),
        "source_facts": "copied from verified catalog", "additional_values": "model-authored analysis; scientifically assessed separately"}


def _agrees(expected, observed):
    expected = _cell(expected)
    if expected == "" and observed is None:
        return True  # XLSX represents an empty string as an empty cell.
    if isinstance(expected, (int, float)) and not isinstance(expected, bool):
        return (isinstance(observed, (int, float)) and not isinstance(observed, bool)
                and math.isclose(expected, observed, rel_tol=1e-12, abs_tol=1e-15))
    return type(expected) is type(observed) and expected == observed


def validate_design_workbooks(payload, result, required_outputs):
    """Check actual file contents, not merely the writer's completion claim."""
    required = {Path(p).resolve() for p in payload.get("source_bound_design_workbooks") or []}
    required &= set(required_outputs)
    if not required:
        return
    from openpyxl import load_workbook
    tables = result.get("design_tables")
    if not isinstance(tables, list) or not tables:
        raise ValueError("Design Excel exports require source-bound design_tables. Use aso_agents.design_workbooks.DesignWorkbook, "
                         "add_table(), and include the dict returned by save() in sandbox_result.json. "
                         "Copy source rows; do not reconstruct targets from overlapping fragments.")
    rows, sources = _catalog(payload["declared_inputs"])
    indexed = {row["binding_id"]: row for row in rows}
    declared = set()
    books = {}
    checked = 0
    try:
        for table in tables:
            path = Path(table["path"]).resolve()
            if path not in required:
                raise ValueError("Design table path is outside its requested output scope")
            if (path, table["sheet"]) in declared:
                raise ValueError("Duplicate design table declaration")
            declared.add((path, table["sheet"]))
            if path not in books:
                books[path] = load_workbook(path, read_only=True, data_only=False)
            book = books[path]
            if table["sheet"] not in book.sheetnames:
                raise ValueError("Declared design table sheet is missing")
            values = list(book[table["sheet"]].values)
            if not values:
                raise ValueError("Design table is empty")
            headers = list(values[0])
            columns = table.get("columns") or {}
            if any(columns.get(name) != name for name in CORE_COLUMNS):
                raise ValueError("Design table must bind every core source column")
            if any(headers.count(name) != 1 for name in columns):
                raise ValueError("Missing or duplicate bound design column")
            ids = table["row_ids"]
            if len(ids) != len(set(ids)) or len(values) - 1 != len(ids):
                raise ValueError("Design table row count or identity differs from its binding")
            for key, cells in zip(ids, values[1:]):
                if key not in indexed:
                    raise ValueError("Unknown source binding in design export")
                row = indexed[key]
                if table.get("task_id") is not None and row["task_id"] != table["task_id"]:
                    raise ValueError("Design row borrowed from another producer")
                for header, field in columns.items():
                    expected = _source_value(row, field)
                    if not _agrees(expected, cells[headers.index(header)]):
                        raise ValueError(f"Design export contradicts source: {path.name}/{table['sheet']} {key} {header}; "
                                         "use the source-bound writer and keep derived values in extra columns")
                checked += 1
            coverage = table.get("coverage")
            if coverage in {"all_raw", "all_final"}:
                stage = "raw_generated" if coverage == "all_raw" else "final_output"
                expected = {row["binding_id"] for row in rows if stage in row["stage_roles"]
                            and (table.get("task_id") is None or row["task_id"] == table["task_id"])}
                if set(ids) != expected:
                    raise ValueError("Design export does not cover its declared complete source stage")
            elif coverage != "subset":
                raise ValueError("Design table coverage must be explicit")
        if required != {path for path, _ in declared}:
            raise ValueError("Every requested design workbook needs a source-bound candidate table")
        for path, book in books.items():
            if PREDICTION_SHEET not in book.sheetnames:
                raise ValueError("Design workbook lost its recorded prediction evidence sheet")
            actual = list(book[PREDICTION_SHEET].values)
            expected = [list(PREDICTION_COLUMNS), *_prediction_rows(indexed, [table for table in tables
                        if Path(table["path"]).resolve() == path])]
            if len(actual) != len(expected) or any(len(a) != len(b) or any(not _agrees(x, y)
                    for x, y in zip(a, b)) for a, b in zip(expected, actual)):
                raise ValueError("Design workbook changed or omitted recorded site-specific predictions")
    except (KeyError, TypeError, AttributeError) as exc:
        raise ValueError("Malformed source-bound design table declaration: " + str(exc)) from exc
    finally:
        for book in books.values():
            book.close()
    result["design_fact_validation"] = {"rows_checked": checked, "source_catalogs": sources,
        "scope": "Bound source columns and declared stage coverage only; derived columns, selection and prose are not scientifically certified."}
