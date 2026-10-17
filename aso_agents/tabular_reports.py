"""Render declared report tables directly from exported observations.

The analysis program supplies semantic column roles. This layer does not infer
intent from prose or recompute a scientific ranking.
"""
import csv
import hashlib
import json
import math
from copy import deepcopy
from pathlib import Path


def _workbook_report_table(path, sheet_name, header_row):
    """Read an explicitly identified rectangular sheet without evaluating formulas."""
    if not isinstance(sheet_name, str) or not sheet_name:
        raise ValueError("A workbook report view needs an explicit nonempty source_sheet.")
    if type(header_row) is not int or header_row < 1:
        raise ValueError("Workbook header_row must be a positive integer.")
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=False)
    try:
        if sheet_name not in book.sheetnames:
            raise ValueError("Report source worksheet does not exist: " + sheet_name)
        sheet = book[sheet_name]
        if header_row > sheet.max_row:
            raise ValueError("Workbook header_row is outside the source worksheet.")
        iterator = sheet.iter_rows(min_row=header_row)
        header = list(next(iterator))
        while header and header[-1].value in (None, ""):
            header.pop()
        if not header or any(not isinstance(cell.value, str) or not cell.value.strip()
                             for cell in header):
            raise ValueError("Workbook table headers must be nonempty column names.")
        fields = [cell.value for cell in header]
        if len(set(fields)) != len(fields):
            raise ValueError("Workbook table headers must be unique.")
        if any(cell.data_type == "f" for cell in header):
            raise ValueError("A workbook report view cannot evaluate a formula header.")
        rows = []
        for cells in iterator:
            if not any(cell.value not in (None, "") for cell in cells):
                continue
            if any(cell.value not in (None, "") for cell in cells[len(fields):]):
                raise ValueError("Workbook data extends beyond its declared header columns.")
            if any(cell.data_type == "f" for cell in cells[:len(fields)]):
                raise ValueError(
                    "A workbook report view cannot evaluate formula cells. "
                    "Export computed values in the requested source sheet."
                )
            rows.append({name: "" if cell.value is None else str(cell.value)
                         for name, cell in zip(fields, cells)})
        return fields, rows
    finally:
        book.close()


def _repeated_export_table_columns(text, source_columns, source_rows, alternative_sources=()):
    """Find a hand-written export table that is not faithful to its source."""
    columns = set(source_columns)
    key = source_columns[0] if source_columns else None
    by_key = {row[key]: row for row in source_rows} if key else {}
    unique_key = len(by_key) == len(source_rows)
    lines = text.splitlines()
    for index in range(len(lines) - 1):
        header, separator = lines[index:index + 2]
        if not (header.lstrip().startswith("|") and separator.lstrip().startswith("|")):
            continue
        names = [cell.strip().strip("`") for cell in header.strip().strip("|").split("|")]
        separators = [cell.strip() for cell in separator.strip().strip("|").split("|")]
        if len(names) != len(separators) or not all(
            cell and set(cell) <= {"-", ":"} and "-" in cell for cell in separators
        ):
            continue
        repeated = columns.intersection(names)
        if key in repeated and len(repeated) >= 3:
            if not unique_key or len(names) != len(set(names)):
                return sorted(repeated)
            shown_rows = []
            for line in lines[index + 2:]:
                if not line.lstrip().startswith("|"):
                    break
                values = [cell.strip().replace("\\|", "|") for cell in line.strip().strip("|").split("|")]
                if len(values) != len(names):
                    return sorted(repeated)
                shown_rows.append(dict(zip(names, values)))
            if not shown_rows:
                return sorted(repeated)
            sources = [(source_columns, source_rows), *alternative_sources]
            def matches(candidate_columns, candidate_rows):
                candidate_columns = set(candidate_columns)
                if key not in candidate_columns or not repeated <= candidate_columns:
                    return False
                indexed = {row[key]: row for row in candidate_rows}
                if len(indexed) != len(candidate_rows):
                    return False
                return all(
                    (source := indexed.get(displayed[key])) is not None
                    and all(displayed[name] == source[name]
                            or (source[name] == "" and displayed[name] in {"-", "—", "–"})
                            for name in repeated)
                    for displayed in shown_rows
                )
            if not any(matches(candidate_columns, candidate_rows)
                       for candidate_columns, candidate_rows in sources):
                return sorted(repeated)
    return []


def _manual_markdown_table_locations(text):
    """Locate populated pipe tables using the existing validation boundary."""
    lines = text.splitlines()
    for index in range(len(lines) - 2):
        header, separator, first_row = lines[index:index + 3]
        if not (header.lstrip().startswith("|") and separator.lstrip().startswith("|")
                and first_row.lstrip().startswith("|")):
            continue
        cells = [cell.strip() for cell in separator.strip().strip("|").split("|")]
        if len(cells) >= 2 and all(cell and set(cell) <= {"-", ":"} and "-" in cell for cell in cells):
            end = index + 2
            while end < len(lines) and lines[end].lstrip().startswith("|"):
                end += 1
            yield {"header_line_1based": index + 1, "header_preview": header[:280],
                   "row_count": end - index - 2}


def _has_manual_markdown_table(text):
    """Detect a populated pipe table written without a bound source marker."""
    return next(_manual_markdown_table_locations(text), None) is not None


def _unbound_markdown_table_error(texts):
    locations = [{"path": str(path), **location}
                 for path, text in texts.items()
                 for location in _manual_markdown_table_locations(text)]
    # Keep feedback bounded for large reports while locating multiple files
    # in the same repair attempt. No scientific values or source are inferred.
    diagnostic = {"tables": locations[:20], "total_unbound_tables": len(locations),
                  "unlisted_tables": max(0, len(locations) - 20)}
    return ValueError(
        "A Markdown result table remains unbound to a declared source table. "
        "Locations in the render input: " + json.dumps(diagnostic, ensure_ascii=False) + ". "
        "Repair every listed destination, including requested Markdown files as well as the summary. "
        "Replace each manual table with a {{table:id}} marker and declare its report_tables view "
        "using an available source_path and its exact columns. Keep the requested scientific "
        "content and outputs; changing only the summary does not repair another report."
    )


def _replace_candidate_tables_with_marker(text, marker, source_row_count):
    """Replace only a candidate-shaped hand table with one source-bound view."""
    lines = text.splitlines()
    output = []
    replaced = 0
    index = 0
    while index < len(lines):
        if index + 1 < len(lines) and lines[index].lstrip().startswith("|") and lines[index + 1].lstrip().startswith("|"):
            cells = [cell.strip() for cell in lines[index + 1].strip().strip("|").split("|")]
            header = lines[index].lower()
            if len(cells) >= 2 and all(cell and set(cell) <= {"-", ":"} and "-" in cell for cell in cells):
                end = index + 2
                while end < len(lines) and lines[end].lstrip().startswith("|"):
                    end += 1
                rows = end - index - 2
                candidate_header = any(word in header for word in ("antisense", "target", "candidate"))
                if candidate_header and rows in {0, source_row_count}:
                    if not replaced:
                        output.append(marker)
                    replaced += 1
                    index = end
                    continue
        output.append(lines[index])
        index += 1
    return "\n".join(output) + ("\n" if text.endswith("\n") else ""), replaced


def _validate_eligibility_reason_columns(path):
    """A reason for failure cannot accompany an observed pass or unknown result."""
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle, delimiter="\t" if path.suffix.lower() == ".tsv" else ",")
        if not {"eligibility", "fail_reasons"}.issubset(reader.fieldnames or []):
            return
        for number, row in enumerate(reader, 2):
            status = str(row["eligibility"] or "").strip().lower()
            reasons = str(row["fail_reasons"] or "").strip()
            if status in {"eligible", "pass", "passed", "indeterminate", "unknown"} and reasons:
                raise ValueError(
                    f"CSV row {number} has eligibility={status} but nonempty fail_reasons. "
                    "Put successful checks and missing requirements in separate fields; "
                    "fail_reasons may contain only observed failed criteria."
                )


def _markdown_tables(text):
    """Read simple pipe tables for structural cross-artifact checks."""
    lines = text.splitlines()
    for index in range(len(lines) - 1):
        header, separator = lines[index:index + 2]
        if not (header.lstrip().startswith("|") and separator.lstrip().startswith("|")):
            continue
        names = [cell.strip().lower() for cell in header.strip().strip("|").split("|")]
        rules = [cell.strip() for cell in separator.strip().strip("|").split("|")]
        if len(names) != len(rules) or not all(
            rule and set(rule) <= {"-", ":"} and "-" in rule for rule in rules
        ):
            continue
        rows = []
        for line in lines[index + 2:]:
            if not line.lstrip().startswith("|"):
                break
            cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
            if len(cells) == len(names):
                rows.append(dict(zip(names, cells)))
        yield names, rows


def _validate_selection_report_views(texts, export_tables):
    """Keep selection counts and walk decisions bound to delivered tables.

    This applies only when the outputs themselves contain an eligible ranked
    universe and a strict selected subset with matching molecule/set keys.
    It never infers a scientific selection from a name or prose claim.
    """
    candidates = []
    for path in export_tables.values():
        with path.open(newline="", encoding="utf-8-sig") as handle:
            rows = list(csv.DictReader(handle, delimiter="\t" if path.suffix.lower() == ".tsv" else ","))
        if rows and {"id", "set", "rank"}.issubset(rows[0]):
            keys = [(row["id"], row["set"]) for row in rows]
            if len(set(keys)) == len(keys):
                candidates.append((set(keys), rows))
    pairs = [(full, chosen) for full_keys, full in candidates
             for chosen_keys, chosen in candidates
             if chosen_keys < full_keys]
    if len(pairs) != 1:
        return
    eligible, selected = pairs[0]
    eligible_keys = {(row["id"], row["set"]) for row in eligible}
    selected_keys = {(row["id"], row["set"]) for row in selected}
    eligible_counts = {name: sum(row["set"] == name for row in eligible)
                       for name in {row["set"] for row in eligible}}
    selected_counts = {name: sum(row["set"] == name for row in selected)
                       for name in eligible_counts}
    for destination, text in texts.items():
        for names, rows in _markdown_tables(text):
            columns = set(names)
            if {"set", "eligible", "selected"}.issubset(columns):
                for row in rows:
                    name = row["set"]
                    if name in eligible_counts and (
                        row["eligible"] != str(eligible_counts[name])
                        or row["selected"] != str(selected_counts[name])
                    ):
                        raise ValueError(
                            f"Selection count table in {Path(destination).name} disagrees "
                            "with the delivered ranked-universe and selected-panel CSV rows. "
                            "Derive counts directly from those outputs."
                        )
            if {"id", "set", "accept?"}.issubset(columns):
                for row in rows:
                    key = row["id"], row["set"]
                    if key not in eligible_keys:
                        continue
                    actual = row["accept?"].strip().lower()
                    if actual in {"yes", "true", "accept", "accepted"}:
                        claimed = True
                    elif actual in {"no", "false", "reject", "rejected"}:
                        claimed = False
                    else:
                        continue
                    if claimed != (key in selected_keys):
                        raise ValueError(
                            f"Selection walk in {Path(destination).name} disagrees "
                            "with the delivered selected-panel CSV. Derive each decision "
                            "and cumulative count from the same selected rows."
                        )


def _canonicalize_selection_report_views(texts, export_tables):
    """Replace contradictory count/walk cells using two delivered CSVs.

    The operation is deliberately narrow: both tables must carry unique
    molecule/set keys and numeric ranks; one must be a strict subset of the
    other. A Markdown table is changed only when its set or molecule keys
    exactly cover that recorded universe. Other prose and tables are untouched.
    """
    candidates = []
    for path in export_tables.values():
        with path.open(newline="", encoding="utf-8-sig") as handle:
            rows = list(csv.DictReader(handle, delimiter="\t" if path.suffix.lower() == ".tsv" else ","))
        if rows and {"id", "set", "rank"}.issubset(rows[0]):
            keys = [(row["id"], row["set"]) for row in rows]
            if len(set(keys)) == len(keys) and all(str(row["rank"]).isdigit() for row in rows):
                candidates.append((set(keys), rows, path))
    pairs = [(full, chosen, full_path, chosen_path)
             for full_keys, full, full_path in candidates
             for chosen_keys, chosen, chosen_path in candidates
             if chosen_keys < full_keys]
    if len(pairs) != 1:
        return texts, []
    eligible, selected, eligible_path, selected_path = pairs[0]
    eligible_keys = {(row["id"], row["set"]) for row in eligible}
    selected_keys = {(row["id"], row["set"]) for row in selected}
    sets = {row["set"] for row in eligible}
    eligible_counts = {name: sum(row["set"] == name for row in eligible) for name in sets}
    selected_counts = {name: sum(row["set"] == name for row in selected) for name in sets}
    amended, receipts = {}, []
    for destination, text in texts.items():
        lines = text.splitlines()
        index = 0
        changes = []
        while index + 1 < len(lines):
            header, separator = lines[index:index + 2]
            if not (header.lstrip().startswith("|") and separator.lstrip().startswith("|")):
                index += 1
                continue
            names = [cell.strip().lower() for cell in header.strip().strip("|").split("|")]
            rules = [cell.strip() for cell in separator.strip().strip("|").split("|")]
            if len(names) != len(rules) or not all(
                rule and set(rule) <= {"-", ":"} and "-" in rule for rule in rules
            ):
                index += 1
                continue
            end = index + 2
            while end < len(lines) and lines[end].lstrip().startswith("|"):
                end += 1
            cells = [[cell.strip() for cell in line.strip().strip("|").split("|")]
                     for line in lines[index + 2:end]]
            if all(len(row) == len(names) for row in cells):
                data = [dict(zip(names, row)) for row in cells]
                if ({"set", "eligible", "selected"}.issubset(names)
                        and {row["set"] for row in data} == sets
                        and len(data) == len(sets)):
                    for offset, row in enumerate(data, index + 2):
                        values = cells[offset - index - 2]
                        values[names.index("eligible")] = str(eligible_counts[row["set"]])
                        values[names.index("selected")] = str(selected_counts[row["set"]])
                        lines[offset] = "| " + " | ".join(values) + " |"
                    changes.append("per_set_counts")
                if ({"id", "set", "accept?"}.issubset(names)
                        and {(row["id"], row["set"]) for row in data} == eligible_keys
                        and len(data) == len(eligible_keys)):
                    cumulative = {}
                    for offset, row in enumerate(data, index + 2):
                        values = cells[offset - index - 2]
                        accepted = (row["id"], row["set"]) in selected_keys
                        if accepted:
                            cumulative[row["set"]] = cumulative.get(row["set"], 0) + 1
                        values[names.index("accept?")] = "yes" if accepted else "no"
                        if "set counts after" in names:
                            values[names.index("set counts after")] = json.dumps(cumulative, sort_keys=True)
                        lines[offset] = "| " + " | ".join(values) + " |"
                    changes.append("selection_walk")
            index = end
        amended[destination] = "\n".join(lines) + ("\n" if text.endswith("\n") else "")
        if changes:
            receipts.append({"destination": destination, "sections": changes,
                             "eligible_source": str(eligible_path),
                             "selected_source": str(selected_path),
                             "eligible_sha256": hashlib.sha256(eligible_path.read_bytes()).hexdigest(),
                             "selected_sha256": hashlib.sha256(selected_path.read_bytes()).hexdigest()})
    return amended, receipts


def render_report_tables(result, required_outputs, summary_path, *, required=False, source_tables=(), computed_tables=()):
    outputs = {str(Path(path).resolve()): Path(path) for path in required_outputs}
    export_tables = {key: path for key, path in outputs.items() if path.suffix.lower() in {".csv", ".tsv"}}
    workbook_exports = {key: path for key, path in outputs.items() if path.suffix.lower() == ".xlsx"}
    tables = dict(export_tables)
    if computed_tables and (export_tables or workbook_exports):
        raise ValueError("Intermediate computed tables cannot replace views of requested tabular outputs.")
    for source in computed_tables:
        path = Path(source).resolve()
        if path.suffix.lower() in {".csv", ".tsv"}:
            tables[str(path)] = path
        elif path.suffix.lower() == ".xlsx":
            workbook_exports[str(path)] = path
        else:
            raise ValueError("Computed report tables must be CSV, TSV or XLSX files.")
    for source in source_tables:
        path = Path(source).resolve()
        if path.suffix.lower() in {".csv", ".tsv"} and path.is_file():
            tables.setdefault(str(path), path)
    for path in export_tables.values():
        _validate_eligibility_reason_columns(path)
    def resolve_source(declared_source):
        source = Path(declared_source or "")
        if not source.is_absolute():
            source = Path(summary_path).resolve().parent / source
        return str(source.resolve())
    def file_digest(path):
        digest = hashlib.sha256()
        with path.open("rb") as handle:
            for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(chunk)
        return digest.digest()

    def table_key(declared_source):
        key = resolve_source(declared_source)
        if key in tables:
            return key
        # An exact staging copy may live beside the analysis program. It is
        # still bound to an actual requested export by relative name and bytes.
        staged_root = Path(summary_path).resolve().parent / "output"
        source = Path(key)
        try:
            relative = source.relative_to(staged_root)
        except ValueError:
            return key
        if not source.is_file():
            return key
        source_digest = file_digest(source)
        matching = [
            output_key for output_key, output_path in {**tables, **workbook_exports}.items()
            if output_path.name == relative.name
            and file_digest(output_path) == source_digest
        ]
        return matching[0] if len(matching) == 1 else key
    # Some programs serialize their result inside the standard sandbox_result
    # envelope. Its explicit table declarations have the same meaning; their
    # paths and columns still undergo all validation below. Never guess a source
    # from a marker's prose or replace conflicting declarations.
    nested = result.get("sandbox_result")
    nested_views = nested.get("report_tables") if isinstance(nested, dict) else None
    direct_views = result.get("report_tables")
    if direct_views and nested_views and direct_views != nested_views:
        raise ValueError("Conflicting report_tables declarations in the result and sandbox_result envelope.")
    views = deepcopy(direct_views or nested_views or [])
    if not isinstance(views, list):
        raise ValueError("report_tables must be an array of table declarations.")
    # A workbook sheet can contain a different population or computation from
    # an upstream native panel. Preserve its explicit identity; never replace
    # it with the only available CSV merely because no CSV export was requested.
    workbook_sources = {}
    for view in views:
        if not isinstance(view, dict):
            continue  # The regular declaration validation below reports this.
        key = table_key(view.get("source_path"))
        if Path(key).suffix.lower() == ".xlsx" and key not in workbook_exports:
            raise ValueError("A workbook report source must be an actual requested XLSX output.")
        if key in workbook_exports:
            sheet_name, header_row = view.get("source_sheet"), view.get("header_row", 1)
            data = _workbook_report_table(workbook_exports[key], sheet_name, header_row)
            workbook_sources[(key, sheet_name, header_row)] = data
    if not tables and not views:
        markdown = [Path(summary_path), *(path for path in outputs.values()
                                         if path.suffix.lower() == ".md")]
        if any("{{table:" in path.read_text() for path in markdown):
            raise ValueError(
                "A workbook report marker needs a report_tables declaration with "
                "the requested source_path and exact source_sheet."
            )
        return {}  # Workbook-only narrative reports have no implicit sheet choice.
    declaration_path = "report_tables" if direct_views else "sandbox_result.report_tables" if nested_views else None
    texts = {str(Path(summary_path).resolve()): Path(summary_path).read_text()}
    texts.update({key: path.read_text() for key, path in outputs.items() if path.suffix.lower() == ".md"})
    from .metric_report_integrity import canonicalize_metric_notes
    texts, metric_repairs = canonicalize_metric_notes(texts, export_tables)
    if metric_repairs:
        result["metric_report_canonicalization"] = metric_repairs
    texts, selection_repairs = _canonicalize_selection_report_views(texts, export_tables)
    if selection_repairs:
        result["selection_report_canonicalization"] = selection_repairs
    _validate_selection_report_views(texts, export_tables)
    original_model_texts = dict(texts)
    canonicalization = None
    if len(source_tables) == 1 and not export_tables and not workbook_sources and not computed_tables:
        source = Path(source_tables[0]).resolve()
        with source.open(newline="", encoding="utf-8-sig") as handle:
            reader = csv.DictReader(handle, delimiter="\t" if source.suffix.lower() == ".tsv" else ",")
            fields = reader.fieldnames or []
            source_row_count = sum(1 for _ in reader)
        preferred = ("candidate", "antisense_5to3", "target_5to3", "sequence",
                     "transcript_accession", "transcript_cdna_start", "transcript_cdna_end",
                     "exon_ordinal", "complete_coding_exon_site",
                     "genomic_chrom", "genomic_start_1based", "genomic_end_1based",
                     "target_genomic_strand", "chemistry_description", "length", "gc_fraction",
                     "max_homopolymer_run_nt", "native_rank", "rank",
                     "off_target_mismatch_budget", "off_target_exact_count", "off_target_exact_count_unit")
        columns = [field for field in preferred if field in fields] or fields[:10]
        original_views = direct_views or nested_views or []
        view_id = (original_views[0].get("id") if isinstance(original_views, list)
                   and len(original_views) == 1 and isinstance(original_views[0], dict)
                   and isinstance(original_views[0].get("id"), str)
                   and original_views[0]["id"] else "native_candidate_panel")
        marker = "{{table:" + view_id + "}}"
        replaced_total = 0
        for destination, body in list(texts.items()):
            revised, replaced = _replace_candidate_tables_with_marker(body, marker, source_row_count)
            texts[destination] = revised
            replaced_total += replaced
        if any(_has_manual_markdown_table(text) for text in texts.values()):
            raise _unbound_markdown_table_error(texts)
        views = [{"id": view_id, "source_path": str(source),
                  "columns": columns,
                  "rank_column": "native_rank" if "native_rank" in fields else "rank" if "rank" in fields else None,
                  "preview_limit": None}]
        declaration_path = "canonical_native_source"
        canonicalization = {"source_path": str(source),
            "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
            "replaced_manual_tables": replaced_total,
            "discarded_model_declarations": deepcopy(direct_views or nested_views or [])}
    elif source_tables and any(_has_manual_markdown_table(text) for text in texts.values()):
        raise _unbound_markdown_table_error(texts)
    originals = dict(texts)
    if required:
        # Rendering an actual exported table does not require another model to
        # describe its columns or transcribe its rows. Preserve source order when
        # no semantic rank is declared; never infer a scientific ordering.
        if not views and any("{{table:" in text for text in texts.values()):
            # An exact exported filename in a marker already identifies its
            # table. Bind that identity without asking the analysis model to
            # restate a redundant JSON declaration. Arbitrary marker names and
            # duplicate basenames remain ambiguous and are rejected.
            markers = {
                tail.split("}}", 1)[0]
                for body in texts.values()
                for tail in body.split("{{table:")[1:]
                if "}}" in tail
            }
            matches = {
                marker: [key for key, path in tables.items() if path.name == marker]
                for marker in markers
            }
            if not markers or any(len(paths) != 1 for paths in matches.values()):
                raise ValueError("Declare report_tables for explicit table markers; their source identity is ambiguous.")
            views.extend({"id": marker, "source_path": paths[0],
                          "rank_column": None, "preview_limit": None,
                          "native_source_order_view": True}
                         for marker, paths in sorted(matches.items()))
        existing = {table_key(view.get("source_path")) for view in views if isinstance(view, dict)}
        automatic_sources = export_tables if workbook_sources or computed_tables else tables
        for key, path in automatic_sources.items():
            if key in existing:
                continue
            view_id = "export_" + hashlib.sha256(key.encode()).hexdigest()[:16]
            views.append({"id": view_id, "source_path": key, "rank_column": None,
                          "preview_limit": None, "native_source_order_view": True})
        # A source-bound table in the execution summary does not make the
        # specifically requested Markdown export complete. With one exported
        # table and one named Markdown report, include that exact table there
        # too even if the model put its marker only in the summary.
        named_markdown = [key for key, path in outputs.items() if path.suffix.lower() == ".md"]
        if len(export_tables) == len(views) == len(named_markdown) == 1:
            view = views[0]
            if isinstance(view, dict) and isinstance(view.get("id"), str):
                marker = "{{table:" + view["id"] + "}}"
                destination = named_markdown[0]
                if marker not in texts[destination]:
                    texts[destination] += "\n\n### Complete exported source table\n\n" + marker + "\n"
        for view in views:
            if not isinstance(view, dict) or not isinstance(view.get("id"), str):
                continue  # Normal validation below reports malformed views.
            marker = "{{table:" + view["id"] + "}}"
            if not any(marker in text for text in texts.values()):
                for destination in texts:
                    texts[destination] += "\n\n### " + Path(view.get("source_path") or "").name + "\n\n" + marker + "\n"
    # Typed audits describe eligibility and discrepancies from the delivered
    # table, rather than free-form illustrative row anecdotes. Bind the same
    # exact table keys used by the normal renderer; retain historical templates.
    from .audit_reconciliation_report import canonical_audit_texts
    audit_views = [{**view, "source_path": table_key(view.get("source_path"))}
                   for view in views if isinstance(view, dict)]
    texts, audit_repair = canonical_audit_texts(result, texts, export_tables, audit_views)
    if audit_repair:
        result.setdefault("historical_model_report_templates", original_model_texts)
        result["audit_reconciliation_canonicalization"] = audit_repair
        original_model_texts = dict(texts)
        originals = dict(texts)
    source_data = {}
    for source_key, source_path in tables.items():
        with source_path.open(newline="", encoding="utf-8-sig") as handle:
            reader = csv.DictReader(handle, delimiter="\t" if source_path.suffix.lower() == ".tsv" else ",")
            source_data[source_key] = (reader.fieldnames or [], list(reader))
    source_data.update(workbook_sources)
    declared, ids, receipts = set(), set(), []
    for view in views:
        if not isinstance(view, dict) or not isinstance(view.get("id"), str) or not view["id"]:
            raise ValueError("Each report table needs a distinct nonempty id.")
        view_id = view["id"]
        if view_id in ids:
            raise ValueError("Duplicate report table id: " + view_id)
        ids.add(view_id)
        declared_source = str(view.get("source_path") or "")
        path_key = table_key(declared_source)
        if path_key != declared_source:
            view["source_path"] = path_key
        if path_key not in tables and path_key not in workbook_exports:
            raise ValueError(
                "A report table source_path must be an actual requested CSV/TSV output, "
                "a requested XLSX output, "
                "or an exact declared read-only report_source_tables path. "
                "For XLSX use source_sheet and optional header_row (default 1). "
                "Do not create a competing unrequested table."
            )
        if path_key in workbook_exports:
            path = workbook_exports[path_key]
            key = (path_key, view["source_sheet"], view.get("header_row", 1))
        else:
            path, key = tables[path_key], path_key
        fields, rows = source_data[key]
        columns = view.get("columns") or fields
        if not isinstance(columns, list) or not columns or any(column not in fields for column in columns):
            raise ValueError("Report columns must refer to the source table's actual columns.")
        if required:
            for destination, text in originals.items():
                repeated = _repeated_export_table_columns(
                    text, fields, rows,
                    alternative_sources=[value for source_key, value in source_data.items()
                                         if source_key != key],
                )
                if repeated:
                    raise ValueError(
                        "A Markdown report independently transcribes the declared CSV/TSV row table "
                        f"with values that cannot be verified against the export ({Path(destination).name}: "
                        f"{', '.join(repeated)}). Use the declared {{table:ID}} marker or correct the "
                        "manual table so measured zeros and missing values match the exported data."
                    )
        rank_column = view.get("rank_column")
        if rank_column is not None:
            if rank_column not in fields:
                raise ValueError("rank_column is absent from the exported table.")
            missing = view.get("missing_rank_values", ["", "NA", "NaN", "null", "None"])
            def rank(item):
                index, row = item
                value = row[rank_column]
                if value in missing:
                    return float("inf"), index
                try:
                    number = float(value)
                except (TypeError, ValueError):
                    raise ValueError("Rank values must be positive integers or declared missing markers.")
                if not math.isfinite(number) or not number.is_integer() or number <= 0:
                    raise ValueError("Rank values must be positive integers or declared missing markers.")
                return number, index
            rows = [row for _, row in sorted(enumerate(rows), key=rank)]
        limit = view.get("preview_limit")
        if limit is not None and (not isinstance(limit, int) or isinstance(limit, bool) or limit <= 0):
            raise ValueError("preview_limit must be a positive integer or null for the complete table.")
        shown = rows if limit is None else rows[:limit]
        def cell(value):
            return str(value).replace("|", "\\|").replace("\n", " ")
        lines = ["| " + " | ".join(cell(column) for column in columns) + " |",
                 "| " + " | ".join("---" for _ in columns) + " |"]
        lines.extend("| " + " | ".join(cell(row[column]) for column in columns) + " |" for row in shown)
        lines.extend(["", f"Shown: {len(shown)} of {len(rows)} exported rows. " +
                      ("Ordered by the recorded numeric rank; missing ranks remain unranked." if rank_column else "Recorded source order; no additional ranking.")])
        marker = "{{table:" + view_id + "}}"
        if not any(marker in text for text in texts.values()):
            raise ValueError("Missing report-table marker: " + marker)
        for destination, text in list(texts.items()):
            texts[destination] = text.replace(marker, "\n".join(lines))
        declared.add(key)
        receipt = {"id": view_id, "source_path": path_key, "declared_source_path": declared_source,
                         "source_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                         "rank_column": rank_column, "row_count": len(rows), "shown_rows": len(shown),
                         "native_source_order_view": bool(view.get("native_source_order_view"))}
        if path_key in workbook_exports:
            receipt.update(source_sheet=view["source_sheet"], header_row=view.get("header_row", 1))
        receipts.append(receipt)
    if required and not set(export_tables).issubset(declared):
        raise ValueError("Every requested CSV/TSV output needs a report_tables view; missing: " + ", ".join(sorted(set(export_tables) - declared)))
    if any("{{table:" in text for text in texts.values()):
        raise ValueError("An unknown report-table marker remains unresolved.")
    if views:
        result["report_tables"] = views
        result["report_table_declarations_path"] = declaration_path
        result["report_table_rendering"] = receipts
        result["original_report_templates"] = original_model_texts
        if canonicalization:
            result["source_table_canonicalization"] = canonicalization
        for destination, text in texts.items():
            Path(destination).write_text(text)
    return {"views": receipts}
