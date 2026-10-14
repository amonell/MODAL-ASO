"""Versioned recipient exports from attested chemistry; no model or reranking.

Mixed PS/PO molecules can be represented and preserved here. Native design
selection still rejects them until its complete scientific route supports them.
"""
import csv
import hashlib
import html
import json
from pathlib import Path
from aso_molecule import (ASSUMPTION_STATUSES, molecule_record, parse_ordering_notation,
                         validate_parsed_chemistry, validate_molecule_record)

SCHEMA_VERSION = "modal-aso-recipient/1.2"
LEGACY_FIELDS = ("candidate_id", "native_rank", "antisense_5to3", "architecture", "linkages_5to3", "chemistry_assumption_status", "nucleotide_chemistry_5to3", "ordering_notation")
FIELDS = (*LEGACY_FIELDS, "molecule_specification")
STATUS = ASSUMPTION_STATUSES


def attest_record(candidate_id, sequence, architecture, notation, *, linkages, assumption_status, native_rank=None):
    if not candidate_id:
        raise ValueError("Candidate identity is required.")
    molecule = molecule_record(sequence, architecture, notation, linkages=linkages,
                               assumption_status=assumption_status)
    values = (candidate_id, native_rank, sequence, molecule["architecture"], molecule["linkages_5to3"],
              assumption_status, molecule["nucleotide_chemistry_5to3"], notation or "", molecule)
    return dict(zip(FIELDS, values))


def _flat(value):
    return value if isinstance(value, str) else json.dumps(value, separators=(",", ":"))


def _validate_records(records, version=SCHEMA_VERSION):
    if not records or len({row["candidate_id"] for row in records}) != len(records):
        raise ValueError("Recipient records must have unique nonempty identities.")
    for row in records:
        molecule = row.get("molecule_specification")
        links = (None if molecule and molecule.get("ordering_status") != "attested"
                 else row["linkages_5to3"])
        rebuilt = attest_record(row["candidate_id"], row["antisense_5to3"], row["architecture"], row["ordering_notation"],
                                linkages=links, assumption_status=row["chemistry_assumption_status"], native_rank=row["native_rank"])
        if version in {"modal-aso-recipient/1.0", "modal-aso-recipient/1.1"}:
            rebuilt = {key: rebuilt[key] for key in LEGACY_FIELDS}
            if version == "modal-aso-recipient/1.0":
                rebuilt.pop("nucleotide_chemistry_5to3", None)
            else:
                # Legacy 1.1 collapsed the MOE-MeC base modification into sugar.
                # Validate its historical projection while separately re-attesting
                # the actual token chemistry; never rewrite the source molecule.
                rebuilt["nucleotide_chemistry_5to3"] = [
                    "2MOE" if item == "2MOE_methylC" else item
                    for item in rebuilt["nucleotide_chemistry_5to3"]]
        if row != rebuilt:
            raise ValueError("Recipient record has unsupported, contradictory or missing fields.")


def _recipient_markdown(payload):
    def cell(value):
        return html.escape(str(value)).replace('|', '&#124;').replace('\n', ' ')
    lines = [f"# {payload['title']}", '', payload['context'], '',
             'ASO sequences are shown 5′→3′. Recorded ranks are preserved. Chemistry assumption status describes its provenance; it does not establish biological activity.', '',
             '[Candidate workbook](candidates.xlsx) · [Candidate CSV](candidates.csv) · [Complete molecule records](candidates.json)', '',
             '| Candidate | Recorded rank | ASO (5′→3′) | Architecture | Chemistry assumption | Ordering status |',
             '|---|---:|---|---|---|---|']
    for row in payload['records']:
        molecule = row['molecule_specification']
        values = [row['candidate_id'], row['native_rank'] if row['native_rank'] is not None else 'Not recorded',
                  row['antisense_5to3'], row['architecture'], row['chemistry_assumption_status'], molecule['ordering_status']]
        lines.append('| ' + ' | '.join(cell(value) for value in values) + ' |')
    lines.extend(['', '## Per-candidate chemistry', '',
                  'Linkages run 5′→3′; entry 1 joins nucleotides 1 and 2. Sugar identity, base modifications, linkages and ordering notation are separate. Unsupported or unspecified ordering chemistry is not replaced with a different chemistry.', ''])
    for row in payload['records']:
        molecule = row['molecule_specification']
        lines.extend([f"### {row['candidate_id']}", '',
                      'Nucleotide chemistry (5′→3′): ' + ', '.join(row['nucleotide_chemistry_5to3']) + '.', '',
                      'Linkages (5′→3′): ' + ', '.join(row['linkages_5to3']) + '.', '',
                      'Ordering notation:', '', '```text', row['ordering_notation'] or 'Not available for this molecule.', '```', ''])
    # Exact, independently verifiable projection remains available in Markdown.
    lines.extend(['## Complete molecule records', '', '```json', json.dumps(payload, indent=2, ensure_ascii=False), '```', ''])
    return '\n'.join(lines)


def _recipient_html(payload, report, directory):
    from aso_agents.report_presentation import TABLE_CSS, markdown_html
    # Full records stay downloadable; the HTML opens on the scientific summary,
    # not a many-page JSON dump. Embed the exact payload for integrity checking.
    readable = report.split('## Complete molecule records\n', 1)[0]
    body = markdown_html(readable, directory)
    evidence = json.dumps(payload, ensure_ascii=True).replace('</', '<\\/')
    return ('<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
            '<title>' + html.escape(payload['title']) + '</title><style>'
            'body{font-family:Arial,sans-serif;color:#153a64;max-width:1200px;margin:32px auto;padding:0 20px;background:#fff}h1{font-size:2rem}h2{margin-top:32px}'
            + TABLE_CSS + '</style></head><body><main class="report-content">' + body
            + '</main><script type="application/json" id="molecule-records">' + evidence + '</script></body></html>')


def write_recipient_report(directory, records, *, title, context, provenance):
    """Create new outputs only; every view is derived from the same records."""
    from openpyxl import Workbook
    _validate_records(records)
    directory = Path(directory)
    if directory.exists():
        raise FileExistsError("Recipient exports require a new directory; originals are immutable.")
    directory.mkdir(parents=True)
    payload = {"schema_version": SCHEMA_VERSION, "title": title, "context": context, "provenance": provenance, "records": records}
    (directory / "candidates.json").write_text(json.dumps(payload, indent=2, ensure_ascii=False))
    flat = [[_flat(row[key]) for key in FIELDS] for row in records]
    if any(len(value) > 32767 for row in flat for value in row):
        raise ValueError("Workbook cell would truncate attested evidence; use a separate evidence artifact.")
    with (directory / "candidates.csv").open("w", newline="") as stream:
        writer = csv.writer(stream); writer.writerow(FIELDS); writer.writerows(flat)
    workbook = Workbook(); sheet = workbook.active; sheet.title = "Candidates"; sheet.append(FIELDS)
    for row in flat:
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            cell.data_type = "s"
    sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions; workbook.save(directory / "candidates.xlsx")
    # HTML escapes vendor slash/star strings. JSON block in Markdown preserves
    # exact notation, including stars, without a second authored chemistry view.
    report = _recipient_markdown(payload)
    (directory / "report.md").write_text(report)
    (directory / "report.html").write_text(_recipient_html(payload, report, directory))
    validate_recipient_report(directory)
    manifest = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in directory.iterdir() if p.is_file()}
    (directory / "checksums.json").write_text(json.dumps(manifest, indent=2))
    return payload


def validate_recipient_report(directory):
    from openpyxl import load_workbook
    directory = Path(directory); payload = json.loads((directory / "candidates.json").read_text())
    version = payload.get("schema_version")
    if version not in {SCHEMA_VERSION, "modal-aso-recipient/1.0", "modal-aso-recipient/1.1"}:
        raise ValueError("Unsupported recipient schema version.")
    records = payload["records"]; _validate_records(records, version)
    fields = FIELDS if version == SCHEMA_VERSION else tuple(k for k in LEGACY_FIELDS if version != "modal-aso-recipient/1.0" or k != "nucleotide_chemistry_5to3")
    expected = [list(fields)] + [[_flat(row[key]) for key in fields] for row in records]
    with (directory / "candidates.csv").open(newline="") as stream:
        if list(csv.reader(stream)) != expected: raise ValueError("CSV chemistry view diverges from JSON.")
    book = load_workbook(directory / "candidates.xlsx", read_only=True)
    try:
        if [["" if value is None else value for value in row] for row in book.active.values] != expected: raise ValueError("Workbook chemistry view diverges from JSON.")
    finally: book.close()
    report = (directory / "report.md").read_text()
    try: projected = json.loads(report.split("```json\n", 1)[1].split("\n```", 1)[0])
    except (IndexError, ValueError) as exc: raise ValueError("Report lacks its attested chemistry projection.") from exc
    if projected != payload: raise ValueError("Report chemistry view diverges from JSON.")
    actual_html = (directory / "report.html").read_text()
    # Historical exports remain valid; newly generated views must reproduce the
    # canonical payload in their readable content as well as their embedded data.
    if 'id="molecule-records"' in actual_html and report != _recipient_markdown(payload):
        raise ValueError("Readable Markdown chemistry view diverges from JSON.")
    expected_html = _recipient_html(payload, report, directory) if 'id="molecule-records"' in actual_html else '<!doctype html><meta charset="utf-8"><pre>'+html.escape(report)+'</pre>'
    if actual_html != expected_html:
        raise ValueError("HTML chemistry view diverges from Markdown.")
    return {"schema_version": version, "records_checked": len(records), "views_checked": ["JSON", "CSV", "XLSX", "Markdown", "HTML"]}
