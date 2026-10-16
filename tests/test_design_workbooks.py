import copy
import json
from pathlib import Path

import openpyxl
import pytest

from aso_agents import pipeline_tools as native
from aso_agents.design_workbooks import CORE_COLUMNS, DesignWorkbook, validate_design_workbooks
from test_design_observation_tables import example, project, write


def context(root, *, second=False, conflict=False):
    row, goal, source = example(root)
    row.update(primary_delta_deletion=0.0, primary_delta_replacement=-0.25,
               off_target_search_status="complete", Off_target_mismatch_budget=1,
               splice_site_deltas=[{"label": "annotated", "site_kind": "donor",
                                    "delta_donor_deletion": 0.0, "delta_donor_replacement": None}])
    rows = [row]
    if second:
        goal["design_programs"].append({"program_id": "other", "transcript_id": "TX.2"})
        write(root, "structured_goal.json", goal)
        rows.append({**row, "program_id": "other"})
    source.write_text(json.dumps({"schema": {"table": "aso_tiles"}, "rows": rows}))
    if conflict:
        scored = write(root, "scored_tiles.json", {"schema": {"table": "aso_tiles"},
                       "rows": [{**row, "Off_target_exact_count": 7, "primary_delta_deletion": -0.7}]})
        write(root, "generated_filter_trace.json", {"sources": {
            "raw_generated": [str(source)], "scored": [str(scored)]}})
    meta, _ = project(root)
    declared = {"completed_run": {"design_evidence": [{"task_id": "design", "observations": meta}]}}
    path = root / "requested.xlsx"
    payload = {"declared_inputs": declared, "source_bound_design_workbooks": [str(path)]}
    return declared, path, payload


def make_book(declared, path, *, stage="raw_generated", coverage="all_raw"):
    writer = DesignWorkbook(declared)
    rows = writer.rows(stage=stage, task_id="design")
    writer.add_table("Candidates", rows, coverage=coverage, task_id="design",
                     extra_columns=["ID", "eligible"],
                     extra_values=[{"ID": f"candidate-{i}", "eligible": None} for i in range(len(rows))],
                     source_columns={"Unknown measurement": "recorded:/Off_target_2mm_count"})
    return writer, writer.save(path)


def test_copies_sequences_exact_metrics_zero_geometry_and_all_prediction_records(tmp_path):
    declared, path, payload = context(tmp_path)
    writer, result = make_book(declared, path)
    validate_design_workbooks(payload, result, [path])
    book = openpyxl.load_workbook(path, data_only=True)
    values = list(book["Candidates"].values)
    row = dict(zip(values[0], values[1]))
    assert row["target_5to3"] == "AACCGTAACCGTAACCG"
    assert row["antisense_5to3"] == "CGGTTACGGTTACGGTT"
    assert row["off_target_exact_count"] == 0 and row["Unknown measurement"] is None
    assert row["gc_bases"] == 9 and row["length_nt"] == 17
    assert row["gc_fraction"] == pytest.approx(9 / 17)
    assert json.loads(row["physical_splice_elements"])[0]["exonic_overlap_bp"] == 3
    assert row["primary_delta_deletion"] == 0 and row["primary_delta_replacement"] == -0.25
    assert book["Recorded splice evidence"].max_row == 3
    assert result["design_fact_validation"]["rows_checked"] == 1


@pytest.mark.parametrize("column,bad", [
    ("target_5to3", "T" * 17), ("antisense_5to3", "A" * 17),
    ("off_target_exact_count", None), ("gc_fraction", 0.53),
    ("physical_splice_elements", "[]"), ("genomic_blocks_bed0_half_open", "[]"),
    ("longest_homopolymer", 7), ("primary_delta_replacement", None)])
def test_actual_cell_corruption_is_rejected(tmp_path, column, bad):
    declared, path, payload = context(tmp_path)
    _, result = make_book(declared, path)
    book = openpyxl.load_workbook(path)
    headers = [c.value for c in book["Candidates"][1]]
    book["Candidates"].cell(2, headers.index(column) + 1).value = bad
    book.save(path)
    with pytest.raises(ValueError, match="contradicts source"):
        validate_design_workbooks(payload, result, [path])


def test_same_sequence_different_programs_remains_distinct_and_complete_coverage_is_checked(tmp_path):
    declared, path, payload = context(tmp_path, second=True)
    writer = DesignWorkbook(declared)
    rows = writer.rows(stage="raw_generated")
    assert len(rows) == 2 and rows[0]["binding_id"] != rows[1]["binding_id"]
    writer.add_table("Truncated", rows[:1], coverage="all_raw", task_id="design")
    result = writer.save(path)
    with pytest.raises(ValueError, match="complete source stage"):
        validate_design_workbooks(payload, result, [path])
    result["design_tables"][0]["coverage"] = "subset"
    validate_design_workbooks(payload, result, [path])
    result["design_tables"][0]["task_id"] = "different-task"
    with pytest.raises(ValueError, match="another producer"):
        validate_design_workbooks(payload, result, [path])


def test_conflicting_measurements_stay_unresolved_with_both_sources(tmp_path):
    declared, path, payload = context(tmp_path, conflict=True)
    writer, result = make_book(declared, path)
    row, = writer.rows()
    assert row["off_target_exact_count"] is None and row["primary_delta_deletion"] is None
    assert row["evidence_conflicts"]["off_target_exact_count"] == [0, 7]
    assert len(row["splice_predictions"]) == 4
    validate_design_workbooks(payload, result, [path])


@pytest.mark.parametrize("corruption", ["row", "sheet", "value"])
def test_missing_or_changed_prediction_evidence_is_rejected(tmp_path, corruption):
    declared, path, payload = context(tmp_path)
    _, result = make_book(declared, path)
    book = openpyxl.load_workbook(path)
    sheet = book["Recorded splice evidence"]
    if corruption == "row":
        sheet.delete_rows(3)
    elif corruption == "sheet":
        book.remove(sheet)
    else:
        headers = [c.value for c in sheet[1]]
        sheet.cell(2, headers.index("primary_delta_deletion") + 1).value = None
    book.save(path)
    with pytest.raises(ValueError, match="prediction"):
        validate_design_workbooks(payload, result, [path])


def test_source_mutation_and_unbound_export_are_rejected(tmp_path):
    declared, path, payload = context(tmp_path)
    _, result = make_book(declared, path)
    with pytest.raises(ValueError, match="source-bound design_tables"):
        validate_design_workbooks(payload, {}, [path])
    source = Path(declared["completed_run"]["design_evidence"][0]["observations"]["path"])
    source.write_text(source.read_text() + " ")
    with pytest.raises(ValueError, match="hash changed"):
        validate_design_workbooks(payload, result, [path])


def test_writer_rejects_changed_source_rows_but_accepts_derived_columns(tmp_path):
    declared, path, payload = context(tmp_path)
    writer = DesignWorkbook(declared)
    rows = writer.rows()
    rows[0]["target_5to3"] = "T" * 17
    with pytest.raises(ValueError, match="row was changed"):
        writer.add_table("Candidate", rows)


def test_native_custom_workbook_uses_source_bound_writer_and_validates_bytes(tmp_path, monkeypatch):
    example(tmp_path)
    destination = tmp_path / "library.xlsx"
    state = {"request": "Export all candidates.", "request_envelope": {
        "schema_version": "v3", "tasks": [{"task_id": "design", "task_type": "design_new_asos"}],
        "final_outputs": [{"kind": "custom", "task_id": "design", "path": str(destination)}]}}
    def generate(payload):
        assert payload["source_bound_design_workbooks"] == [str(destination)]
        declared = json.loads((tmp_path / "requested_artifacts/declared_inputs.json").read_text())
        assert declared["request_envelope"] == state["request_envelope"]
        assert declared["request_text"] == state["request"]
        return {"python_code": "from pathlib import Path\nimport json\n"
                "from aso_agents.design_workbooks import DesignWorkbook\n"
                "w=DesignWorkbook(json.loads(Path('declared_inputs.json').read_text()))\n"
                "w.add_table('Candidates',w.rows(stage='raw_generated'),coverage='all_raw')\n"
                f"result=w.save({str(destination)!r})\n"
                "Path('sandbox_result.json').write_text(json.dumps(result))\n"
                "Path('sandbox_summary.md').write_text('Recorded candidate facts exported.')\n"}
    monkeypatch.setattr(native, "_run_sandbox_analysis_sync", generate)
    result = native._materialize_requested_final_outputs(state, str(tmp_path))
    assert result["ok"], result
    saved = json.loads((tmp_path / "requested_artifacts/sandbox_result.json").read_text())
    assert saved["design_fact_validation"]["rows_checked"] == 1


def test_unrelated_workbooks_are_outside_design_validation_scope(tmp_path):
    validate_design_workbooks({"source_bound_design_workbooks": []}, {}, [tmp_path / "analysis.xlsx"])


def test_exact_original_native_rows_resolve_to_their_verified_catalog_binding(tmp_path):
    declared, path, payload = context(tmp_path)
    raw = json.loads(native._run_path(tmp_path, "tiles.json").read_text())["rows"]
    writer = DesignWorkbook(declared)
    writer.add_table("Native records", raw, task_id="design", coverage="all_raw")
    result = writer.save(path)
    validate_design_workbooks(payload, result, [path])
    assert result["design_fact_validation"]["rows_checked"] == 1


def test_native_binding_id_column_does_not_impersonate_the_catalog_schema(tmp_path):
    row, goal, raw = example(tmp_path)
    row["binding_id"] = "native-label"
    raw.write_text(json.dumps({"schema": {"table": "aso_tiles"}, "rows": [row]}))
    meta, _ = project(tmp_path)
    declared = {"completed_run": {"design_evidence": [{"task_id": "design", "observations": meta}]}}
    writer = DesignWorkbook(declared)
    writer.add_table("Original rows", [row], task_id="design")
    path = tmp_path / "native.xlsx"
    result = writer.save(path)
    validate_design_workbooks({"declared_inputs": declared, "source_bound_design_workbooks": [str(path)]}, result, [path])


@pytest.mark.parametrize("change", ["missing_field", "changed_zero", "extra_field"])
def test_native_row_adapter_requires_the_complete_unchanged_source_record(tmp_path, change):
    declared, path, payload = context(tmp_path)
    row = json.loads(native._run_path(tmp_path, "tiles.json").read_text())["rows"][0]
    if change == "missing_field":
        row.pop("IDT")
    elif change == "changed_zero":
        row["Off_target_exact_count"] = None
    else:
        row["rank"] = 1
    with pytest.raises(ValueError, match="exactly one unchanged"):
        DesignWorkbook(declared).add_table("Candidates", [row], task_id="design")


def test_identical_native_row_in_different_tasks_requires_unambiguous_scope(tmp_path):
    declared, path, payload = context(tmp_path)
    meta = declared["completed_run"]["design_evidence"][0]["observations"]
    data = json.loads(Path(meta["path"]).read_text())
    data["task_id"] = "other"
    data["rows"][0]["binding_identity"]["task_id"] = "other"
    data["rows"][0]["binding_match_key"] = "different-task-binding"
    copied = tmp_path / "other-observations.json"
    copied.write_text(json.dumps(data))
    import hashlib
    declared["completed_run"]["design_evidence"].append({"task_id": "other", "observations": {
        "path": str(copied), "sha256": hashlib.sha256(copied.read_bytes()).hexdigest()}})
    raw = json.loads(native._run_path(tmp_path, "tiles.json").read_text())["rows"]
    writer = DesignWorkbook(declared)
    with pytest.raises(ValueError, match="exactly one unchanged"):
        writer.add_table("Ambiguous", raw)
    writer.add_table("Scoped", raw, task_id="design")
    result = writer.save(path)
    validate_design_workbooks(payload, result, [path])
