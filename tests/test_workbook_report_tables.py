import hashlib
from pathlib import Path

import openpyxl
import pytest

from aso_agents.tabular_reports import render_report_tables


def workbook(path, sheets):
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    book.save(path)
    book.close()
    return path


def reports(tmp_path, markers):
    summary = tmp_path / "sandbox_summary.md"
    report = tmp_path / "report.md"
    body = "# Calculations\n\n" + "\n\n".join("{{table:" + name + "}}" for name in markers)
    summary.write_text(body)
    report.write_text(body)
    return summary, report


@pytest.mark.parametrize("native_source", [False, True])
def test_requested_workbook_view_retains_computed_population(tmp_path, native_source):
    export = workbook(tmp_path / "results.xlsx", {
        "Observations": [["id", "measurement", "rank"],
                         ["new-c", 0, 3], ["new-a", 0.125, 1], ["new-b", None, 2]],
        "Methods": [["Method", "Description"], ["calculation", "from supplied observations"]],
    })
    summary, report = reports(tmp_path, ["observations"])
    upstream = tmp_path / "upstream.csv"
    upstream.write_text("candidate,sequence\nold-panel-row,ACGT\n")
    original = export.read_bytes()
    result = {"report_tables": [{"id": "observations", "source_path": str(export),
                                  "source_sheet": "Observations", "rank_column": "rank"}]}
    render_report_tables(result, [export, report], summary, required=True,
                         source_tables=[upstream] if native_source else [])
    for destination in (summary, report):
        text = destination.read_text()
        assert "{{table:" not in text
        assert "| new-a | 0.125 | 1 |" in text
        assert "| new-c | 0 | 3 |" in text
        assert "| new-b |  | 2 |" in text
        assert text.index("new-a") < text.index("new-b") < text.index("new-c")
        assert "old-panel-row" not in text
        assert "Shown: 3 of 3 exported rows." in text
    assert export.read_bytes() == original
    receipt, = result["report_table_rendering"]
    assert receipt["source_path"] == str(export)
    assert receipt["source_sheet"] == "Observations"
    assert receipt["header_row"] == 1
    assert receipt["source_sha256"] == hashlib.sha256(original).hexdigest()
    assert receipt["row_count"] == receipt["shown_rows"] == 3


def test_multiple_workbook_sheets_have_distinct_identity_and_projections(tmp_path):
    export = workbook(tmp_path / "results.xlsx", {
        "Rows": [["title"], ["id", "rank", "keep"], ["r2", 2, False],
                 ["r1", 1, True], ["unranked", None, None]],
        "Pairs": [["id", "rank", "count"], ["pair-b", 2, 0], ["pair-a", 1, 7]],
    })
    summary, report = reports(tmp_path, ["rows", "pairs"])
    result = {"report_tables": [
        {"id": "rows", "source_path": "results.xlsx", "source_sheet": "Rows",
         "header_row": 2, "columns": ["id", "keep"], "rank_column": "rank"},
        {"id": "pairs", "source_path": "results.xlsx", "source_sheet": "Pairs",
         "columns": ["id", "count"], "rank_column": "rank", "preview_limit": 1},
    ]}
    render_report_tables(result, [export, report], summary, required=True)
    text = report.read_text()
    assert "| r1 | True |" in text and "| r2 | False |" in text
    assert text.index("r1") < text.index("r2") < text.index("unranked")
    assert "| pair-a | 7 |" in text and "pair-b" not in text
    assert "Shown: 1 of 2 exported rows." in text
    assert [(r["source_sheet"], r["header_row"], r["row_count"])
            for r in result["report_table_rendering"]] == [("Rows", 2, 3), ("Pairs", 1, 2)]


@pytest.mark.parametrize("change,message", [
    ({"source_sheet": "absent"}, "worksheet"),
    ({"source_sheet": None}, "source_sheet"),
    ({"header_row": True}, "header_row"),
    ({"header_row": 0}, "header_row"),
    ({"header_row": 20}, "header_row"),
    ({"columns": ["invented"]}, "actual columns"),
])
def test_workbook_view_rejects_unestablished_schema(tmp_path, change, message):
    export = workbook(tmp_path / "results.xlsx", {"Rows": [["id", "value"], ["a", 2]]})
    summary, report = reports(tmp_path, ["rows"])
    view = {"id": "rows", "source_path": str(export), "source_sheet": "Rows", **change}
    original = (summary.read_bytes(), report.read_bytes(), export.read_bytes())
    with pytest.raises(ValueError, match=message):
        render_report_tables({"report_tables": [view]}, [export, report], summary, required=True)
    assert (summary.read_bytes(), report.read_bytes(), export.read_bytes()) == original


@pytest.mark.parametrize("rows,message", [
    ([["id", "id"], ["a", "b"]], "unique"),
    ([["id", None, "value"], ["a", None, 1]], "nonempty"),
    ([["id", "value"], ["a", "=1+1"]], "formula"),
    ([["id", "value"], ["a", 1, "unheaded"]], "header"),
])
def test_workbook_view_rejects_ambiguous_or_unevaluated_cells(tmp_path, rows, message):
    export = workbook(tmp_path / "results.xlsx", {"Rows": rows})
    summary, report = reports(tmp_path, ["rows"])
    with pytest.raises(ValueError, match=message):
        render_report_tables({"report_tables": [{"id": "rows", "source_path": str(export),
                                                  "source_sheet": "Rows"}]},
                             [export, report], summary, required=True)


def test_unrequested_workbook_cannot_be_a_report_source(tmp_path):
    export = workbook(tmp_path / "requested.xlsx", {"Rows": [["id"], ["right"]]})
    other = workbook(tmp_path / "other.xlsx", {"Rows": [["id"], ["wrong"]]})
    summary, report = reports(tmp_path, ["rows"])
    with pytest.raises(ValueError, match="requested"):
        render_report_tables({"report_tables": [{"id": "rows", "source_path": str(other),
                                                  "source_sheet": "Rows"}]},
                             [export, report], summary, required=True)


def test_workbook_marker_cannot_silently_survive_without_a_sheet_binding(tmp_path):
    export = workbook(tmp_path / "results.xlsx", {"Rows": [["id"], ["row"]]})
    summary, report = reports(tmp_path, ["unbound"])
    with pytest.raises(ValueError, match="source_sheet"):
        render_report_tables({}, [export, report], summary, required=True)
    summary.write_text("See the delivered workbook for the observations.")
    report.write_text(summary.read_text())
    assert render_report_tables({}, [export, report], summary, required=True) == {}


def test_csv_and_workbook_views_share_validation_without_mixing_rows(tmp_path):
    export = workbook(tmp_path / "results.xlsx", {"Rows": [["id", "count"], ["book", 0]]})
    csv = tmp_path / "counts.csv"
    csv.write_text("id,count\ntext,8\n")
    summary, report = reports(tmp_path, ["book", "text"])
    result = {"report_tables": [
        {"id": "book", "source_path": str(export), "source_sheet": "Rows"},
        {"id": "text", "source_path": str(csv)},
    ]}
    render_report_tables(result, [export, csv, report], summary, required=True)
    assert "| book | 0 |" in report.read_text()
    assert "| text | 8 |" in report.read_text()
    assert len(result["report_table_rendering"]) == 2


@pytest.mark.parametrize("identical", [True, False])
def test_workbook_staging_copy_requires_exact_delivered_bytes(tmp_path, identical):
    export = workbook(tmp_path / "results.xlsx", {"Rows": [["id"], ["delivered"]]})
    staged = tmp_path / "output" / "results.xlsx"
    staged.parent.mkdir()
    if identical:
        staged.write_bytes(export.read_bytes())
    else:
        workbook(staged, {"Rows": [["id"], ["different"]]})
    summary, report = reports(tmp_path, ["rows"])
    result = {"report_tables": [{"id": "rows", "source_path": str(staged),
                                  "source_sheet": "Rows"}]}
    if not identical:
        with pytest.raises(ValueError, match="actual requested XLSX"):
            render_report_tables(result, [export, report], summary, required=True)
    else:
        render_report_tables(result, [export, report], summary, required=True)
        assert "| delivered |" in report.read_text()
        receipt, = result["report_table_rendering"]
        assert receipt["source_path"] == str(export)
        assert receipt["declared_source_path"] == str(staged)


def test_analysis_execution_renders_requested_workbook_and_markdown(tmp_path):
    from aso_agents.analysis_execution import execute_analysis_program

    export, report = tmp_path / "results.xlsx", tmp_path / "report.md"
    code = f'''
import json
from pathlib import Path
from openpyxl import Workbook
b = Workbook()
s = b.active
s.title = "Values"
s.append(["id", "count"])
s.append(["observation", 0])
b.save({str(export)!r})
b.close()
Path("sandbox_summary.md").write_text("Results\\n{{{{table:values}}}}\\n")
Path({str(report)!r}).write_text("Results\\n{{{{table:values}}}}\\n")
Path("sandbox_result.json").write_text(json.dumps({{
    "report_tables": [{{"id": "values", "source_path": {str(export)!r},
                        "source_sheet": "Values"}}]
}}))
'''
    result = execute_analysis_program(directory=tmp_path,
        payload={"require_tabular_report_views": True},
        generate=lambda _: {"python_code": code}, required_outputs=[export, report], max_attempts=1)
    assert result["ok"], result
    assert "| observation | 0 |" in report.read_text()
    assert "{{table:" not in report.read_text()
