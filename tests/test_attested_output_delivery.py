"""Existing custom exports must pass delivery, while verified bytes are reusable."""

import hashlib
import json
import os
from pathlib import Path

import openpyxl
import pytest

from aso_agents import pipeline_tools as pipeline


def delivery_state(root, destination):
    goal = root / "artifacts/intake/structured_goal.json"
    goal.parent.mkdir(parents=True, exist_ok=True)
    goal.write_text("{}")
    return {"request": "Export the supplied observations with their measured values.",
            "request_envelope": {"schema_version": "v3",
                "tasks": [{"task_id": "evaluation", "task_type": "evaluate_existing_aso_set"}],
                "final_outputs": [{"kind": "file", "task_id": "evaluation",
                                   "path": str(destination)}]}}


def install_writer(monkeypatch, *, workbook=False):
    calls, reviews = [], []

    def generate(payload):
        calls.append(payload)
        contract = payload["execution_contract"]
        code = "from pathlib import Path\n"
        code += f"Path({contract['result_json']!r}).write_text('{{}}')\n"
        code += f"Path({contract['summary_markdown']!r}).write_text('Recorded observations exported.')\n"
        for path in payload["requested_output_paths"]:
            code += f"assert not Path({path!r}).exists()\n"
            if workbook:
                code += "from openpyxl import Workbook\nw=Workbook()\nw.active.title='Observations'\n"
                code += "w.active.append(['id', 'measured'])\nw.active.append(['A', 3])\n"
                code += f"w.save({path!r})\n"
            else:
                code += f"Path({path!r}).write_text('{{\"measured\": 3}}')\n"
        return {"python_code": code}

    def review(payload):
        reviews.append(payload)
        return {"program_sha256": payload["program_sha256"], "checks": [{
            "requirement_quote": payload["request_text"],
            "obligation": "Unit fixture output is exported",
            "status": "unverified", "implementation_evidence": "Isolated review boundary",
            "problem": None, "suggested_fix": None}]}

    monkeypatch.setattr(pipeline, "_run_sandbox_analysis_sync", generate)
    monkeypatch.setattr(pipeline, "_run_analysis_contract_review_sync", review)
    return calls, reviews


def test_generic_precreated_workbook_enters_review_and_preserves_original(tmp_path, monkeypatch):
    destination = tmp_path / "requested.xlsx"
    written, error = pipeline._write_v2_rows_excel(
        [{"id": "A", "measured": 3}], destination, sheet_groups=pipeline._V2_EXCEL_SHEETS)
    assert written and not error
    original = destination.read_bytes()
    calls, reviews = install_writer(monkeypatch, workbook=True)
    result = pipeline._materialize_requested_final_outputs(
        delivery_state(tmp_path, destination), str(tmp_path))
    assert result["ok"] and len(calls) == len(reviews) == 1
    book = openpyxl.load_workbook(destination)
    assert book.sheetnames == ["Observations"]
    assert list(book.active.values) == [("id", "measured"), ("A", 3)]
    book.close()
    archived = list(tmp_path.rglob("previous_outputs/*.json"))
    metadata = next(json.loads(p.read_text()) for p in archived
                    if json.loads(p.read_text())["path"] == str(destination))
    assert Path(metadata["archive"]).read_bytes() == original
    assert result["output_sha256"][str(destination)] == hashlib.sha256(destination.read_bytes()).hexdigest()


@pytest.mark.parametrize("suffix", [".json", ".csv", ".md"])
@pytest.mark.parametrize("legacy_receipt", [False, True])
def test_unattested_existing_custom_files_do_not_bypass_writer(tmp_path, monkeypatch, suffix, legacy_receipt):
    destination = tmp_path / ("requested" + suffix)
    destination.write_text('{"measured": 0}')
    if legacy_receipt:
        directory = pipeline._write_run_path(str(tmp_path), "requested_artifacts")
        directory.mkdir(parents=True)
        (directory / "delivery.json").write_text(json.dumps({"ok": True, "paths": [str(destination)]}))
    calls, reviews = install_writer(monkeypatch)
    result = pipeline._materialize_requested_final_outputs(
        delivery_state(tmp_path, destination), str(tmp_path))
    assert result["ok"] and len(calls) == len(reviews) == 1
    assert json.loads(destination.read_text()) == {"measured": 3}


def test_verified_delivery_is_reused_until_bytes_change_even_at_same_timestamp(tmp_path, monkeypatch):
    destination = tmp_path / "requested.json"
    state = delivery_state(tmp_path, destination)
    calls, reviews = install_writer(monkeypatch)
    assert pipeline._materialize_requested_final_outputs(state, str(tmp_path))["ok"]
    assert pipeline._materialize_requested_final_outputs(state, str(tmp_path))["already_present"]
    assert len(calls) == len(reviews) == 1
    stat = destination.stat()
    destination.write_text(destination.read_text().replace("3", "9"))
    os.utime(destination, ns=(stat.st_atime_ns, stat.st_mtime_ns))
    assert destination.stat().st_size == stat.st_size
    assert pipeline._materialize_requested_final_outputs(state, str(tmp_path))["ok"]
    assert len(calls) == len(reviews) == 2
    assert json.loads(destination.read_text()) == {"measured": 3}


def test_prior_success_does_not_attest_changed_request(tmp_path, monkeypatch):
    destination = tmp_path / "requested.json"
    state = delivery_state(tmp_path, destination)
    calls, _ = install_writer(monkeypatch)
    assert pipeline._materialize_requested_final_outputs(state, str(tmp_path))["ok"]
    state["request"] += " Include a units column."
    assert pipeline._materialize_requested_final_outputs(state, str(tmp_path))["ok"]
    assert len(calls) == 2
    assert calls[-1]["request_text"] == state["request"]


def test_explicit_native_copy_repairs_preexisting_wrong_bytes_without_generation(tmp_path, monkeypatch):
    source, destination = tmp_path / "native.csv", tmp_path / "requested.csv"
    source.write_text("id,measured\nA,3\n")
    destination.write_text("id,measured\nA,9\n")
    state = delivery_state(tmp_path, destination)
    state["request_envelope"]["final_outputs"][0]["kind"] = "comparison_table_csv"
    monkeypatch.setattr(pipeline, "_native_export_catalog",
                        lambda *args: {"comparison_table_csv": str(source)})
    monkeypatch.setattr(pipeline, "_run_sandbox_analysis_sync",
                        lambda payload: pytest.fail("Explicit native copies need no generated writer"))
    receipt = pipeline._materialize_requested_final_outputs(state, str(tmp_path))
    assert receipt["ok"] and destination.read_bytes() == source.read_bytes()
    assert receipt["execution_attempts"] == []
    assert pipeline._materialize_requested_final_outputs(state, str(tmp_path))["already_present"]
