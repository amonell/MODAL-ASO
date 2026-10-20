import csv
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from aso_agents import pipeline_tools as pipeline


def artifact_task(tmp_path, *, suffix, values):
    table=tmp_path/'measurements.csv'
    with table.open('w',newline='') as f:
        writer=csv.writer(f);writer.writerow(['id','value'])
        writer.writerows(zip(['row1','row1','row2'],values))
    output=tmp_path/('calculated'+suffix) if suffix else None
    request=(f'Calculate the arithmetic mean and number of observations, retaining zero values and repeated IDs. Write the actual results to {output}.'
             if output else 'Which files contain the input measurements?')
    task={'task_id':'calculate','task_type':'analyze_artifact','label':'Analyze supplied measurements',
          'native':True,'context':{'artifact_refs':[{'path':str(table),'artifact_type':'csv'}]}}
    raw={'schema_version':'v3','request_text':request,'request_mode':'single_task','tasks':[task],
         'final_outputs':([{'task_id':'calculate','kind':'file','format':suffix[1:],'path':str(output)}] if output else [])}
    envelope=pipeline._build_request_envelope(request_text=request,request_id='artifact_export_control',structured_goal={},raw_envelope=raw)
    state={'request':request,'status':'request_parsed_ready','phase':'request_parsed_ready',
           'run_dir':str(tmp_path),'request_envelope':envelope,'last_error':None,
           'clarification':{},'pending_questions':[],'clarification_qa_history':[]}
    pipeline._ensure_v2_request_metadata(state)
    path=tmp_path/'state.json';path.write_text(json.dumps(state))
    return path,table,output,request


@pytest.mark.parametrize('suffix',['.xlsx','.csv','.md'])
@pytest.mark.parametrize('values',[[0,4,8],[2,3,13]])
def test_file_producing_artifact_analysis_executes_real_calculation(tmp_path,monkeypatch,suffix,values):
    state,table,output,request=artifact_task(tmp_path,suffix=suffix,values=values)
    source_bytes=table.read_bytes();generated=[]
    def generate(payload):
        generated.append(payload)
        refs=payload['declared_inputs']['artifact_refs']
        assert {'path':str(table),'exists':True} in refs
        assert payload['request_text']==request
        return {'python_code':
            'import csv,json\nfrom pathlib import Path\nfrom openpyxl import Workbook\n'
            +f'with Path({str(table)!r}).open() as f:\n    values=[float(r["value"]) for r in csv.DictReader(f)]\n'
            +'result={"mean":sum(values)/len(values),"n_used":len(values)}\n'
            +f'output=Path({str(output)!r})\n'
            +'if output.suffix==".xlsx":\n    book=Workbook();sheet=book.active;sheet.append(list(result));sheet.append(list(result.values()));book.save(output)\n'
            +'elif output.suffix==".csv":\n    with output.open("w",newline="") as f:\n        writer=csv.DictWriter(f,fieldnames=list(result));writer.writeheader();writer.writerow(result)\n'
            +'else:\n    output.write_text(json.dumps(result))\n'
            +'Path("sandbox_result.json").write_text(json.dumps(result))\n'
            +'Path("sandbox_summary.md").write_text("Computed the mean and count from every actual observation.")\n'}
    monkeypatch.setattr(pipeline,'_run_sandbox_analysis_sync',generate)
    def forbid_explanation(*args,**kwargs):
        raise AssertionError('A prose explanation cannot create the declared analysis export')
    monkeypatch.setattr(pipeline,'_run_workbench_response_sync',forbid_explanation)
    result=pipeline.tool_analyze_artifact(str(state))
    assert result['status']=='ok'
    assert generated and output.is_file()
    if suffix=='.xlsx':
        book=load_workbook(output,data_only=True)
        records=list(book.active.values);actual=dict(zip(records[0],records[1]));book.close()
    elif suffix=='.csv':
        with output.open() as f:actual=next(csv.DictReader(f))
    else:actual=json.loads(output.read_text())
    assert float(actual['mean'])==sum(values)/len(values)
    assert int(actual['n_used'])==len(values)
    assert table.read_bytes()==source_bytes
    saved=json.loads(state.read_text())
    assert saved['task_statuses']['calculate']['status']=='completed'


def test_missing_declared_file_is_not_completed_by_an_inventory(tmp_path,monkeypatch):
    state,table,output,request=artifact_task(tmp_path,suffix='.csv',values=[0,4,8])
    def omit_export(payload):
        return {'python_code':'from pathlib import Path\nPath("sandbox_result.json").write_text("{}")\nPath("sandbox_summary.md").write_text("No calculation export was written.")\n'}
    monkeypatch.setattr(pipeline,'_run_sandbox_analysis_sync',omit_export)
    monkeypatch.setattr(pipeline,'_run_workbench_response_sync',lambda payload:{'markdown':'A prose response without the requested export.'})
    result=pipeline.tool_analyze_artifact(str(state))
    assert result['status']!='ok'
    assert not output.exists()
    saved=json.loads(state.read_text())
    assert saved['task_statuses']['calculate']['status']!='completed'


def test_inventory_without_declared_export_keeps_native_inspection(tmp_path,monkeypatch):
    state,table,output,request=artifact_task(tmp_path,suffix='',values=[0,4,8])
    def unexpected_generation(payload):
        raise AssertionError('A plain file inventory does not need a generated program')
    monkeypatch.setattr(pipeline,'_run_sandbox_analysis_sync',unexpected_generation)
    result=pipeline.tool_analyze_artifact(str(state))
    assert result['status']=='ok'
    response=Path(result['artifacts'].get('json') or result['artifacts']['response_json'])
    assert json.loads(response.read_text())['analysis_method']=='deterministic_artifact_index'
