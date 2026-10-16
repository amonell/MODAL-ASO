"""Serialization and edit failures must not discard recorded evidence."""
import json
from pathlib import Path

import openpyxl
import pytest

from aso_agents.analysis_execution import execute_analysis_program
from aso_agents.analysis_runtime import structured_excel_cells


def test_structured_cells_round_trip_without_changing_scalars(tmp_path):
    from openpyxl.cell.cell import Cell
    original = Cell._bind_value
    values = [0, None, False, 2.5, "001", [{"coordinate": 17, "count": 0, "missing": None}],
              {"sequence": "ACGT", "observations": [False, 0, "β"]}, [], {}]
    path = tmp_path / "evidence.xlsx"
    with structured_excel_cells() as receipt:
        book = openpyxl.Workbook()
        book.active.append(values)
        book.save(path)
    assert Cell._bind_value is original
    loaded = list(openpyxl.load_workbook(path, data_only=True).active.values)[0]
    assert loaded[:5] == tuple(values[:5])
    assert [json.loads(value) for value in loaded[5:]] == values[5:]
    assert receipt["converted_cells"] == 4
    assert receipt["source_types"] == {"list": 2, "dict": 2}


@pytest.mark.parametrize("length,accepted", [(32763, True), (32764, False)])
def test_structured_cell_limit_never_silently_truncates(length, accepted):
    from openpyxl.cell.cell import Cell
    original = Cell._bind_value
    value = ["x" * length]
    with structured_excel_cells():
        sheet = openpyxl.Workbook().active
        if accepted:
            sheet["A1"] = value
            assert json.loads(sheet["A1"].value) == value
        else:
            with pytest.raises(ValueError, match="preserve the complete value"):
                sheet["A1"] = value
    assert Cell._bind_value is original


def test_nonfinite_nested_value_remains_an_error_and_patch_restores():
    from openpyxl.cell.cell import Cell
    original = Cell._bind_value
    with pytest.raises(ValueError, match="JSON compliant"):
        with structured_excel_cells():
            openpyxl.Workbook().active["A1"] = [float("nan")]
    assert Cell._bind_value is original


def test_actual_analysis_runtime_records_serialization_without_leaf_values(tmp_path):
    output = tmp_path / "evidence.xlsx"
    code = (
        "import json\nfrom pathlib import Path\nimport openpyxl\n"
        "book=openpyxl.Workbook()\nbook.active.append(['observation',0,[{'locus':'private-value','count':0}]])\n"
        f"book.save({str(output)!r})\n"
        "Path('sandbox_result.json').write_text('{}')\nPath('sandbox_summary.md').write_text('Recorded evidence exported.')\n"
    )
    result = execute_analysis_program(directory=tmp_path, payload={},
        generate=lambda request: {"python_code": code}, required_outputs=[output], max_attempts=1)
    assert result["ok"], result
    receipt = result["attempts"][0]["serialization"]
    assert receipt == {"format": "json", "converted_cells": 1, "source_types": {"list": 1}}
    assert "private-value" not in json.dumps(receipt)
    assert json.loads(openpyxl.load_workbook(output).active["C1"].value) == [{"locus": "private-value", "count": 0}]


def test_rejected_edit_retains_last_program_plan_and_error_for_next_repair(tmp_path):
    first = "raise ValueError('precise_original_failure')"
    plan = {"python_code": first, "analysis_plan": "Use unchanged supplied observations.", "limitations": ["Activity is unmeasured."]}
    calls = []
    def generate(request):
        calls.append(request)
        if len(calls) == 1:
            return plan
        if len(calls) == 2:
            return {"python_code": "", "error": {"error": "invalid_structured_output", "reason": "bad edit span"}}
        repair = request["execution_repair"]
        assert repair["previous_code"] == first and repair["previous_plan"] == plan
        assert "precise_original_failure" in repair["error"]
        assert "bad edit span" in repair["latest_generation_error"]
        assert len(repair["previous_attempts"]) == 2
        return {"python_code": "from pathlib import Path\nPath('sandbox_result.json').write_text('{}')\nPath('sandbox_summary.md').write_text('Repaired without changing observations.')"}
    result = execute_analysis_program(directory=tmp_path, payload={}, generate=generate, required_outputs=[], max_attempts=3)
    assert result["ok"] and len(calls) == 3
    assert (tmp_path / "execution_attempts/01/generated_code.py").read_text() == first
    assert (tmp_path / "execution_attempts/02/generated_code.py").read_text() == ""
    assert json.loads((tmp_path / "execution_attempts/02/plan.json").read_text())["error"]["reason"] == "bad edit span"
