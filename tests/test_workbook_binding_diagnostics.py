import copy
import hashlib
import json
from pathlib import Path

import openpyxl
import pytest

from aso_agents.analysis_execution import execute_analysis_program
from aso_agents.design_workbooks import render_design_workbook_plans
from test_design_workbook_rendering import prepared


@pytest.mark.parametrize("unknown", ["foreign-display-id", "x" * 10000])
def test_unknown_binding_identifies_exact_plan_row_without_writing_a_workbook(tmp_path, unknown):
    declared, path, payload, plan, rows = prepared(tmp_path)
    extra = copy.deepcopy(plan["design_workbooks"][0]["tables"][0])
    extra.update(title="Selected", coverage="subset")
    extra["rows"].append({"binding_id": unknown, "values": {"ID": "display-2"}})
    plan["design_workbooks"][0]["tables"].append(extra)
    original = copy.deepcopy(plan)
    with pytest.raises(ValueError) as caught:
        render_design_workbook_plans(payload, plan, [path])
    error = str(caught.value)
    assert "binding_id is absent" in error
    assert "design_workbooks[0].tables[1].rows[1].binding_id" in error
    assert "Table 'Selected', producer 'design'" in error
    assert "['design_catalog']['path']" in error
    assert "display IDs" in error
    assert "notes entries of {title, rows}" in error
    assert len(error) < 1400
    assert not path.exists()
    assert plan == original


def test_binding_error_reaches_repair_and_summary_moves_to_notes_without_changing_candidates(tmp_path):
    declared, path, payload, plan, rows = prepared(tmp_path, second=True)
    catalog = Path(declared["completed_run"]["design_catalog"]["path"])
    catalog_hash = hashlib.sha256(catalog.read_bytes()).hexdigest()
    broken = copy.deepcopy(plan)
    broken["design_workbooks"][0]["tables"].append({
        "title": "ReferenceChecks", "task_id": "design", "coverage": "subset",
        "extra_columns": ["Check", "Count"],
        "rows": [{"binding_id": "check-summary", "values": {"Check": "Conflicts", "Count": 0}}],
    })
    calls = []

    def generate(request):
        calls.append(request)
        if len(calls) == 1:
            current = broken
        else:
            assert len(calls) == 2
            assert not path.exists()
            error = request["execution_repair"]["error"]
            assert "design_workbooks[0].tables[1].rows[0].binding_id" in error
            assert "'check-summary'" in error and "notes entries" in error
            current = copy.deepcopy(plan)
            current["design_workbooks"][0]["notes"].append({
                "title": "ReferenceChecks", "rows": [["Check", "Count"], ["Conflicts", 0]],
            })
        return {"python_code": "from pathlib import Path\nimport json\n"
                f"Path('sandbox_result.json').write_text(json.dumps({current!r}))\n"
                "Path('sandbox_summary.md').write_text('All source candidates and separate reference checks.')\n"}

    result = execute_analysis_program(directory=tmp_path, payload=payload, required_outputs=[path],
                                      generate=generate, max_attempts=2)
    assert result["ok"], result
    assert len(calls) == len(result["attempts"]) == 2
    assert "check-summary" in result["attempts"][0]["error"]
    assert result["attempts"][1]["error"] is None
    assert hashlib.sha256(catalog.read_bytes()).hexdigest() == catalog_hash
    book = openpyxl.load_workbook(path, data_only=True)
    try:
        values = list(book["Library"].values)
        actual = [dict(zip(values[0], row)) for row in values[1:]]
        assert [row["binding_id"] for row in actual] == [row["binding_id"] for row in rows]
        assert [row["antisense_5to3"] for row in actual] == [row["antisense_5to3"] for row in rows]
        assert list(book["ReferenceChecks"].values) == [("Check", "Count"), ("Conflicts", 0)]
    finally:
        book.close()
