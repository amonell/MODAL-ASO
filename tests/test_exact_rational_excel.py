"""Exact arithmetic must survive the workbook serialization boundary."""
from fractions import Fraction
import json

import openpyxl
import pytest

from aso_agents.analysis_execution import execute_analysis_program
from aso_agents.analysis_runtime import structured_excel_cells


@pytest.mark.parametrize("value", [Fraction(17, 13), Fraction(-11, 4), Fraction(9), Fraction(0)])
def test_exact_rational_cell_round_trip_without_rounding(tmp_path, value):
    from openpyxl.cell.cell import Cell
    original = Cell._bind_value
    output = tmp_path / "ratios.xlsx"
    with structured_excel_cells() as receipt:
        workbook = openpyxl.Workbook()
        workbook.active.append([value, 0, False, "001", 2.5])
        workbook.save(output)
    assert Cell._bind_value is original
    row = next(openpyxl.load_workbook(output).active.values)
    assert isinstance(row[0], str)
    assert Fraction(row[0]) == value
    assert row[1:] == (0, False, "001", 2.5)
    assert receipt["converted_cells"] == 1
    assert receipt["source_types"] == {"Fraction": 1}


def test_nested_exact_ratios_preserve_input_values_and_ordering(tmp_path):
    values = [Fraction(5, 17), Fraction(2, 7), Fraction(-3, 8)]
    original_order = sorted(values)
    source = {"ratios": values, "missing": None, "zero": 0}
    output = tmp_path / "nested.xlsx"
    with structured_excel_cells():
        workbook = openpyxl.Workbook()
        workbook.active.append([source])
        workbook.save(output)
    encoded = json.loads(openpyxl.load_workbook(output).active["A1"].value)
    assert [Fraction(x) for x in encoded["ratios"]] == values
    assert encoded["missing"] is None and encoded["zero"] == 0
    assert all(isinstance(x, Fraction) for x in source["ratios"])
    assert sorted(source["ratios"]) == original_order


def test_actual_runtime_exports_exact_ratios_without_program_repair(tmp_path):
    output = tmp_path / "calculation.xlsx"
    code = (
        "from fractions import Fraction\nfrom pathlib import Path\nimport openpyxl\n"
        "workbook=openpyxl.Workbook()\n"
        "workbook.active.append([Fraction(-13, 8), {'ratio': Fraction(5, 19)}])\n"
        f"workbook.save({str(output)!r})\n"
        "Path('sandbox_result.json').write_text('{}')\n"
        "Path('sandbox_summary.md').write_text('Exact ratios exported.')\n"
    )
    result = execute_analysis_program(
        directory=tmp_path, payload={}, generate=lambda request: {"python_code": code},
        required_outputs=[output], max_attempts=1,
    )
    assert result["ok"], result
    assert len(result["attempts"]) == 1
    row = next(openpyxl.load_workbook(output).active.values)
    assert Fraction(row[0]) == Fraction(-13, 8)
    assert Fraction(json.loads(row[1])["ratio"]) == Fraction(5, 19)
    assert result["attempts"][0]["serialization"]["source_types"] == {"Fraction": 1, "dict": 1}
