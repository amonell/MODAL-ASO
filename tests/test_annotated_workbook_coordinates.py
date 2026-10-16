import hashlib
import json
from pathlib import Path

import openpyxl
import pytest

from aso_agents.design_workbooks import render_design_workbook_plans, validate_design_workbooks
from test_design_observation_tables import example, project
from test_design_workbook_rendering import prepared


@pytest.mark.parametrize("strand,start,end,relative,exonic,intronic", [
    ("+", 146, 163, [-2, 14], 3, 14),
    ("-", 185, 202, [-2, 14], 3, 14),
    ("+", 159, 176, [11, 27], 0, 17),
    ("-", 173, 190, [10, 26], 0, 17),
    ("+", 110, 127, [-38, -22], 17, 0),
    ("-", 210, 227, [-27, -11], 17, 0),
])
def test_reference_context_covers_crossing_and_wholly_exonic_or_intronic_targets(
        tmp_path, strand, start, end, relative, exonic, intronic):
    row, _, raw = example(tmp_path, strand)
    row.update(Start=start, End=end)
    raw.write_text(json.dumps({"schema": {"table": "aso_tiles"}, "rows": [row]}))
    _, observations = project(tmp_path)
    observation = observations["rows"][0]
    context, = [x for x in observation["annotated_splice_contexts"]
                if x["transcript_id"] == "TX.2" and x["element"] == "donor" and x["exon_ordinal"] == 1]
    assert context["relative_interval_inclusive"] == relative
    assert context["exonic_overlap_bp"] == exonic
    assert context["intronic_overlap_bp"] == intronic
    assert context["target_interval_1based_inclusive"] == [start + 1, end]
    assert bool(observation["physical_splice_elements"]) == bool(exonic and intronic)


def selector(field="relative_interval_inclusive", **changes):
    return {"kind": "splice_context", "transcript_id": "TX.2", "exon_ordinal": 1,
            "element": "donor", "field": field, **changes}


def test_host_projects_requested_coordinates_counts_and_aliases_and_detects_corruption(tmp_path):
    declared, path, payload, plan, rows = prepared(tmp_path)
    plan["design_workbooks"][0]["tables"][0]["source_columns"] = {
        "Relative interval": selector(), "Exonic bases": selector("exonic_overlap_bp"),
        "Intronic bases": selector("intronic_overlap_bp"),
        "Genome intervals": "genomic_blocks_1based_inclusive", "ASO": "antisense_5to3"}
    render_design_workbook_plans(payload, plan, [path])
    validate_design_workbooks(payload, plan, [path])
    book = openpyxl.load_workbook(path)
    cells = list(book["Library"].values)
    row = dict(zip(cells[0], cells[1]))
    assert json.loads(row["Relative interval"]) == [-2, 14]
    assert row["Exonic bases"] == 3 and row["Intronic bases"] == 14
    assert json.loads(row["Genome intervals"]) == [["chr7", 147, 163]]
    assert row["ASO"] == "CGGTTACGGTTACGGTT"
    book["Library"].cell(2, cells[0].index("Relative interval") + 1).value = "[-1, 15]"
    book.save(path)
    with pytest.raises(ValueError, match="contradicts source"):
        validate_design_workbooks(payload, plan, [path])


@pytest.mark.parametrize("changes", [{"transcript_id": "TX.3"}, {"exon_ordinal": 2}, {"element": "acceptor"}])
def test_missing_reference_frame_remains_unknown(tmp_path, changes):
    _, path, payload, plan, _ = prepared(tmp_path)
    plan["design_workbooks"][0]["tables"][0]["source_columns"] = {"Relative": selector(**changes)}
    render_design_workbook_plans(payload, plan, [path])
    validate_design_workbooks(payload, plan, [path])
    values = list(openpyxl.load_workbook(path, data_only=True)["Library"].values)
    assert dict(zip(values[0], values[1]))["Relative"] is None


@pytest.mark.parametrize("changes", [{"boundary_1based": 150}, {"exon_ordinal": "1"},
                                     {"field": "invented"}, {"kind": "predicted_response_site"}])
def test_selector_cannot_supply_an_anchor_or_change_reference_semantics(tmp_path, changes):
    _, path, payload, plan, _ = prepared(tmp_path)
    plan["design_workbooks"][0]["tables"][0]["source_columns"] = {"Relative": selector(**changes)}
    with pytest.raises(ValueError, match="Invalid annotated"):
        render_design_workbook_plans(payload, plan, [path])


def test_multiple_target_blocks_are_not_silently_collapsed_to_one_interval(tmp_path):
    declared, path, payload, plan, _ = prepared(tmp_path)
    meta = declared["completed_run"]["design_evidence"][0]["observations"]
    source = Path(meta["path"])
    observations = json.loads(source.read_text())
    context = next(x for x in observations["rows"][0]["annotated_splice_contexts"] if x["element"] == "donor")
    observations["rows"][0]["annotated_splice_contexts"].append(
        {**context, "target_interval_1based_inclusive": [160, 176], "relative_interval_inclusive": [11, 27]})
    source.write_text(json.dumps(observations))
    meta["sha256"] = hashlib.sha256(source.read_bytes()).hexdigest()
    plan["design_workbooks"][0]["tables"][0]["source_columns"] = {"Relative": selector()}
    with pytest.raises(ValueError, match="ambiguous across target blocks"):
        render_design_workbook_plans(payload, plan, [path])
