from __future__ import annotations

import asyncio
import importlib
import pytest
import json
from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import run_goal_intake
from Bio import SeqIO
from Bio.Seq import Seq
from Bio.SeqFeature import FeatureLocation, SeqFeature
from Bio.SeqRecord import SeqRecord
from aso_agents.parameter_constraints import (
    apply_parameter_constraints_to_goal,
    evaluate_row_against_constraints,
    get_binding_count_from_row,
    get_gc_bounds,
    get_gc_inclusive_flags,
    get_homopolymer_per_base_limits,
    parse_parameter_constraints,
    sequence_passes_homopolymer_limits,
)
from aso_agents import pipeline_tools
from aso_agents.run_layout import resolve_run_path

splicing_agent_module = importlib.import_module("aso_agents.splicing_agent")
splice_strategy_router_module = importlib.import_module(
    "aso_agents.splice_strategy_router_agent"
)


def test_splice_strategy_agents_delegate_output_budget_to_provider() -> None:
    settings = (
        splice_strategy_router_module.PLANNER_MODEL_SETTINGS,
        splice_strategy_router_module.EXECUTOR_MODEL_SETTINGS,
        splice_strategy_router_module.ROUTE_RANKER_MODEL_SETTINGS,
    )

    for model_settings in settings:
        assert model_settings.max_output_tokens is None
        assert "max_output_tokens" not in model_settings.extra_body


def test_route_ranker_repairs_noncanonical_candidate_direction(monkeypatch) -> None:
    calls: list[tuple[str, dict[str, Any]]] = []

    def fake_call(agent, payload, *, max_steps):
        calls.append((agent.name, payload))
        if len(calls) == 1:
            parsed = {
                "ranking_summary": "Initial ordering.",
                "ranked_routes": [
                    {
                        "rank": 1,
                        "label": "acceptor@100",
                        "direction": "decrease",
                        "rationale": "Strongest measured candidate.",
                        "confidence": "high",
                    }
                ],
            }
        else:
            parsed = {
                "ranking_summary": "Contract-repaired ordering.",
                "ranked_routes": [
                    {
                        "rank": 1,
                        "label": "acceptor@100",
                        "direction": "increase",
                        "rationale": "Strongest measured candidate.",
                        "confidence": "high",
                    }
                ],
            }
        return {
            "parsed": parsed,
            "responses": [f"response-{len(calls)}"],
            "raw_output": json.dumps(parsed),
        }

    monkeypatch.setattr(splice_strategy_router_module, "_call_json_agent", fake_call)

    result = splice_strategy_router_module.rank_route_candidates_with_llm(
        {
            "job": {"desired_delta": "decrease"},
            "candidates": [
                {"label": "acceptor@100", "direction": "increase"}
            ],
        }
    )

    assert result["ranked_routes"][0]["direction"] == "increase"
    assert result["_repair_applied"] is True
    assert len(calls) == 2
    assert calls[1][0] == "ASO Strategy Route Ranking Contract Repair"
    assert calls[1][1]["canonical_candidate_pairs"] == [
        {"label": "acceptor@100", "direction": "increase"}
    ]
    assert "label/direction" in calls[1][1]["contract_errors"][0]


def test_route_ranker_prompt_separates_job_outcome_from_site_action() -> None:
    instructions = str(
        splice_strategy_router_module.route_ranker_agent.instructions
    )

    assert "candidates[].direction is the canonical site-action direction" in instructions
    assert "job.desired_delta describes the requested program-level outcome" in instructions


def _write_test_genbank(path: Path, features: list[SeqFeature]) -> None:
    record = SeqRecord(Seq("ACGT" * 200), id="LMNA_chr1", name="LMNA", description="LMNA genomic segment")
    record.annotations["molecule_type"] = "DNA"
    record.features = features
    SeqIO.write(record, str(path), "genbank")


def _state_from_raw_envelope(raw_envelope: dict[str, Any], tmp_path: Path) -> dict[str, Any]:
    request_text = str(raw_envelope.get("request_text") or "Composite ASO request").strip()
    envelope = pipeline_tools._build_request_envelope(
        request_text=request_text,
        request_id="unit_large_graph",
        structured_goal={},
        raw_envelope=raw_envelope,
    )
    state = {
        "request": request_text,
        "status": "request_parsed_ready",
        "phase": "request_parsed_ready",
        "run_dir": str(tmp_path),
        "request_envelope": envelope,
        "last_error": None,
        "clarification": {"isoform_scope_confirmed": False},
        "pending_questions": [],
        "clarification_qa_history": [],
    }
    pipeline_tools._ensure_v2_request_metadata(state)
    return state


def _mark_completed(state: dict[str, Any], *task_ids: str) -> None:
    statuses = state.setdefault("task_statuses", {})
    for task_id in task_ids:
        entry = statuses.setdefault(task_id, {})
        entry["status"] = "completed"


def _next_tool(planning: dict[str, Any]) -> str:
    actions = planning.get("next_actions") if isinstance(planning.get("next_actions"), list) else []
    return str(actions[0].get("tool") or "") if actions and isinstance(actions[0], dict) else ""


def test_source_backed_clarification_preserves_and_resolves_evidence(
    tmp_path: Path,
) -> None:
    source_path = tmp_path / "source_context.json"
    source_path.write_text(
        json.dumps(
            {
                "scope": "reviewed input",
                "region": {"chrom": "chr1", "start": 100, "end": 140},
            }
        ),
        encoding="utf-8",
    )
    normalized = pipeline_tools._normalize_clarification_answers(
        [
            {
                "question": "Which supplied source should be used?",
                "answer": "Use the reviewed source artifact.",
                "evidence_refs": [str(source_path)],
            }
        ]
    )

    assert normalized[0]["evidence_refs"] == [str(source_path)]
    history = pipeline_tools._clarification_qa_history(
        {"clarification_history": [{"answers": normalized}]}
    )
    assert history[0]["evidence_refs"] == [str(source_path)]
    rendered = pipeline_tools._render_request_workbench_history(history)
    assert str(source_path) in rendered
    assert "chr1" not in rendered
    assert "not additional request intent" in rendered

    audit_sources = pipeline_tools._goal_audit_evidence_sources(
        "Use the supplied reviewed source.",
        history,
    )
    source_evidence = [
        item
        for item in audit_sources
        if item["kind"] == "clarification_source_evidence"
    ]
    assert len(source_evidence) == 1
    assert '"start\\": 100' not in source_evidence[0]["text"]
    assert '"question": "Which supplied source should be used?"' in source_evidence[0]["text"]
    assert '"answer": "Use the reviewed source artifact."' in source_evidence[0]["text"]
    assert "not an additional user request" in source_evidence[0]["text"]


def test_existing_aso_orientation_interpretation_marks_target_site_sequence() -> None:
    entry = {
        "name": "aso_1",
        "sequence_role_confirmation": {"role": "target_site_sequence"},
        "sequence": "TTTAGCTTGCGGTCTTTAGA",
        "reverse_complement_sequence": "TCTAAAGACCGCAAGCTAAA",
        "mapping_hits": [
            {
                "match_orientation": "input_sequence",
                "matched_target_sequence": "TTTAGCTTGCGGTCTTTAGA",
                "mapping_scope": "pre_mrna_locus",
            }
        ],
    }

    interpretation = pipeline_tools._provided_aso_orientation_interpretation(entry)
    assert interpretation["provided_sequence_role"] == "target_site_sequence"
    assert interpretation["matched_target_sequence"] == "TTTAGCTTGCGGTCTTTAGA"
    assert interpretation["synthesized_antisense_sequence"] == "TCTAAAGACCGCAAGCTAAA"

    markdown = pipeline_tools._provided_aso_orientation_markdown({"provided_asos": [{"name": "aso_1", **entry, "orientation_interpretation": interpretation}]})
    assert "Sequence Orientation" in markdown
    assert "target site sequence" in markdown
    assert "TCTAAAGACCGCAAGCTAAA" in markdown


def test_provided_aso_files_load_no_header_tsv(tmp_path: Path) -> None:
    aso_path = tmp_path / "otud4_asos.tsv"
    aso_path.write_text(
        "OTUD4_45\tTCCAAATGAGAACAGAGTGC\t/52MOErT/*/i2MOErC/*/32MOErC/\n"
        "OTUD4_46\tCAAATGAGAACAGAGTGCAA\t/52MOErC/*/i2MOErA/*/32MOErA/\n",
        encoding="utf-8",
    )

    rows = pipeline_tools._normalize_provided_asos(
        {"provided_aso_files": [{"path": str(aso_path), "file_format": "tsv"}]}
    )

    assert [row["name"] for row in rows] == ["OTUD4_45", "OTUD4_46"]
    assert rows[0]["sequence"] == "TCCAAATGAGAACAGAGTGC"
    assert rows[0]["chemistry"] == "/52MOErT/*/i2MOErC/*/32MOErC/"
    assert rows[0]["source_file"] == str(aso_path)
    assert rows[0]["source_row"] == 1


def test_evaluate_existing_aso_set_accepts_aso_table_file_context(tmp_path: Path) -> None:
    aso_path = tmp_path / "lab_asos.csv"
    aso_path.write_text(
        "ASO Name,Sequence,Chemistry\n"
        "lab_aso_1,ACGTACGTACGTACGTACGT,2MOE PS\n",
        encoding="utf-8",
    )
    state = _state_from_raw_envelope(
        {
            "schema_version": "v3",
            "request_text": f"Evaluate the ASOs in {aso_path} for human MAPT.",
            "request_mode": "single_task",
            "tasks": [
                {
                    "task_id": "evaluate_lab_asos",
                    "task_type": "evaluate_existing_aso_set",
                    "context": {
                        "gene_symbol": "MAPT",
                        "species": "human",
                        "provided_aso_files": [{"path": str(aso_path), "file_format": "csv"}],
                    },
                }
            ],
        },
        tmp_path,
    )

    validation = pipeline_tools._request_validation_summary(state["request_envelope"])
    task_validation = validation["tasks"]["evaluate_lab_asos"]
    assert "provided_asos_missing" not in task_validation["issues"]
    assert task_validation["status"] == "runnable"


def test_gapmer_uniqueness_context_defaults_gene_wide_for_standard_knockdown() -> None:
    context = run_goal_intake._resolve_gapmer_uniqueness_context(
        {"intent": "knockdown", "gene_symbol": "OTUD4"},
        "canonical",
    )

    assert context["mode"] == "gene_wide"
    assert context["target_label"] is None
    assert context["target_labels"] == []


def test_gapmer_uniqueness_context_honors_explicit_isoform_scope() -> None:
    context = run_goal_intake._resolve_gapmer_uniqueness_context(
        {"intent": "knockdown", "target_scope": "isoform_specific"},
        "canonical",
    )

    assert context["mode"] == "target_isoform_specific"
    assert context["target_label"] == "canonical"


def test_gapmer_uniqueness_context_honors_explicit_gene_wide_scope() -> None:
    context = run_goal_intake._resolve_gapmer_uniqueness_context(
        {"intent": "knockdown", "target_scope": "gene_wide"},
        "canonical",
    )

    assert context["mode"] == "gene_wide"
    assert context["target_label"] is None


def test_gapmer_uniqueness_context_uses_pathogenic_isoform_objective() -> None:
    context = run_goal_intake._resolve_gapmer_uniqueness_context(
        {
            "intent": "splice_modulation",
            "splice_isoform_objectives": [
                {"isoform_label": "canonical", "desired_abundance": "decrease", "role": "pathogenic"}
            ],
        },
        "canonical",
    )

    assert context["mode"] == "target_isoform_specific"
    assert context["target_label"] == "canonical"
    assert context["pathogenic_labels"] == ["canonical"]


def test_splice_nearest_site_helper_identifies_nearest_boundary() -> None:
    nearest = splicing_agent_module._nearest_splice_site_for_mid(
        490,
        [
            {"label": "exon_1", "acceptor": {"position": 100}, "donor": {"position": 250}},
            {"canonical_label": "exon_2", "acceptor": {"position": 500}, "donor": {"position": 700}},
        ],
    )

    assert nearest == {
        "distance_bp": 10,
        "site_kind": "acceptor",
        "position": 500,
        "label": "exon_2",
        "exon_index": None,
        "isoform_label": None,
    }


def test_requested_splice_site_helper_honors_site_kind_for_large_exons() -> None:
    site = {
        "canonical_label": "donor@100",
        "site_kind": "donor",
        "acceptor": {"position": 20000},
        "donor": {"position": 100},
        "exon_index": 326,
        "isoform_label": "TTN-full-length",
    }

    near_acceptor = splicing_agent_module._requested_splice_site_for_mid(19995, site)
    near_donor = splicing_agent_module._requested_splice_site_for_mid(105, site)

    assert near_acceptor is not None
    assert near_acceptor["site_kind"] == "donor"
    assert near_acceptor["position"] == 100
    assert near_acceptor["distance_bp"] == 19895
    assert near_donor is not None
    assert near_donor["distance_bp"] == 5


def test_artifact_analysis_ranks_generated_aso_set_rows(tmp_path: Path) -> None:
    run_dir = tmp_path / "goal_test"
    run_dir.mkdir()
    (run_dir / "generated_aso_set.json").write_text(
        json.dumps(
            {
                "rows": [
                    {
                        "sequence": "AAAACCCCGGGGTTTTAAAA",
                        "reverse_complement_sequence": "TTTTAAAACCCCGGGGTTTT",
                        "route_label": "donor@100",
                        "route_rank": 1,
                        "chrom": "chr1",
                        "start": "80",
                        "end": "100",
                        "strand": "+",
                        "gc_fraction": 0.4,
                        "repeat_risk_class": "none",
                        "source_annotations": {
                            "primary_direction": "decrease",
                            "primary_delta_deletion": "-0.91",
                            "primary_delta_replacement": "-0.90",
                            "off_target_total_upto_mismatch_count": "0",
                            "off_target_weighted_risk_score": "0",
                            "selected_transcript_exon_labels": "ENST_TEST:exon2",
                            "tile_score": "800",
                            "no_homopolymer": "True",
                            "no_secondary_structure": "True",
                            "safeguards_status": "pass",
                        },
                    },
                    {
                        "sequence": "CCCCAAAATTTTGGGGCCCC",
                        "reverse_complement_sequence": "GGGGCCCCAAAATTTTGGGG",
                        "route_label": "acceptor@200",
                        "route_rank": 2,
                        "chrom": "chr1",
                        "start": "200",
                        "end": "220",
                        "strand": "+",
                        "gc_fraction": 0.5,
                        "repeat_risk_class": "none",
                        "source_annotations": {
                            "primary_direction": "decrease",
                            "primary_delta_deletion": "-0.82",
                            "primary_delta_replacement": "-0.81",
                            "off_target_total_upto_mismatch_count": "0",
                            "off_target_weighted_risk_score": "0",
                            "selected_transcript_exon_labels": "ENST_TEST:exon3",
                            "tile_score": "900",
                            "no_homopolymer": "True",
                            "no_secondary_structure": "True",
                            "safeguards_status": "pass",
                        },
                    },
                ]
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    request_envelope = {"artifact_refs": [{"path": str(run_dir)}], "request_text": "Return top ASOs by route."}
    artifact_context = pipeline_tools._build_artifact_context(request_envelope)
    response = pipeline_tools._build_generated_aso_artifact_analysis_response(
        artifact_context=artifact_context,
        response_meta={"request_text": "Return top ASOs by route."},
        request_envelope=request_envelope,
    )

    assert response is not None
    assert response["analysis_method"] == "deterministic_generated_aso_ranker"
    assert response["top_ranked_asos"][0]["target_site_sequence"] == "AAAACCCCGGGGTTTTAAAA"
    assert "AAAACCCCGGGGTTTTAAAA" in response["markdown"]
    assert "cannot output a full top-10" not in response["markdown"].lower()


def test_analysis_only_artifact_request_overrides_misclassified_design_envelope(tmp_path: Path) -> None:
    run_dir = tmp_path / "goal_existing_ttn"
    run_dir.mkdir()
    (run_dir / "generated_aso_set.json").write_text('{"rows": []}', encoding="utf-8")
    request = (
        f"Evaluate the MODAL-ASO output in {run_dir}. Do not design new ASOs and do not rerun tiling. "
        "From the existing output artifacts, return the top 10 ASOs most likely to work."
    )
    raw_envelope = {
        "workflow_name": "design_new_asos",
        "tasks": [
            {
                "task_id": "bad_design_task",
                "task_type": "design_new_asos",
                "context": {
                    "gene_symbol": "TTN",
                    "species": "human",
                    "structured_goal": {"gene_symbol": "TTN", "species": "human"},
                },
            }
        ],
    }

    envelope = pipeline_tools._build_request_envelope(
        request_text=request,
        request_id="artifact_eval",
        structured_goal={"gene_symbol": "TTN", "species": "human"},
        raw_envelope=raw_envelope,
    )

    assert envelope["workflow_name"] == "analyze_artifact"
    assert envelope["response_mode"] == "analysis"
    tasks = [task for task in envelope["tasks"] if isinstance(task, dict)]
    assert [task["task_type"] for task in tasks] == ["analyze_artifact"]
    assert tasks[0]["context"]["artifact_refs"] == [{"path": str(run_dir)}]
    assert not pipeline_tools._request_envelope_validation_issues(envelope)


def test_artifact_route_ranker_diversifies_strategy_routes_with_scoped_prompt() -> None:
    records = []
    for idx in range(6):
        records.append(
            {
                "route_rank": 1,
                "route_label": "donor@100",
                "target_site_sequence": f"AAAAACCCCCAAAAA{idx:04d}",
                "synthesis_antisense_sequence": f"TTTTTGGGGGTTTTT{idx:04d}",
                "directional_effect": 0.99 - idx * 0.001,
                "composite_rank_score": 990 - idx,
                "off_target_total_upto_mismatch_count": 0,
                "passes_basic_filters": True,
                "source_path": "/tmp/run/strategy_routes/donor.route.csv",
            }
        )
    for idx in range(3):
        records.append(
            {
                "route_rank": 2,
                "route_label": "acceptor@200",
                "target_site_sequence": f"CCCCCAAAAACCCCC{idx:04d}",
                "synthesis_antisense_sequence": f"GGGGGTTTTTGGGGG{idx:04d}",
                "directional_effect": 0.95 - idx * 0.001,
                "composite_rank_score": 950 - idx,
                "off_target_total_upto_mismatch_count": 0,
                "passes_basic_filters": True,
                "source_path": "/tmp/run/strategy_routes/acceptor.route.csv",
            }
        )

    selected = pipeline_tools._select_top_artifact_asos(
        records,
        request_text="Analyze existing ASO output artifacts",
        limit=4,
    )

    assert [item["route_label"] for item in selected] == [
        "donor@100",
        "donor@100",
        "acceptor@200",
        "acceptor@200",
    ]


def test_artifact_context_indexes_tables_and_outputs(tmp_path: Path) -> None:
    run_dir = tmp_path / "goal_table_index"
    routes_dir = run_dir / "strategy_routes"
    routes_dir.mkdir(parents=True)
    (run_dir / "report.md").write_text("# Run Report\n\nTop candidates are in generated outputs.\n", encoding="utf-8")
    (run_dir / "generated_aso_set.json").write_text(
        json.dumps(
            {
                "rows": [
                    {
                        "sequence": "AAAACCCCGGGGTTTTAAAA",
                        "reverse_complement_sequence": "TTTTAAAACCCCGGGGTTTT",
                        "route_label": "donor@100",
                        "off_target_total_upto_mismatch_count": 0,
                    }
                ]
            }
        ),
        encoding="utf-8",
    )
    (routes_dir / "donor_rank_001.route.csv").write_text(
        "primary_label,sequence,chemistry_synthesis,primary_delta_deletion,off_target_total_upto_mismatch_count\n"
        "donor@100,AAAACCCCGGGGTTTTAAAA,TTTTAAAACCCCGGGGTTTT,-0.91,0\n",
        encoding="utf-8",
    )

    artifact_context = pipeline_tools._build_artifact_context(
        {
            "request_text": f"Where should I look in {run_dir}?",
            "artifact_refs": [{"path": str(run_dir)}],
        }
    )
    artifact_index = artifact_context["artifact_index"]
    aggregate = artifact_index["aggregate"]
    tables = pipeline_tools._artifact_index_key_tables(artifact_index)
    table_by_rel = {str(table["rel_path"]): table for table in tables}

    assert aggregate["table_count"] >= 2
    assert aggregate["by_kind"]["generated_aso_set"] == 1
    assert aggregate["by_kind"]["strategy_route_table"] == 1
    assert table_by_rel["generated_aso_set.json"]["row_count"] == 1
    assert "sequence" in table_by_rel["generated_aso_set.json"]["columns"]
    assert table_by_rel["strategy_routes/donor_rank_001.route.csv"]["sample_rows"][0]["primary_label"] == "donor@100"

    ranking_response = pipeline_tools._build_generated_aso_artifact_analysis_response(
        artifact_context=artifact_context,
        response_meta={"request_text": "Where should I look to evaluate performance?"},
        request_envelope={"request_text": "Where should I look to evaluate performance?"},
    )
    response = pipeline_tools._build_artifact_index_analysis_response(
        artifact_context=artifact_context,
        response_meta={"request_text": "Where should I look to evaluate performance?"},
        request_envelope={"request_text": "Where should I look to evaluate performance?"},
    )

    assert ranking_response is None
    assert response is not None
    assert response["analysis_method"] == "deterministic_artifact_index"
    assert "generated_aso_set.json" in response["markdown"]
    assert "strategy_routes/donor_rank_001.route.csv" in response["markdown"]


def test_splice_scored_row_skip_status_is_preserved() -> None:
    scored_row = {
        "splice_eval_status": "skipped_no_nearby_splice_site",
        "splice_eval_skip_reason": "no_splice_sites_within_2000bp",
        "nearest_splice_site_distance_bp": 2430,
        "nearest_splice_site_label": "exon_21_acceptor",
    }

    assert pipeline_tools._splice_stage_from_scored_row(scored_row) == {
        "status": "skipped",
        "reason_code": "skipped_no_nearby_splice_site",
        "details": "no_splice_sites_within_2000bp",
    }
    summary = pipeline_tools._splice_summary_from_scored_row(scored_row)
    assert summary["splice_eval_status"] == "skipped_no_nearby_splice_site"
    assert summary["nearest_splice_site_distance_bp"] == 2430


def test_chooser_focused_v2_markdown_summarizes_provided_aso_evaluation(tmp_path: Path) -> None:
    run_dir = tmp_path / "goal_eval"
    scoring_dir = run_dir / "v2_tasks" / "normal_scoring" / "provided_aso_scoring"
    splice_dir = run_dir / "v2_tasks" / "splice_eval"
    scoring_dir.mkdir(parents=True)
    splice_dir.mkdir(parents=True)
    scoring_path = scoring_dir / "provided_aso_scored_uniqueness.json"
    scoring_path.write_text(
        json.dumps(
            {
                "rows": [
                    {
                        "name": "OTUD4_46",
                        "sequence": "CAAATGAGAACAGAGTGCAA",
                        "selection_rank": 1,
                        "isoform_uniqueness_score": 1.0,
                        "isoform_uniqueness_class": "gene_wide",
                        "off_target_gene_count": 0,
                    }
                ]
            }
        ),
        encoding="utf-8",
    )
    (splice_dir / "scored_tiles.json").write_text(
        json.dumps(
            {
                "rows": [
                    {
                        "name": "OTUD4_46",
                        "sequence": "CAAATGAGAACAGAGTGCAA",
                        "delta_inclusion_deletion": -0.82,
                        "delta_inclusion_replacement": 0.12,
                    },
                    {
                        "name": "OTUD4_47",
                        "sequence": "TGAGAACAGAGTGCAACATG",
                        "splice_eval_status": "skipped_no_nearby_splice_site",
                        "splice_eval_skip_reason": "no_splice_sites_within_2000bp",
                    },
                ]
            }
        ),
        encoding="utf-8",
    )
    payload = {
        "evaluation_summary": {
            "context_summary": {
                "gene_symbol": "OTUD4",
                "species": "human",
                "target_context": "canonical transcript in neurons",
            },
            "provided_aso_scoring": {"scored_tiles_path": str(scoring_path)},
            "provided_asos": [
                {
                    "name": "OTUD4_46",
                    "sequence": "CAAATGAGAACAGAGTGCAA",
                    "selection_rank": 1,
                    "isoform_uniqueness_score": 1.0,
                    "isoform_uniqueness_class": "gene_wide",
                    "off_target_gene_count": 0,
                    "mapping_status": "mapped",
                    "splice_eval_summary": {
                        "nearest_splice_site_distance_bp": 38,
                        "delta_inclusion_deletion": -0.82,
                        "delta_inclusion_replacement": 0.12,
                    },
                },
                {
                    "name": "OTUD4_47",
                    "sequence": "TGAGAACAGAGTGCAACATG",
                    "selection_rank": 2,
                    "isoform_uniqueness_score": 0.5,
                    "isoform_uniqueness_class": "gene_wide",
                    "off_target_gene_count": 1,
                    "mapping_status": "mapped",
                },
            ],
            "limitations": ["No direct PTC/NMD consequence model was run."],
        }
    }

    markdown = pipeline_tools._chooser_focused_v2_markdown(
        run_dir=str(run_dir),
        markdown_text="# Sequence Orientation\nVerbose orientation dump",
        response_payload=payload,
        response_basename="existing_aso_evaluation",
    )

    assert markdown.startswith("# Answer")
    assert "What To Do Next" in markdown
    assert markdown.find("# Answer") < markdown.find("## What To Do Next")
    assert markdown.find("## What To Do Next") < markdown.find("## Main Files To Open")
    assert "use two shortlists" not in markdown
    assert "The top-ranked sequence is `CAAATGAGAACAGAGTGCAA`" in markdown
    assert "OTUD4_46" in markdown
    assert "Top Normal-Ranking Candidates" in markdown
    assert "Strongest Splice-Perturbation Candidates" in markdown
    assert "v2_tasks/normal_scoring/provided_aso_scoring/provided_aso_scored_uniqueness.json" in markdown
    assert "v2_tasks/splice_eval/scored_tiles.json" in markdown
    assert "2/2 ASO(s) were scored as `gene_wide`" in markdown
    # The sibling score file is linked separately; its unbound skip status
    # cannot fill the missing status in this evaluation payload.
    assert "Splice evaluation scored 1/2 row(s) and skipped 0/2 row(s)." in markdown
    assert "1/2 row(s) have neither a usable splice delta nor a recorded skip status in this evaluation." in markdown
    assert "| 1 | OTUD4_46 | CAAATGAGAACAGAGTGCAA | gene_wide | 1.0000 | NA / NA | NA | 0 |" in markdown
    assert "Sequence Orientation" not in markdown
    assert str(run_dir) not in markdown

    artifacts = pipeline_tools._write_v2_response_outputs(
        run_dir=str(run_dir),
        markdown_text="# Sequence Orientation\nVerbose orientation dump",
        response_payload=payload,
        response_basename="existing_aso_evaluation",
        title="ASO Evaluation",
    )
    written_report = Path(artifacts["report_md"]).read_text(encoding="utf-8")
    written_payload = json.loads(Path(artifacts["response_json"]).read_text(encoding="utf-8"))
    assert written_report == "# Sequence Orientation\nVerbose orientation dump"
    assert written_payload["markdown"] == written_report
    assert written_payload.get("answer_summary") == payload.get("answer_summary")
    written_html = Path(artifacts["report_html"]).read_text(encoding="utf-8")
    assert "<h1>Sequence Orientation</h1>" in written_html
    assert "<pre>" not in written_html
    assert (resolve_run_path(run_dir, "existing_aso_evaluation.xlsx")).exists()
    assert scoring_path.with_suffix(".xlsx").exists()
    assert (splice_dir / "scored_tiles.xlsx").exists()
    assert artifacts["existing_aso_evaluation_xlsx"].endswith("existing_aso_evaluation.xlsx")
    assert "excel_outputs" in written_payload["evaluation_summary"]
    from openpyxl import load_workbook

    workbook = load_workbook(resolve_run_path(run_dir, "existing_aso_evaluation.xlsx"), read_only=True, data_only=True)
    assert "Splice" in workbook.sheetnames
    splice_sheet = workbook["Splice"]
    splice_headers = [cell.value for cell in next(splice_sheet.iter_rows(min_row=1, max_row=1))]
    assert "name" in splice_headers
    assert "splice_eval_status" in splice_headers
    assert "nearest_splice_site_distance_bp" in splice_headers
    assert "delta_inclusion_deletion" in splice_headers
    first_splice_row = [cell.value for cell in next(splice_sheet.iter_rows(min_row=2, max_row=2))]
    splice_values = dict(zip(splice_headers, first_splice_row))
    assert splice_values["name"] == "OTUD4_46"
    assert splice_values["splice_eval_status"] == "scored"
    assert splice_values["nearest_splice_site_distance_bp"] == 38
    assert splice_values["delta_inclusion_deletion"] == -0.82
    normal_workbook = load_workbook(scoring_path.with_suffix(".xlsx"), read_only=True, data_only=True)
    assert "Splice" not in normal_workbook.sheetnames
    assert (
        pipeline_tools._summary_rel_path(
            "runs/goal_eval/v2_tasks/eval_otud4_user_asos/scored_tiles.json",
            "runs/goal_eval",
        )
        == "v2_tasks/eval_otud4_user_asos/scored_tiles.json"
    )


def test_chooser_focused_v2_markdown_answers_provided_splice_evaluation(tmp_path: Path) -> None:
    run_dir = tmp_path / "goal_splice_eval"
    splice_dir = run_dir / "v2_tasks" / "eval_aso_1"
    splice_dir.mkdir(parents=True)
    (run_dir / "request.txt").write_text(
        "Evaluate TTTAGCTTGCGGTCTTTAGA for exon 27 skipping on the requested transcript. "
        "Return whether it decreases inclusion, the qualitative effect size, and the acceptor rationale.",
        encoding="utf-8",
    )
    (splice_dir / "scored_tiles.json").write_text(
        json.dumps(
            {
                "rows": [
                    {
                        "name": "user_aso_1",
                        "Sequence": "TTTAGCTTGCGGTCTTTAGA",
                        "delta_inclusion_deletion": None,
                        "delta_inclusion_replacement": None,
                        "delta_names": "acceptor@40303948",
                        "delta_deletion": "-0.9985",
                        "delta_replacement": "-0.9983",
                    }
                ]
            }
        ),
        encoding="utf-8",
    )
    payload = {
        "assumptions": ["The intended modality is a steric-blocking splice-switching oligo."],
        "limitations": ["No experimental splicing readout was run."],
        "suggested_followups": ["Design additional acceptor-tiling SSOs to empirically rank skippers."],
        "evaluation_summary": {
            "splice_related_request": True,
            "context_summary": {
                "gene_symbol": "TESTGENE",
                "species": "human",
                "transcript_id": "ENST00000000001",
                "target_context": "human iPSC-derived dopaminergic neurons",
            },
            "evaluation_targets": [
                {
                    "transcript_id": "ENST00000000001",
                    "desired_effect": "skip_exon",
                    "target_type": "splicing",
                    "site_kind": "acceptor",
                    "exon_label": 27,
                }
            ],
            "provided_asos": [
                {
                    "name": "user_aso_1",
                    "sequence": "TTTAGCTTGCGGTCTTTAGA",
                    "chemistry": "steric non-RNase H splice-switching ASO",
                    "mapping_status": "mapped",
                    "transcript_id": "ENST00000000001",
                    "genomic_span": {"chrom": "12", "start": 40303943, "end": 40303962},
                    "reverse_complement_sequence": "TCTAAAGACCGCAAGCTAAA",
                    "off_target_gene_count": 0,
                    "isoform_uniqueness_class": "non_target",
                    "isoform_uniqueness_score": 0.0,
                    "selection_rank": 1,
                    "orientation_interpretation": {
                        "provided_sequence_role": "target_site_sequence",
                        "synthesized_antisense_sequence": "TCTAAAGACCGCAAGCTAAA",
                    },
                    "splice_eval_summary": {
                        "delta_inclusion_deletion": None,
                        "delta_inclusion_replacement": None,
                        "delta_names": "acceptor@40303948",
                        "delta_deletion": "-0.9985",
                        "delta_replacement": "-0.9983",
                    },
                }
            ],
            "limitations": [],
        },
    }

    markdown = pipeline_tools._chooser_focused_v2_markdown(
        run_dir=str(run_dir),
        markdown_text="# Sequence Orientation\nVerbose orientation dump",
        response_payload=payload,
        response_basename="existing_aso_evaluation",
    )

    assert markdown.startswith("# Answer")
    assert markdown.find("# Answer") < markdown.find("## What To Do Next")
    assert markdown.find("## What To Do Next") < markdown.find("## Main Files To Open")
    assert "Yes: it is predicted to decrease exon 27 inclusion and increase skipping" in markdown
    assert "very strong" in markdown
    assert "delta_deletion` = -0.9985" in markdown
    assert "acceptor junction" in markdown
    assert "TCTAAAGACCGCAAGCTAAA" in markdown
    assert "target-site (sense) sequence" in markdown
    assert "Exact other-gene off-targets" in markdown
    assert "Splice evaluation scored 1/1 row(s) and skipped 0/1 row(s)." in markdown
    assert "use two shortlists" not in markdown
    assert "Top Normal-Ranking Candidates" not in markdown
    assert "Sequence Orientation" not in markdown
    assert "Design additional acceptor-tiling SSOs" in markdown
    assert "No experimental splicing readout was run." in markdown


def test_chooser_focused_v2_markdown_summarizes_composite_outputs(tmp_path: Path) -> None:
    run_dir = tmp_path / "goal_composite"
    child_report = run_dir / "child_runs" / "design_otud4" / "report.md"
    child_report.parent.mkdir(parents=True)
    child_report.write_text("# Child Report\n", encoding="utf-8")
    payload = {
        "answer_summary": "Completed one design task.",
        "report_status": "complete",
        "completed_tasks": [
            {
                "task_id": "design_otud4",
                "task_type": "design_new_asos",
                "artifacts": {"response_md": str(child_report)},
            }
        ],
        "limitations": ["Final ASO selection still needs experimental review."],
    }

    markdown = pipeline_tools._chooser_focused_v2_markdown(
        run_dir=str(run_dir),
        markdown_text=f"Original detailed model response at {child_report}.",
        response_payload=payload,
        response_basename="composite_report",
    )

    assert markdown.startswith("# Run Output Summary")
    assert "## Task Outcomes" in markdown
    assert "child_runs/design_otud4/report.md" in markdown
    assert "Final ASO selection still needs experimental review." in markdown
    assert "## Detailed Response" in markdown
    assert "Original detailed model response at child_runs/design_otud4/report.md." in markdown
    assert "child_runs/design_otud4/report.md" in markdown
    assert str(run_dir) not in markdown

    regenerated = pipeline_tools._chooser_focused_v2_markdown(
        run_dir=str(run_dir),
        markdown_text=markdown,
        response_payload=payload,
        response_basename="composite_report",
    )
    assert regenerated.count("# Run Output Summary") == 1
    assert regenerated.count("## Detailed Response") == 1
    assert "Original detailed model response at child_runs/design_otud4/report.md." in regenerated
    assert str(run_dir) not in regenerated


def test_chooser_focused_v2_markdown_uses_design_outcomes_for_composite(tmp_path: Path) -> None:
    run_dir = tmp_path / "goal_multi"
    child_dir = run_dir / "child_runs" / "task_design_gene1"
    gapmer_dir = child_dir / "gapmer_tiling"
    gapmer_dir.mkdir(parents=True)
    (gapmer_dir / "gapmer_scored_uniqueness.xlsx").write_text("placeholder", encoding="utf-8")
    panel_dir = child_dir / "artifacts" / "candidates"
    panel_dir.mkdir(parents=True)
    (panel_dir / "design_panel.md").write_text(
        "# Candidate panel\n\nSelected candidate rows: 5.\n\n"
        "## Off-target search references\n\nGenomic index and mismatch budget recorded.\n",
        encoding="utf-8",
    )
    payload = {
        "answer_summary": "One design succeeded and one produced no candidates.",
        "report_status": "complete",
        "completed_tasks": [{"task_id": "t1", "task_type": "design_new_asos"}],
        "deterministic_summary": {
            "design_outcomes": [
                {
                    "gene": "GENE1",
                    "label": "Design GENE1",
                    "child_run_dir": str(child_dir),
                    "design_outcome": {
                        "outcome_status": "completed_with_candidates",
                        "candidate_count": 5,
                        "unique_sequence_count": 5,
                        "evidence_details": {
                            "cross_species": {
                                "checked_candidate_count": 20,
                                "passing_candidate_count": 5,
                                "requested_species": ["human", "mouse"],
                            }
                        },
                    },
                },
                {
                    "gene": "GENE2",
                    "label": "Design GENE2",
                    "child_run_dir": str(run_dir / "child_runs" / "task_design_gene2"),
                    "design_outcome": {
                        "outcome_status": "completed_no_candidates",
                        "candidate_count": 0,
                        "unique_sequence_count": 0,
                        "no_candidate_reason": "No exact cross-species candidates.",
                    },
                },
            ]
        },
    }

    markdown = pipeline_tools._chooser_focused_v2_markdown(
        run_dir=str(run_dir),
        markdown_text="Detailed model text.",
        response_payload=payload,
        response_basename="composite_report",
    )

    assert "## Design Run Status" in markdown
    assert "| GENE1 | `completed_with_candidates` | 5 | 5 | human, mouse | 5 exact / 20 checked (recorded filter scope) | `child_runs/task_design_gene1/gapmer_tiling/gapmer_scored_uniqueness.xlsx` |" in markdown
    assert "GENE2: `completed_no_candidates`. No exact cross-species candidates." in markdown
    assert "## Selected physical candidate panels" in markdown
    assert "### Off-target search references" in markdown
    assert "Genomic index and mismatch budget recorded." in markdown
    assert "see task artifacts" not in markdown
    assert str(run_dir) not in markdown


def test_provided_aso_scoring_merges_duplicate_sequences_by_row_index(tmp_path: Path) -> None:
    class FakeRuntime:
        def _annotate_tile_rows_with_rnafold(self, rows):
            for row in rows:
                row["No_secondary_structure"] = True
            return len(rows)

        def _resolve_gene_id(self, structured_goal, final_payload, *, allow_interactive=True):
            structured_goal["gene_id"] = "ENSGTARGET"
            final_payload["gene_id"] = "ENSGTARGET"
            return "ENSGTARGET"

        def _annotate_gapmer_off_targets(self, tiles_json_path, structured_goal, parameter_constraints):
            payload = json.loads(Path(tiles_json_path).read_text(encoding="utf-8"))
            for row in payload["rows"]:
                idx = int(row["provided_row_index"])
                row["Off_target_gene_ids"] = [f"ENSGOFF{idx}"]
                row["Off_target_genes"] = [f"OFF{idx}"]
                row["off_target_gene_count"] = idx + 1
            Path(tiles_json_path).write_text(json.dumps(payload), encoding="utf-8")
            return {"rows": len(payload["rows"]), "queries": 1, "rows_with_hits": 0}

        def _build_gapmer_kmer_index_by_len(self, cdna_map, lengths):
            return {4: {"TTTT": {"tx1", "tx2"}}}

        def _resolve_gapmer_uniqueness_context(self, structured_goal, iso_label_primary):
            return {"target_label": "tx1", "beneficial_labels": [], "pathogenic_labels": []}

        def _annotate_gapmer_uniqueness(
            self,
            tiles_json_path,
            isoform_kmers_by_len,
            target_isoform_label,
            total_isoforms,
            *,
            beneficial_isoform_labels=None,
            pathogenic_isoform_labels=None,
            output_name="provided_aso_scored_uniqueness.json",
            transcript_catalog=None,
            resolved_transcripts=None,
        ):
            payload = json.loads(Path(tiles_json_path).read_text(encoding="utf-8"))
            for row in payload["rows"]:
                labels = sorted(isoform_kmers_by_len[len(row["Sequence"])][row["Sequence"]])
                row["isoform_uniqueness_labels"] = labels
                row["isoform_uniqueness_count"] = len(labels)
                row["isoform_uniqueness_score"] = 1 / len(labels)
                row["isoform_uniqueness_class"] = "target_without_beneficial"
            out_path = Path(tiles_json_path).with_name(output_name)
            out_path.write_text(json.dumps(payload), encoding="utf-8")
            return str(out_path)

        def _annotate_gapmer_selection_ranks(self, tiles_json_path):
            payload = json.loads(Path(tiles_json_path).read_text(encoding="utf-8"))
            for row in payload["rows"]:
                row["selection_rank"] = int(row["provided_row_index"]) + 10
            Path(tiles_json_path).write_text(json.dumps(payload), encoding="utf-8")
            return {"status": "complete", "row_count": len(payload["rows"])}

    rows = [
        {"name": "duplicate", "sequence": "AAAA", "comparison_sequence": "TTTT", "chemistry": "steric-hindrance", "stage_results": {}},
        {"name": "duplicate", "sequence": "AAAA", "comparison_sequence": "TTTT", "chemistry": "steric-hindrance", "stage_results": {}},
    ]

    summary = pipeline_tools._score_provided_asos_with_existing_pipeline(
        rows=rows,
        eval_goal={"gene_symbol": "TEST", "species": "human"},
        cdna_map={"tx1": "TTTT", "tx2": "TTTT"},
        iso_label="tx1",
        runtime=FakeRuntime(),
        run_dir=str(tmp_path),
        task_id="eval",
    )

    assert summary["status"] == "complete"
    assert rows[0]["Off_target_gene_ids"] == ["ENSGOFF0"]
    assert rows[1]["Off_target_gene_ids"] == ["ENSGOFF1"]
    assert rows[0]["selection_rank"] == 10
    assert rows[1]["selection_rank"] == 11
    assert rows[0]["stage_results"]["ranking"]["status"] == "pass"


def test_existing_aso_orientation_interpretation_marks_antisense_sequence() -> None:
    entry = {
        "name": "aso_1",
        "sequence": "TCTAAAGACCGCAAGCTAAA",
        "reverse_complement_sequence": "TTTAGCTTGCGGTCTTTAGA",
        "mapping_hits": [
            {
                "match_orientation": "reverse_complement",
                "matched_target_sequence": "TTTAGCTTGCGGTCTTTAGA",
                "mapping_scope": "pre_mrna_locus",
            }
        ],
    }

    interpretation = pipeline_tools._provided_aso_orientation_interpretation(entry)
    assert interpretation["provided_sequence_role"] == "antisense_oligo_sequence"
    assert interpretation["matched_target_sequence"] == "TTTAGCTTGCGGTCTTTAGA"
    assert interpretation["synthesized_antisense_sequence"] == "TCTAAAGACCGCAAGCTAAA"


def test_repairs_empty_splice_site_plan_from_resolved_isoforms(monkeypatch) -> None:
    class Runtime:
        @staticmethod
        def _compute_splice_site_plan(objectives, isoform_transcripts, structured_goal):
            assert objectives[0]["desired_abundance"] == "decrease"
            assert "progerin" in isoform_transcripts
            return (
                [
                    {
                        "chrom": "1",
                        "position": 100,
                        "site": "donor",
                        "delta_direction": "decrease",
                        "present_in": ["progerin"],
                        "absent_in": ["lamin_a"],
                    }
                ],
                [{"chrom": "1", "position": 100, "site": "donor"}],
            )

    monkeypatch.setattr(pipeline_tools, "_goal_runtime_module", lambda: Runtime())
    goal = {
        "splice_site_plan": [],
        "splice_isoform_objectives": [
            {"isoform_label": "progerin", "desired_abundance": "decrease"},
            {"isoform_label": "lamin_a", "desired_abundance": "increase"},
        ],
        "isoform_transcripts": {"progerin": {"sites": []}, "lamin_a": {"sites": []}},
    }

    repaired = pipeline_tools._repair_splice_site_plan_from_resolved_isoforms(goal)

    assert repaired[0]["position"] == 100
    assert goal["splice_site_plan"] == repaired
    assert goal["splice_site_comparisons"][0]["site"] == "donor"


def test_runtime_splice_site_plan_uses_canonical_desired_abundance() -> None:
    plan, _comparisons = run_goal_intake._compute_splice_site_plan(
        [
            {"isoform_label": "pathogenic", "desired_abundance": "decrease"},
            {"isoform_label": "beneficial", "desired_abundance": "increase"},
        ],
        {
            "pathogenic": {
                "sites": [
                    {"chrom": "chr1", "donor": 100, "exon_index": 1, "strand": "+", "coding": True},
                    {"chrom": "chr1", "acceptor": 200, "exon_index": 2, "strand": "+", "coding": True},
                ]
            },
            "beneficial": {
                "sites": [
                    {"chrom": "chr1", "acceptor": 200, "exon_index": 2, "strand": "+", "coding": True},
                    {"chrom": "chr1", "donor": 300, "exon_index": 2, "strand": "+", "coding": True},
                ]
            },
        },
        {"splice_isoform_objectives": [
            {"isoform_label": "pathogenic", "desired_abundance": "decrease"},
            {"isoform_label": "beneficial", "desired_abundance": "increase"},
        ]},
    )

    by_label = {f"{entry['site']}@{entry['position']}": entry for entry in plan}
    assert by_label["donor@100"]["delta_direction"] == "decrease"
    assert by_label["donor@300"]["delta_direction"] == "increase"
    assert by_label["acceptor@200"]["delta_per_isoform"] == {
        "beneficial": "increase",
        "pathogenic": "decrease",
    }
    assert by_label["acceptor@200"]["isoform_abundance_objectives"] == {
        "beneficial": "increase",
        "pathogenic": "decrease",
    }


def test_design_goal_contract_reports_exact_canonical_paths_and_rejects_direction_aliases() -> None:
    issues = pipeline_tools._canonical_design_goal_contract_issues(
        {
            "design_programs": [
                {
                    "program_id": "branch",
                    "label": "branch",
                    "modality": "steric",
                    "intent": "splice_modulation",
                    "objective": "change the supplied target",
                    "direction": "decrease",
                }
            ],
            "evaluation_targets": [
                {
                    "objective_type": "primary_outcome",
                    "direction": "decrease",
                }
            ],
            "splice_isoform_objectives": [
                {"isoform_label": "supplied isoform", "direction": "decrease"}
            ],
        }
    )

    assert issues == [
        "context.structured_goal.design_programs[0].desired_delta",
        "context.structured_goal.evaluation_targets[0].desired_delta",
        "context.structured_goal.splice_isoform_objectives[0].desired_abundance",
    ]


def test_design_goal_contract_scopes_splice_execution_fields_to_programs() -> None:
    issues = pipeline_tools._canonical_design_goal_contract_issues(
        {
            "design_programs": [
                {
                    "program_id": "branch",
                    "label": "branch",
                    "modality": "steric",
                    "intent": "knockdown",
                    "objective": "change the supplied exon-level target",
                    "desired_delta": "decrease",
                    "desired_effect": "skipping",
                    "exon_label": 8,
                }
            ],
            "evaluation_targets": [
                {
                    "objective_type": "route_target",
                    "desired_delta": "decrease",
                    "desired_effect": "skipping",
                    "exon_label": 8,
                }
            ],
        }
    )

    assert issues == [
        "context.structured_goal.design_programs[0].run_splice_router",
        "context.structured_goal.design_programs[0].run_splice_eval",
        "context.structured_goal.design_programs[0].site_kind",
    ]


def test_named_splice_route_requires_matching_typed_isoform_objective() -> None:
    goal = {
        "design_programs": [
            {
                "program_id": "named_route",
                "label": "Named route",
                "modality": "steric",
                "intent": "splice_modulation",
                "objective": "Change the named isoform abundance.",
                "desired_delta": "increase",
                "isoform_label": "beneficial isoform",
            }
        ],
        "evaluation_targets": [
            {
                "objective_type": "route_target",
                "isoform_label": "beneficial isoform",
                "desired_delta": "increase",
            }
        ],
        "splice_isoform_objectives": [],
    }

    assert pipeline_tools._canonical_design_goal_contract_issues(goal) == [
        "context.structured_goal.splice_isoform_objectives"
    ]

    goal["splice_isoform_objectives"] = [
        {
            "objective_scope": "default_exon_context",
            "isoform_label": "beneficial isoform",
            "desired_abundance": "increase",
        }
    ]
    assert pipeline_tools._canonical_design_goal_contract_issues(goal) == [
        "context.structured_goal.splice_isoform_objectives[0].objective_scope"
    ]

    goal["splice_isoform_objectives"][0]["objective_scope"] = "specific_isoform"
    assert pipeline_tools._canonical_design_goal_contract_issues(goal) == []


def test_transcript_primary_outcome_does_not_require_splice_site_kind() -> None:
    issues = pipeline_tools._canonical_design_goal_contract_issues(
        {
            "design_programs": [
                {
                    "program_id": "branch",
                    "label": "branch",
                    "modality": "steric",
                    "intent": "splice_modulation",
                    "objective": "change the supplied transcript outcome",
                    "desired_delta": "decrease",
                }
            ],
            "evaluation_targets": [
                {
                    "target_type": "transcript",
                    "objective_type": "primary_outcome",
                    "desired_delta": "decrease",
                    "desired_effect": "inclusion",
                    "transcript_id": "TX-SUPPLIED",
                }
            ],
        }
    )

    assert (
        "context.structured_goal.evaluation_targets[0].site_kind"
        not in issues
    )


def test_splice_relative_window_anchor_requires_matching_program_site_kind() -> None:
    goal = {
        "design_programs": [
            {
                "program_id": "regulatory_window",
                "label": "Regulatory window",
                "modality": "steric",
                "intent": "splice_modulation",
                "objective": "Modulate splicing through a junction-relative window.",
                "desired_delta": "increase",
                "region_constraints": [
                    {
                        "source_type": "splice_relative_window",
                        "splice_relative_window": {
                            "anchor_site": "donor",
                            "relative_start_nt": 5,
                            "relative_end_nt": 25,
                        },
                    }
                ],
            }
        ]
    }

    assert pipeline_tools._canonical_design_goal_contract_issues(goal) == [
        "context.structured_goal.design_programs[0].site_kind"
    ]

    goal["design_programs"][0]["site_kind"] = "donor"
    assert pipeline_tools._canonical_design_goal_contract_issues(goal) == []


def test_design_goal_contract_reports_cross_field_contradictions() -> None:
    issues = pipeline_tools._canonical_design_goal_contract_issues(
        {
            "intent": "splice_modulation",
            "design_programs": [
                {
                    "program_id": "branch",
                    "label": "branch",
                    "modality": "steric",
                    "intent": "binding",
                    "objective": "block the supplied target",
                    "desired_delta": "increase",
                }
            ],
            "evaluation_targets": [
                {
                    "objective_type": "primary_outcome",
                    "desired_delta": "decrease",
                }
            ],
        }
    )

    assert issues == ["context.structured_goal.intent"]


def test_design_goal_contract_allows_independent_outcome_directions() -> None:
    issues = pipeline_tools._canonical_design_goal_contract_issues(
        {
            "design_programs": [
                {
                    "program_id": "branch",
                    "label": "branch",
                    "modality": "steric",
                    "intent": "splice_modulation",
                    "objective": "shift the requested isoform balance",
                    "desired_delta": "decrease",
                }
            ],
            "evaluation_targets": [
                {
                    "objective_type": "primary_outcome",
                    "desired_delta": "decrease",
                },
                {
                    "objective_type": "primary_outcome",
                    "desired_delta": "increase",
                },
            ],
        }
    )

    assert issues == []


def test_design_goal_contract_rejects_splice_effect_direction_conflicts() -> None:
    goal = {
        "intent": "splice_modulation",
        "design_programs": [
            {
                "program_id": "splice_branch",
                "label": "Change exon usage",
                "modality": "steric",
                "intent": "splice_modulation",
                "objective": "Increase skipping of the requested exon.",
                "desired_effect": "skipping",
                "desired_delta": "increase",
                "site_kind": "both",
                "run_splice_router": True,
                "run_splice_eval": True,
            }
        ],
        "evaluation_targets": [
            {
                "objective_type": "primary_outcome",
                "target_type": "exon",
                "exon_label": 8,
                "desired_effect": "inclusion",
                "desired_delta": "decrease",
            }
        ],
        "splice_isoform_objectives": [
            {
                "objective_scope": "default_exon_context",
                "desired_effect": "skipping",
                "desired_abundance": "increase",
            }
        ],
    }

    assert pipeline_tools._canonical_design_goal_contract_issues(goal) == [
        "context.structured_goal.design_programs[0].desired_delta",
        "context.structured_goal.evaluation_targets[0].desired_delta",
        "context.structured_goal.splice_isoform_objectives[0].desired_abundance",
    ]

    goal["design_programs"][0]["desired_delta"] = "decrease"
    goal["evaluation_targets"][0]["desired_delta"] = "increase"
    goal["splice_isoform_objectives"][0]["desired_abundance"] = "decrease"

    assert pipeline_tools._canonical_design_goal_contract_issues(goal) == []


def test_design_goal_contract_allows_route_direction_to_oppose_exon_outcome() -> None:
    goal = {
        "intent": "splice_modulation",
        "design_programs": [
            {
                "program_id": "alternate_site_branch",
                "label": "Reduce the competing splice site",
                "modality": "steric",
                "intent": "splice_modulation",
                "objective": "Reduce a competing acceptor to increase exon inclusion.",
                "desired_effect": "inclusion",
                "desired_delta": "decrease",
                "site_kind": "acceptor",
                "isoform_label": "alternate transcript",
                "run_splice_router": True,
                "run_splice_eval": True,
            }
        ],
    }

    assert pipeline_tools._canonical_design_goal_contract_issues(goal) == []


def test_design_goal_contract_scopes_splice_direction_to_target_identity() -> None:
    goal = {
        "design_programs": [
            {
                "program_id": "splice_branch",
                "label": "Change exon usage",
                "modality": "steric",
                "intent": "splice_modulation",
                "objective": "Decrease usage of the requested exon.",
                "desired_effect": "skipping",
                "desired_delta": "decrease",
                "site_kind": "both",
                "run_splice_router": True,
                "run_splice_eval": True,
            }
        ],
        "evaluation_targets": [
            {
                "objective_type": "primary_outcome",
                "target_type": "isoform_cohort",
                "isoform_label": "productive transcript cohort",
                "desired_effect": "skipping",
                "desired_delta": "increase",
            },
            {
                "objective_type": "route_target",
                "target_type": "exon",
                "exon_label": "target exon",
                "desired_effect": "skipping",
                "desired_delta": "decrease",
            },
        ],
    }

    assert pipeline_tools._canonical_design_goal_contract_issues(goal) == []

    goal["evaluation_targets"][1]["desired_delta"] = "increase"
    assert pipeline_tools._canonical_design_goal_contract_issues(goal) == [
        "context.structured_goal.evaluation_targets[1].desired_delta"
    ]


def test_design_goal_contract_allows_isoform_abundance_opposite_splice_event() -> None:
    goal = {
        "design_programs": [
            {
                "program_id": "splice_branch",
                "label": "Shift the supplied isoform balance",
                "modality": "steric",
                "intent": "splice_modulation",
                "objective": "Shift the supplied isoform balance",
                "desired_delta": "decrease",
            }
        ],
        "splice_isoform_objectives": [
            {
                "objective_scope": "specific_transcript",
                "transcript_id": "ENST00000000001.1",
                "desired_effect": "skipping",
                "desired_abundance": "increase",
            },
            {
                "objective_scope": "specific_isoform",
                "isoform_label": "included isoform",
                "desired_effect": "inclusion",
                "desired_abundance": "decrease",
            },
        ]
    }

    assert pipeline_tools._canonical_design_goal_contract_issues(goal) == []


def test_design_goal_contract_supports_nondirectional_monitoring_targets() -> None:
    goal = {
        "design_programs": [
            {
                "program_id": "branch",
                "label": "branch",
                "modality": "steric",
                "intent": "binding",
                "objective": "block the supplied site",
                "desired_delta": "decrease",
            }
        ],
        "evaluation_targets": [
            {
                "objective_type": "monitoring_safeguard",
                "label": "orthogonal readout",
            }
        ],
    }

    assert pipeline_tools._canonical_design_goal_contract_issues(goal) == []

    goal["evaluation_targets"][0]["desired_delta"] = "increase"
    assert pipeline_tools._canonical_design_goal_contract_issues(goal) == [
        "context.structured_goal.evaluation_targets[0].desired_delta"
    ]


def test_design_goal_contract_separates_directional_outcomes_from_identity_anchors() -> None:
    goal = {
        "design_programs": [
            {
                "program_id": "junction_branch",
                "label": "Junction-targeted knockdown",
                "modality": "gapmer",
                "intent": "knockdown",
                "objective": "Decrease the requested isoform through its defining junction.",
                "desired_delta": "decrease",
                "exon_labels": [2, 4],
            }
        ],
        "evaluation_targets": [
            {
                "objective_type": "primary_outcome",
                "target_type": "isoform",
                "isoform_label": "requested isoform",
                "desired_delta": "decrease",
            },
            {
                "objective_type": "identity_anchor",
                "target_type": "exon",
                "exon_label": 2,
            },
            {
                "objective_type": "identity_anchor",
                "target_type": "exon",
                "exon_label": 4,
            },
        ],
    }

    assert pipeline_tools._canonical_design_goal_contract_issues(goal) == []

    goal["evaluation_targets"][1]["desired_delta"] = "decrease"
    assert pipeline_tools._canonical_design_goal_contract_issues(goal) == [
        "context.structured_goal.evaluation_targets[1].desired_delta"
    ]

    del goal["evaluation_targets"][1]["desired_delta"]
    del goal["evaluation_targets"][1]["objective_type"]
    assert pipeline_tools._canonical_design_goal_contract_issues(goal) == [
        "context.structured_goal.evaluation_targets[1].objective_type"
    ]


def test_binding_only_programs_do_not_require_optional_target_directions() -> None:
    goal = {
        "design_programs": [
            {
                "program_id": "binding_branch",
                "label": "binding branch",
                "modality": "steric",
                "intent": "binding",
                "objective": "block the supplied binding sites",
            }
        ],
        "evaluation_targets": [
                {
                    "objective_type": "identity_anchor",
                    "target_type": "binding_site",
                    "label": "supplied site",
                    "chrom": "chr1",
                "start": 10,
                "end": 20,
            }
        ],
    }

    assert pipeline_tools._canonical_design_goal_contract_issues(goal) == []


def test_allele_specific_goal_contract_requires_typed_variant_and_overlap_policy() -> None:
    goal = {
        "design_programs": [
            {
                "program_id": "variant_branch",
                "label": "Variant branch",
                "modality": "gapmer",
                "intent": "knockdown",
                "objective": "Decrease the requested mutant allele.",
                "desired_delta": "decrease",
            }
        ],
        "evaluation_targets": [
            {
                "target_type": "mutant_variant",
                "objective_type": "primary_outcome",
                "desired_delta": "decrease",
            },
            {"target_type": "wildtype_variant", "objective_type": "monitoring_safeguard"},
        ],
    }

    assert pipeline_tools._canonical_design_goal_contract_issues(goal) == [
        "context.structured_goal.patient_variant.notation",
        "context.structured_goal.patient_variant.kind",
        "context.structured_goal.constraints.allele_specific.target_allele",
        "context.structured_goal.constraints.allele_specific.non_target_allele",
        "context.structured_goal.constraints.allele_specific.require_variant_overlap",
    ]

    goal["patient_variant"] = {
        "notation": "NC_TEST:g.101A>C",
        "kind": "genomic",
        "chr": "chr1",
        "pos": 101,
        "ref": "A",
        "alt": "C",
    }
    goal["constraints"] = {
        "allele_specific": {
            "target_allele": "C",
            "non_target_allele": "A",
            "require_variant_overlap": True,
        }
    }

    assert pipeline_tools._canonical_design_goal_contract_issues(goal) == []


def test_patient_sequence_context_does_not_imply_allele_specific_policy() -> None:
    goal = {
        "design_programs": [
            {
                "program_id": "context_branch",
                "label": "Sequence-context branch",
                "modality": "gapmer",
                "intent": "knockdown",
                "objective": "Decrease the requested RNA in the supplied context.",
                "desired_delta": "decrease",
            }
        ],
        "patient_variant": {
            "notation": "TEST:c.10A>C",
            "kind": "cdna",
            "ref": "A",
            "alt": "C",
        },
    }

    assert pipeline_tools._canonical_design_goal_contract_issues(goal) == []


def test_patient_variant_resolution_uses_source_backed_gene_and_refseq_mapping(
    monkeypatch: Any,
) -> None:
    def fake_ensembl_json(url: str) -> Any:
        if "/vep/" in url:
            return [
                {
                    "assembly_name": "GRCh38",
                    "seq_region_name": "7",
                    "start": 501,
                    "strand": -1,
                    "allele_string": "G/A",
                    "transcript_consequences": [
                        {
                            "transcript_id": "ENST_TEST_LOCKED",
                            "gene_id": "ENSG_TEST",
                            "gene_symbol": "GENE_TEST",
                            "biotype": "protein_coding",
                            "strand": -1,
                            "cdna_start": 44,
                            "cds_start": 12,
                            "variant_allele": "A",
                        },
                        {
                            "transcript_id": "ENST_TEST_OTHER",
                            "gene_id": "ENSG_TEST",
                            "gene_symbol": "GENE_TEST",
                            "biotype": "protein_coding",
                            "strand": -1,
                            "cdna_start": 51,
                            "cds_start": 12,
                            "variant_allele": "A",
                        },
                    ],
                }
            ]
        if "ENST_TEST_LOCKED" in url:
            return [
                {
                    "display_id": "NM_TEST.4",
                    "primary_id": "NM_TEST",
                    "version": "4",
                    "dbname": "RefSeq_mRNA",
                }
            ]
        return []

    monkeypatch.setattr(run_goal_intake, "_ensembl_variant_json", fake_ensembl_json)
    goal = {
        "species": "human",
        "transcript_policy": "user_locked",
        "transcript_lock": {
            "kind": "reference_accession",
            "value": "NM_TEST.4",
            "note": "User supplied this accession.",
        },
        "patient_variant": {
            "notation": "NM_TEST.4:c.12G>A",
            "kind": "genomic",
            "reference_accession": "NM_TEST.4",
            "ref": "G",
            "alt": "A",
        },
    }
    final_payload: dict[str, Any] = {}

    result = run_goal_intake._resolve_patient_variant_context(goal, final_payload)

    assert result["status"] == "resolved"
    assert goal["gene_symbol"] == "GENE_TEST"
    assert goal["gene_id"] == "ENSG_TEST"
    assert goal["patient_variant"]["chr"] == "chr7"
    assert goal["patient_variant"]["genomic_ref"] == "C"
    assert goal["patient_variant"]["genomic_alt"] == "T"
    assert goal["patient_variant"]["kind"] == "cdna"
    assert goal["patient_variant"]["allele_orientation"] == "transcript"
    assert goal["transcript_lock"]["value"] == "ENST_TEST_LOCKED"
    assert result["reference_accession_resolution"]["confidence"] == "high"
    assert final_payload["patient_variant"]["resolution"]["source"] == "ensembl_vep"


def test_gene_qualified_cdna_variant_uses_source_backed_canonical_transcript(
    monkeypatch: Any,
) -> None:
    requested_urls: list[str] = []

    def fake_ensembl_json(url: str) -> Any:
        requested_urls.append(url)
        if "/lookup/symbol/" in url or "/lookup/id/" in url:
            return {"canonical_transcript": "ENST_CANONICAL.2"}
        if "/vep/" in url:
            return [
                {
                    "assembly_name": "GRCh38",
                    "seq_region_name": "1",
                    "start": 101,
                    "allele_string": "A/C",
                    "transcript_consequences": [
                        {
                            "transcript_id": "ENST_CANONICAL",
                            "gene_id": "ENSG_TEST",
                            "gene_symbol": "GENE_TEST",
                            "biotype": "protein_coding",
                            "strand": 1,
                            "cdna_start": 10,
                            "variant_allele": "C",
                        }
                    ],
                }
            ]
        return None

    monkeypatch.setattr(run_goal_intake, "_ensembl_variant_json", fake_ensembl_json)
    goal = {
        "species": "human",
        "gene_symbol": "GENE_TEST",
        "transcript_policy": "agent_choice",
        "patient_variant": {
            "notation": "GENE_TEST c.10A>C",
            "kind": "cdna",
            "ref": "A",
            "alt": "C",
        },
    }
    final_payload: dict[str, Any] = {}

    result = run_goal_intake._resolve_patient_variant_context(goal, final_payload)

    assert result["status"] == "resolved"
    assert result["resolved_transcript_id"] == "ENST_CANONICAL"
    assert result["hgvs_query_resolution"]["method"] == (
        "ensembl_canonical_transcript"
    )
    assert any(
        "/vep/homo_sapiens/hgvs/ENST_CANONICAL:c.10A%3EC" in url
        for url in requested_urls
    )
    assert goal["patient_variant"]["notation"] == "GENE_TEST c.10A>C"


def test_unresolvable_variant_query_is_semantic_not_infrastructure(
    monkeypatch: Any,
) -> None:
    def reject_query(_url: str) -> Any:
        raise run_goal_intake.EnsemblServiceError(
            "invalid variant query",
            retryable=False,
            status_code=400,
        )

    monkeypatch.setattr(run_goal_intake, "_ensembl_variant_json", reject_query)
    goal = {
        "species": "human",
        "gene_symbol": "GENE_TEST",
        "patient_variant": {
            "notation": "NM_TEST.1:c.invalid",
            "kind": "cdna",
        },
    }

    result = run_goal_intake._resolve_patient_variant_context(goal, {})

    assert result["status"] == "unresolved"
    assert result["reason"] == "patient_variant_not_resolved_by_ensembl_vep"


def test_negative_strand_genomic_variant_projects_complemented_transcript_alleles() -> None:
    goal = {
        "patient_variant": {
            "notation": "NC_TEST:g.101T>G",
            "kind": "genomic",
            "chr": "chr1",
            "pos": 101,
            "ref": "T",
            "alt": "G",
            "allele_orientation": "genomic",
            "resolution": {"status": "resolved"},
            "transcript_consequences": {
                "ENST_TEST": {
                    "transcript_id": "ENST_TEST",
                    "cdna_start": 3,
                    "strand": -1,
                }
            },
        }
    }
    entry = {"transcript_id": "ENST_TEST", "strand": "-", "exons_tx": []}

    projection = run_goal_intake._patient_variant_cdna_projection(
        goal,
        entry,
        "ENST_TEST",
        "CCAC",
    )

    assert projection == {
        "status": "resolved",
        "transcript_id": "ENST_TEST",
        "cdna_position": 3,
        "wildtype_allele": "A",
        "mutant_allele": "C",
        "wildtype_cdna": "CCAC",
        "mutant_cdna": "CCCC",
        "source": "ensembl_vep_validated_patient_variant",
    }


def test_hgvs_cdna_notation_overrides_incorrect_declared_genomic_kind() -> None:
    goal = {
        "patient_variant": {
            "notation": "NM_TEST.1:c.3G>A",
            "kind": "genomic",
            "ref": "G",
            "alt": "A",
            "allele_orientation": "genomic",
            "resolution": {"status": "resolved"},
            "transcript_consequences": {
                "ENST_TEST": {
                    "transcript_id": "ENST_TEST",
                    "cdna_start": 3,
                    "strand": -1,
                }
            },
        }
    }
    entry = {"transcript_id": "ENST_TEST", "strand": "-", "exons_tx": []}

    projection = run_goal_intake._patient_variant_cdna_projection(
        goal,
        entry,
        "ENST_TEST",
        "CCGC",
    )

    assert projection["wildtype_allele"] == "G"
    assert projection["mutant_allele"] == "A"
    assert projection["mutant_cdna"] == "CCAC"
    assert goal["patient_variant"]["kind"] == "cdna"
    assert goal["patient_variant"]["allele_orientation"] == "transcript"
    assert goal["patient_variant"]["genomic_ref"] == "C"
    assert goal["patient_variant"]["genomic_alt"] == "T"


def test_allele_specific_gapmer_tiling_keeps_only_variant_overlapping_windows(
    monkeypatch: Any,
    tmp_path: Path,
) -> None:
    wildtype_cdna = "ACGT" * 12 + "AC"
    variant_position = 25
    assert wildtype_cdna[variant_position - 1] == "A"
    goal = {
        "species": "human",
        "chemistry": "5-10-5",
        "constraints": {
            "allele_specific": {
                "target_allele": "C",
                "non_target_allele": "A",
                "require_variant_overlap": True,
            }
        },
        "patient_variant": {
            "notation": "NM_TEST.1:c.25A>C",
            "kind": "cdna",
            "ref": "A",
            "alt": "C",
            "allele_orientation": "transcript",
            "resolution": {"status": "resolved"},
            "transcript_consequences": {
                "ENST_TEST": {
                    "transcript_id": "ENST_TEST",
                    "cdna_start": variant_position,
                    "strand": 1,
                }
            },
        },
    }
    transcript = {
        "transcript_id": "ENST_TEST",
        "strand": "+",
        "chrom": "chr1",
        "exons_tx": [{"index": 1, "start": 1, "end": len(wildtype_cdna)}],
    }
    monkeypatch.setattr(
        run_goal_intake,
        "_build_isoform_cdna_sequences",
        lambda *_args, **_kwargs: {"ENST_TEST": wildtype_cdna},
    )
    monkeypatch.setattr(
        run_goal_intake,
        "_annotate_tile_rows_with_rnafold",
        lambda _rows: None,
    )

    result = run_goal_intake._build_cdna_gapmer_tiles(
        goal,
        {"ENST_TEST": transcript},
        "ENST_TEST",
        {"ENST_TEST": transcript},
        str(tmp_path),
    )

    assert result.get("error") is None
    payload = json.loads(Path(result["json_path"]).read_text(encoding="utf-8"))
    rows = payload["rows"]
    assert len(rows) == 20
    assert all(row["allele_specific"] is True for row in rows)
    assert all(row["wildtype_mismatch_count"] == 1 for row in rows)
    assert all(row["mutant_target_sequence"] != row["wildtype_target_sequence"] for row in rows)
    assert all(1 <= row["variant_position_in_aso"] <= 20 for row in rows)
    assert sum(bool(row["variant_in_gap"]) for row in rows) == 10


def test_allele_specific_gapmer_ranking_prioritizes_central_gap_mismatch(
    tmp_path: Path,
) -> None:
    tiles_path = tmp_path / "allele_specific_tiles.json"
    shared = {
        "allele_specific": True,
        "wildtype_mismatch_count": 1,
        "variant_in_gap": True,
        "isoform_uniqueness_class": "gene_wide",
        "isoform_uniqueness_score": 1.0,
        "off_target_gene_count": 0,
        "GC": 0.5,
    }
    rows = [
        {
            **shared,
            "Name": "gap-edge",
            "Sequence": "A" * 20,
            "mutant_aso_sequence": "T" * 20,
            "variant_position_in_aso": 6,
        },
        {
            **shared,
            "Name": "gap-center",
            "Sequence": "C" * 20,
            "mutant_aso_sequence": "G" * 20,
            "variant_position_in_aso": 10,
        },
        {
            **shared,
            "Name": "wing",
            "Sequence": "G" * 20,
            "mutant_aso_sequence": "C" * 20,
            "variant_position_in_aso": 2,
            "variant_in_gap": False,
        },
    ]
    tiles_path.write_text(json.dumps({"rows": rows}), encoding="utf-8")

    result = run_goal_intake._annotate_gapmer_selection_ranks(str(tiles_path))
    ranked = json.loads(tiles_path.read_text(encoding="utf-8"))["rows"]

    assert result["status"] == "complete"
    assert [row["Name"] for row in ranked] == ["gap-center", "gap-edge", "wing"]
    assert ranked[0]["selection_score_components"]["variant_in_gap"] is True


def test_gapmer_constraints_rematerialize_stale_selection_ranking_metadata(
    tmp_path: Path,
) -> None:
    scored_path = tmp_path / "gapmer_scored_uniqueness.json"
    scored_path.write_text(
        json.dumps(
            {
                "rows": [
                    {
                        "Name": "wing",
                        "Sequence": "A" * 20,
                        "mutant_aso_sequence": "T" * 20,
                        "allele_specific": True,
                        "wildtype_mismatch_count": 1,
                        "variant_position_in_aso": 2,
                        "variant_in_gap": False,
                    },
                    {
                        "Name": "center",
                        "Sequence": "C" * 20,
                        "mutant_aso_sequence": "G" * 20,
                        "allele_specific": True,
                        "wildtype_mismatch_count": 1,
                        "variant_position_in_aso": 10,
                        "variant_in_gap": True,
                    },
                ]
            }
        ),
        encoding="utf-8",
    )
    runtime = SimpleNamespace(
        _annotate_gapmer_selection_ranks=run_goal_intake._annotate_gapmer_selection_ranks,
        filter_tiles_json=lambda path, constraints: {
            "applied": False,
            "output_path": path,
        },
    )

    result = pipeline_tools._apply_gapmer_constraints_to_result(
        gapmer_results={
            "scored_tiles_path": str(scored_path),
            "selection_ranking": {
                "status": "complete",
                "row_count": 2,
                "path": str(scored_path),
                "method": "superseded",
            },
        },
        structured_goal={},
        runtime=runtime,
    )
    ranked = json.loads(scored_path.read_text(encoding="utf-8"))["rows"]

    assert "error" not in result
    assert [row["Name"] for row in ranked] == ["center", "wing"]
    assert [row["selection_rank"] for row in ranked] == [1, 2]
    assert result["gapmer_results"]["selection_ranking"]["method"].startswith(
        "lexicographic_required_target_evidence"
    )


def test_request_validation_requires_child_runtime_species_at_canonical_path() -> None:
    envelope = {
        "tasks": [
            {
                "task_id": "design_a",
                "task_type": "design_new_asos",
                "context": {
                    "gene_symbol": "GENE_A",
                    "target_context": "primary cells",
                    "structured_goal": {
                        "gene_symbol": "GENE_A",
                        "target_context": "primary cells",
                        "intent": "knockdown",
                        "design_programs": [
                            {
                                "program_id": "program_a",
                                "label": "Program A",
                                "modality": "gapmer",
                                "intent": "knockdown",
                                "objective": "Decrease the requested RNA",
                                "desired_delta": "decrease",
                            }
                        ],
                    },
                },
            }
        ]
    }

    validation = pipeline_tools._request_validation_summary(envelope)

    assert validation["tasks"]["design_a"]["status"] == "runnable"
    assert validation["tasks"]["design_a"]["repairable_issues"] == [
        "context.structured_goal.species"
    ]


def test_request_validation_accepts_structured_species_set_for_anchor_materialization() -> None:
    envelope = {
        "tasks": [
            {
                "task_id": "design_a",
                "task_type": "design_new_asos",
                "context": {
                    "gene_symbol": "GENE_A",
                    "target_context": "primary cells",
                    "structured_goal": {
                        "gene_symbol": "GENE_A",
                        "target_context": "primary cells",
                        "intent": "knockdown",
                        "species_targets": [
                            {"species": "organism_a"},
                            {"species": "organism_b"},
                        ],
                        "design_programs": [
                            {
                                "program_id": "program_a",
                                "label": "Program A",
                                "modality": "gapmer",
                                "intent": "knockdown",
                                "objective": "Decrease the requested RNA",
                                "desired_delta": "decrease",
                            }
                        ],
                    },
                },
            }
        ]
    }

    validation = pipeline_tools._request_validation_summary(envelope)

    assert validation["tasks"]["design_a"]["repairable_issues"] == []


def test_isoform_objective_explicit_abundance_precedes_splice_effect() -> None:
    normalized = pipeline_tools._normalized_design_objective_preseed(
        {
            "splice_isoform_objectives": [
                {
                    "isoform_label": "supplied isoform",
                    "desired_abundance": "increase",
                    "desired_effect": "skipping",
                }
            ]
        }
    )["splice_isoform_objectives"]

    assert normalized[0]["desired_abundance"] == "increase"


def test_semantic_goal_audit_replaces_unsupported_value_from_request_evidence(
    monkeypatch: Any,
) -> None:
    request_text = "Use an 18-nt steric design for the supplied target in reviewed cells."
    envelope = {
        "schema_version": "v3",
        "request_text": request_text,
        "tasks": [
            {
                "task_id": "design",
                "task_type": "design_new_asos",
                "context": {
                    "target_context": "reviewed cells",
                    "structured_goal": {
                        "target_context": "reviewed cells",
                        "design_programs": [
                            {
                                "program_id": "branch",
                                "label": "branch",
                                "modality": "steric",
                                "intent": "binding",
                                "objective": "bind the supplied target",
                                "desired_delta": "decrease",
                                "length": 40,
                            }
                        ],
                    },
                },
            }
        ],
    }

    calls = {"count": 0}

    async def fake_audit(prompt: str) -> dict[str, Any]:
        calls["count"] += 1
        assert request_text in prompt
        if calls["count"] == 2:
            assert '"audit_mode": "post_repair_semantic_verification"' in prompt
            return {"status": "valid", "patches": [], "goal_replacements": [], "unresolved": []}
        return {
            "status": "repairable",
            "patches": [
                {
                    "op": "replace",
                    "path": "/tasks/0/context/structured_goal/design_programs/0/length",
                    "value": 18,
                    "evidence_refs": ["request"],
                    "evidence_quote": "18-nt",
                    "reason": "The request supplies the oligo length.",
                }
            ],
            "unresolved": [],
        }

    monkeypatch.setattr(
        pipeline_tools,
        "_run_goal_contract_audit_agent_once",
        fake_audit,
    )
    result = pipeline_tools._repair_request_envelope_if_needed(
        request_text=request_text,
        run_dir="",
        structured_goal={},
        request_envelope=envelope,
        qa_history=[],
        semantic_audit=True,
    )

    program = result["request_envelope"]["tasks"][0]["context"]["structured_goal"][
        "design_programs"
    ][0]
    assert result["status"] == "ready"
    assert calls["count"] == 2
    assert program["length"] == 18


def test_parent_envelope_preseed_is_reused_only_before_child_clarification() -> None:
    envelope = {
        "schema_version": "v3",
        "tasks": [
            {
                "task_id": "design",
                "task_type": "design_new_asos",
                "context": {
                    "structured_goal": {
                        "target_context": "supplied cells",
                        "design_programs": [],
                    }
                },
            }
        ],
    }
    state = {"request_envelope_preseed": envelope}

    parsed = pipeline_tools._request_parse_output_from_envelope_preseed(
        state,
        [],
    )

    assert parsed == {
        "status": "ready",
        "questions": [],
        "reason": None,
        "request_envelope": envelope,
    }
    assert parsed["request_envelope"] is not envelope
    assert (
        pipeline_tools._request_parse_output_from_envelope_preseed(
            state,
            [{"question": "Which scope?", "answer": "Use the supplied scope."}],
        )
        == {}
    )


def test_internal_envelope_preseed_uses_internal_not_llm_transport_contract() -> None:
    envelope = {
        "schema_version": "v3",
        "request_id": "internal_request",
        "request_text": "Execute the already validated task.",
        "request_mode": "single_task",
        "request_summary": "Execute the validated task.",
        "request_classifications": ["design_new_asos"],
        "workflow_name": "design_new_asos",
        "workflow_stage_order": ["intake", "pipeline", "report"],
        "default_context": {},
        "tasks": [
            {
                "task_id": "design",
                "task_type": "design_new_asos",
                "label": "Validated design",
                "native": True,
                "can_run_in_parallel": False,
                "context": {},
            }
        ],
        "task_dependencies": [],
        "final_outputs": [
            {"kind": "design_report", "task_id": "design", "label": None}
        ],
        "response_mode": "report",
        "requires_execution": True,
        "assumptions": [],
    }
    payload = {
        "status": "ready",
        "questions": [],
        "reason": None,
        "request_envelope": envelope,
    }

    assert pipeline_tools._request_workbench_contract_errors(payload)
    assert pipeline_tools._request_parse_contract_errors(
        payload,
        reused_envelope_preseed=True,
    ) == []


def test_child_preseed_uses_audited_canonical_not_runtime_projection() -> None:
    canonical_task = {
        "task_id": "design",
        "task_type": "design_new_asos",
        "label": "Validated design",
        "native": True,
        "can_run_in_parallel": False,
        "context": {
            "structured_goal": {
                "design_programs": [
                    {
                        "program_id": "primary",
                        "label": "primary",
                        "modality": "steric",
                        "intent": "binding",
                        "objective": "bind the supplied target",
                        "desired_delta": "decrease",
                    }
                ]
            }
        },
    }
    canonical_envelope = {
        "schema_version": "v3",
        "tasks": [canonical_task],
        "final_outputs": [
            {"kind": "design_report", "task_id": "design", "label": None}
        ],
    }
    runtime_envelope = deepcopy(canonical_envelope)
    runtime_goal = runtime_envelope["tasks"][0]["context"]["structured_goal"]
    runtime_goal["parameter_constraints"] = {
        "runtime_projection": True
    }

    scoped = pipeline_tools._task_scoped_child_preseed_envelope(
        {"canonical_request_envelope": canonical_envelope},
        runtime_envelope,
        runtime_envelope["tasks"][0],
    )

    scoped_goal = scoped["tasks"][0]["context"]["structured_goal"]
    assert "parameter_constraints" not in scoped_goal
    assert scoped_goal == canonical_task["context"]["structured_goal"]

    evidence_task, evidence_context = pipeline_tools._child_design_evidence_scope(
        scoped,
        runtime_envelope["tasks"][0],
    )
    child_request = pipeline_tools._build_scoped_design_request_text(
        "Execute the validated design.",
        evidence_task,
        scoped,
        [],
    )

    assert "runtime_projection" not in evidence_context["structured_goal"]
    assert "runtime_projection" not in child_request
    assert evidence_task["context"] == canonical_task["context"]


def test_composite_design_envelope_is_goal_ready_only_when_every_goal_exists() -> None:
    tasks = [
        {
            "task_id": "design_a",
            "task_type": "design_new_asos",
            "context": {"structured_goal": {"design_programs": [{}]}},
        },
        {
            "task_id": "design_b",
            "task_type": "design_new_asos",
            "context": {"structured_goal": {"design_programs": [{}]}},
        },
        {
            "task_id": "report",
            "task_type": "build_composite_report",
            "context": {},
        },
    ]
    envelope = {"tasks": tasks}

    assert pipeline_tools._request_envelope_has_complete_design_goals(envelope)

    tasks[1]["context"].pop("structured_goal")
    assert not pipeline_tools._request_envelope_has_complete_design_goals(envelope)


def test_auditor_can_remove_unsupported_optional_field_that_triggers_dependencies(
    monkeypatch: Any,
) -> None:
    request_text = "Create a binding-only steric design in reviewed cells."
    envelope = {
        "schema_version": "v3",
        "request_text": request_text,
        "tasks": [
            {
                "task_id": "design",
                "task_type": "design_new_asos",
                "context": {
                    "target_context": "reviewed cells",
                    "structured_goal": {
                        "target_context": "reviewed cells",
                        "intent": "binding",
                        "design_programs": [
                            {
                                "program_id": "primary",
                                "label": "Primary design",
                                "modality": "steric",
                                "intent": "binding",
                                "objective": "Bind the supplied target.",
                                "desired_delta": "increase",
                                "desired_effect": "inclusion",
                            }
                        ],
                    },
                },
            }
        ],
    }
    calls = {"count": 0}

    async def fake_audit(prompt: str) -> dict[str, Any]:
        calls["count"] += 1
        if calls["count"] == 2:
            assert '"audit_mode": "post_repair_semantic_verification"' in prompt
            return {
                "status": "valid",
                "patches": [],
                "goal_replacements": [],
                "field_dispositions": [],
                "unresolved": [],
            }
        assert "run_splice_router" in prompt
        assert "run_splice_eval" in prompt
        return {
            "status": "repairable",
            "patches": [
                {
                    "op": "remove",
                    "path": "/tasks/0/context/structured_goal/design_programs/0/desired_effect",
                    "value": None,
                    "evidence_refs": ["request"],
                    "evidence_quote": "binding-only steric design",
                    "reason": "The evidence supports binding, not a splice effect.",
                }
            ],
            "goal_replacements": [],
            "field_dispositions": [],
            "unresolved": [],
        }

    monkeypatch.setattr(
        pipeline_tools,
        "_run_goal_contract_audit_agent_once",
        fake_audit,
    )
    result = pipeline_tools._repair_request_envelope_if_needed(
        request_text=request_text,
        run_dir="",
        structured_goal={},
        request_envelope=envelope,
        qa_history=[],
        semantic_audit=True,
    )

    program = result["request_envelope"]["tasks"][0]["context"][
        "structured_goal"
    ]["design_programs"][0]
    assert result["status"] == "ready"
    assert calls["count"] == 2
    assert "desired_effect" not in program
    assert "run_splice_router" not in program
    assert "run_splice_eval" not in program


def test_semantic_audit_reserves_verification_and_exposes_repair_cycles(
    monkeypatch: Any,
) -> None:
    request_text = (
        "Design one binding branch in supplied cells; supplied context alpha and "
        "supplied annotation are source labels."
    )
    envelope = {
        "schema_version": "v3",
        "request_text": request_text,
        "tasks": [
            {
                "task_id": "design",
                "task_type": "design_new_asos",
                "context": {
                    "target_context": "supplied cells",
                    "structured_goal": {
                        "target_context": "supplied cells",
                        "design_programs": [
                            {
                                "program_id": "primary",
                                "label": "binding branch",
                                "modality": "steric",
                                "intent": "binding",
                                "objective": "bind the supplied target",
                                "desired_delta": "decrease",
                                "length": 18,
                            }
                        ],
                    },
                },
            }
        ],
    }
    calls = {"count": 0}

    async def fake_audit(prompt: str) -> dict[str, Any]:
        calls["count"] += 1
        audit_input = pipeline_tools._goal_audit_protocol_payload(prompt)
        program = audit_input["current_request_envelope"]["tasks"][0][
            "context"
        ]["structured_goal"]["design_programs"][0]
        if calls["count"] == 1:
            assert audit_input["repair_cycle_history"] == []
            return {
                "status": "repairable",
                "patches": [
                    {
                        "op": "replace",
                        "path": "/tasks/0/context/structured_goal/design_programs/0/length",
                        "value": 20,
                        "evidence_refs": ["request"],
                        "evidence_quote": "supplied annotation",
                        "reason": "Apply the first evidence-backed correction.",
                    }
                ],
                "goal_replacements": [],
                "field_dispositions": [],
                "unresolved": [],
            }
        if calls["count"] == 2:
            assert program["length"] == 20
            return {
                "status": "repairable",
                "patches": [
                    {
                        "op": "replace",
                        "path": "/tasks/0/context/structured_goal/design_programs/0/length",
                        "value": 18,
                        "evidence_refs": ["request"],
                        "evidence_quote": "binding branch",
                        "reason": "Restore the original evidence-backed value.",
                    }
                ],
                "goal_replacements": [],
                "field_dispositions": [],
                "unresolved": [],
            }
        if calls["count"] == 3:
            history = audit_input["repair_cycle_history"]
            assert len(history) == 2
            assert [
                repair["op"]
                for transition in history
                for repair in transition["applied_repairs"]
            ] == ["replace", "replace"]
            return {
                "status": "repairable",
                "patches": [
                    {
                        "op": "replace",
                        "path": "/tasks/0/context/structured_goal/design_programs/0/length",
                        "value": 19,
                        "evidence_refs": ["request"],
                        "evidence_quote": "supplied annotation",
                        "reason": "Use the evidence-supported canonical destination.",
                    }
                ],
                "goal_replacements": [],
                "field_dispositions": [],
                "unresolved": [],
            }
        assert calls["count"] == 4
        assert audit_input["audit_mode"] == "post_repair_semantic_verification"
        assert program["length"] == 19
        return {
            "status": "valid",
            "patches": [],
            "goal_replacements": [],
            "field_dispositions": [],
            "unresolved": [],
        }

    monkeypatch.setenv("ASO_GOAL_CONTRACT_AUDIT_MAX_ROUNDS", "1")
    monkeypatch.setattr(
        pipeline_tools,
        "_run_goal_contract_audit_agent_once",
        fake_audit,
    )
    monkeypatch.setattr(
        pipeline_tools,
        "_build_request_envelope",
        lambda *, raw_envelope, **_kwargs: deepcopy(raw_envelope),
    )

    def apply_test_repairs(
        request_envelope: dict[str, Any],
        patches: Any,
        **_kwargs: Any,
    ) -> dict[str, Any]:
        updated = deepcopy(request_envelope)
        applied: list[dict[str, Any]] = []
        for patch in patches:
            updated["tasks"][0]["context"]["structured_goal"][
                "design_programs"
            ][0]["length"] = patch["value"]
            applied.append({"op": patch["op"], "path": patch["path"]})
        return {
            "request_envelope": updated,
            "applied": applied,
            "rejected": [],
        }

    monkeypatch.setattr(
        pipeline_tools,
        "_apply_goal_audit_patches",
        apply_test_repairs,
    )
    monkeypatch.setattr(
        pipeline_tools,
        "_resolved_declared_contract_blocks",
        lambda request_envelope: (deepcopy(request_envelope), False),
    )

    result = pipeline_tools._repair_request_envelope_if_needed(
        request_text=request_text,
        run_dir="",
        structured_goal={},
        request_envelope=envelope,
        qa_history=[],
        semantic_audit=True,
    )

    assert result["status"] == "ready", result
    assert calls["count"] == 4


def test_audit_remove_patch_cannot_delete_required_canonical_field() -> None:
    request_text = "Create a binding-only steric design in reviewed cells."
    envelope = {
        "tasks": [
            {
                "task_id": "design",
                "task_type": "design_new_asos",
                "context": {
                    "structured_goal": {
                        "design_programs": [
                            {
                                "program_id": "primary",
                                "label": "Primary design",
                                "modality": "steric",
                                "intent": "binding",
                                "objective": "Bind the supplied target.",
                                "desired_delta": "increase",
                            }
                        ]
                    }
                },
            }
        ]
    }

    result = pipeline_tools._apply_goal_audit_patches(
        envelope,
        [
            {
                "op": "remove",
                "path": "/tasks/0/context/structured_goal/design_programs/0/program_id",
                "value": None,
                "evidence_refs": ["request"],
                "evidence_quote": "binding-only steric design",
                "reason": "Attempt to remove a required field.",
            }
        ],
        issue_records=[],
        evidence_sources=[
            {
                "evidence_id": "request",
                "kind": "user_request",
                "text": request_text,
            }
        ],
    )

    assert result["applied"] == []
    assert result["rejected"] == [
        {
            "path": "/tasks/0/context/structured_goal/design_programs/0/program_id",
            "reason": "remove_target_not_optional_canonical_field",
        }
    ]


def test_audit_remove_patch_can_delete_optional_non_design_context_field() -> None:
    request_text = "Analyze the two supplied sequences."
    envelope = {
        "tasks": [
            {
                "task_id": "analysis",
                "task_type": "sandbox_analysis",
                "context": {
                    "gene_symbol": ".",
                    "provided_asos": ["ACGT", "TGCA"],
                },
            }
        ]
    }

    result = pipeline_tools._apply_goal_audit_patches(
        envelope,
        [
            {
                "op": "remove",
                "path": "/tasks/0/context/gene_symbol",
                "value": None,
                "evidence_refs": ["request"],
                "evidence_quote": "two supplied sequences",
                "reason": "The generated placeholder is unsupported.",
            }
        ],
        issue_records=[],
        evidence_sources=[
            {
                "evidence_id": "request",
                "kind": "user_request",
                "text": request_text,
            }
        ],
    )

    assert result["rejected"] == []
    assert result["request_envelope"]["tasks"][0]["context"] == {
        "provided_asos": ["ACGT", "TGCA"]
    }


def test_audit_patch_validates_request_envelope_output_schema() -> None:
    request_text = "Return the requested machine-readable design artifact."
    envelope = {
        "tasks": [],
        "final_outputs": ["report_markdown"],
    }
    evidence_sources = [
        {
            "evidence_id": "request",
            "kind": "user_request",
            "text": request_text,
        }
    ]
    invalid = pipeline_tools._apply_goal_audit_patches(
        envelope,
        [
            {
                "op": "replace",
                "path": "/final_outputs",
                "value": [{"label": "requested artifact"}],
                "evidence_refs": ["request"],
                "evidence_quote": "machine-readable design artifact",
                "reason": "Preserve the requested output intent.",
            }
        ],
        issue_records=[],
        evidence_sources=evidence_sources,
    )

    assert invalid["applied"] == []
    assert invalid["rejected"] == [
        {
            "path": "/final_outputs",
            "reason": (
                "canonical_request_envelope_value_invalid:"
                "item_0:value_matches_no_schema_branch"
            ),
        }
    ]
    feedback = pipeline_tools._goal_audit_repair_feedback(invalid["rejected"])
    assert feedback[0]["path"] == "/final_outputs"
    assert feedback[0]["reason"] == invalid["rejected"][0]["reason"]
    output_schemas = feedback[0]["expected_schema"]["items"]["anyOf"]
    assert output_schemas[0] == {"type": "string"}
    assert output_schemas[1]["required"] == ["kind"]
    assert output_schemas[1]["additionalProperties"] is False

    valid = pipeline_tools._apply_goal_audit_patches(
        envelope,
        [
            {
                "op": "replace",
                "path": "/final_outputs",
                "value": ["requested_machine_readable_artifact"],
                "evidence_refs": ["request"],
                "evidence_quote": "machine-readable design artifact",
                "reason": "Preserve the requested output intent.",
            }
        ],
        issue_records=[],
        evidence_sources=evidence_sources,
    )

    assert valid["rejected"] == []
    assert valid["request_envelope"]["final_outputs"] == [
        "requested_machine_readable_artifact"
    ]


def test_binding_identity_anchor_does_not_require_unsupported_abundance_direction() -> None:
    goal = {
        "intent": "binding",
        "design_programs": [
            {
                "program_id": "blocker",
                "label": "Block the supplied sites",
                "modality": "steric",
                "intent": "binding",
                "objective": "Block the supplied binding sites.",
            }
        ],
        "evaluation_targets": [
            {
                "label": "Supplied binding site",
                "objective_type": "identity_anchor",
                "target_type": "binding_site",
            }
        ],
    }

    assert pipeline_tools._canonical_design_goal_contract_issues(goal) == []


def test_splice_program_still_requires_biological_direction() -> None:
    goal = {
        "intent": "splice_modulation",
        "design_programs": [
            {
                "program_id": "splice",
                "label": "Change exon usage",
                "modality": "steric",
                "intent": "splice_modulation",
                "objective": "Change use of the supplied exon.",
            }
        ],
    }

    assert (
        "context.structured_goal.design_programs[0].desired_delta"
        in pipeline_tools._canonical_design_goal_contract_issues(goal)
    )


def test_goal_intake_clarification_defers_to_complete_verified_intake_restart(
    tmp_path: Path,
    monkeypatch: Any,
) -> None:
    request_text = "Use an 18-nt steric design for the supplied target."
    envelope = {
        "schema_version": "v3",
        "request_text": request_text,
        "tasks": [
            {
                "task_id": "design",
                "task_type": "design_new_asos",
                "context": {
                    "structured_goal": {
                        "design_programs": [
                            {
                                "program_id": "branch",
                                "label": "branch",
                                "modality": "steric",
                                "intent": "binding",
                                "objective": "bind the supplied target",
                                "desired_delta": "decrease",
                                "length": 40,
                            }
                        ]
                    }
                },
            }
        ],
    }
    question = "What biological target context should be used?"
    state = {
        "status": "clarification_answers_ready",
        "phase": "clarification_answers_ready",
        "request": request_text,
        "request_envelope": envelope,
        "clarification": {
            "kind": "goal_intake",
            "pending_apply": True,
            "restart_stage": "intake",
            "last_answers": [{"question": question, "answer": "reviewed cells"}],
        },
        "pending_questions": [],
        "clarification_qa_history": [],
        "force_rerun_stages": [],
    }
    state_path = tmp_path / "run_state.json"
    state_path.write_text(json.dumps(state), encoding="utf-8")

    async def fail_legacy_patch_audit(_prompt: str) -> dict[str, Any]:
        raise AssertionError(
            "intake clarification must be rebuilt by the complete verified workbench"
        )

    monkeypatch.setattr(
        pipeline_tools,
        "_run_goal_contract_audit_agent_once",
        fail_legacy_patch_audit,
    )

    result = pipeline_tools.apply_clarification_to_goal(str(state_path))
    updated = json.loads(state_path.read_text(encoding="utf-8"))
    clarified_goal = updated["request_envelope"]["tasks"][0]["context"][
        "structured_goal"
    ]

    assert result["status"] == "ok"
    assert updated["status"] == "initialized"
    assert updated["phase"] == "restart_intake_pending"
    assert clarified_goal["target_context"] == "reviewed cells"
    assert clarified_goal["design_programs"][0]["length"] == 40


def test_contract_repair_applies_evidence_backed_patch_to_exact_missing_path(
    monkeypatch: Any,
) -> None:
    envelope = {
        "schema_version": "v3",
        "request_text": "Decrease the supplied splice outcome with a steric design in reviewed cells.",
        "tasks": [
            {
                "task_id": "design",
                "task_type": "design_new_asos",
                "context": {
                    "target_context": "reviewed cells",
                    "structured_goal": {
                        "target_context": "reviewed cells",
                        "design_programs": [
                            {
                                "program_id": "branch",
                                "label": "branch",
                                "modality": "steric",
                                "intent": "splice_modulation",
                                "objective": "decrease the supplied splice outcome",
                            }
                        ],
                    },
                },
            }
        ],
    }

    async def fake_repair(_prompt: str) -> dict[str, Any]:
        return {
            "status": "repairable",
            "patches": [
                {
                    "op": "add",
                    "path": "/tasks/0/context/structured_goal/design_programs/0/desired_delta",
                    "value": "decrease",
                    "evidence_refs": ["request"],
                    "evidence_quote": "Decrease",
                    "reason": "The request supplies the direction.",
                }
            ],
            "unresolved": [],
        }

    monkeypatch.setattr(
        pipeline_tools,
        "_run_goal_contract_audit_agent_once",
        fake_repair,
    )

    result = pipeline_tools._repair_request_envelope_if_needed(
        request_text=envelope["request_text"],
        run_dir="",
        structured_goal=deepcopy(
            envelope["tasks"][0]["context"]["structured_goal"]
        ),
        request_envelope=envelope,
        qa_history=[],
    )

    assert result["status"] == "ready"
    assert "blocked_tasks" not in result["request_envelope"]
    assert (
        result["request_envelope"]["tasks"][0]["context"]["structured_goal"][
            "design_programs"
        ][0]["desired_delta"]
        == "decrease"
    )


def test_contract_repair_fills_all_validator_paths_without_user_clarification(
    monkeypatch: Any,
) -> None:
    request_text = "Increase the requested splice outcome at each supplied target interval."
    envelope = {
        "schema_version": "v3",
        "request_text": request_text,
        "tasks": [
            {
                "task_id": "design",
                "task_type": "design_new_asos",
                "context": {
                    "target_context": "reviewed cells",
                    "evaluation_targets": [
                        {"label": "first supplied interval"},
                        {"label": "second supplied interval"},
                    ],
                    "structured_goal": {
                        "target_context": "reviewed cells",
                        "intent": "splice_modulation",
                        "design_programs": [
                            {
                                "program_id": "branch",
                                "label": "branch",
                                "modality": "steric",
                                "intent": "splice_modulation",
                                "objective": "increase the splice outcome at the supplied target intervals",
                                "desired_delta": "increase",
                            }
                        ],
                        "evaluation_targets": [
                            {
                                "label": "first supplied interval",
                                "objective_type": "primary_outcome",
                            },
                            {
                                "label": "second supplied interval",
                                "objective_type": "primary_outcome",
                            },
                        ],
                    },
                },
            }
        ],
    }

    async def fake_audit(prompt: str) -> dict[str, Any]:
        if '"audit_mode": "post_repair_semantic_verification"' in prompt:
            return {"status": "valid", "patches": [], "unresolved": []}
        assert '"audit_mode": "contract_repair"' in prompt
        assert "/tasks/0/context/structured_goal/evaluation_targets/0/desired_delta" in prompt
        assert "/tasks/0/context/structured_goal/evaluation_targets/1/desired_delta" in prompt
        return {
            "status": "repairable",
            "patches": [
                {
                    "op": "add",
                    "path": f"/tasks/0/context/structured_goal/evaluation_targets/{index}/desired_delta",
                    "value": "increase",
                    "evidence_refs": ["request"],
                    "evidence_quote": "Increase",
                    "reason": "The request supplies one direction for every target.",
                }
                for index in range(2)
            ],
            "unresolved": [],
        }

    monkeypatch.setattr(
        pipeline_tools,
        "_run_goal_contract_audit_agent_once",
        fake_audit,
    )

    result = pipeline_tools._repair_request_envelope_if_needed(
        request_text=request_text,
        run_dir="",
        structured_goal={},
        request_envelope=envelope,
        qa_history=[],
        semantic_audit=True,
    )

    targets = result["request_envelope"]["tasks"][0]["context"][
        "structured_goal"
    ]["evaluation_targets"]
    assert result["status"] == "ready"
    assert [target["desired_delta"] for target in targets] == [
        "increase",
        "increase",
    ]


def test_semantic_audit_builds_missing_design_goal_from_evidence(
    monkeypatch: Any,
) -> None:
    request_text = "Design a 20 nucleotide steric blocker to increase output in reviewed cells."
    envelope = {
        "schema_version": "v3",
        "request_text": request_text,
        "tasks": [
            {
                "task_id": "design",
                "task_type": "design_new_asos",
                "context": {"target_context": "reviewed cells"},
            }
        ],
    }

    calls = {"count": 0}

    async def fake_audit(prompt: str) -> dict[str, Any]:
        calls["count"] += 1
        assert '"current_request_envelope"' in prompt
        if calls["count"] == 2:
            assert '"audit_mode": "post_repair_semantic_verification"' in prompt
            assert '"structured_goal"' in prompt
            assert '"design_goal_locations"' in prompt
            assert '"json_pointer": "/tasks/0/context/structured_goal"' in prompt
            return {"status": "valid", "patches": [], "goal_replacements": [], "unresolved": []}
        assert '"audit_mode": "semantic_completeness_and_contract_repair"' in prompt
        return {
            "status": "repairable",
            "patches": [],
            "goal_replacements": [
                {
                    "task_index": 0,
                    "task_id": "design",
                    "structured_goal": {
                        "target_context": "reviewed cells",
                        "intent": "binding",
                        "design_programs": [
                            {
                                "program_id": "primary",
                                "label": "primary",
                                "modality": "steric",
                                "intent": "binding",
                                "objective": "increase output by steric blockade",
                                "desired_delta": "increase",
                                "length": 20,
                            }
                        ],
                    },
                    "evidence_refs": ["request"],
                    "evidence_quote": "20 nucleotide steric blocker to increase output",
                    "reason": "The request supplies the executable design goal.",
                }
            ],
            "unresolved": [],
        }

    monkeypatch.setattr(
        pipeline_tools,
        "_run_goal_contract_audit_agent_once",
        fake_audit,
    )

    result = pipeline_tools._repair_request_envelope_if_needed(
        request_text=request_text,
        run_dir="",
        structured_goal={},
        request_envelope=envelope,
        qa_history=[],
        semantic_audit=True,
    )

    assert result["status"] == "ready"
    assert calls["count"] == 2
    goal = result["request_envelope"]["tasks"][0]["context"]["structured_goal"]
    assert goal["design_programs"][0]["desired_delta"] == "increase"
    assert goal["design_programs"][0]["length"] == 20


def test_staged_request_envelope_defers_design_goal_defaults_until_after_audit() -> None:
    envelope = pipeline_tools._build_request_envelope(
        request_text="Design ASOs for the supplied target context.",
        request_id="staged_design",
        structured_goal=None,
        raw_envelope={
            "tasks": [
                {
                    "task_id": "design",
                    "task_type": "design_new_asos",
                    "context": {
                        "gene_symbol": "GENE1",
                        "species": "human",
                        "target_context": "reviewed cells",
                        "chemistry": "steric",
                    },
                }
            ]
        },
        finalize_design_contract=False,
    )

    context = envelope["tasks"][0]["context"]
    assert "structured_goal" not in context
    assert "intent" not in context
    serialized = json.dumps(envelope)
    assert "request_contract_default_program" not in serialized
    assert "design_programs" not in serialized
    assert "explicit_splice_targets" not in serialized
    assert "focus_exons_only" not in serialized
    assert "homopolymer" not in serialized


def test_goal_replacement_channel_checks_task_identity_and_source_evidence(
    monkeypatch: Any,
) -> None:
    request_text = "Build a steric design for the supplied target in reviewed cells."
    envelope = {
        "schema_version": "v3",
        "request_text": request_text,
        "tasks": [
            {
                "task_id": "design",
                "task_type": "design_new_asos",
                "context": {"target_context": "reviewed cells"},
            }
        ],
    }

    async def fake_audit(_prompt: str) -> dict[str, Any]:
        return {
            "status": "repairable",
            "patches": [],
            "goal_replacements": [
                {
                    "task_index": 0,
                    "task_id": "different_task",
                    "structured_goal": {
                        "target_context": "reviewed cells",
                        "intent": "binding",
                        "design_programs": [
                            {
                                "program_id": "primary",
                                "label": "primary",
                                "modality": "steric",
                                "intent": "binding",
                                "objective": "bind the supplied target",
                                "desired_delta": "increase",
                            }
                        ],
                    },
                    "evidence_refs": ["request"],
                    "evidence_quote": "supplied target",
                    "reason": "Build the supported design goal.",
                }
            ],
            "unresolved": [],
        }

    monkeypatch.setattr(
        pipeline_tools,
        "_run_goal_contract_audit_agent_once",
        fake_audit,
    )

    result = pipeline_tools._repair_request_envelope_if_needed(
        request_text=request_text,
        run_dir="",
        structured_goal={},
        request_envelope=envelope,
        qa_history=[],
        semantic_audit=True,
    )

    assert result["status"] == "model_failure"
    assert result["reason"] == "goal_contract_audit_rejected_patch"
    assert result["rejected_patches"] == [
        {
            "task_id": "different_task",
            "reason": "goal_replacement_task_identity_mismatch",
        }
    ]


def test_goal_replacement_schema_uses_only_canonical_registry_keys() -> None:
    from aso_agents.goal_contract import (
        CANONICAL_GOAL_COLLECTION_FIELDS,
        CANONICAL_GOAL_FIELD_SCHEMAS,
        CANONICAL_REGION_CONSTRAINT_FIELDS,
        CANONICAL_SEQUENCE_CONSTRAINT_FIELDS,
        CANONICAL_STRUCTURED_GOAL_FIELDS,
        CANONICAL_TASK_CONTEXT_FIELDS,
        CANONICAL_TASK_TYPES,
    )
    from aso_agents.structured_outputs import (
        GOAL_CONTRACT_AUDIT_OUTPUT_SCHEMA,
        REQUEST_WORKBENCH_OUTPUT_SCHEMA,
    )

    replacement = GOAL_CONTRACT_AUDIT_OUTPUT_SCHEMA["properties"][
        "goal_replacements"
    ]["items"]
    goal_schema = replacement["properties"]["structured_goal"]

    assert goal_schema["additionalProperties"] is False
    assert replacement["properties"]["evidence_quote"]["maxLength"] == 160
    region_schema = CANONICAL_GOAL_FIELD_SCHEMAS["region_constraints"]["items"]
    assert region_schema["additionalProperties"] is False
    assert set(region_schema["properties"]) == set(
        CANONICAL_REGION_CONSTRAINT_FIELDS
    )
    sequence_schema = CANONICAL_GOAL_FIELD_SCHEMAS["constraints"]
    assert sequence_schema["additionalProperties"] is False
    assert set(sequence_schema["properties"]) == set(
        CANONICAL_SEQUENCE_CONSTRAINT_FIELDS
    )
    assert set(goal_schema["properties"]) == set(
        CANONICAL_STRUCTURED_GOAL_FIELDS
    )
    for collection, canonical_fields in CANONICAL_GOAL_COLLECTION_FIELDS.items():
        item_schema = goal_schema["properties"][collection]["items"]
        assert item_schema["additionalProperties"] is False
        assert set(item_schema["properties"]) == set(canonical_fields)

    task_schema = REQUEST_WORKBENCH_OUTPUT_SCHEMA["properties"][
        "request_envelope"
    ]["anyOf"][1]["properties"]["tasks"]["items"]
    context_schema = task_schema["properties"]["context"]
    assert task_schema["additionalProperties"] is False
    assert set(task_schema["properties"]["task_type"]["enum"]) == set(
        CANONICAL_TASK_TYPES
    )
    assert context_schema["additionalProperties"] is False
    assert set(context_schema["properties"]) == set(CANONICAL_TASK_CONTEXT_FIELDS)
    assert set(task_schema["properties"]) == {
        "task_id",
        "task_type",
        "label",
        "native",
        "can_run_in_parallel",
        "context",
    }


def test_contract_issue_records_expose_unsupported_task_type_schema() -> None:
    from aso_agents.goal_contract import CANONICAL_TASK_TYPES

    envelope = {
        "tasks": [
            {
                "task_id": "task",
                "task_type": "noncanonical_type",
                "context": {},
            }
        ]
    }
    validation = pipeline_tools._request_validation_summary(envelope)

    records = pipeline_tools._goal_contract_issue_records(envelope, validation)

    assert records == [
        {
                "task_id": "task",
                "validation_issue": "unsupported_task_type",
                "canonical_path": "task_type",
                "json_pointer": "/tasks/0/task_type",
                "required_operation": "replace",
                "expected_schema": {
                "type": "string",
                "enum": sorted(CANONICAL_TASK_TYPES),
            },
        }
    ]


def test_validation_issue_localization_builds_schema_derived_issue_record() -> None:
    envelope = {
        "tasks": [
            {
                "task_id": "design",
                "task_type": "design_new_asos",
                "context": {
                    "structured_goal": {
                        "intent": "knockdown",
                        "design_programs": [
                            {
                                "program_id": "primary",
                                "label": "Primary design",
                                "modality": "gapmer",
                                "intent": "knockdown",
                                "objective": "Decrease the requested RNA.",
                                "desired_delta": "decrease",
                            }
                        ],
                    }
                },
            }
        ]
    }
    validation = pipeline_tools._request_validation_summary(envelope)
    exact_records = pipeline_tools._goal_contract_issue_records(
        envelope,
        validation,
    )
    issue_instances = pipeline_tools._goal_validation_issue_instances(
        validation,
        exact_records,
    )
    candidates = pipeline_tools._goal_validation_candidate_destinations(
        envelope,
        issue_instances,
    )
    target_candidate = next(
        candidate
        for candidate in candidates
        if candidate["canonical_path"]
        == "context.structured_goal.target_context"
    )

    localized_records, ledger = (
        pipeline_tools._goal_validation_localization_records(
            envelope,
            issue_instances,
            candidates,
            {
                "verdicts": [
                    {
                        "issue_index": 0,
                        "resolution": "localized",
                        "candidate_indexes": [
                            target_candidate["candidate_index"]
                        ],
                        "reason": "The destination semantics match the validator requirement.",
                    }
                ]
            },
        )
    )

    assert [item["validation_issue"] for item in issue_instances] == [
        "target_context_missing"
    ]
    assert localized_records == [
        {
            "task_id": "design",
            "validation_issue": "target_context_missing",
            "canonical_path": "context.structured_goal.target_context",
            "json_pointer": (
                "/tasks/0/context/structured_goal/target_context"
            ),
            "required_operation": "add",
            "expected_schema": {"type": "string"},
        }
    ]
    assert ledger[0]["localization"] == "localized"
    assert ledger[0]["localized_json_pointers"] == [
        "/tasks/0/context/structured_goal/target_context"
    ]


def test_contract_audit_rejects_noncanonical_task_type_patch() -> None:
    from aso_agents.goal_contract import CANONICAL_TASK_TYPES

    envelope = {
        "tasks": [
            {
                "task_id": "task",
                "task_type": "answer_question",
                "context": {},
            }
        ]
    }
    patch_result = pipeline_tools._apply_goal_audit_patches(
        envelope,
        [
            {
                "op": "replace",
                "path": "/tasks/0/task_type",
                "value": "noncanonical_type",
                "evidence_refs": ["request"],
                "evidence_quote": "Design steric oligonucleotides",
                "reason": "Correct the requested operation.",
            }
        ],
        issue_records=[],
        evidence_sources=[
            {
                "evidence_id": "request",
                "kind": "user_request",
                "text": "Design steric oligonucleotides for the target.",
            }
        ],
    )

    assert patch_result["applied"] == []
    assert patch_result["rejected"] == [
        {
            "path": "/tasks/0/task_type",
            "reason": (
                "canonical_task_type_required:"
                + ",".join(sorted(CANONICAL_TASK_TYPES))
            ),
        }
    ]
    assert envelope["tasks"][0]["task_type"] == "answer_question"


def test_semantic_audit_retries_rejected_task_type_with_canonical_registry(
    monkeypatch: Any,
) -> None:
    request_text = (
        "Design steric oligonucleotides to increase GENE1 output in human "
        "reviewed cells."
    )
    envelope = {
        "schema_version": "v3",
        "request_text": request_text,
        "request_mode": "single_task",
        "request_summary": "Answer a question about GENE1.",
        "request_classifications": ["question"],
        "tasks": [
            {
                "task_id": "task",
                "task_type": "answer_question",
                "label": "GENE1 question",
                "native": True,
                "can_run_in_parallel": False,
                "context": {
                    "gene_symbol": "GENE1",
                    "species": "human",
                    "target_context": "reviewed cells",
                },
            }
        ],
        "task_dependencies": [],
        "final_outputs": ["answer"],
        "response_mode": "answer",
        "requires_execution": True,
        "assumptions": [],
    }
    calls = {"count": 0}

    async def fake_audit(prompt: str) -> dict[str, Any]:
        calls["count"] += 1
        assert '"canonical_task_types"' in prompt
        assert '"design_new_asos"' in prompt
        if calls["count"] == 1:
            return {
                "status": "repairable",
                "patches": [
                    {
                        "op": "replace",
                        "path": "/tasks/0/task_type",
                        "value": "noncanonical_type",
                        "evidence_refs": ["request"],
                        "evidence_quote": "Design steric oligonucleotides",
                        "reason": "Correct the requested operation.",
                    }
                ],
                "goal_replacements": [],
                "field_dispositions": [],
                "unresolved": [],
            }
        if calls["count"] == 2:
            assert '"prior_rejected_patches"' in prompt
            assert '"canonical_task_type_required:' in prompt
            return {
                "status": "repairable",
                "patches": [
                    {
                        "op": "replace",
                        "path": "/tasks/0/task_type",
                        "value": "design_new_asos",
                        "evidence_refs": ["request"],
                        "evidence_quote": "Design steric oligonucleotides",
                        "reason": "Use the canonical design operation.",
                    }
                ],
                "goal_replacements": [
                    {
                        "task_index": 0,
                        "task_id": "task",
                        "structured_goal": {
                            "gene_symbol": "GENE1",
                            "species": "human",
                            "target_context": "reviewed cells",
                            "intent": "binding",
                            "design_programs": [
                                {
                                    "program_id": "primary",
                                    "label": "primary",
                                    "modality": "steric",
                                    "intent": "binding",
                                    "objective": "increase GENE1 output",
                                    "desired_delta": "increase",
                                }
                            ],
                        },
                        "evidence_refs": ["request"],
                        "evidence_quote": "increase GENE1 output",
                        "reason": "Represent the requested executable design.",
                    }
                ],
                "field_dispositions": [],
                "unresolved": [],
            }
        return {
            "status": "valid",
            "patches": [],
            "goal_replacements": [],
            "field_dispositions": [],
            "unresolved": [],
        }

    monkeypatch.setattr(
        pipeline_tools,
        "_run_goal_contract_audit_agent_once",
        fake_audit,
    )

    result = pipeline_tools._repair_request_envelope_if_needed(
        request_text=request_text,
        run_dir="",
        structured_goal={},
        request_envelope=envelope,
        qa_history=[],
        semantic_audit=True,
    )

    assert result["status"] == "ready"
    assert calls["count"] == 3
    task = result["request_envelope"]["tasks"][0]
    assert task["task_type"] == "design_new_asos"
    assert task["context"]["structured_goal"]["gene_symbol"] == "GENE1"


def test_contract_audit_rejects_patch_without_verified_source_evidence(
    monkeypatch: Any,
) -> None:
    request_text = "Use the supplied steric design in reviewed cells."
    envelope = {
        "schema_version": "v3",
        "request_text": request_text,
        "tasks": [
            {
                "task_id": "design",
                "task_type": "design_new_asos",
                "context": {
                    "target_context": "reviewed cells",
                    "structured_goal": {
                        "target_context": "reviewed cells",
                        "intent": "binding",
                        "design_programs": [
                            {
                                "program_id": "branch",
                                "label": "branch",
                                "modality": "steric",
                                "intent": "binding",
                                "objective": "bind the supplied target",
                                "desired_delta": "decrease",
                                "length": 20,
                            }
                        ],
                    },
                },
            }
        ],
    }

    async def fake_audit(_prompt: str) -> dict[str, Any]:
        return {
            "status": "repairable",
            "patches": [
                {
                    "op": "replace",
                    "path": "/tasks/0/context/structured_goal/design_programs/0/length",
                    "value": 99,
                    "evidence_refs": ["request"],
                    "evidence_quote": "99 nucleotide",
                    "reason": "Unsupported model assertion.",
                }
            ],
            "unresolved": [],
        }

    monkeypatch.setattr(
        pipeline_tools,
        "_run_goal_contract_audit_agent_once",
        fake_audit,
    )

    result = pipeline_tools._repair_request_envelope_if_needed(
        request_text=request_text,
        run_dir="",
        structured_goal={},
        request_envelope=envelope,
        qa_history=[],
        semantic_audit=True,
    )

    assert result["status"] == "model_failure"
    assert result["reason"] == "goal_contract_audit_rejected_patch"
    assert result["rejected_patches"][0]["reason"] == "evidence_not_verified"
    assert (
        envelope["tasks"][0]["context"]["structured_goal"]["design_programs"]
        [0]["length"]
        == 20
    )


def test_contract_audit_rejects_noncanonical_collection_item_keys() -> None:
    envelope = {
        "tasks": [
            {
                "task_id": "design",
                "context": {"structured_goal": {}},
            }
        ]
    }
    patch_result = pipeline_tools._apply_goal_audit_patches(
        envelope,
        [
            {
                "op": "add",
                "path": "/tasks/0/context/structured_goal/evaluation_targets",
                "value": [
                    {
                        "label": "supplied interval",
                        "chrom": "chr1",
                        "start": 100,
                        "end": 140,
                        "desired_delta": "increase",
                        "alternate_target_key": "unsupported",
                    }
                ],
                "evidence_refs": ["request"],
                "evidence_quote": "supplied interval",
                "reason": "Add the supplied target.",
            }
        ],
        issue_records=[],
        evidence_sources=[
            {
                "evidence_id": "request",
                "kind": "user_request",
                "text": "Use the supplied interval.",
            }
        ],
    )

    assert patch_result["applied"] == []
    assert patch_result["rejected"] == [
        {
            "path": "/tasks/0/context/structured_goal/evaluation_targets",
            "reason": "noncanonical_patch_value_keys:alternate_target_key",
        }
    ]
    assert envelope["tasks"][0]["context"]["structured_goal"] == {}


def test_contract_audit_adds_canonical_collection_items_at_numeric_indices() -> None:
    envelope = {
        "tasks": [
            {
                "task_id": "design",
                "context": {
                    "structured_goal": {
                        "evaluation_targets": [
                            {
                                "label": "existing target",
                                "target_type": "region",
                                "desired_delta": "decrease",
                            }
                        ]
                    }
                },
            }
        ]
    }
    patch_result = pipeline_tools._apply_goal_audit_patches(
        envelope,
        [
            {
                "op": "add",
                "path": "/tasks/0/context/structured_goal/evaluation_targets/1",
                "value": {
                    "label": "appended target",
                    "target_type": "splice_site",
                    "desired_delta": "increase",
                    "site_kind": "acceptor",
                },
                "evidence_refs": ["request"],
                "evidence_quote": "appended target",
                "reason": "Add the independently requested target.",
            },
            {
                "op": "add",
                "path": "/tasks/0/context/structured_goal/evaluation_targets/1",
                "value": {
                    "label": "inserted target",
                    "target_type": "splice_site",
                    "desired_delta": "decrease",
                    "site_kind": "donor",
                },
                "evidence_refs": ["request"],
                "evidence_quote": "inserted target",
                "reason": "Insert the independently requested target.",
            },
        ],
        issue_records=[],
        evidence_sources=[
            {
                "evidence_id": "request",
                "kind": "user_request",
                "text": "Use the appended target and inserted target.",
            }
        ],
    )

    assert patch_result["rejected"] == []
    assert [
        item["label"]
        for item in patch_result["request_envelope"]["tasks"][0]["context"][
            "structured_goal"
        ]["evaluation_targets"]
    ] == ["existing target", "inserted target", "appended target"]


def test_contract_audit_materializes_missing_canonical_object_parents() -> None:
    request_text = (
        "Target the mutant allele, protect the wildtype allele, and require "
        "every candidate to overlap the variant."
    )
    envelope = {
        "tasks": [
            {
                "task_id": "design",
                "context": {"structured_goal": {}},
            }
        ]
    }
    paths_and_values = [
        (
            "/tasks/0/context/structured_goal/constraints/allele_specific/target_allele",
            "mutant",
        ),
        (
            "/tasks/0/context/structured_goal/constraints/allele_specific/non_target_allele",
            "wildtype",
        ),
        (
            "/tasks/0/context/structured_goal/constraints/allele_specific/require_variant_overlap",
            True,
        ),
    ]

    patch_result = pipeline_tools._apply_goal_audit_patches(
        envelope,
        [
            {
                "op": "add",
                "path": path,
                "value": value,
                "evidence_refs": ["request"],
                "evidence_quote": request_text,
                "reason": "Restore the explicitly requested constraint.",
            }
            for path, value in paths_and_values
        ],
        issue_records=[
            {"json_pointer": path}
            for path, _value in paths_and_values
        ],
        evidence_sources=[
            {
                "evidence_id": "request",
                "kind": "user_request",
                "text": request_text,
            }
        ],
    )

    assert patch_result["rejected"] == []
    assert patch_result["request_envelope"]["tasks"][0]["context"][
        "structured_goal"
    ]["constraints"] == {
        "allele_specific": {
            "target_allele": "mutant",
            "non_target_allele": "wildtype",
            "require_variant_overlap": True,
        }
    }


def test_contract_audit_does_not_materialize_missing_array_parents() -> None:
    request_text = "Use the requested target."
    path = "/tasks/0/context/structured_goal/evaluation_targets/0/label"
    envelope = {
        "tasks": [
            {
                "task_id": "design",
                "context": {"structured_goal": {}},
            }
        ]
    }

    patch_result = pipeline_tools._apply_goal_audit_patches(
        envelope,
        [
            {
                "op": "add",
                "path": path,
                "value": "requested target",
                "evidence_refs": ["request"],
                "evidence_quote": request_text,
                "reason": "Restore the explicitly requested target.",
            }
        ],
        issue_records=[{"json_pointer": path}],
        evidence_sources=[
            {
                "evidence_id": "request",
                "kind": "user_request",
                "text": request_text,
            }
        ],
    )

    assert patch_result["applied"] == []
    assert patch_result["rejected"] == [
        {"path": path, "reason": "patch_parent_missing"}
    ]
    assert patch_result["request_envelope"] == envelope


def test_contract_audit_rejects_numeric_list_add_beyond_current_length() -> None:
    envelope = {
        "tasks": [
            {
                "task_id": "design",
                "context": {"structured_goal": {"evaluation_targets": []}},
            }
        ]
    }
    patch_result = pipeline_tools._apply_goal_audit_patches(
        envelope,
        [
            {
                "op": "add",
                "path": "/tasks/0/context/structured_goal/evaluation_targets/1",
                "value": {
                    "label": "target",
                    "target_type": "region",
                    "desired_delta": "decrease",
                },
                "evidence_refs": ["request"],
                "evidence_quote": "target",
                "reason": "Add the requested target.",
            }
        ],
        issue_records=[],
        evidence_sources=[
            {
                "evidence_id": "request",
                "kind": "user_request",
                "text": "Use the target.",
            }
        ],
    )

    assert patch_result["applied"] == []
    assert patch_result["rejected"] == [
        {
            "path": "/tasks/0/context/structured_goal/evaluation_targets/1",
            "reason": "invalid_add_target",
        }
    ]


def test_contract_audit_rejects_noncanonical_keys_in_full_goal_patch() -> None:
    envelope = {
        "tasks": [
            {
                "task_id": "design",
                "context": {},
            }
        ]
    }
    patch_result = pipeline_tools._apply_goal_audit_patches(
        envelope,
        [
            {
                "op": "add",
                "path": "/tasks/0/context/structured_goal",
                "value": {
                    "intent": "binding",
                    "unsupported_goal_field": "unsupported",
                    "design_programs": [
                        {
                            "program_id": "primary",
                            "label": "primary",
                            "modality": "steric",
                            "intent": "binding",
                            "objective": "bind the supplied target",
                            "desired_delta": "increase",
                        }
                    ],
                },
                "evidence_refs": ["request"],
                "evidence_quote": "supplied target",
                "reason": "Build the supported design goal.",
            }
        ],
        issue_records=[],
        evidence_sources=[
            {
                "evidence_id": "request",
                "kind": "user_request",
                "text": "Bind the supplied target.",
            }
        ],
    )

    assert patch_result["applied"] == []
    assert patch_result["rejected"] == [
        {
            "path": "/tasks/0/context/structured_goal",
            "reason": "noncanonical_structured_goal_keys:unsupported_goal_field",
        }
    ]
    assert "structured_goal" not in envelope["tasks"][0]["context"]


def test_contract_audit_rejects_invalid_canonical_goal_value_types() -> None:
    envelope = {
        "tasks": [
            {
                "task_id": "design",
                "context": {},
            }
        ]
    }
    patch_result = pipeline_tools._apply_goal_audit_patches(
        envelope,
        [
            {
                "op": "add",
                "path": "/tasks/0/context/structured_goal",
                "value": {
                    "intent": "binding",
                    "region_constraints": {
                        "source_type": "explicit_rows",
                    },
                    "design_programs": [
                        {
                            "program_id": "primary",
                            "label": "primary",
                            "modality": "steric",
                            "intent": "binding",
                            "objective": "bind the supplied target",
                            "desired_delta": "increase",
                        }
                    ],
                },
                "evidence_refs": ["request"],
                "evidence_quote": "supplied target",
                "reason": "Build the supported design goal.",
            }
        ],
        issue_records=[],
        evidence_sources=[
            {
                "evidence_id": "request",
                "kind": "user_request",
                "text": "Bind the supplied target.",
            }
        ],
    )

    assert patch_result["applied"] == []
    assert patch_result["rejected"] == [
        {
            "path": "/tasks/0/context/structured_goal",
            "reason": (
                "canonical_field_value_invalid:region_constraints:expected_array"
            ),
        }
    ]


def test_contract_audit_repair_feedback_exposes_exact_destination_schema() -> None:
    feedback = pipeline_tools._goal_audit_repair_feedback(
        [
            {
                "path": "/tasks/0/context/structured_goal/isoform_scope_source",
                "reason": "canonical_field_value_invalid:value_not_in_enum",
            }
        ]
    )

    assert feedback == [
        {
            "path": "/tasks/0/context/structured_goal/isoform_scope_source",
            "reason": "canonical_field_value_invalid:value_not_in_enum",
            "expected_schema": pipeline_tools.CANONICAL_GOAL_FIELD_SCHEMAS[
                "isoform_scope_source"
            ],
        }
    ]


def test_targeted_contract_audit_omits_unneeded_semantic_catalogs() -> None:
    prompt = pipeline_tools._goal_audit_prompt(
        request_envelope={
            "tasks": [
                {
                    "task_id": "design",
                    "task_type": "design_new_asos",
                    "context": {"structured_goal": {}},
                }
            ]
        },
        validation_issues=["missing canonical value"],
        issue_records=[
            {
                "task_id": "design",
                "canonical_path": "isoform_scope_source",
                "json_pointer": (
                    "/tasks/0/context/structured_goal/isoform_scope_source"
                ),
                "expected_schema": pipeline_tools.CANONICAL_GOAL_FIELD_SCHEMAS[
                    "isoform_scope_source"
                ],
            }
        ],
        evidence_sources=[
            {
                "evidence_id": "request",
                "kind": "user_request",
                "text": "Use the explicitly supplied isoform scope.",
            }
        ],
        semantic_audit=False,
    )
    audit_input = json.loads(prompt.split("\n\n", 1)[1])

    assert audit_input["canonical_destination_catalog"] == []
    assert audit_input["canonical_task_context_registry"] == []
    assert audit_input["canonical_contract_issues"][0][
        "expected_schema"
    ] == pipeline_tools.CANONICAL_GOAL_FIELD_SCHEMAS["isoform_scope_source"]


def test_semantic_audit_bounds_response_to_patch_and_unresolved_channels() -> None:
    prompt = pipeline_tools._goal_audit_prompt(
        request_envelope={
            "tasks": [
                {
                    "task_id": "design",
                    "task_type": "design_new_asos",
                    "context": {
                        "structured_goal": {
                            "intent": "binding",
                            "design_programs": [
                                {
                                    "program_id": "primary",
                                    "label": "primary",
                                    "modality": "steric",
                                    "intent": "binding",
                                    "objective": "bind the supplied target",
                                    "desired_delta": "decrease",
                                }
                            ],
                        }
                    },
                }
            ]
        },
        validation_issues=[],
        issue_records=[],
        evidence_sources=[
            {
                "evidence_id": "request",
                "kind": "user_request",
                "text": "Bind the supplied target.",
            }
        ],
        semantic_audit=True,
        post_repair_verification=True,
    )
    audit_input = pipeline_tools._goal_audit_protocol_payload(prompt)
    schema = pipeline_tools._goal_contract_audit_schema_for_prompt(prompt)

    assert audit_input["allowed_nonempty_channels"] == [
        "unresolved",
        "patches",
    ]
    assert audit_input["task_context_locations"] == []
    assert audit_input["design_goal_locations"] == [
        {
            "task_index": 0,
            "task_id": "design",
            "json_pointer": "/tasks/0/context/structured_goal",
        }
    ]
    assert "maxItems" not in schema["properties"]["patches"]
    assert "maxItems" not in schema["properties"]["unresolved"]
    assert schema["properties"]["goal_replacements"]["maxItems"] == 0
    assert schema["properties"]["field_dispositions"]["maxItems"] == 0
    assert "structured_goal" not in schema["properties"]["goal_replacements"][
        "items"
    ]["properties"]

    invalid_audit = {
        "status": "repairable",
        "patches": [
            {
                "op": "replace",
                "path": "/tasks/0/context/structured_goal/evaluation_targets",
                "value": [{"name": "candidate", "sequence": "ACGT"}],
                "evidence_refs": ["request"],
                "evidence_quote": "Bind the supplied target.",
                "reason": "Invalid task-context value shape.",
            }
        ],
        "goal_replacements": [],
        "field_dispositions": [],
        "unresolved": [],
    }
    assert pipeline_tools._json_schema_validation_errors(invalid_audit, schema)


def test_semantic_audit_binds_each_patch_path_to_its_exact_value_schema() -> None:
    prompt = pipeline_tools._goal_audit_prompt(
        request_envelope={
            "tasks": [
                {
                    "task_id": "design",
                    "task_type": "design_new_asos",
                    "context": {
                        "structured_goal": {
                            "candidate_limit": 12,
                            "design_programs": [
                                {
                                    "program_id": "primary",
                                    "label": "Primary route",
                                    "modality": "steric",
                                    "intent": "splice_modulation",
                                    "objective": "Change the requested splice outcome.",
                                }
                            ],
                            "splice_isoform_objectives": [
                                {
                                    "objective_scope": "specific_isoform",
                                    "isoform_label": "requested isoform",
                                    "desired_abundance": "increase",
                                }
                            ],
                        }
                    },
                }
            ]
        },
        validation_issues=[],
        issue_records=[],
        evidence_sources=[
            {
                "evidence_id": "request",
                "kind": "user_request",
                "text": "Preserve both requested splice outcomes.",
            }
        ],
        semantic_audit=True,
    )
    schema = pipeline_tools._goal_contract_audit_schema_for_prompt(prompt)
    branches = schema["properties"]["patches"]["items"]["anyOf"]
    def _branch_for(operation: str, path: str) -> dict[str, object]:
        return next(
            branch["properties"]
            for branch in branches
            if branch["properties"]["op"]["enum"] == [operation]
            and path in branch["properties"]["path"]["enum"]
        )

    collection_path = (
        "/tasks/0/context/structured_goal/splice_isoform_objectives"
    )
    append_path = f"{collection_path}/-"
    abundance_path = f"{collection_path}/0/desired_abundance"
    per_site_limit_path = (
        "/tasks/0/context/structured_goal/candidate_limit_per_site"
    )

    assert not any(
        collection_path in branch["properties"]["path"]["enum"]
        and branch["properties"]["op"]["enum"] == ["replace"]
        for branch in branches
    )
    append_branch = _branch_for("add", append_path)
    abundance_branch = _branch_for("replace", abundance_path)
    per_site_limit_branch = _branch_for("add", per_site_limit_path)
    assert "objective_id" not in append_branch["value"]["properties"]
    assert abundance_branch["value"]["enum"] == ["increase", "decrease"]
    assert per_site_limit_branch["value"] == {"type": "integer", "minimum": 1}

    valid_audit = {
        "status": "repairable",
        "patches": [
            {
                "op": "add",
                "path": append_path,
                "value": {
                    "objective_scope": "default_exon_context",
                    "desired_abundance": "decrease",
                },
                "evidence_refs": ["request"],
                "evidence_quote": "Preserve both requested splice outcomes.",
                "reason": "Append the independently requested outcome.",
            }
        ],
        "goal_replacements": [],
        "field_dispositions": [],
        "unresolved": [],
    }
    assert not pipeline_tools._json_schema_validation_errors(valid_audit, schema)

    invalid_audit = deepcopy(valid_audit)
    invalid_audit["patches"][0]["value"]["objective_id"] = "invented"
    assert pipeline_tools._json_schema_validation_errors(invalid_audit, schema)


def test_semantic_audit_does_not_treat_source_coordinates_as_overlap_policy() -> None:
    prompt = pipeline_tools._goal_audit_prompt(
        request_envelope={
            "tasks": [
                {
                    "task_id": "design",
                    "task_type": "design_new_asos",
                    "context": {
                        "structured_goal": {
                            "design_programs": [
                                {
                                    "program_id": "primary",
                                    "label": "primary",
                                    "modality": "steric",
                                    "intent": "binding",
                                    "objective": "bind the supplied target",
                                    "desired_delta": "decrease",
                                }
                            ],
                            "evaluation_targets": [
                                {
                                    "target_type": "region",
                                    "chrom": "supplied_contig",
                                    "start": 10,
                                    "end": 30,
                                    "strand": "+",
                                }
                            ],
                        }
                    },
                }
            ]
        },
        validation_issues=[],
        issue_records=[],
        evidence_sources=[
            {
                "evidence_id": "source.0",
                "kind": "declared_task_artifact",
                "text": "A supplied source contains a labeled interval.",
            }
        ],
        semantic_audit=True,
        post_repair_verification=True,
    )

    assert "Source coordinates or interval rows do not alone prove" in prompt
    assert "Never invent a least-restrictive overlap policy" in prompt
    assert "optional region_constraints entry" in prompt


def test_goal_requirement_catalog_reports_schema_requiredness() -> None:
    catalog = pipeline_tools.canonical_goal_requirement_catalog()

    assert catalog["context.structured_goal"] == ["design_programs"]
    assert catalog["context.structured_goal.design_programs[*]"] == [
        "intent",
        "label",
        "modality",
        "objective",
        "program_id",
    ]
    assert catalog["context.structured_goal.evaluation_targets[*]"] == [
        "objective_type"
    ]


def test_semantic_audit_applies_repairs_before_optional_unresolved_question(
    monkeypatch: Any,
) -> None:
    request_text = (
        "Design one steric branch and measure transcript TX-OUTCOME as the global "
        "biological outcome in the supplied cell model."
    )
    envelope = {
        "schema_version": "v3",
        "request_text": request_text,
        "tasks": [
            {
                "task_id": "design",
                "task_type": "design_new_asos",
                "context": {
                    "target_context": "supplied cell model",
                    "structured_goal": {
                        "target_context": "supplied cell model",
                        "design_programs": [
                            {
                                "program_id": "primary",
                                "label": "steric branch",
                                "modality": "steric",
                                "intent": "binding",
                                "objective": "bind the requested target",
                                "desired_delta": "decrease",
                                "transcript_id": "TX-OUTCOME",
                            }
                        ],
                        "evaluation_targets": [
                            {
                                "label": "global transcript outcome",
                                "target_type": "transcript",
                                "transcript_id": "TX-OUTCOME",
                                "desired_delta": "decrease",
                                "objective_type": "primary_outcome",
                            }
                        ],
                    }
                },
            }
        ],
    }
    calls = {"count": 0}

    async def fake_audit(prompt: str) -> dict[str, Any]:
        calls["count"] += 1
        audit_input = pipeline_tools._goal_audit_protocol_payload(prompt)
        assert audit_input["canonical_requirement_catalog"]
        program = audit_input["current_request_envelope"]["tasks"][0][
            "context"
        ]["structured_goal"]["design_programs"][0]
        if calls["count"] == 2:
            assert "transcript_id" not in program
            return {
                "status": "valid",
                "patches": [],
                "goal_replacements": [],
                "field_dispositions": [],
                "unresolved": [],
            }
        return {
            "status": "needs_clarification",
            "patches": [
                {
                    "op": "remove",
                    "path": (
                        "/tasks/0/context/structured_goal/design_programs/0/"
                        "transcript_id"
                    ),
                    "value": None,
                    "evidence_refs": ["request"],
                    "evidence_quote": "global biological outcome",
                    "reason": "The outcome identifier was not assigned to the program.",
                }
            ],
            "goal_replacements": [],
            "field_dispositions": [],
            "unresolved": [
                {
                    "path": (
                        "/tasks/0/context/structured_goal/design_programs/0/"
                        "transcript_id"
                    ),
                    "question": "Which outcome identifier belongs to this branch?",
                    "reason": "The optional program mapping is absent.",
                    "resolution": "user_input",
                }
            ],
        }

    monkeypatch.setattr(
        pipeline_tools,
        "_run_goal_contract_audit_agent_once",
        fake_audit,
    )
    monkeypatch.setattr(
        pipeline_tools,
        "_run_goal_unresolved_adjudication_sync",
        lambda _payload: {
            "verdicts": [
                {
                    "unresolved_index": 0,
                    "disposition": "nonblocking",
                    "canonical_path": None,
                    "reason": "The accompanying repair removes the optional mapping.",
                }
            ]
        },
    )

    result = pipeline_tools._repair_request_envelope_if_needed(
        request_text=request_text,
        run_dir="",
        structured_goal={},
        request_envelope=envelope,
        qa_history=[],
        semantic_audit=True,
    )

    assert result["status"] == "ready", result
    assert calls["count"] == 2


def test_semantic_audit_treats_noop_only_verdict_as_verified(
    monkeypatch: Any,
) -> None:
    request_text = "Design one steric branch in the supplied cell model."
    envelope = {
        "schema_version": "v3",
        "request_text": request_text,
        "tasks": [
            {
                "task_id": "design",
                "task_type": "design_new_asos",
                "context": {
                    "target_context": "supplied cell model",
                    "structured_goal": {
                        "target_context": "supplied cell model",
                        "design_programs": [
                            {
                                "program_id": "primary",
                                "label": "steric branch",
                                "modality": "steric",
                                "intent": "binding",
                                "objective": "bind the requested target",
                                "desired_delta": "decrease",
                            }
                        ],
                    },
                },
            }
        ],
    }

    async def fake_audit(_prompt: str) -> dict[str, Any]:
        return {
            "status": "repairable",
            "patches": [
                {
                    "op": "remove",
                    "path": (
                        "/tasks/0/context/structured_goal/design_programs/0/"
                        "transcript_id"
                    ),
                    "value": None,
                    "evidence_refs": ["request"],
                    "evidence_quote": "one steric branch",
                    "reason": "The optional field is absent.",
                }
            ],
            "goal_replacements": [],
            "field_dispositions": [],
            "unresolved": [],
        }

    monkeypatch.setattr(
        pipeline_tools,
        "_run_goal_contract_audit_agent_once",
        fake_audit,
    )

    result = pipeline_tools._repair_request_envelope_if_needed(
        request_text=request_text,
        run_dir="",
        structured_goal={},
        request_envelope=envelope,
        qa_history=[],
        semantic_audit=True,
    )

    assert result["status"] == "ready"


def test_noncanonical_field_audit_exposes_only_disposition_channel() -> None:
    prompt = pipeline_tools._goal_audit_prompt(
        request_envelope={
            "tasks": [
                {
                    "task_id": "design",
                    "task_type": "design_new_asos",
                    "context": {
                        "structured_goal": {
                            "intent": "binding",
                            "invented_requirement": "supplied meaning",
                        }
                    },
                }
            ]
        },
        validation_issues=[],
        issue_records=[],
        evidence_sources=[
            {
                "evidence_id": "request",
                "kind": "user_request",
                "text": "Use the supplied meaning.",
            }
        ],
        semantic_audit=True,
    )
    audit_input = pipeline_tools._goal_audit_protocol_payload(prompt)
    schema = pipeline_tools._goal_contract_audit_schema_for_prompt(prompt)

    assert audit_input["allowed_nonempty_channels"] == [
        "unresolved",
        "field_dispositions",
    ]
    assert "maxItems" not in schema["properties"]["field_dispositions"]
    assert schema["properties"]["patches"]["maxItems"] == 0
    assert schema["properties"]["goal_replacements"]["maxItems"] == 0


def test_contract_audit_rejects_noncanonical_nested_region_keys() -> None:
    issue = pipeline_tools._goal_contract_value_schema_issue(
        [
            {
                "source_type": "explicit_rows",
                "coordinate_space": "genomic",
                "overlap_policy": "any_overlap",
            }
        ],
        pipeline_tools.CANONICAL_GOAL_FIELD_SCHEMAS["region_constraints"],
    )

    assert issue == "item_0:noncanonical_keys:overlap_policy"


def test_staged_normalization_preserves_audited_canonical_constraints() -> None:
    goal = {
        "intent": "binding",
        "aso_overlap_requirement": "any_overlap",
        "minimum_overlap_bp": 1,
        "constraints": {
            "gc_content": {
                "min": 0.25,
                "max": 0.75,
                "min_inclusive": False,
                "max_inclusive": False,
            }
        },
        "region_constraints": [
            {
                "source_type": "explicit_rows",
                "coordinate_space": "genomic",
                "aso_overlap_requirement": "any_overlap",
                "minimum_overlap_bp": 1,
                "intervals_bed0_half_open": [
                    {
                        "chrom": "chr1",
                        "start": 100,
                        "end": 120,
                        "label": "supplied interval",
                        "strand": "+",
                    }
                ],
            }
        ],
        "design_programs": [
            {
                "program_id": "primary",
                "label": "primary",
                "modality": "steric",
                "intent": "binding",
                "objective": "bind the supplied interval",
                "desired_delta": "increase",
            }
        ],
    }
    envelope = pipeline_tools._build_request_envelope(
        request_text="Use the supplied interval in reviewed cells.",
        request_id="canonical_constraints",
        structured_goal={},
        raw_envelope={
            "tasks": [
                {
                    "task_id": "design",
                    "task_type": "design_new_asos",
                    "context": {"structured_goal": goal},
                }
            ]
        },
        finalize_design_contract=False,
    )

    normalized_goal = envelope["tasks"][0]["context"]["structured_goal"]
    assert normalized_goal["constraints"] == goal["constraints"]
    assert "aso_overlap_requirement" not in normalized_goal["design_programs"][0]
    assert "minimum_overlap_bp" not in normalized_goal["design_programs"][0]


def test_contract_audit_surfaces_exhausted_incomplete_response_without_clarification(
    monkeypatch: Any,
) -> None:
    request_text = "Use the supplied steric design in reviewed cells."
    envelope = {
        "schema_version": "v3",
        "request_text": request_text,
        "tasks": [
            {
                "task_id": "design",
                "task_type": "design_new_asos",
                "context": {
                    "target_context": "reviewed cells",
                    "structured_goal": {
                        "target_context": "reviewed cells",
                        "intent": "binding",
                        "design_programs": [
                            {
                                "program_id": "branch",
                                "label": "branch",
                                "modality": "steric",
                                "intent": "binding",
                                "objective": "bind the supplied target",
                                "desired_delta": "decrease",
                            }
                        ],
                    },
                },
            }
        ],
    }
    async def fake_audit(_prompt: str) -> dict[str, Any]:
        return {
            "error": "incomplete_response",
            "status": "incomplete",
            "reason": "max_output_tokens",
        }

    monkeypatch.setattr(
        pipeline_tools,
        "_run_goal_contract_audit_agent_once",
        fake_audit,
    )

    result = pipeline_tools._repair_request_envelope_if_needed(
        request_text=request_text,
        run_dir="",
        structured_goal={},
        request_envelope=envelope,
        qa_history=[],
        semantic_audit=True,
    )

    assert result["status"] == "model_failure"
    assert result["reason"] == "goal_contract_audit_incomplete_response"
    assert result["model_reason"] == "max_output_tokens"
    assert "questions" not in result


def test_contract_audit_uses_validator_feedback_to_correct_rejected_patch(
    monkeypatch: Any,
) -> None:
    request_text = "Decrease the requested splice outcome with the supplied steric design."
    envelope = {
        "schema_version": "v3",
        "request_text": request_text,
        "tasks": [
            {
                "task_id": "design",
                "task_type": "design_new_asos",
                "context": {
                    "target_context": "reviewed cells",
                    "structured_goal": {
                        "target_context": "reviewed cells",
                        "intent": "splice_modulation",
                        "design_programs": [
                            {
                                "program_id": "branch",
                                "label": "branch",
                                "modality": "steric",
                                "intent": "splice_modulation",
                                "objective": "modulate the supplied splice target",
                            }
                        ],
                    },
                },
            }
        ],
    }
    calls = {"count": 0}

    async def fake_audit(prompt: str) -> dict[str, Any]:
        calls["count"] += 1
        if calls["count"] == 2:
            assert "replace_target_missing" in prompt
        return {
            "status": "repairable",
            "patches": [
                {
                    "op": "replace" if calls["count"] == 1 else "add",
                    "path": "/tasks/0/context/structured_goal/design_programs/0/desired_delta",
                    "value": "decrease",
                    "evidence_refs": ["request"],
                    "evidence_quote": "Decrease",
                    "reason": "The request supplies the direction.",
                }
            ],
            "unresolved": [],
        }

    monkeypatch.setattr(
        pipeline_tools,
        "_run_goal_contract_audit_agent_once",
        fake_audit,
    )

    result = pipeline_tools._repair_request_envelope_if_needed(
        request_text=request_text,
        run_dir="",
        structured_goal={},
        request_envelope=envelope,
        qa_history=[],
    )

    assert calls["count"] == 2
    assert result["status"] == "ready"
    assert (
        result["request_envelope"]["tasks"][0]["context"]["structured_goal"]
        ["design_programs"][0]["desired_delta"]
        == "decrease"
    )


def test_route_json_annotation_emits_one_canonical_route_contract(tmp_path: Path) -> None:
    route_path = tmp_path / "route.json"
    route_path.write_text(
        json.dumps(
            {
                "primary_targets": [
                    {"label": "donor@200", "direction": "decrease"}
                ]
            }
        ),
        encoding="utf-8",
    )
    goal = {
        "splice_site_plan": [
            {
                "chrom": "chr1",
                "position": 200,
                "site": "donor",
                "exon_indices": {"TX1": 6},
                "present_in": ["TX1"],
            }
        ]
    }

    pipeline_tools._annotate_route_json_with_site_context(
        json_path=str(route_path),
        structured_goal=goal,
    )

    payload = json.loads(route_path.read_text(encoding="utf-8"))
    assert payload["routes"] == [
        {
            "label": "donor@200",
            "direction": "decrease",
            "site_kind": "donor",
            "position": 200,
            "exon_label": 6,
        }
    ]


def test_runtime_splice_site_plan_uses_inclusion_direction_for_shared_focus_exon() -> None:
    plan, _comparisons = run_goal_intake._compute_splice_site_plan(
        [
            {
                "isoform_label": "skipped",
                "desired_abundance": "decrease",
                "desired_effect": "inclusion",
                "exon_label": "exon 7",
            },
            {
                "isoform_label": "included",
                "desired_abundance": "increase",
                "desired_effect": "inclusion",
                "exon_label": "exon 7",
            },
        ],
        {
            "skipped": {
                "sites": [
                    {"chrom": "chr1", "acceptor": 200, "donor": 250, "exon_index": 7, "strand": "+", "coding": True},
                ],
                "exons_tx": [{"index": idx} for idx in range(1, 9)],
            },
            "included": {
                "sites": [
                    {"chrom": "chr1", "acceptor": 200, "donor": 250, "exon_index": 7, "strand": "+", "coding": True},
                ],
                "exons_tx": [{"index": idx} for idx in range(1, 9)],
            },
        },
        {
            "splice_isoform_objectives": [
                {
                    "isoform_label": "skipped",
                    "desired_abundance": "decrease",
                    "desired_effect": "inclusion",
                    "exon_label": "exon 7",
                },
                {
                    "isoform_label": "included",
                    "desired_abundance": "increase",
                    "desired_effect": "inclusion",
                    "exon_label": "exon 7",
                },
            ],
            "desired_effect": "inclusion",
            "focus_exons_only": True,
        },
    )

    by_label = {f"{entry['site']}@{entry['position']}": entry for entry in plan}
    assert by_label["acceptor@200"]["delta_direction"] == "increase"
    assert by_label["donor@250"]["delta_direction"] == "increase"


@pytest.mark.parametrize('exon_label', [14, '14x (poison exon)'])
def test_explicit_isoform_abundance_is_not_overwritten_by_effect_prose(exon_label) -> None:
    raw_goal = {
        "splice_isoform_objectives": [
            {
                "isoform_label": "ENST00000376285",
                "transcript_id": "ENST00000376285",
                "desired_abundance": "increase",
                "objective": "promote_skipping",
                "intended_effects": ["decrease_poison_exon_inclusion", "increase_productive_mRNA"],
                "exon_label": exon_label,
            }
        ],
        "focus_exons_only": True,
    }
    normalized = pipeline_tools._normalized_design_objective_preseed(raw_goal)["splice_isoform_objectives"]
    assert normalized[0]["desired_abundance"] == "increase"

    plan, _comparisons = run_goal_intake._compute_splice_site_plan(
        normalized,
        {
            "ENST00000376285": {
                "sites": [
                    {"chrom": "chr13", "acceptor": 1300, "donor": 1350, "exon_index": 13, "strand": "+", "coding": True},
                    {"chrom": "chr13", "acceptor": 1400, "donor": 1450, "exon_index": 14, "strand": "+", "coding": True},
                    {"chrom": "chr13", "acceptor": 1500, "donor": 1550, "exon_index": 15, "strand": "+", "coding": True},
                ],
                "exons_tx": [{"index": idx} for idx in range(1, 16)],
            }
        },
        {"splice_isoform_objectives": normalized, "focus_exons_only": True},
    )

    by_label = {f"{entry['site']}@{entry['position']}": entry for entry in plan}
    if isinstance(exon_label, str):
        # The old assertion silently converted a named noncanonical exon to 14.
        # Its abundance remains intact, but no site can be assigned until the
        # named feature has a resolved reference span.
        assert by_label == {}
        return
    assert sorted(by_label) == ["acceptor@1400", "donor@1450"]
    assert by_label["acceptor@1400"]["delta_direction"] == "increase"
    assert by_label["donor@1450"]["delta_per_isoform"] == {"ENST00000376285": "increase"}


def test_isoform_conversion_objective_expands_to_transcript_objectives() -> None:
    raw_goal = {
        "splice_isoform_objectives": [
            {
                "event": "exon_inclusion",
                "direction": "increase_inclusion",
                "exon_label": "exon 7",
                "exon_numbering_transcript_id": "ENST00000380743",
                "from_transcript_id": "ENST00000626847",
                "to_transcript_id": "ENST00000380743",
            }
        ]
    }

    objectives = pipeline_tools._normalized_design_objective_preseed(raw_goal)["splice_isoform_objectives"]
    by_tx = {obj["transcript_id"]: obj for obj in objectives}

    assert set(by_tx) == {"ENST00000626847", "ENST00000380743"}
    assert by_tx["ENST00000626847"]["desired_abundance"] == "decrease"
    assert by_tx["ENST00000626847"]["conversion_role"] == "from"
    assert by_tx["ENST00000380743"]["desired_abundance"] == "increase"
    assert by_tx["ENST00000380743"]["conversion_role"] == "to"
    assert by_tx["ENST00000380743"]["exon_label"] == "exon 7"


def test_isoform_conversion_objective_accepts_reference_alternative_transcript_ids() -> None:
    raw_goal = {
        "splice_isoform_objectives": [
            {
                "objective_type": "increase_exon_inclusion",
                "exon_label": "exon 7",
                "exon_numbering_transcript_id": "ENST00000380743",
                "reference_isoform_transcript_id": "ENST00000380743",
                "alternative_isoform_transcript_id": "ENST00000626847",
                "desired_direction": "increase",
            }
        ]
    }

    objectives = pipeline_tools._normalized_design_objective_preseed(raw_goal)["splice_isoform_objectives"]
    by_tx = {obj["transcript_id"]: obj for obj in objectives}

    assert set(by_tx) == {"ENST00000626847", "ENST00000380743"}
    assert by_tx["ENST00000626847"]["desired_abundance"] == "decrease"
    assert by_tx["ENST00000626847"]["conversion_role"] == "from"
    assert by_tx["ENST00000380743"]["desired_abundance"] == "increase"
    assert by_tx["ENST00000380743"]["conversion_role"] == "to"


def test_isoform_conversion_objective_accepts_nested_inclusion_skipping_ids() -> None:
    raw_goal = {
        "splice_isoform_objectives": [
            {
                "objective_type": "increase_exon_inclusion",
                "exon_label": "7",
                "exon_numbering_transcript_id": "ENST00000380743",
                "isoform_definitions": {
                    "inclusion_isoform_transcript_id": "ENST00000380743",
                    "skipping_isoform_transcript_id": "ENST00000626847",
                },
                "direction": "shift_toward_inclusion_isoform",
            }
        ]
    }

    objectives = pipeline_tools._normalized_design_objective_preseed(raw_goal)["splice_isoform_objectives"]
    by_tx = {obj["transcript_id"]: obj for obj in objectives}

    assert set(by_tx) == {"ENST00000626847", "ENST00000380743"}
    assert by_tx["ENST00000626847"]["desired_abundance"] == "decrease"
    assert by_tx["ENST00000626847"]["conversion_role"] == "from"
    assert by_tx["ENST00000380743"]["desired_abundance"] == "increase"
    assert by_tx["ENST00000380743"]["conversion_role"] == "to"


def test_isoform_conversion_objective_accepts_nested_isoform_definition_objects() -> None:
    raw_goal = {
        "splice_isoform_objectives": [
            {
                "objective_kind": "increase_exon_inclusion",
                "direction": "shift_toward_included_isoform",
                "exon_label": "exon 7",
                "exon_numbering_transcript_id": "ENST00000380743",
                "isoform_definitions": {
                    "included_isoform": {
                        "label": "full_length",
                        "transcript_id": "ENST00000380743",
                    },
                    "skipped_isoform": {
                        "label": "exon7_skipped",
                        "transcript_id": "ENST00000626847",
                    },
                },
            }
        ]
    }

    objectives = pipeline_tools._normalized_design_objective_preseed(raw_goal)["splice_isoform_objectives"]
    by_tx = {obj["transcript_id"]: obj for obj in objectives}

    assert set(by_tx) == {"ENST00000626847", "ENST00000380743"}
    assert by_tx["ENST00000626847"]["desired_abundance"] == "decrease"
    assert by_tx["ENST00000626847"]["conversion_role"] == "from"
    assert by_tx["ENST00000380743"]["desired_abundance"] == "increase"
    assert by_tx["ENST00000380743"]["conversion_role"] == "to"


def test_isoform_conversion_objective_accepts_desired_undesired_isoform_objects() -> None:
    raw_goal = {
        "splice_isoform_objectives": [
            {
                "objective_kind": "increase_exon_inclusion",
                "exon_reference_transcript_id": "ENST00000380743",
                "exon_label": "exon 7",
                "isoform_definitions": {
                    "desired_isoform": {
                        "transcript_id": "ENST00000380743",
                        "label": "full-length protein",
                    },
                    "undesired_isoform": {
                        "transcript_id": "ENST00000626847",
                        "label": "exon 7 skipped",
                    },
                },
                "directionality": {
                    "increase": ["ENST00000380743"],
                    "decrease": ["ENST00000626847"],
                },
            }
        ]
    }

    objectives = pipeline_tools._normalized_design_objective_preseed(raw_goal)["splice_isoform_objectives"]
    by_tx = {obj["transcript_id"]: obj for obj in objectives}

    assert set(by_tx) == {"ENST00000626847", "ENST00000380743"}
    assert by_tx["ENST00000626847"]["desired_abundance"] == "decrease"
    assert by_tx["ENST00000626847"]["conversion_role"] == "from"
    assert by_tx["ENST00000380743"]["desired_abundance"] == "increase"
    assert by_tx["ENST00000380743"]["conversion_role"] == "to"


def test_isoform_conversion_preseed_extracts_convert_toward_text() -> None:
    preseed = pipeline_tools._request_isoform_conversion_preseed_from_text(
        "Design ASOs to convert Ensembl transcript ENST00000626847 (exon 7 skipped) "
        "toward Ensembl transcript ENST00000380743 by increasing exon 7 inclusion."
    )
    objectives = pipeline_tools._normalized_design_objective_preseed(preseed)["splice_isoform_objectives"]
    by_tx = {obj["transcript_id"]: obj for obj in objectives}

    assert set(by_tx) == {"ENST00000626847", "ENST00000380743"}
    assert by_tx["ENST00000626847"]["desired_abundance"] == "decrease"
    assert by_tx["ENST00000380743"]["desired_abundance"] == "increase"
    assert by_tx["ENST00000380743"]["exon_label"] == "exon 7"


def test_isoform_conversion_preseed_extracts_skipped_and_included_text() -> None:
    preseed = pipeline_tools._request_isoform_conversion_preseed_from_text(
        "Use ENST00000626847 as the exon 7 skipped isoform to decrease and "
        "ENST00000380743 as the exon 7 included isoform to increase."
    )
    objectives = pipeline_tools._normalized_design_objective_preseed(preseed)["splice_isoform_objectives"]
    by_tx = {obj["transcript_id"]: obj for obj in objectives}

    assert set(by_tx) == {"ENST00000626847", "ENST00000380743"}
    assert by_tx["ENST00000626847"]["desired_abundance"] == "decrease"
    assert by_tx["ENST00000380743"]["desired_abundance"] == "increase"
    assert by_tx["ENST00000380743"]["exon_label"] == "exon 7"


def test_isoform_conversion_preseed_accepts_named_increase_with_all_other_decrease_text() -> None:
    preseed = pipeline_tools._request_isoform_conversion_preseed_from_text(
        "Which transcript or isoform IDs should define the SMN2 exon 7 splice objective? "
        "Ensembl ID ENST00000380743 is the isoform which needs to increase. "
        "All other isoforms of SMN2 can decrease."
    )
    objectives = pipeline_tools._normalized_design_objective_preseed(preseed)["splice_isoform_objectives"]

    assert len(objectives) == 1
    objective = objectives[0]
    assert objective["transcript_id"] == "ENST00000380743"
    assert objective["desired_abundance"] == "increase"
    assert objective["conversion_role"] == "to"
    assert objective["conversion_pair"]["to_transcript_id"] == "ENST00000380743"
    assert objective["conversion_pair"]["from_transcript_id"] is None
    assert objective["exon_label"] == "exon 7"


def test_isoform_conversion_preseed_accepts_named_increase_then_other_decrease_text() -> None:
    preseed = pipeline_tools._request_isoform_conversion_preseed_from_text(
        "Which transcript or isoform IDs should define the SMN2 exon 7 splice objective? "
        "I want to increase the long isoform (ENST00000380743.9) and decrease the "
        "other isoforms that dont have exon 7 present in this isoform."
    )
    objectives = pipeline_tools._normalized_design_objective_preseed(preseed)["splice_isoform_objectives"]

    assert len(objectives) == 1
    objective = objectives[0]
    assert objective["transcript_id"] == "ENST00000380743.9"
    assert objective["desired_abundance"] == "increase"
    assert objective["conversion_role"] == "to"
    assert objective["conversion_pair"]["to_transcript_id"] == "ENST00000380743.9"
    assert objective["conversion_pair"]["from_transcript_id"] is None
    assert objective["exon_label"] == "exon 7"


def test_isoform_conversion_preseed_accepts_colon_labeled_increase_and_other_decrease_text() -> None:
    preseed = pipeline_tools._request_isoform_conversion_preseed_from_text(
        "Increase: ENST00000380743.9 (promote exon 7 inclusion). "
        "Decrease: all other SMN2 isoforms/transcripts that do not include exon 7."
    )
    objectives = pipeline_tools._normalized_design_objective_preseed(preseed)["splice_isoform_objectives"]

    assert len(objectives) == 1
    objective = objectives[0]
    assert objective["transcript_id"] == "ENST00000380743.9"
    assert objective["desired_abundance"] == "increase"
    assert objective["conversion_role"] == "to"
    assert objective["conversion_pair"]["to_transcript_id"] == "ENST00000380743.9"
    assert objective["conversion_pair"]["from_transcript_id"] is None
    assert objective["exon_label"] == "exon 7"


def test_normalized_design_preseed_accepts_transcript_id_increase_field() -> None:
    raw_goal = {
        "splice_isoform_objectives": [
            {
                "event_type": "exon_inclusion",
                "exon_label": "7",
                "transcript_id_increase": "ENST00000380743.9",
                "transcripts_to_decrease_description": "other SMN2 isoforms that do not include exon 7",
            }
        ]
    }
    objectives = pipeline_tools._normalized_design_objective_preseed(raw_goal)["splice_isoform_objectives"]

    assert len(objectives) == 1
    objective = objectives[0]
    assert objective["transcript_id"] == "ENST00000380743.9"
    assert objective["desired_abundance"] == "increase"
    assert objective["conversion_role"] == "to"
    assert objective["conversion_pair"]["to_transcript_id"] == "ENST00000380743.9"
    assert objective["conversion_pair"]["from_transcript_id"] is None


def test_normalized_design_preseed_preserves_structured_aggregate_objective_without_label() -> None:
    raw_goal = {
        "splice_isoform_objectives": [
            {
                "objective_scope": "aggregate_class",
                "aggregate_class": {
                    "basis": "exon_presence",
                    "exon_label": 7,
                    "presence": "excluded",
                },
                "desired_abundance": "decrease",
                "role": "pathogenic",
            }
        ]
    }

    objectives = pipeline_tools._normalized_design_objective_preseed(raw_goal)["splice_isoform_objectives"]

    assert len(objectives) == 1
    assert objectives[0]["objective_scope"] == "aggregate_class"
    assert objectives[0]["aggregate_class"]["basis"] == "exon_presence"
    assert objectives[0]["aggregate_class"]["presence"] == "excluded"
    assert "isoform_label" not in objectives[0]


def test_design_preseed_preserves_aggregate_objectives_without_labels() -> None:
    aggregate_objective = {
        "objective_scope": "aggregate_class",
        "aggregate_class": {
            "basis": "exon_presence",
            "exon_label": 7,
            "presence": "excluded",
        },
        "desired_abundance": "decrease",
        "role": "pathogenic",
        "rationale": "Decrease isoforms that do not include exon 7.",
    }
    task = {
        "context": {
            "gene_symbol": "SMN2",
            "species": "human",
            "structured_goal": {
                "gene_symbol": "SMN2",
                "species": "human",
                "intent": "splice_modulation",
                "splice_isoform_objectives": [
                    {
                        "objective_scope": "specific_transcript",
                        "transcript_id": "ENST00000380743.9",
                        "isoform_label": "ENST00000380743.9",
                        "desired_abundance": "increase",
                        "role": "beneficial",
                    },
                    aggregate_objective,
                ],
            },
        }
    }

    preseed = pipeline_tools._prepare_design_objective_preseed(
        task["context"]["structured_goal"],
        task["context"],
    )

    objectives = preseed["splice_isoform_objectives"]
    assert len(objectives) == 2
    assert objectives[1]["objective_scope"] == "aggregate_class"
    assert objectives[1]["aggregate_class"] == aggregate_objective["aggregate_class"]


def test_restore_preseed_merges_aggregate_objectives_without_labels() -> None:
    specific_objective = {
        "objective_scope": "specific_transcript",
        "transcript_id": "ENST00000380743.9",
        "isoform_label": "ENST00000380743.9",
        "desired_abundance": "increase",
        "role": "beneficial",
    }
    aggregate_objective = {
        "objective_scope": "aggregate_class",
        "aggregate_class": {
            "basis": "exon_presence",
            "exon_label": 7,
            "presence": "excluded",
        },
        "desired_abundance": "decrease",
        "role": "pathogenic",
    }
    structured_goal = {
        "gene_symbol": "SMN2",
        "species": "human",
        "intent": "splice_modulation",
        "splice_isoform_objectives": [specific_objective],
    }
    final_payload = {"structured_goal": deepcopy(structured_goal)}
    preseed = {"splice_isoform_objectives": [specific_objective, aggregate_objective]}

    changed = pipeline_tools._restore_design_objective_preseed_to_payload(
        structured_goal,
        final_payload,
        preseed,
    )

    assert changed == 1
    objectives = structured_goal["splice_isoform_objectives"]
    assert len(objectives) == 2
    assert objectives[1]["objective_scope"] == "aggregate_class"
    assert final_payload["splice_isoform_objectives"] == objectives


def test_restore_preseed_merges_source_provenance_into_matching_explicit_target() -> None:
    structured_goal = {
        "explicit_splice_targets": [
            {
                "chrom": "chr2",
                "start": 300,
                "end": 350,
                "strand": "+",
                "label": "descriptive parsed target",
                "exon_label": 4,
                "notes": "Richer model-generated biological detail.",
            }
        ]
    }
    final_payload = {"structured_goal": deepcopy(structured_goal)}
    preseed = {
        "structured_goal_overlay": {
            "explicit_splice_targets": [
                {
                    "chrom": "2",
                    "start": 300,
                    "end": 350,
                    "strand": "+",
                    "transcript_id": "TX0001.2",
                    "source": "clarification_answers",
                    "evidence_refs": ["reviewed_source.json"],
                }
            ]
        }
    }

    changed = pipeline_tools._restore_design_objective_preseed_to_payload(
        structured_goal,
        final_payload,
        preseed,
    )

    assert changed == 1
    assert len(structured_goal["explicit_splice_targets"]) == 1
    target = structured_goal["explicit_splice_targets"][0]
    assert target["label"] == "descriptive parsed target"
    assert target["exon_label"] == 4
    assert target["notes"] == "Richer model-generated biological detail."
    assert target["transcript_id"] == "TX0001.2"
    assert target["source"] == "clarification_answers"
    assert target["evidence_refs"] == ["reviewed_source.json"]


def test_restore_preseed_does_not_bind_provenance_to_unrelated_explicit_target() -> None:
    structured_goal = {
        "explicit_splice_targets": [
            {
                "chrom": "chr2",
                "start": 500,
                "end": 550,
                "strand": "+",
                "transcript_id": "TX0002",
                "label": "different parsed target",
            }
        ]
    }
    final_payload = {"structured_goal": deepcopy(structured_goal)}
    preseed = {
        "structured_goal_overlay": {
            "explicit_splice_targets": [
                {
                    "chrom": "chr2",
                    "start": 300,
                    "end": 350,
                    "strand": "+",
                    "transcript_id": "TX0001",
                    "source": "clarification_answers",
                    "evidence_refs": ["reviewed_source.json"],
                }
            ]
        }
    }

    pipeline_tools._restore_design_objective_preseed_to_payload(
        structured_goal,
        final_payload,
        preseed,
    )

    targets = structured_goal["explicit_splice_targets"]
    assert len(targets) == 2
    assert targets[0]["transcript_id"] == "TX0002"
    assert "evidence_refs" not in targets[0]
    assert targets[1]["transcript_id"] == "TX0001"
    assert targets[1]["evidence_refs"] == ["reviewed_source.json"]


def test_restore_preseed_does_not_merge_conflicting_transcript_identities() -> None:
    structured_goal = {
        "explicit_splice_targets": [
            {
                "chrom": "chr2",
                "start": 300,
                "end": 350,
                "strand": "+",
                "transcript_id": "TX0002",
            }
        ]
    }
    final_payload = {"structured_goal": deepcopy(structured_goal)}
    preseed = {
        "structured_goal_overlay": {
            "explicit_splice_targets": [
                {
                    "chrom": "chr2",
                    "start": 300,
                    "end": 350,
                    "strand": "+",
                    "transcript_id": "TX0001",
                    "evidence_refs": ["reviewed_source.json"],
                }
            ]
        }
    }

    pipeline_tools._restore_design_objective_preseed_to_payload(
        structured_goal,
        final_payload,
        preseed,
    )

    targets = structured_goal["explicit_splice_targets"]
    assert [target["transcript_id"] for target in targets] == ["TX0002", "TX0001"]
    assert "evidence_refs" not in targets[0]


def test_scoped_design_request_excludes_sibling_task_text() -> None:
    parent_request = (
        "Run five tasks: design human SMN2 ASOs to convert ENST00000626847 toward "
        "ENST00000380743 by increasing exon 7 inclusion; design mouse Mapt ASOs for "
        "exon 10 inclusion; download missing reference assets to the default location."
    )
    task = {
        "task_id": "task_design_new_asos_mapt_exon10",
        "task_type": "design_new_asos",
        "label": "Design mouse Mapt splice-switching ASOs for exon 10 inclusion",
        "context": {
            "gene_symbol": "Mapt",
            "species": "mouse",
            "structured_goal": {
                "gene_symbol": "Mapt",
                "species": "mouse",
                "intent": "splice_modulation",
                "exon_label": 10,
                "design_programs": [
                    {
                        "program_id": "steric_splice_switch",
                        "modality": "steric",
                        "intent": "splice_switching",
                        "exon_label": 10,
                    }
                ],
                "splice_isoform_objectives": [
                    {"objective": "increase_exon_inclusion", "exon_label": "exon 10"}
                ],
            },
        },
    }
    request = pipeline_tools._build_scoped_design_request_text(
        parent_request,
        task,
        {"tasks": [task]},
        [],
    )

    assert "Original user request" not in request
    assert "ENST00000626847" not in request
    assert "ENST00000380743" not in request
    assert "SMN2" not in request
    assert "Mapt" in request
    assert "exon 10" in request
    assert "download them to the default location" in request
    assert pipeline_tools._request_isoform_conversion_preseed_from_text(request) == {}


def test_scoped_design_request_inherits_completed_dependency_results(tmp_path: Path) -> None:
    analysis_path = tmp_path / "artifact_analysis.json"
    analysis_path.write_text(
        json.dumps(
            {
                "answer_summary": "The supplied reference defines chr2:1201-1280 on the minus strand.",
                "evidence_refs": ["inputs/reference_context.json"],
                "assumptions": [],
                "limitations": [],
            }
        ),
        encoding="utf-8",
    )
    analyze_task = {
        "task_id": "analyze_reference",
        "task_type": "analyze_artifact",
        "label": "Analyze supplied reference",
        "context": {"artifact_refs": [{"path": "inputs/reference_context.json"}]},
    }
    design_task = {
        "task_id": "design_target",
        "task_type": "design_new_asos",
        "label": "Design target ASOs",
        "depends_on": ["analyze_reference"],
        "context": {
            "gene_symbol": "GENE1",
            "species": "human",
            "target_context": "human cell model",
        },
    }
    state = {
        "task_graph": {
            "tasks": [analyze_task, design_task],
            "dependencies": [
                {
                    "upstream_task_id": "analyze_reference",
                    "downstream_task_id": "design_target",
                }
            ],
        },
        "task_statuses": {"analyze_reference": {"status": "completed"}},
        "task_artifacts": {"analyze_reference": {"json": str(analysis_path)}},
    }

    dependency_results = pipeline_tools._completed_dependency_results_for_task(
        state,
        design_task,
    )
    request = pipeline_tools._build_scoped_design_request_text(
        "Design ASOs using the supplied reference.",
        design_task,
        {"tasks": [analyze_task, design_task]},
        [],
        dependency_results,
    )

    assert len(dependency_results) == 1
    assert dependency_results[0]["upstream_task_id"] == "analyze_reference"
    assert dependency_results[0]["output_payload"]["evidence_refs"] == [
        "inputs/reference_context.json"
    ]
    assert "Completed upstream dependency results" in request
    assert "chr2:1201-1280" in request
    assert "Do not ask for a fact that is already supplied" in request


def test_scoped_design_request_inherits_indexed_direct_artifact_context(
    tmp_path: Path,
) -> None:
    artifact_path = tmp_path / "reference_context.json"
    artifact_path.write_text(
        json.dumps(
            {
                "genome_build": "GRCh38",
                "region": {
                    "chrom": "chr6",
                    "start": 4401,
                    "end": 4480,
                    "strand": "+",
                },
            }
        ),
        encoding="utf-8",
    )
    task = {
        "task_id": "design_target",
        "task_type": "design_new_asos",
        "label": "Design target ASOs",
        "context": {
            "gene_symbol": "GENE1",
            "species": "human",
            "target_context": "human cell model",
            "structured_goal": {
                "artifact_refs": [{"path": str(artifact_path)}],
            },
        },
    }
    context = pipeline_tools._task_context_from_request_envelope(
        {"tasks": [task]},
        task,
    )
    refs = pipeline_tools._task_context_artifact_refs(context)
    artifact_context = pipeline_tools._build_artifact_context(
        {"artifact_refs": refs}
    )
    request = pipeline_tools._build_scoped_design_request_text(
        "Design ASOs using the supplied reference artifact.",
        task,
        {"tasks": [task]},
        [],
        direct_artifact_context=artifact_context,
    )

    assert refs == [str(artifact_path)]
    assert artifact_context["entries"][0]["exists"] is True
    assert "Indexed direct-artifact context" in request
    assert "chr6" in request
    assert "4401" in request
    assert "Do not ask for a fact already present" in request


def test_artifact_reference_merge_uses_declared_path_identity() -> None:
    merged = pipeline_tools._merge_artifact_refs(
        [
            {
                "path": "inputs/reference_context.json",
                "artifact_type": "frozen_sequence_context",
            }
        ],
        [
            {
                "path": "inputs/reference_context.json",
                "label": "sample-specific sequence context",
            }
        ],
    )

    assert merged == [
        {
            "path": "inputs/reference_context.json",
            "artifact_type": "frozen_sequence_context",
            "label": "sample-specific sequence context",
        }
    ]


def test_goal_audit_evidence_loads_declared_local_artifact_contents(
    tmp_path: Path,
) -> None:
    artifact_path = tmp_path / "reference_context.json"
    artifact_path.write_text(
        json.dumps(
            {
                "genome_build": "GRCh38",
                "region": {
                    "chrom": "chr8",
                    "start": 5101,
                    "end": 5180,
                    "strand": "-",
                },
            }
        ),
        encoding="utf-8",
    )
    envelope = {
        "tasks": [
            {
                "task_id": "design_target",
                "task_type": "design_new_asos",
                "context": {
                    "structured_goal": {
                        "artifact_refs": [{"path": str(artifact_path)}]
                    }
                },
            }
        ]
    }
    evidence = [
        {
            "evidence_id": "request",
            "kind": "user_request",
            "text": "Use the supplied reference context.",
        }
    ]

    pipeline_tools._extend_goal_audit_evidence_from_declared_artifacts(
        evidence,
        envelope,
    )
    pipeline_tools._extend_goal_audit_evidence_from_declared_artifacts(
        evidence,
        envelope,
    )

    artifact_sources = [
        item
        for item in evidence
        if item["kind"] == "declared_task_artifact"
    ]
    assert len(artifact_sources) == 1
    assert artifact_sources[0]["evidence_id"] == "artifact.0"
    assert str(artifact_path) in artifact_sources[0]["text"]
    assert "chr8" in artifact_sources[0]["text"]
    assert "5101" in artifact_sources[0]["text"]


def test_scoped_design_request_filters_parent_history_to_task_scope() -> None:
    mapt_task = {
        "task_id": "task_design_new_asos_mapt_mouse_exon10",
        "task_type": "design_new_asos",
        "label": "Task 5: Mouse Mapt splice-switching ASOs to promote exon 10 inclusion",
        "context": {
            "gene_symbol": "Mapt",
            "species": "mouse",
            "structured_goal": {
                "gene_symbol": "Mapt",
                "species": "mouse",
                "intent": "splice_modulation",
                "exon_label": 10,
                "splice_isoform_objectives": [
                    {
                        "objective_kind": "splice_event",
                        "exon_label": "10",
                        "desired_delta": "increase_inclusion",
                        "notes": "Shift tau isoform balance toward exon 10 inclusion",
                    }
                ],
            },
        },
    }
    smn2_task = {
        "task_id": "task_design_new_asos_smn2_baseline_exon7",
        "task_type": "design_new_asos",
        "label": "Task 3: Human SMN2 splice-switching ASOs",
        "context": {
            "gene_symbol": "SMN2",
            "species": "human",
            "structured_goal": {
                "gene_symbol": "SMN2",
                "species": "human",
                "intent": "splice_modulation",
                "exon_label": 7,
                "splice_isoform_objectives": [
                    {
                        "transcript_id": "ENST00000380743.9",
                        "isoform_label": "ENST00000380743.9",
                        "desired_abundance": "increase",
                    }
                ],
            },
        },
    }
    request = pipeline_tools._build_scoped_design_request_text(
        "Run mouse Mapt exon 10 and human SMN2 exon 7 ASO tasks.",
        mapt_task,
        {"tasks": [smn2_task, mapt_task]},
        [
            {
                "question": "For the mouse Mapt splice-switching design, which exon should be included?",
                "answer": "Increase inclusion of mouse Mapt exon 10 to shift the tau isoform balance toward exon 10 inclusion.",
            },
            {
                "question": "For isoform other_isoforms_without_exon_7, should I use ENST00000940690?",
                "answer": "Skip it; use ENST00000380743.9 as the transcript to increase for the SMN2 task.",
            },
        ],
    )

    assert "Increase inclusion of mouse Mapt exon 10" in request
    assert "ENST00000940690" not in request
    assert "ENST00000380743.9" not in request
    assert pipeline_tools._request_isoform_conversion_preseed_from_text(request) == {}


def test_per_site_request_requires_target_site_clarification_without_site_definitions() -> None:
    request = (
        "Design steric splice-switching ASOs for human SMN2 in spinal cord motor neurons "
        "to convert ENST00000626847 toward ENST00000380743 by increasing exon 7 inclusion. "
        "Return at least 30 ASOs per site."
    )

    questions = pipeline_tools._per_site_targeting_clarification_questions(
        {"request": request},
        {"gene_symbol": "SMN2", "candidate_limit_per_site": 30},
    )

    assert len(questions) == 3


def test_per_site_request_accepts_clarified_target_site_definitions() -> None:
    state = {
        "request": "Return at least 30 ASOs per site.",
        "clarification_qa_history": [
            {
                "question": "What should count as distinct target sites?",
                "answer": "Use the intron 7 ISS-N1 region, donor junction, and acceptor junction.",
            }
        ],
    }

    assert pipeline_tools._per_site_targeting_clarification_questions(state, {}) == []


def test_scoped_per_site_request_ignores_task_metadata_when_site_scope_missing() -> None:
    scoped_request = (
        "Create a scoped ASO design run for exactly one task from a larger user request.\n\n"
        "Original user request:\n"
        "Design steric splice-switching ASOs for human SMN2 in spinal cord motor neurons "
        "to convert ENST00000626847 toward ENST00000380743 by increasing exon 7 inclusion. "
        "Return at least 30 ASOs per site.\n\n"
        "Scoped design task:\n"
        '{"task_id":"design_smn2_ex7_acceptor","context":{"structured_goal":'
        '{"design_programs":[{"objective":"Design against splice-site or nearby splice-regulatory regions"}]}}}'
    )

    questions = pipeline_tools._per_site_targeting_clarification_questions(
        {"request": scoped_request},
        {"gene_symbol": "SMN2", "candidate_limit_per_site": 30},
    )

    assert len(questions) == 3


def test_per_site_clarification_builds_splice_relative_site_programs() -> None:
    structured_goal = {
        "gene_symbol": "SMN2",
        "target_context": "spinal cord motor neurons",
        "per_site_candidate_limit": 30,
        "splice_isoform_objectives": [
            {
                "conversion_role": "to",
                "transcript_id": "ENST00000380743",
                "exon_label": "exon 7",
            }
        ],
    }
    text = (
        "Use biologically distinct exon 7 inclusion-relevant steric target regions. "
        "Primary site: the intron 7 ISS-N1 / donor-proximal downstream silencer region. "
        "Also include exon 7 donor-proximal and acceptor-proximal junction regions if targetable. "
        "For the acceptor site, use 25 nt upstream in intron 6 plus the first 25 nt of exon 7. "
        "For the donor/ISS-N1 site, use the last 10 nt of exon 7 plus 60 nt downstream into intron 7. "
        "If ISS-N1 exact coordinates are available use motif span plus padding; otherwise use donor +1 to +60 intron 7 window."
    )

    programs = pipeline_tools._request_splice_relative_window_programs_from_text(text, structured_goal)
    windows = [
        program["region_constraints"][0]["splice_relative_window"]
        for program in programs
    ]

    assert len(programs) == 3
    assert {
        (window["anchor_site"], window["relative_start_nt"], window["relative_end_nt"], window.get("feature_name"))
        for window in windows
    } == {
        ("acceptor", -25, 24, None),
        ("donor", -9, 60, None),
        ("donor", 1, 60, "ISS-N1"),
    }
    assert all(program["candidate_limit"] == 30 for program in programs)


def test_runtime_route_program_prefers_direct_splice_plan_site_over_shared_context() -> None:
    goal = {
        "splice_site_plan": [
            {
                "chrom": "chr17",
                "position": 100,
                "site": "acceptor",
                "delta_direction": "decrease",
                "covers_all_isoforms": False,
                "exon_indices": {"pathogenic": 7},
            },
            {
                "chrom": "chr17",
                "position": 200,
                "site": "donor",
                "delta_direction": "decrease",
                "covers_all_isoforms": False,
                "exon_indices": {"pathogenic": 7},
            },
            {
                "chrom": "chr17",
                "position": 300,
                "site": "acceptor",
                "delta_direction": "decrease",
                "covers_all_isoforms": True,
                "exon_indices": {"pathogenic": 10, "productive": 9},
            },
        ]
    }

    acceptor_entries = run_goal_intake._collect_program_splice_plan_entries(
        goal,
        {"site_kind": "acceptor", "exon_label": 10, "label": "acceptor route"},
    )
    donor_entries = run_goal_intake._collect_program_splice_plan_entries(
        goal,
        {"site_kind": "donor", "exon_label": 10, "label": "donor route"},
    )

    assert [entry["position"] for entry in acceptor_entries] == [100]
    assert [entry["position"] for entry in donor_entries] == [200]


def test_exploratory_derivation_preserves_requested_route_site_kind() -> None:
    structured_goal = {
        "splice_site_plan": [
            {
                "chrom": "chr17",
                "position": 100,
                "site": "acceptor",
                "strand": "+",
                "delta_direction": "decrease",
                "covers_all_isoforms": False,
                "present_in": ["pathogenic"],
                "absent_in": ["productive"],
            },
            {
                "chrom": "chr17",
                "position": 190,
                "site": "donor",
                "strand": "+",
                "delta_direction": "decrease",
                "covers_all_isoforms": False,
                "present_in": ["pathogenic"],
                "absent_in": ["productive"],
            },
        ],
        "splice_isoform_objectives": [
            {"isoform_label": "pathogenic", "desired_abundance": "decrease"}
        ],
    }

    donor_programs, _meta = pipeline_tools._derive_exploratory_steric_programs(
        structured_goal=structured_goal,
        base_program={
            "program_id": "route_donor_window",
            "label": "donor route",
            "modality": "steric",
            "intent": "splice_modulation",
            "objective": "change the requested isoform balance",
            "desired_delta": "decrease",
            "length": 17,
            "isoform_label": "pathogenic",
            "site_kind": "donor",
            "run_splice_eval": True,
            "run_splice_router": True,
        },
    )

    assert len(donor_programs) == 1
    assert donor_programs[0]["site_kind"] == "donor"
    assert donor_programs[0]["candidate_sites"] == ["donor@190"]
    assert donor_programs[0]["routing_candidate_sites"] == ["donor@190"]
    assert donor_programs[0]["window_summaries"][0]["site_kind"] == "donor"
    assert donor_programs[0]["desired_delta"] == "decrease"
    assert donor_programs[0]["length"] == 17
    assert donor_programs[0]["isoform_label"] == "pathogenic"


def test_exploratory_derivation_scopes_named_routes_by_identity_and_direction() -> None:
    structured_goal = {
        "splice_site_plan": [
            {
                "chrom": "chr1",
                "position": 100,
                "site": "acceptor",
                "strand": "+",
                "delta_direction": "increase",
                "delta_per_isoform": {"beneficial": "increase"},
                "present_in": ["beneficial"],
                "absent_in": ["pathogenic"],
                "covers_all_isoforms": False,
            },
            {
                "chrom": "chr1",
                "position": 500,
                "site": "acceptor",
                "strand": "+",
                "delta_direction": "decrease",
                "delta_per_isoform": {"pathogenic": "decrease"},
                "present_in": ["pathogenic"],
                "absent_in": ["beneficial"],
                "covers_all_isoforms": False,
            },
        ],
        "splice_isoform_objectives": [
            {"isoform_label": "beneficial", "desired_abundance": "increase"},
            {"isoform_label": "pathogenic", "desired_abundance": "decrease"},
        ],
        "isoform_transcript_ids": {
            "beneficial": "TX_BENEFICIAL",
            "pathogenic": "TX_PATHOGENIC",
        },
    }
    # Strategy analysis proposes candidates across every route. That global
    # pool must not restore sites removed by the typed per-route scope filter.
    structured_goal["splice_strategy_analysis"] = {
        "context_payload": {
            "steric_site_candidates": [
                {**entry, "label": f"acceptor@{entry['position']}"}
                for entry in structured_goal["splice_site_plan"]
            ]
        }
    }

    beneficial, _meta = pipeline_tools._derive_exploratory_steric_programs(
        structured_goal=structured_goal,
        base_program={
            "program_id": "beneficial_route",
            "modality": "steric",
            "intent": "splice_modulation",
            "isoform_label": "beneficial",
            "desired_delta": "increase",
            "site_kind": "acceptor",
        },
    )
    pathogenic, _meta = pipeline_tools._derive_exploratory_steric_programs(
        structured_goal=structured_goal,
        base_program={
            "program_id": "pathogenic_route",
            "modality": "steric",
            "intent": "splice_modulation",
            "transcript_id": "TX_PATHOGENIC",
            "desired_delta": "decrease",
            "site_kind": "acceptor",
        },
    )

    assert beneficial[0]["anchor_sites"] == ["acceptor@100"]
    assert pathogenic[0]["anchor_sites"] == ["acceptor@500"]
    assert pipeline_tools._opposing_named_splice_route_collisions(
        structured_goal, beneficial + pathogenic,
    ) == []
    assert beneficial[0]["anchor_site_records"] == [
        {
            "chrom": "chr1",
            "position": 100,
            "site_kind": "acceptor",
            "strand": "+",
        }
    ]


def test_exploratory_derivation_reports_unresolved_named_route_scope() -> None:
    structured_goal = {
        "splice_site_plan": [
            {
                "chrom": "chr2",
                "position": 210,
                "site": "donor",
                "strand": "+",
                "delta_per_isoform": {"beneficial": "increase"},
                "present_in": ["beneficial"],
            }
        ],
        "isoform_transcript_ids": {"pathogenic": "TX_PATHOGENIC"},
    }

    programs, metadata = pipeline_tools._derive_exploratory_steric_programs(
        structured_goal=structured_goal,
        base_program={
            "program_id": "pathogenic_route",
            "modality": "steric",
            "intent": "splice_modulation",
            "isoform_label": "pathogenic",
            "desired_delta": "decrease",
            "site_kind": "donor",
        },
    )

    assert programs == []
    assert metadata["status"] == "unresolved_named_route_scope"
    assert metadata["program_id"] == "pathogenic_route"
    assert metadata["identities"] == ["pathogenic", "tx_pathogenic"]


def test_route_integrity_detects_opposing_named_routes_on_same_anchor() -> None:
    goal = {
        "isoform_transcript_ids": {
            "beneficial": "TX_BENEFICIAL",
            "pathogenic": "TX_PATHOGENIC",
        }
    }
    programs = [
        {
            "program_id": "increase_route",
            "modality": "steric",
            "isoform_label": "beneficial",
            "desired_delta": "increase",
            "anchor_site_records": [
                {
                    "chrom": "chr3",
                    "position": 300,
                    "site_kind": "acceptor",
                    "strand": "+",
                }
            ],
        },
        {
            "program_id": "decrease_route",
            "modality": "steric",
            "transcript_id": "TX_PATHOGENIC",
            "desired_delta": "decrease",
            "anchor_site_records": [
                {
                    "chrom": "3",
                    "position": 300,
                    "site_kind": "acceptor",
                    "strand": "+",
                }
            ],
        },
    ]

    collisions = pipeline_tools._opposing_named_splice_route_collisions(
        goal,
        programs,
    )

    assert len(collisions) == 1
    assert collisions[0]["left_program_id"] == "increase_route"
    assert collisions[0]["right_program_id"] == "decrease_route"


def test_route_integrity_accepts_distinct_anchors_for_opposing_named_routes() -> None:
    goal = {
        "isoform_transcript_ids": {
            "beneficial": "TX_BENEFICIAL",
            "pathogenic": "TX_PATHOGENIC",
        }
    }
    programs = [
        {
            "program_id": "increase_route",
            "modality": "steric",
            "isoform_label": "beneficial",
            "desired_delta": "increase",
            "anchor_site_records": [
                {
                    "chrom": "chr3",
                    "position": 300,
                    "site_kind": "acceptor",
                    "strand": "+",
                }
            ],
        },
        {
            "program_id": "decrease_route",
            "modality": "steric",
            "isoform_label": "pathogenic",
            "desired_delta": "decrease",
            "anchor_site_records": [
                {
                    "chrom": "chr3",
                    "position": 900,
                    "site_kind": "acceptor",
                    "strand": "+",
                }
            ],
        },
    ]

    assert pipeline_tools._opposing_named_splice_route_collisions(goal, programs) == []


def test_same_gene_site_split_tasks_retain_independent_counts_and_outputs() -> None:
    base_goal = {
        "candidate_limit": 20,
        "splice_isoform_objectives": [
            {
                "objective_scope": "default_exon_context",
                "exon_label": 7,
                "desired_abundance": "increase",
                "desired_effect": "inclusion",
                "role": "beneficial",
            }
        ],
        "gene_symbol": "SMN2",
        "species": "human",
        "target_context": "Human spinal cord motor neurons in an SMA in vitro model.",
        "intent": "splice_modulation",
    }

    def scoped_goal(program_id: str, site_kind: str, site_context: str) -> dict[str, Any]:
        goal = deepcopy(base_goal)
        goal["design_programs"] = [
            {
                "program_id": program_id,
                "label": f"{site_kind} {site_context}",
                "modality": "steric",
                "intent": "splice_modulation",
                "objective": "increase exon inclusion at the supplied route",
                "desired_delta": "increase",
                "length": 17,
                "exon_label": 7,
                "site_kind": site_kind,
                "site_context": site_context,
                "focus_exons_only": True,
                "candidate_limit": 20,
            }
        ]
        goal["evaluation_targets"] = [
            {
                "exon_label": 7,
                "site_kind": site_kind,
                "site_context": site_context,
                "desired_delta": "increase",
            }
        ]
        return goal

    raw_envelope = {
        "request_mode": "composite",
        "workflow_name": "composite_workflow",
        "final_outputs": [
            {
                "kind": "acceptor_candidates",
                "task_id": "task_design_new_asos_acceptor",
            },
            {
                "kind": "donor_candidates",
                "task_id": "task_design_new_asos_donor",
            },
        ],
        "tasks": [
            {
                "task_id": "task_design_new_asos_acceptor",
                "task_type": "design_new_asos",
                "label": "Design SMN2 exon 7 acceptor-proximal junction window",
                "context": {
                    "structured_goal": scoped_goal(
                        "steric_acceptor",
                        "acceptor",
                        "site_window",
                    )
                },
            },
            {
                "task_id": "task_design_new_asos_donor",
                "task_type": "design_new_asos",
                "label": "Design SMN2 exon 7 donor-proximal junction window",
                "context": {
                    "structured_goal": scoped_goal(
                        "steric_donor",
                        "donor",
                        "site_window",
                    )
                },
            },
            {
                "task_id": "task_design_new_asos_issn1",
                "task_type": "design_new_asos",
                "label": "Design SMN2 exon 7 ISS-N1 downstream silencer window",
                "context": {
                    "structured_goal": scoped_goal(
                        "steric_downstream",
                        "donor",
                        "downstream_intron",
                    )
                },
            },
            {
                "task_id": "task_build_composite_report",
                "task_type": "build_composite_report",
                "label": "Combine per-region outputs",
                "depends_on": [
                    "task_design_new_asos_acceptor",
                    "task_design_new_asos_donor",
                    "task_design_new_asos_issn1",
                ],
            },
        ],
    }

    envelope = pipeline_tools._build_request_envelope(
        request_text="Design SMN2 exon 7 ASOs and return 20 candidates per site.",
        request_id="request_smn2_sites",
        structured_goal=None,
        raw_envelope=raw_envelope,
    )

    assert envelope["request_mode"] == "composite"
    assert envelope["workflow_name"] == "composite_workflow"
    design_tasks = [task for task in envelope["tasks"] if task["task_type"] == "design_new_asos"]
    assert {task["task_id"] for task in design_tasks} == {
        "task_design_new_asos_acceptor",
        "task_design_new_asos_donor",
        "task_design_new_asos_issn1",
    }
    goals = [task["context"]["structured_goal"] for task in design_tasks]
    assert [goal["design_programs"][0]["program_id"] for goal in goals] == [
        "steric_acceptor",
        "steric_donor",
        "steric_downstream",
    ]
    for goal in goals:
        assert goal["candidate_limit"] == 20
        assert len(goal["evaluation_targets"]) == 1
        assert len(goal["design_programs"]) == 1
        assert goal["design_programs"][0]["length"] == 17
        assert goal["design_programs"][0]["candidate_limit"] == 20
    assert {output["task_id"] for output in envelope["final_outputs"]} == {
        "task_design_new_asos_acceptor", "task_design_new_asos_donor",
    }
    report = next(task for task in envelope["tasks"] if task["task_type"] == "build_composite_report")
    assert set(report["depends_on"]) == {task["task_id"] for task in design_tasks}
    task = design_tasks[0]
    scoped = pipeline_tools._task_scoped_request_envelope(envelope, task)
    assert {output["task_id"] for output in scoped["final_outputs"]} == {
        task["task_id"]
    }


def test_splice_relative_windows_use_route_signal_exon_when_transcript_indices_are_offset(tmp_path: Path) -> None:
    structured_goal = {
        "gene_symbol": "SMN2",
        "focus_exons_only": True,
        "splice_isoform_objectives": [
            {
                "isoform_label": "ENST00000626847",
                "transcript_id": "ENST00000626847",
                "desired_abundance": "decrease",
                "conversion_role": "from",
                "exon_label": "exon 7",
            },
            {
                "isoform_label": "ENST00000380743",
                "transcript_id": "ENST00000380743",
                "desired_abundance": "increase",
                "conversion_role": "to",
                "exon_label": "exon 7",
            },
        ],
        "isoform_transcripts": {
            "ENST00000626847": {
                "transcript_id": "ENST00000626847",
                "chrom": "chr5",
                "strand": "+",
                "exons_tx": [
                    {"index": 7, "start": 70070641, "end": 70070751, "seq_region_name": "5"},
                    {"index": 8, "start": 70077019, "end": 70077595, "seq_region_name": "5"},
                ],
            },
            "ENST00000380743": {
                "transcript_id": "ENST00000380743",
                "chrom": "chr5",
                "strand": "+",
                "exons_tx": [
                    {"index": 7, "start": 70070641, "end": 70070751, "seq_region_name": "5"},
                    {"index": 8, "start": 70076521, "end": 70076574, "seq_region_name": "5"},
                ],
            },
        },
        "splice_site_plan": [
            {
                "chrom": "chr5",
                "position": 70070641,
                "site": "acceptor",
                "delta_direction": "increase",
                "present_in": ["ENST00000380743", "ENST00000626847"],
                "exon_indices": {"ENST00000380743": 7, "ENST00000626847": 7},
                "uniqueness": "shared",
                "covers_all_isoforms": True,
                "delta_per_isoform": {"ENST00000380743": "increase", "ENST00000626847": "decrease"},
            },
            {
                "chrom": "chr5",
                "position": 70076521,
                "site": "acceptor",
                "delta_direction": "increase",
                "present_in": ["ENST00000380743"],
                "absent_in": ["ENST00000626847"],
                "exon_indices": {"ENST00000380743": 8},
                "uniqueness": "unique",
                "covers_all_isoforms": False,
                "delta_per_isoform": {"ENST00000380743": "increase"},
            },
            {
                "chrom": "chr5",
                "position": 70077019,
                "site": "acceptor",
                "delta_direction": "decrease",
                "present_in": ["ENST00000626847"],
                "absent_in": ["ENST00000380743"],
                "exon_indices": {"ENST00000626847": 8},
                "uniqueness": "unique",
                "covers_all_isoforms": False,
                "delta_per_isoform": {"ENST00000626847": "decrease"},
            },
        ],
        "design_programs": [
            {
                "program_id": "steric_exon7_acceptor_window",
                "modality": "steric",
                "region_constraints": [
                    {
                        "source_type": "splice_relative_window",
                        "label": "exon7_acceptor_window",
                        "exon_numbering_transcript_id": "ENST00000380743",
                        "splice_relative_window": {
                            "exon_label": 7,
                            "anchor_site": "acceptor",
                            "relative_start_nt": -25,
                            "relative_end_nt": 24,
                            "reference_transcript_id": "ENST00000380743",
                        },
                    }
                ],
            },
            {
                "program_id": "steric_exon7_iss_n1_window",
                "modality": "steric",
                "region_constraints": [
                    {
                        "source_type": "splice_relative_window",
                        "label": "exon7_iss_n1_window",
                        "exon_numbering_transcript_id": "ENST00000380743",
                        "aso_overlap_requirement": "any_overlap",
                        "splice_relative_window": {
                            "exon_label": 7,
                            "anchor_site": "donor",
                            "relative_start_nt": 1,
                            "relative_end_nt": 60,
                            "reference_transcript_id": "ENST00000380743",
                            "feature_name": "ISS-N1",
                        },
                    }
                ],
            },
        ],
    }
    final_payload: dict[str, Any] = {}

    run_goal_intake._materialize_region_constraints(structured_goal, final_payload, str(tmp_path))
    by_program = structured_goal["materialized_region_constraints"]["by_program"]

    acceptor = by_program["steric_exon7_acceptor_window"][0]
    iss_n1 = by_program["steric_exon7_iss_n1_window"][0]
    assert acceptor["materialized_rows"] == [
        ["chr5", 70076495, 70076545, "exon7_acceptor_window", 0, "+"]
    ]
    assert iss_n1["materialized_rows"] == [
        ["chr5", 70076574, 70076634, "exon7_iss_n1_window", 0, "+"]
    ]
    assert acceptor["materialization_summary"]["anchor_resolution"]["source"] == "splice_site_plan_route_signal"
    assert iss_n1["materialization_summary"]["resolved_exon_label"] == 8

    program = structured_goal["design_programs"][0]
    program_rows, program_meta = run_goal_intake._build_steric_program_region_rows(structured_goal, program)
    constrained_rows, constraint_meta = run_goal_intake._apply_region_constraints_to_rows(
        structured_goal,
        program_rows,
        output_space="genomic",
        program=program,
    )
    assert program_meta["direct_region_constraints_seed"] is True
    assert constrained_rows == [["chr5", 70076495, 70076545, "exon7_acceptor_window", 0, "+"]]
    assert constraint_meta["constrained_row_count"] == 1


def test_splice_planning_uses_verified_biological_label_mapping_without_renaming_request() -> None:
    transcript = {
        "transcript_id": "TRANSCRIPT_WITH_VERIFIED_SOURCE",
        "chrom": "chr2",
        "strand": "+",
        "exons_tx": [
            {"index": 2, "start": 1000, "end": 1050},
            {"index": 3, "start": 2000, "end": 2060},
            {"index": 4, "start": 3000, "end": 3060},  # Keep the mapped exon internal.
        ],
        "sites": [
            {
                "exon_index": 2,
                "chrom": "chr2",
                "strand": "+",
                "acceptor": 1000,
                "donor": 1050,
                "coding": True,
            },
            {
                "exon_index": 3,
                "chrom": "chr2",
                "strand": "+",
                "acceptor": 2000,
                "donor": 2060,
                "coding": True,
            },
        ],
        "metadata": {
            "exon_nomenclature_resolutions": {
                "2": {
                    "status": "verified_alternate_mapping",
                    "source": "authoritative_exon_nomenclature_mapping",
                    "requested_exon_label": "2",
                    "resolved_transcript_ordinal": 3,
                    "ordered_biological_labels": ["1", "2a", "2"],
                    "evidence_quote": "The source defines the complete order.",
                    "evidence_refs": ["https://example.invalid/record"],
                    "source_sha256": "a" * 64,
                    "checker_reason": "The mapping is supported.",
                }
            }
        },
    }
    goal = {
        "gene_symbol": "GENE",
        "exon_label": 2,
        "focus_exons_only": True,
        "desired_effect": "inclusion",
        "splice_isoform_objectives": [
            {
                "isoform_label": "reference transcript",
                "exon_label": 2,
                "desired_abundance": "increase",
            }
        ],
    }

    plan, _comparisons = run_goal_intake._compute_splice_site_plan(
        goal["splice_isoform_objectives"],
        {"reference transcript": transcript},
        goal,
    )

    assert goal["exon_label"] == 2
    assert {entry["position"] for entry in plan} == {2000, 2060}
    assert {entry["exon_label"] for entry in plan} == {3}

    resolved, metadata = run_goal_intake._resolve_splice_relative_effective_exon(
        constraint={},
        structured_goal=goal,
        iso_label="reference transcript",
        transcript_entry=transcript,
        isoform_transcripts={"reference transcript": transcript},
        requested_exon_label=2,
        requested_site="acceptor",
    )
    assert resolved == 3
    assert metadata["source"] == "authoritative_exon_nomenclature_mapping"
    assert metadata["authoritative_relation_mapping"] is True


def test_authoritative_exon_mapping_cache_survives_in_transcript_metadata(
    monkeypatch,
) -> None:
    calls = []

    def resolve_mapping(**kwargs):
        calls.append(kwargs)
        return {
            "status": "verified_alternate_mapping",
            "source": "authoritative_exon_nomenclature_mapping",
            "requested_exon_label": "1",
            "resolved_transcript_ordinal": 2,
        }

    monkeypatch.setattr(
        run_goal_intake,
        "resolve_transcript_exon_nomenclature_sync",
        resolve_mapping,
    )
    transcript = {
        "transcript_id": "ENST00000000001",
        "metadata": {"id": "ENST00000000001"},
        "exons_tx": [
            {"index": 1, "start": 100, "end": 150},
            {"index": 2, "start": 200, "end": 250},
        ],
    }

    first = run_goal_intake._cached_authoritative_exon_nomenclature_mapping(
        transcript,
        1,
    )
    second = run_goal_intake._cached_authoritative_exon_nomenclature_mapping(
        transcript,
        1,
    )

    assert first == second
    assert len(calls) == 1
    assert transcript["metadata"]["exon_nomenclature_resolutions"]["1"] == first
    assert "exon_nomenclature_resolutions" not in transcript


def test_transcript_identity_takes_scope_precedence_over_descriptive_isoform_label() -> None:
    goal = {
        "intent": "splice_modulation",
        "design_programs": [
            {
                "program_id": "route_a",
                "label": "Named splice route",
                "modality": "steric",
                "intent": "splice_modulation",
                "objective": "Increase the transcript-defined product.",
                "desired_delta": "increase",
                "desired_effect": "inclusion",
                "site_kind": "acceptor",
                "run_splice_router": True,
                "run_splice_eval": True,
                "transcript_id": "ENST00000000001",
                "isoform_label": "descriptive product label",
            }
        ],
        "evaluation_targets": [
            {
                "label": "Transcript-defined route target",
                "objective_type": "route_target",
                "desired_delta": "increase",
                "target_type": "transcript",
                "transcript_id": "ENST00000000001",
                "isoform_label": "descriptive product label",
            }
        ],
        "splice_isoform_objectives": [
            {
                "objective_scope": "specific_transcript",
                "transcript_id": "ENST00000000001",
                "isoform_label": "descriptive product label",
                "desired_abundance": "increase",
                "desired_effect": "inclusion",
            }
        ],
    }

    issues = pipeline_tools._canonical_design_goal_contract_issues(goal)

    assert (
        "context.structured_goal.splice_isoform_objectives[0].objective_scope"
        not in issues
    )


def test_isoform_only_route_still_requires_specific_isoform_scope() -> None:
    goal = {
        "intent": "splice_modulation",
        "design_programs": [
            {
                "program_id": "route_a",
                "label": "Named splice route",
                "modality": "steric",
                "intent": "splice_modulation",
                "objective": "Increase the named product.",
                "desired_delta": "increase",
                "desired_effect": "inclusion",
                "site_kind": "acceptor",
                "run_splice_router": True,
                "run_splice_eval": True,
                "isoform_label": "descriptive product label",
            }
        ],
        "evaluation_targets": [
            {
                "label": "Named route target",
                "objective_type": "route_target",
                "desired_delta": "increase",
                "target_type": "isoform",
                "isoform_label": "descriptive product label",
            }
        ],
        "splice_isoform_objectives": [
            {
                "objective_scope": "specific_transcript",
                "isoform_label": "descriptive product label",
                "desired_abundance": "increase",
                "desired_effect": "inclusion",
            }
        ],
    }

    issues = pipeline_tools._canonical_design_goal_contract_issues(goal)

    assert (
        "context.structured_goal.splice_isoform_objectives[0].objective_scope"
        in issues
    )


def test_explicit_target_mapping_precedes_inferred_splice_route_labels() -> None:
    transcript = {
        "transcript_id": "TRANSCRIPT_A",
        "chrom": "chr2",
        "strand": "+",
        "exons_tx": [
            {"index": 4, "start": 1000, "end": 1100},
            {"index": 5, "start": 2000, "end": 2060},
        ],
    }
    goal = {
        "explicit_splice_targets": [
            {
                "transcript_id": "TRANSCRIPT_A",
                "requested_exon_label": 4,
                "chrom": "chr2",
                "start": 1990,
                "end": 2020,
                "strand": "+",
            },
            {
                "transcript_id": "TRANSCRIPT_A",
                "requested_exon_label": 4,
                "chrom": "chr2",
                "start": 2050,
                "end": 2090,
                "strand": "+",
            },
        ],
        "splice_site_plan": [
            {
                "position": 1000,
                "site": "acceptor",
                "delta_direction": "increase",
                "exon_indices": {"synthetic route": 4},
                "uniqueness": "unique",
                "covers_all_isoforms": False,
            }
        ],
    }
    constraint = {
        "reference_transcript_id": "TRANSCRIPT_A",
        "splice_relative_window": {"reference_transcript_id": "TRANSCRIPT_A"},
    }

    resolved, metadata = run_goal_intake._resolve_splice_relative_effective_exon(
        constraint=constraint,
        structured_goal=goal,
        iso_label="included isoform",
        transcript_entry=transcript,
        isoform_transcripts={"included isoform": transcript},
        requested_exon_label=4,
        requested_site="acceptor",
    )

    assert resolved == 5
    assert metadata["source"] == "explicit_splice_target_mapping"
    assert metadata["explicit_target_indexes"] == [0, 1]


def test_explicit_target_coordinates_override_conflicting_transcript_ordinal() -> None:
    goal = {
        "explicit_splice_targets": [
            {
                "chrom": "chr2",
                "start": 2000,
                "end": 2060,
                "strand": "+",
                "isoform_label": "included isoform::explicit target model",
                "base_isoform_label": "included isoform",
                "requested_exon_label": 4,
                "exon_label": 4,
            }
        ]
    }
    transcripts = {
        "included isoform": {
            "transcript_id": "TRANSCRIPT_A",
            "strand": "+",
            "exons_tx": [
                {"index": 4, "start": 1000, "end": 1100},
                {"index": 5, "start": 2000, "end": 2060},
            ],
        }
    }

    normalized = run_goal_intake._normalize_explicit_splice_targets(
        goal,
        transcripts,
    )

    assert normalized[0]["requested_exon_label"] == 4
    assert normalized[0]["exon_label"] == 5
    assert normalized[0]["mapped_start"] == 2000
    assert normalized[0]["mapped_end"] == 2060
    assert normalized[0]["mapping_mode"] == "transcript_exon_overlap"


def test_resolved_target_consistency_payload_binds_materialization_to_source_evidence() -> None:
    goal = {
        "design_programs": [
            {
                "program_id": "site_a",
                "label": "requested feature route",
                "objective": "Target the requested feature.",
                "notes": "Do not substitute a distinct unresolved feature.",
                "transcript_id": "TX0001.2",
                "exon_label": 2,
                "site_kind": "acceptor",
            }
        ],
        "explicit_splice_targets": [
            {
                "chrom": "chr2",
                "start": 300,
                "end": 350,
                "strand": "+",
                "transcript_id": "TX0001.2",
                "source": "clarification_answers",
                "evidence_refs": ["reviewed_source.json"],
            }
        ],
        "isoform_transcripts": {
            "requested isoform": {
                "transcript_id": "TX0001.2",
                "strand": "+",
                "exons_tx": [
                    {"index": 2, "start": 300, "end": 350, "seq_region_name": "2"},
                    {"index": 3, "start": 500, "end": 560, "seq_region_name": "2"},
                ],
            }
        },
        "splice_site_plan": [
            {
                "chrom": "chr2",
                "position": 300,
                "site": "acceptor",
                "strand": "+",
                "exon_indices": {"requested isoform": 2},
            },
            {
                "chrom": "chr2",
                "position": 500,
                "site": "acceptor",
                "strand": "+",
                "source": "explicit_splice_target",
                "exon_indices": {"synthetic target model": 3},
            },
        ],
        "materialized_region_constraints": {
            "global": [],
            "by_program": {
                "site_a": [
                    {
                        "constraint_id": "window_a",
                        "source_type": "splice_relative_window",
                        "label": "requested acceptor window",
                        "splice_relative_window": {
                            "reference_transcript_id": "TX0001",
                            "exon_label": 2,
                            "anchor_site": "acceptor",
                            "relative_start_nt": -10,
                            "relative_end_nt": 10,
                        },
                        "materialized_rows": [
                            ["chr2", 289, 310, "requested acceptor window", 0, "+"]
                        ],
                        "materialization_summary": {
                            "status": "materialized",
                            "transcript_id": "TX0001.2",
                            "exon_label": 2,
                            "resolved_exon_label": 2,
                            "anchor_site": "acceptor",
                            "anchor_position": 300,
                        },
                    }
                ]
            },
        },
    }
    state = {
        "request": "Use the requested feature and exclude a different unresolved feature.",
        "clarification_qa_history": [
            {
                "question": "Which exact feature mapping should be used?",
                "answer": "Use transcript feature index 2.",
                "evidence_refs": ["reviewed_source.json"],
            }
        ],
    }

    payload = pipeline_tools._resolved_target_consistency_review_payload(
        state=state,
        structured_goal=goal,
    )

    assert payload["request_evidence"]["clarification_history"][0]["answer"] == (
        "Use transcript feature index 2."
    )
    assert payload["materialization_records"][0]["record_index"] == 0
    assert payload["materialization_records"][0]["declared_program_target"] == {
        "label": "requested feature route",
        "objective": "Target the requested feature.",
        "notes": "Do not substitute a distinct unresolved feature.",
        "transcript_id": "TX0001.2",
        "exon_label": 2,
        "site_kind": "acceptor",
    }
    assert payload["materialization_records"][0]["resolved_target"]["anchor_position"] == 300
    assert payload["materialization_records"][0]["resolved_target"][
        "evidence_capabilities"
    ] == [
        "ordinal_identity",
        "explicit_positive_mapping",
        "authoritative_relation_mapping",
    ]
    assert payload["materialization_records"][0]["resolved_target"][
        "evidence_bindings"
    ] == [
        {
            "source_type": "explicit_splice_target",
            "target_index": 0,
            "source": "clarification_answers",
            "chrom": "chr2",
            "start": 300,
            "end": 350,
            "strand": "+",
            "transcript_id": "TX0001.2",
            "evidence_refs": ["reviewed_source.json"],
        }
    ]
    assert payload["authoritative_resolution_evidence"]["transcript_features"][0][
        "exons_in_transcript_order"
    ][1]["index"] == 3
    assert len(payload["authoritative_resolution_evidence"]["splice_site_plan"]) == 1


def test_resolved_target_evidence_binding_rejects_a_different_explicit_target() -> None:
    record = {
        "requested_target": {"reference_transcript_id": "TX0001"},
        "resolved_target": {
            "transcript_id": "TX0001.2",
            "anchor_position": 300,
            "materialized_rows": [["chr2", 289, 310, "window", 0, "+"]],
        },
        "declared_program_target": {},
    }

    bindings = pipeline_tools._compact_explicit_target_evidence_bindings(
        record,
        [
            {
                "chrom": "chr2",
                "start": 500,
                "end": 560,
                "strand": "+",
                "transcript_id": "TX0001.2",
                "evidence_refs": ["reviewed_source.json"],
            }
        ],
    )

    assert bindings == []


def test_resolved_target_consistency_review_blocks_task_with_canonical_missing_field(
    tmp_path: Path,
    monkeypatch: Any,
) -> None:
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    goal = {
        "gene_symbol": "GENE1",
        "species": "human",
        "target_context": "reviewed cells",
        "intent": "splice_modulation",
        "run_splice_router": True,
        "run_splice_eval": True,
        "design_programs": [
            {
                "program_id": "site_a",
                "label": "site a",
                "modality": "steric",
                "intent": "splice_modulation",
                "objective": "increase the requested feature",
                "desired_delta": "increase",
            }
        ],
        "materialized_region_constraints": {
            "global": [
                {
                    "constraint_id": "window_a",
                    "source_type": "splice_relative_window",
                    "splice_relative_window": {
                        "reference_transcript_id": "TX0001",
                        "exon_label": 2,
                        "anchor_site": "acceptor",
                        "relative_start_nt": -10,
                        "relative_end_nt": 10,
                    },
                    "materialized_rows": [["chr2", 289, 310, "window", 0, "+"]],
                    "materialization_summary": {
                        "status": "materialized",
                        "transcript_id": "TX0001",
                        "exon_label": 2,
                        "resolved_exon_label": 2,
                        "anchor_site": "acceptor",
                        "anchor_position": 300,
                    },
                }
            ],
            "by_program": {},
        }
    }
    final_payload = {"structured_goal": deepcopy(goal)}
    raw_payload = {"final_payload": final_payload}
    pipeline_tools._save_goal_payload(str(run_dir), raw_payload, goal, final_payload)
    request_text = "Use the named feature but exclude a separately defined alternative."
    state = _state_from_raw_envelope(
        {
            "request_text": request_text,
            "tasks": [
                {
                    "task_id": "design",
                    "task_type": "design_new_asos",
                    "label": "design",
                    "context": {"structured_goal": deepcopy(goal)},
                }
            ],
        },
        tmp_path,
    )
    state["run_dir"] = str(run_dir)
    state["current_task_id"] = "design"
    state["task_statuses"] = {"design": {"status": "in_progress"}}
    state_path = tmp_path / "run_state.json"
    state_path.write_text(json.dumps(state), encoding="utf-8")
    monkeypatch.setattr(pipeline_tools, "provider_credentials_available", lambda: True)
    monkeypatch.setattr(
        pipeline_tools,
        "_run_resolved_target_consistency_sync",
        lambda _payload: {
            "status": "needs_clarification",
            "contradiction_record_indexes": [0],
            "question": "What exact positive feature mapping should be used?",
            "reason": "The supplied exclusion cannot be verified from the resolved target.",
        },
    )

    result = pipeline_tools._review_materialized_target_consistency(
        state=state,
        state_path=str(state_path),
        run_dir=str(run_dir),
        payload=raw_payload,
        structured_goal=goal,
        final_payload=final_payload,
    )
    saved = json.loads(state_path.read_text(encoding="utf-8"))

    assert result is not None and result["needs_user_input"] is True
    assert saved["clarification_queue"][0]["missing_fields"] == [
        "target_feature_mapping_missing"
    ]
    assert saved["pending_questions"] == [
        pipeline_tools.TARGET_FEATURE_MAPPING_QUESTION
    ]


def test_resolved_target_consistency_contract_requires_finite_clarification() -> None:
    schema = pipeline_tools._resolved_target_consistency_schema(1)
    errors = pipeline_tools._resolved_target_consistency_contract_errors(
        {
            "status": "needs_clarification",
            "record_findings": [
                {
                    "record_index": 0,
                    "identity_requirement": "The intended target must be distinguished.",
                    "requirement_kind": "substitution_exclusion",
                    "status": "not_verifiable",
                    "evidence_assessment": "No positive mapping was supplied.",
                }
            ],
            "contradiction_record_indexes": [],
            "question": None,
            "reason": None,
        },
        record_count=1,
        schema=schema,
    )

    assert "needs_clarification review must select at least one materialization record" in errors
    assert "needs_clarification review must return one question" in errors
    assert "needs_clarification review must return one reason" in errors


def test_resolved_target_consistency_contract_rejects_unsupported_relational_claim() -> None:
    schema = pipeline_tools._resolved_target_consistency_schema(1)
    payload = {
        "status": "consistent",
        "record_findings": [
            {
                "record_index": 0,
                "identity_requirement": "A distinct target must not be substituted.",
                "requirement_kind": "substitution_exclusion",
                "status": "satisfied",
                "evidence_assessment": "The selected ordinal matches the requested ordinal.",
            }
        ],
        "contradiction_record_indexes": [],
        "question": None,
        "reason": None,
    }
    records = [
        {
            "record_index": 0,
            "resolved_target": {"evidence_capabilities": ["ordinal_identity"]},
        }
    ]

    errors = pipeline_tools._resolved_target_consistency_contract_errors(
        payload,
        record_count=1,
        schema=schema,
        records=records,
    )

    assert errors == [
        "positive relational finding for record 0 lacks a relation-capable evidence source"
    ]

    recovered = pipeline_tools._recover_resolved_target_consistency_clarification(
        payload,
        records,
        schema,
    )

    assert recovered is not None
    assert recovered["status"] == "needs_clarification"
    assert recovered["contradiction_record_indexes"] == [0]
    assert recovered["question"] == pipeline_tools.TARGET_FEATURE_MAPPING_QUESTION
    assert recovered["record_findings"][0]["status"] == "not_verifiable"
    assert (
        pipeline_tools._resolved_target_consistency_contract_errors(
            recovered,
            record_count=1,
            schema=schema,
            records=records,
        )
        == []
    )


def test_resolved_target_consistency_recovery_rejects_other_contract_errors() -> None:
    schema = pipeline_tools._resolved_target_consistency_schema(1)
    payload = {
        "status": "consistent",
        "record_findings": [],
        "contradiction_record_indexes": [],
        "question": None,
        "reason": None,
    }
    records = [
        {
            "record_index": 0,
            "resolved_target": {"evidence_capabilities": ["ordinal_identity"]},
        }
    ]

    recovered = pipeline_tools._recover_resolved_target_consistency_clarification(
        payload,
        records,
        schema,
    )

    assert recovered is None


def test_resolved_target_consistency_contract_accepts_supported_relational_claim() -> None:
    schema = pipeline_tools._resolved_target_consistency_schema(1)
    payload = {
        "status": "consistent",
        "record_findings": [
            {
                "record_index": 0,
                "identity_requirement": "The requested target has an explicit mapping.",
                "requirement_kind": "substitution_exclusion",
                "status": "satisfied",
                "evidence_assessment": "The resolver supplied an explicit positive mapping.",
            }
        ],
        "contradiction_record_indexes": [],
        "question": None,
        "reason": None,
    }
    records = [
        {
            "record_index": 0,
            "resolved_target": {
                "evidence_capabilities": [
                    "ordinal_identity",
                    "explicit_positive_mapping",
                ]
            },
        }
    ]

    errors = pipeline_tools._resolved_target_consistency_contract_errors(
        payload,
        record_count=1,
        schema=schema,
        records=records,
    )

    assert errors == []


def test_resolved_target_consistency_verification_contract_requires_evidence() -> None:
    schema = pipeline_tools._resolved_target_consistency_verification_schema(1)
    errors = pipeline_tools._resolved_target_consistency_verification_contract_errors(
        {
            "status": "not_verified",
            "record_findings": [
                {
                    "record_index": 0,
                    "identity_requirement": "The intended target must be distinguished.",
                    "requirement_kind": "substitution_exclusion",
                    "status": "not_verified",
                    "evidence_assessment": "No positive mapping was supplied.",
                }
            ],
            "affected_record_indexes": [],
            "reason": "The intended feature is not verifiable.",
        },
        record_count=1,
        schema=schema,
    )

    assert "not_verified must select at least one materialization record" in errors


def test_resolved_target_consistency_verifier_can_overrule_consistent_review(
    tmp_path: Path,
    monkeypatch: Any,
) -> None:
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    goal = {
        "gene_symbol": "GENE1",
        "species": "human",
        "target_context": "reviewed cells",
        "intent": "splice_modulation",
        "run_splice_router": True,
        "run_splice_eval": True,
        "design_programs": [
            {
                "program_id": "site_a",
                "label": "site a",
                "modality": "steric",
                "intent": "splice_modulation",
                "objective": "increase the requested feature",
                "desired_delta": "increase",
            }
        ],
        "materialized_region_constraints": {
            "global": [
                {
                    "constraint_id": "window_a",
                    "source_type": "splice_relative_window",
                    "splice_relative_window": {
                        "reference_transcript_id": "TX0001",
                        "exon_label": 2,
                        "anchor_site": "acceptor",
                        "relative_start_nt": -10,
                        "relative_end_nt": 10,
                    },
                    "materialized_rows": [["chr2", 289, 310, "window", 0, "+"]],
                    "materialization_summary": {
                        "status": "materialized",
                        "transcript_id": "TX0001",
                        "resolved_exon_label": 2,
                        "anchor_site": "acceptor",
                        "anchor_position": 300,
                    },
                }
            ],
            "by_program": {},
        },
    }
    final_payload = {"structured_goal": deepcopy(goal)}
    raw_payload = {"final_payload": final_payload}
    pipeline_tools._save_goal_payload(str(run_dir), raw_payload, goal, final_payload)
    state = _state_from_raw_envelope(
        {
            "request_text": "Use this feature and exclude a separate unresolved alternative.",
            "tasks": [
                {
                    "task_id": "design",
                    "task_type": "design_new_asos",
                    "label": "design",
                    "context": {"structured_goal": deepcopy(goal)},
                }
            ],
        },
        tmp_path,
    )
    state["run_dir"] = str(run_dir)
    state["current_task_id"] = "design"
    state["task_statuses"] = {"design": {"status": "in_progress"}}
    state_path = tmp_path / "run_state.json"
    state_path.write_text(json.dumps(state), encoding="utf-8")
    monkeypatch.setattr(pipeline_tools, "provider_credentials_available", lambda: True)
    monkeypatch.setattr(
        pipeline_tools,
        "_run_resolved_target_consistency_sync",
        lambda _payload: {
            "status": "consistent",
            "contradiction_record_indexes": [],
            "question": None,
            "reason": None,
        },
    )
    monkeypatch.setattr(
        pipeline_tools,
        "_run_resolved_target_consistency_verification_sync",
        lambda _payload: {
            "status": "not_verified",
            "affected_record_indexes": [0],
            "reason": "The explicit exclusion cannot be evaluated from supplied evidence.",
        },
    )
    monkeypatch.setattr(
        pipeline_tools,
        "_run_resolved_target_consistency_adjudication_sync",
        lambda _payload: {
            "status": "needs_clarification",
            "contradiction_record_indexes": [0],
            "question": "What exact positive feature mapping should be used?",
            "reason": "The explicit exclusion still cannot be evaluated from supplied evidence.",
        },
    )

    result = pipeline_tools._review_materialized_target_consistency(
        state=state,
        state_path=str(state_path),
        run_dir=str(run_dir),
        payload=raw_payload,
        structured_goal=goal,
        final_payload=final_payload,
    )
    saved = json.loads(state_path.read_text(encoding="utf-8"))

    assert result is not None and result["needs_user_input"] is True
    assert saved["clarification_queue"][0]["missing_fields"] == [
        "target_feature_mapping_missing"
    ]
    review = goal["resolved_target_consistency_review"]
    assert review["status"] == "needs_clarification"
    assert review["initial_review"]["status"] == "consistent"
    assert review["verification"]["status"] == "not_verified"
    assert review["adjudication"]["status"] == "needs_clarification"


def test_resolved_target_consistency_adjudication_can_reject_false_challenge(
    tmp_path: Path,
    monkeypatch: Any,
) -> None:
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    goal = {
        "gene_symbol": "GENE1",
        "species": "human",
        "target_context": "reviewed cells",
        "intent": "splice_modulation",
        "run_splice_router": True,
        "run_splice_eval": True,
        "design_programs": [
            {
                "program_id": "site_a",
                "label": "named window",
                "modality": "steric",
                "intent": "splice_modulation",
                "objective": "target the named window",
                "desired_delta": "increase",
            }
        ],
        "materialized_region_constraints": {
            "global": [
                {
                    "constraint_id": "window_a",
                    "source_type": "splice_relative_window",
                    "splice_relative_window": {
                        "reference_transcript_id": "TX0001",
                        "exon_label": 2,
                        "anchor_site": "donor",
                        "relative_start_nt": 10,
                        "relative_end_nt": 35,
                        "feature_name": "named window",
                    },
                    "materialized_rows": [["chr2", 309, 335, "named window", 0, "+"]],
                    "materialization_summary": {
                        "status": "materialized",
                        "transcript_id": "TX0001",
                        "resolved_exon_label": 2,
                        "anchor_site": "donor",
                        "anchor_position": 300,
                    },
                }
            ],
            "by_program": {},
        },
    }
    final_payload = {"structured_goal": deepcopy(goal)}
    raw_payload = {"final_payload": final_payload}
    pipeline_tools._save_goal_payload(str(run_dir), raw_payload, goal, final_payload)
    state = _state_from_raw_envelope(
        {
            "request_text": "Define the named window as donor +10 through +35.",
            "tasks": [
                {
                    "task_id": "design",
                    "task_type": "design_new_asos",
                    "label": "design",
                    "context": {"structured_goal": deepcopy(goal)},
                }
            ],
        },
        tmp_path,
    )
    state["run_dir"] = str(run_dir)
    state["current_task_id"] = "design"
    state["task_statuses"] = {"design": {"status": "in_progress"}}
    state_path = tmp_path / "run_state.json"
    state_path.write_text(json.dumps(state), encoding="utf-8")
    monkeypatch.setattr(pipeline_tools, "provider_credentials_available", lambda: True)
    monkeypatch.setattr(
        pipeline_tools,
        "_run_resolved_target_consistency_sync",
        lambda _payload: {
            "status": "consistent",
            "contradiction_record_indexes": [],
            "question": None,
            "reason": None,
        },
    )
    monkeypatch.setattr(
        pipeline_tools,
        "_run_resolved_target_consistency_verification_sync",
        lambda _payload: {
            "status": "not_verified",
            "affected_record_indexes": [0],
            "reason": "No external record repeats the feature name.",
        },
    )
    monkeypatch.setattr(
        pipeline_tools,
        "_run_resolved_target_consistency_adjudication_sync",
        lambda _payload: {
            "status": "consistent",
            "contradiction_record_indexes": [],
            "question": None,
            "reason": None,
        },
    )

    result = pipeline_tools._review_materialized_target_consistency(
        state=state,
        state_path=str(state_path),
        run_dir=str(run_dir),
        payload=raw_payload,
        structured_goal=goal,
        final_payload=final_payload,
    )

    assert result is None
    review = goal["resolved_target_consistency_review"]
    assert review["status"] == "consistent"
    assert review["verification"]["status"] == "not_verified"
    assert review["adjudication"]["status"] == "consistent"


def test_per_site_clarification_accepts_upstream_downstream_site_phrasing() -> None:
    structured_goal = {
        "gene_symbol": "SMN2",
        "target_context": "spinal cord motor neurons",
        "candidate_limit": 30,
        "splice_isoform_objectives": [
            {
                "conversion_role": "to",
                "transcript_id": "ENST00000380743",
                "exon_label": "exon 7",
                "objective_kind": "increase_exon_inclusion",
            }
        ],
    }
    text = (
        "Use three distinct target sites/site classes: SMN2 exon 7 acceptor junction, "
        "SMN2 exon 7 donor junction, and the intronic ISS-N1 silencer region downstream of exon 7. "
        "Use transcript ENST00000380743 exon 7 numbering. "
        "For the exon 7 acceptor, tile ASOs fully contained from 25 nt upstream of the acceptor "
        "through 24 nt downstream of the acceptor. "
        "For the exon 7 donor, tile ASOs fully contained from 9 nt upstream of the donor "
        "through 60 nt downstream of the donor. "
        "For ISS-N1, tile ASOs overlapping the intronic silencer downstream of exon 7, "
        "using the donor-proximal +1 to +60 downstream region with any-overlap allowed for the ISS-N1 motif."
    )

    programs = pipeline_tools._request_splice_relative_window_programs_from_text(text, structured_goal)
    windows = [
        program["region_constraints"][0]["splice_relative_window"]
        for program in programs
    ]

    assert len(programs) == 3
    assert {
        (window["anchor_site"], window["relative_start_nt"], window["relative_end_nt"], window.get("feature_name"))
        for window in windows
    } == {
        ("acceptor", -25, 24, None),
        ("donor", -9, 60, None),
        ("donor", 1, 60, "ISS-N1"),
    }


def test_per_site_clarification_accepts_exon_start_end_and_downstream_feature_phrasing() -> None:
    structured_goal = {
        "gene_symbol": "SMN2",
        "target_context": "spinal cord motor neurons",
        "candidate_limit": 30,
        "splice_isoform_objectives": [
            {
                "conversion_role": "to",
                "transcript_id": "ENST00000380743",
                "exon_label": "exon 7",
                "objective_kind": "increase_exon_inclusion",
            }
        ],
    }
    text = (
        "Use three distinct target site classes: the exon 7 splice acceptor junction, "
        "the exon 7 splice donor junction, and the ISS-N1/downstream intronic splicing "
        "silencer region immediately downstream of exon 7. Use transcript-derived exon 7 "
        "windows on ENST00000380743: acceptor window from 50 nt upstream of the exon 7 "
        "start through 20 nt into exon 7; donor window from 20 nt upstream of the exon 7 "
        "end through 50 nt downstream; ISS-N1/downstream silencer window from +10 to +35 "
        "nt downstream of the exon 7 donor, with enough padding to tile 15-mers fully "
        "across that interval. Intronic pre-mRNA binding is allowed and expected for the "
        "donor/downstream ISS-N1 site."
    )

    programs = pipeline_tools._request_splice_relative_window_programs_from_text(text, structured_goal)
    windows = [
        program["region_constraints"][0]["splice_relative_window"]
        for program in programs
    ]
    overlap_by_feature = {
        window.get("feature_name") or window["anchor_site"]: program["region_constraints"][0]["aso_overlap_requirement"]
        for program, window in zip(programs, windows)
    }

    assert len(programs) == 3
    assert {
        (window["anchor_site"], window["relative_start_nt"], window["relative_end_nt"], window.get("feature_name"))
        for window in windows
    } == {
        ("acceptor", -50, 19, None),
        ("donor", -19, 50, None),
        ("donor", 10, 35, "ISS-N1"),
    }
    assert overlap_by_feature["ISS-N1"] == "any_overlap"
    assert all(program["candidate_limit"] == 30 for program in programs)


def test_scoped_parent_clarification_builds_splice_relative_programs() -> None:
    structured_goal = {
        "gene_symbol": "SMN2",
        "target_context": "spinal cord motor neurons",
        "candidate_limit": 30,
        "splice_isoform_objectives": [
            {
                "conversion_role": "to",
                "transcript_id": "ENST00000380743",
                "exon_label": "exon 7",
            }
        ],
        "design_programs": [{"program_id": "steric_smn2_isoform_switch", "modality": "steric"}],
    }
    scoped_request = (
        "Create a scoped ASO design run for exactly one task from a larger user request.\n\n"
        "Original user request:\n"
        "Design steric splice-switching ASOs for human SMN2 in spinal cord motor neurons "
        "to convert ENST00000626847 toward ENST00000380743 by increasing exon 7 inclusion. "
        "Return at least 30 ASOs per site.\n\n"
        "Scoped design task:\n"
        '{"task_id":"design_smn2_ex7_acceptor"}\n\n'
        "Parent-level clarification history (authoritative inherited inputs):\n"
        "Q: What should count as distinct target sites?\n"
        "- Option B: 30 ASOs per splice-adjacent region (e.g., exon 7 body, intron 6 upstream flank, intron 7 downstream flank)\n"
        "A: Use biologically distinct exon 7 inclusion-relevant steric target regions. "
        "Primary site: the intron 7 ISS-N1 / donor-proximal downstream silencer region. "
        "Also include exon 7 donor-proximal and acceptor-proximal junction regions if targetable.\n"
        "Q: For each site, what targeting window should be used?\n"
        "A: For the acceptor site, use 25 nt upstream in intron 6 plus the first 25 nt of exon 7. "
        "For the donor/ISS-N1 site, use the last 10 nt of exon 7 plus 60 nt downstream into intron 7. "
        "If ISS-N1 exact coordinates are available use motif span plus padding; otherwise use donor +1 to +60 intron 7 window.\n"
        "Q: Are ASOs allowed to bind intronic pre-mRNA sequence?\n"
        "A: Allow both exonic and intronic pre-mRNA binding.\n"
        "Apply every relevant answer to this scoped task."
    )
    final_payload = {}

    changed = pipeline_tools._apply_request_splice_relative_window_programs(
        {"request": scoped_request},
        structured_goal,
        final_payload,
    )

    assert changed is True
    windows = [
        program["region_constraints"][0]["splice_relative_window"]
        for program in structured_goal["design_programs"]
    ]
    assert {
        (window["anchor_site"], window["relative_start_nt"], window["relative_end_nt"], window.get("feature_name"))
        for window in windows
    } == {
        ("acceptor", -25, 24, None),
        ("donor", -9, 60, None),
        ("donor", 1, 60, "ISS-N1"),
    }


def test_clarified_site_windows_replace_generic_scoped_steric_program() -> None:
    structured_goal = {
        "gene_symbol": "SMN2",
        "target_context": "spinal cord motor neurons",
        "candidate_limit": 30,
        "splice_isoform_objectives": [
            {
                "conversion_role": "to",
                "transcript_id": "ENST00000380743",
                "exon_label": "exon 7",
                "objective_kind": "increase_exon_inclusion",
            }
        ],
        "region_constraints": [
            {
                "constraint_id": "constraint_01",
                "label": "ex7_acceptor_junction_window",
                "exon_label": 7,
                "source_type": "exon",
                "aso_overlap_requirement": "full_containment",
            },
            {
                "constraint_id": "constraint_02",
                "label": "ex7_donor_junction_window",
                "exon_label": 7,
                "source_type": "exon",
                "aso_overlap_requirement": "full_containment",
            },
            {
                "constraint_id": "constraint_03",
                "label": "iss_n1_window",
                "exon_label": 7,
                "source_type": "exon",
                "aso_overlap_requirement": "any_overlap",
            },
        ],
        "design_programs": [
            {
                "program_id": "steric_smn2_01",
                "modality": "steric",
                "region_constraints": [
                    {
                        "constraint_id": "constraint_01",
                        "label": "ex7_acceptor_junction_window",
                        "exon_label": 7,
                        "source_type": "exon",
                    },
                    {
                        "constraint_id": "constraint_02",
                        "label": "ex7_donor_junction_window",
                        "exon_label": 7,
                        "source_type": "exon",
                    },
                ],
            }
        ],
    }
    final_payload = {}
    state = {
        "request": "Design SMN2 ASOs. Return at least 30 ASOs per site.",
        "clarification_qa_history": [
            {
                "question": "What should count as distinct target sites?",
                "answer": "Use SMN2 exon 7 acceptor junction, exon 7 donor junction, and ISS-N1.",
            },
            {
                "question": "For each site, what targeting window should be used?",
                "answer": (
                    "For the exon 7 acceptor, tile ASOs fully contained from 25 nt upstream of the acceptor "
                    "through 24 nt downstream of the acceptor. For the exon 7 donor, tile ASOs fully contained "
                    "from 9 nt upstream of the donor through 60 nt downstream of the donor. For ISS-N1, use the "
                    "donor-proximal +1 to +60 downstream region with any-overlap allowed for the ISS-N1 motif."
                ),
            },
        ],
    }

    changed = pipeline_tools._apply_request_splice_relative_window_programs(
        state,
        structured_goal,
        final_payload,
    )

    assert changed is True
    assert len(structured_goal["design_programs"]) == 3
    assert all(
        program["targeting_policy"] == "explicit_user_site_window"
        for program in structured_goal["design_programs"]
    )
    assert all(
        "splice_relative_window" in program["region_constraints"][0]
        for program in structured_goal["design_programs"]
    )
    assert "region_constraints" not in structured_goal
    assert len(structured_goal["superseded_per_site_region_constraints"]) == 3


def test_save_goal_payload_does_not_repromote_superseded_per_site_constraints(tmp_path: Path) -> None:
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    structured_goal = {
        "gene_symbol": "SMN2",
        "per_site_window_program_preseed": {"status": "applied"},
        "design_programs": [
            {
                "program_id": "steric_exon7_acceptor_window",
                "modality": "steric",
                "region_constraints": [
                    {
                        "label": "exon7_acceptor_window",
                        "aso_overlap_requirement": "full_containment",
                        "splice_relative_window": {
                            "exon_label": 7,
                            "anchor_site": "acceptor",
                            "relative_start_nt": -25,
                            "relative_end_nt": 24,
                        },
                    }
                ],
            },
            {
                "program_id": "steric_exon7_iss_n1_window",
                "modality": "steric",
                "region_constraints": [
                    {
                        "label": "exon7_iss_n1_window",
                        "aso_overlap_requirement": "any_overlap",
                        "splice_relative_window": {
                            "exon_label": 7,
                            "anchor_site": "donor",
                            "relative_start_nt": 1,
                            "relative_end_nt": 60,
                        },
                    }
                ],
            },
        ],
    }
    aggregate_constraints = [
        {"label": "acceptor", "source_type": "exon", "aso_overlap_requirement": "full_containment"},
        {"label": "ISS-N1", "source_type": "exon", "aso_overlap_requirement": "any_overlap"},
    ]
    payload = {
        "request": "Design SMN2 ASOs.",
        "request_envelope": {
            "tasks": [
                {
                    "context": {
                        "structured_goal": {
                            "region_constraints": aggregate_constraints,
                        }
                    }
                }
            ]
        },
    }
    final_payload: dict[str, Any] = {}

    pipeline_tools._save_goal_payload(str(run_dir), payload, structured_goal, final_payload)
    saved_goal = json.loads(resolve_run_path(run_dir, "structured_goal.json").read_text(encoding="utf-8"))

    assert "region_constraints" not in saved_goal


def test_requested_aso_length_overrides_chemistry_model_length() -> None:
    structured_goal = {
        "gene_symbol": "SMN2",
        "intent": "splice_modulation",
        "aso_length_nt": 15,
        "design_programs": [
            {
                "program_id": "steric_exon7_donor_window",
                "modality": "steric",
            }
        ],
    }
    final_payload = {
        "chemistry_profiles": {
            "by_program": {
                "steric_exon7_donor_window": {
                    "recommendation": {
                        "CHEMISTRY": "steric-hindrance",
                        "platform": "steric blocker",
                        "length": 20,
                    },
                    "source": "chemistry_agent",
                }
            }
        }
    }

    profiles = pipeline_tools._ensure_program_chemistry_profiles(
        structured_goal=structured_goal,
        final_payload=final_payload,
        modality="steric",
    )

    payload = profiles["by_program"]["steric_exon7_donor_window"]
    assert payload["recommendation"]["length"] == 15
    assert payload["requested_length_override"]["model_recommended_length_nt"] == 20


def test_design_task_contract_preserves_isoform_conversion_ids_in_authoritative_goal() -> None:
    task = {
        "task_id": "design_smn2_ssso",
        "task_type": "design_new_asos",
        "context": {
            "gene_symbol": "SMN2",
            "species": "human",
            "structured_goal": {
                "gene_symbol": "SMN2",
                "species": "human",
                "intent": "splice_modulation",
                "splice_isoform_objectives": [
                    {
                        "event": "exon_inclusion",
                        "direction": "increase_inclusion",
                        "exon_label": "exon 7",
                        "exon_numbering_transcript_id": "ENST00000380743",
                        "from_transcript_id": "ENST00000626847",
                        "to_transcript_id": "ENST00000380743",
                    }
                ],
            },
        },
    }

    finalized = pipeline_tools._finalize_design_task_contract(task, {})
    objectives = finalized["context"]["structured_goal"]["splice_isoform_objectives"]
    by_tx = {obj["transcript_id"]: obj for obj in objectives}

    assert by_tx["ENST00000626847"]["desired_abundance"] == "decrease"
    assert by_tx["ENST00000380743"]["desired_abundance"] == "increase"
    assert "splice_isoform_objectives" not in finalized["context"]


def test_design_task_contract_keeps_evaluation_targets_only_in_authoritative_goal() -> None:
    targets = [
        {
            "label": "Requested transcript",
            "target_type": "transcript",
            "objective_type": "primary_outcome",
            "desired_delta": "decrease",
        }
    ]
    task = {
        "task_id": "design",
        "task_type": "design_new_asos",
        "context": {
            "gene_symbol": "GENE1",
            "species": "human",
            "evaluation_targets": targets,
            "structured_goal": {
                "gene_symbol": "GENE1",
                "species": "human",
                "design_programs": [
                    {
                        "program_id": "program_1",
                        "label": "Create candidates",
                        "modality": "gapmer",
                        "intent": "knockdown",
                        "objective": "Decrease the requested transcript.",
                        "desired_delta": "decrease",
                    }
                ],
            },
        },
    }

    finalized = pipeline_tools._finalize_design_task_contract(task, {})

    assert finalized["context"]["structured_goal"]["evaluation_targets"] == targets
    assert "evaluation_targets" not in finalized["context"]


def test_design_task_contract_does_not_project_target_effect_across_program_scopes() -> None:
    task = {
        "task_id": "design",
        "task_type": "design_new_asos",
        "context": {
            "target_context": "supplied cells",
            "structured_goal": {
                "target_context": "supplied cells",
                "evaluation_targets": [
                    {
                        "exon_label": 6,
                        "desired_effect": "skipping",
                        "desired_delta": "increase",
                        "objective_type": "route_target",
                        "target_type": "exon",
                    }
                ],
                "design_programs": [
                    {
                        "program_id": "splice_branch",
                        "label": "splice branch",
                        "modality": "steric",
                        "intent": "splice_modulation",
                        "objective": "change the requested splice outcome",
                        "desired_delta": "increase",
                        "desired_effect": "skipping",
                    },
                    {
                        "program_id": "cleavage_branch",
                        "label": "cleavage branch",
                        "modality": "gapmer",
                        "intent": "knockdown",
                        "objective": "decrease the requested RNA class",
                        "desired_delta": "decrease",
                    },
                ],
            },
        },
    }

    finalized = pipeline_tools._finalize_design_task_contract(task, {})
    programs = finalized["context"]["structured_goal"]["design_programs"]

    assert programs[0]["desired_effect"] == "skipping"
    assert "desired_effect" not in programs[1]


def test_design_task_contract_uses_structured_species_only_when_primary_is_absent() -> None:
    task = {
        "task_id": "design",
        "task_type": "design_new_asos",
        "context": {
            "structured_goal": {
                "species_targets": [
                    {"species": "mouse"},
                    {"species": "human"},
                ],
                "cross_species_requirement": {
                    "required": True,
                    "species": ["mouse", "human"],
                },
                "design_programs": [
                    {
                        "program_id": "conserved_design",
                        "label": "conserved design",
                        "modality": "steric",
                        "intent": "binding",
                        "objective": "Design against the shared target.",
                        "desired_delta": "decrease",
                    }
                ],
            },
        },
    }

    finalized = pipeline_tools._finalize_design_task_contract(task, {})
    context = finalized["context"]
    goal = context["structured_goal"]

    assert context["species"] == "mouse"
    assert goal["species"] == "mouse"
    assert goal["cross_species_requirement"]["species"] == ["mouse", "human"]
    assert [item["species"] for item in goal["species_targets"]] == [
        "mouse",
        "human",
    ]


def test_design_task_contract_preserves_valid_cross_species_anchor() -> None:
    task = {
        "task_id": "design",
        "task_type": "design_new_asos",
        "context": {
            "species": "human",
            "structured_goal": {
                "species": "human",
                "species_targets": [
                    {"species": "mouse"},
                    {"species": "human"},
                ],
                "cross_species_requirement": {
                    "required": True,
                    "species": ["mouse", "human"],
                },
                "design_programs": [],
            },
        },
    }

    finalized = pipeline_tools._finalize_design_task_contract(task, {})

    assert finalized["context"]["species"] == "human"
    assert finalized["context"]["structured_goal"]["species"] == "human"


def test_clarification_coordinates_are_preserved_as_explicit_splice_targets() -> None:
    state = {
        "request": (
            "Design steric splice-switching ASOs for PCCA to decrease poison exon 14x "
            "isoform ENST00000376285 by promoting exon skipping."
        ),
        "request_envelope": {
            "default_context": {
                "gene_symbol": "PCCA",
                "species": "human",
                "splice_isoform_objectives": [
                    {
                        "isoform_label": "ENST00000376285",
                        "transcript_id": "ENST00000376285",
                        "desired_abundance": "decrease",
                    }
                ],
            }
        },
    }
    preseed = pipeline_tools._explicit_splice_target_preseed_from_clarifications(
        state,
        [
            {
                "question": "Which coordinates define the poison exon?",
                "answer": (
                    "hg38 chr13:100305751-100305834, strand +. This is PCCA poison exon 14x; "
                    "design for decreased exon inclusion."
                ),
                "evidence_refs": ["reviewed_mapping.json"],
            }
        ],
    )

    overlay = preseed["structured_goal_overlay"]
    target = overlay["explicit_splice_targets"][0]
    assert target["chrom"] == "chr13"
    assert target["start"] == 100305751
    assert target["end"] == 100305834
    assert target["strand"] == "+"
    assert target["transcript_id"] == "ENST00000376285"
    assert target["desired_delta"] == "decrease"
    assert target["requested_exon_label"] == "14x"
    assert target["evidence_refs"] == ["reviewed_mapping.json"]
    assert overlay["evaluation_targets"][0]["evidence_refs"] == [
        "reviewed_mapping.json"
    ]
    assert "exon_label" not in target
    assert overlay["transcript_lock"]["value"] == "ENST00000376285"


def test_binding_interval_clarifications_do_not_create_explicit_splice_targets() -> None:
    state = {
        "request": (
            "Design 20-nt steric-blocking target-site blocker ASOs for human CDKN1B "
            "to block PUM1 eCLIP and miR-221/miR-222 seed windows in the 3' UTR."
        ),
        "request_envelope": {
            "default_context": {
                "gene_symbol": "CDKN1B",
                "species": "human",
                "intent": "binding",
                "target_region": "utr3",
            }
        },
    }
    preseed = pipeline_tools._explicit_splice_target_preseed_from_clarifications(
        state,
        [
            {
                "question": "Which sites should be targeted?",
                "answer": (
                    "PUM1/eCLIP site 1: chr12:12721313-12721416 (+). "
                    "miR-221/miR-222 seed window 1: chr12:12721208-12721253 (+)."
                ),
            }
        ],
    )

    assert preseed == {}


def test_binding_interval_request_text_does_not_preseed_splice_targets() -> None:
    state = {
        "request": (
            "Design 20-nt steric-blocking target-site blocker ASOs for human CDKN1B "
            "to increase CDKN1B/p27 by blocking negative sites in the 3' UTR. "
            "PUM1/eCLIP site 1: chr12:12721313-12721416 (+). "
            "miR-221/miR-222 seed window 1: chr12:12721208-12721253 (+)."
        )
    }
    goal = {
        "gene_symbol": "CDKN1B",
        "species": "human",
        "intent": "binding",
        "target_region": "utr3",
        "design_programs": [
            {
                "program_id": "steric_cdkn1b_01",
                "modality": "steric",
                "intent": "binding",
            }
        ],
    }

    assert pipeline_tools._request_coordinate_preseed_if_available(state, goal) == {}


def test_nonstandard_splice_exon_label_requires_coordinates_at_intake() -> None:
    intake = run_goal_intake._normalize_intake_output(
        {
            "status": "ready",
            "structured_goal": {
                "intent": "splice_modulation",
                "gene_symbol": "PCCA",
                "species": "human",
                "splice_isoform_objectives": [
                    {
                        "isoform_label": "ENST00000376285",
                        "transcript_id": "ENST00000376285",
                        "exon_label": "14x",
                        "objective_kind": "exon_skipping",
                    }
                ],
            },
        }
    )

    assert intake["status"] == "need_clarification"
    assert intake["reason"] == "nonstandard_splice_exon_coordinates_missing"
    assert "14x" in intake["questions"][0]


def test_nonstandard_exon_guard_ignores_generated_rationale_and_labels() -> None:
    goal = {
        "intent": "splice_modulation",
        "gene_symbol": "GENE1",
        "species": "human",
        "exon_label": 10,
        "splice_isoform_objectives": [
            {
                "isoform_label": "productive isoform",
                "exon_label": 10,
                "notes": "The annotation rationale also discusses exon 4a.",
            }
        ],
        "design_programs": [
            {
                "label": "Avoid changing exon 4a",
                "objective": "Modulate the typed numeric exon.",
                "exon_label": 10,
            }
        ],
    }

    assert pipeline_tools._nonstandard_splice_exon_labels_from_goal(goal) == []
    intake = run_goal_intake._normalize_intake_output(
        {"status": "ready", "structured_goal": goal}
    )
    assert intake["status"] == "ready"


def test_controller_requests_coordinates_for_nonstandard_exon_in_request_text() -> None:
    state = {
        "request": (
            "Design steric splice-switching ASOs for human PCCA to decrease the poison exon 14x "
            "isoform ENST00000376285 by promoting exon 14x skipping."
        )
    }
    goal = {
        "intent": "splice_modulation",
        "gene_symbol": "PCCA",
        "species": "human",
        "splice_isoform_objectives": [
            {
                "isoform_label": "poison exon 14x isoform",
                "transcript_id": "ENST00000376285",
                "desired_abundance": "decrease",
            }
        ],
    }

    payload = pipeline_tools._nonstandard_splice_coordinate_clarification_payload(state, goal)

    assert payload is not None
    assert payload["reason"] == "nonstandard_splice_exon_coordinates_missing"
    assert "14x" in payload["questions"][0]
    assert "poison exon 14x isoform" not in payload["questions"][0]


def test_controller_extracts_request_coordinates_into_preseed() -> None:
    state = {
        "request": (
            "Design steric splice-switching ASOs for PCCA exon 14x at "
            "hg38 chr13:100305751-100305834, strand + in ENST00000376285."
        )
    }
    goal = {
        "intent": "splice_modulation",
        "gene_symbol": "PCCA",
        "species": "human",
        "splice_isoform_objectives": [
            {
                "isoform_label": "ENST00000376285",
                "transcript_id": "ENST00000376285",
                "desired_abundance": "decrease",
            }
        ],
    }

    preseed = pipeline_tools._request_coordinate_preseed_if_available(state, goal)

    target = preseed["structured_goal_overlay"]["explicit_splice_targets"][0]
    assert target["start"] == 100305751
    assert target["end"] == 100305834
    assert target["transcript_id"] == "ENST00000376285"
    assert target["source"] == "clarification_answers"


def test_scoped_child_does_not_promote_upstream_coordinates_to_user_targets() -> None:
    scoped_request = (
        "Create a scoped ASO design run for exactly one task from a larger user request.\n"
        "Ignore sibling tasks, unrelated genes, and unrelated species.\n\n"
        "Scoped design task:\n"
        '{"task_id":"design_splice","context":{"gene_symbol":"PCCA"}}\n\n'
        "Completed upstream dependency results (authoritative source-backed execution inputs):\n"
        "Reference analysis describes exon 14x at hg38 chr13:100305751-100305834.\n\n"
        "Parent-level clarification history (authoritative inherited inputs):\n"
        "1. Q: What design should I perform?\n"
        "   A: Design steric splice-switching ASOs for PCCA; determine the splice event from the reference.\n"
        "Apply every relevant answer to this scoped task.\n"
    )
    goal = {"intent": "splice_modulation", "gene_symbol": "PCCA", "species": "human"}

    assert pipeline_tools._user_authored_request_text(scoped_request) == ""
    assert pipeline_tools._request_coordinate_preseed_if_available(
        {"request": scoped_request}, goal
    ) == {}


def test_scoped_child_keeps_parent_answer_coordinates_as_user_targets() -> None:
    scoped_request = (
        "Create a scoped ASO design run for exactly one task from a larger user request.\n\n"
        "Scoped design task:\n"
        '{"task_id":"design_splice","context":{"gene_symbol":"PCCA"}}\n\n'
        "Completed upstream dependency results (authoritative source-backed execution inputs):\n"
        "Analysis mentions a separate reference exon at chr13:100400000-100400100.\n\n"
        "Parent-level clarification history (authoritative inherited inputs):\n"
        "1. Q: Which region did you specify?\n"
        "   A: Use PCCA exon 14x at hg38 chr13:100305751-100305834, strand + in ENST00000376285.\n"
        "Apply every relevant answer to this scoped task.\n"
    )
    goal = {"intent": "splice_modulation", "gene_symbol": "PCCA", "species": "human"}

    preseed = pipeline_tools._request_coordinate_preseed_if_available(
        {"request": scoped_request}, goal
    )
    targets = preseed["structured_goal_overlay"]["explicit_splice_targets"]
    assert len(targets) == 1
    assert (targets[0]["start"], targets[0]["end"]) == (100305751, 100305834)


def test_nonstandard_splice_exon_label_with_span_passes_intake_guard() -> None:
    intake = run_goal_intake._normalize_intake_output(
        {
            "status": "ready",
            "structured_goal": {
                "intent": "splice_modulation",
                "gene_symbol": "PCCA",
                "species": "human",
                "evaluation_targets": [
                    {
                        "transcript_id": "ENST00000376285",
                        "chrom": "chr13",
                        "strand": "+",
                        "exon_label": "14x",
                        "start": 100305751,
                        "end": 100305834,
                    }
                ],
                "splice_isoform_objectives": [
                    {
                        "isoform_label": "ENST00000376285",
                        "transcript_id": "ENST00000376285",
                        "exon_label": "14x",
                        "objective_kind": "exon_skipping",
                    }
                ],
            },
        }
    )

    assert intake["status"] == "ready"


def test_evaluation_span_overrides_numeric_exon_focus_for_splice_plan() -> None:
    goal = {
        "gene_symbol": "PCCA",
        "species": "human",
        "intent": "splice_modulation",
        "focus_exons_only": True,
        "exon_label": 14,
        "evaluation_targets": [
            {
                "transcript_id": "ENST00000376285",
                "chrom": "chr13",
                "strand": "+",
                "exon_label": 14,
                "start": 100305751,
                "end": 100305834,
            }
        ],
        "splice_isoform_objectives": [
            {
                "isoform_label": "ENST00000376285",
                "transcript_id": "ENST00000376285",
                "desired_abundance": "decrease",
                "exon_label": "14x (poison exon)",
            }
        ],
        "isoform_transcripts": {
            "ENST00000376285": {
                "transcript_id": "ENST00000376285",
                "metadata": {"id": "ENST00000376285", "seq_region_name": "13", "strand": 1},
                "chrom": "chr13",
                "strand": "+",
                "sites": [
                    {"chrom": "chr13", "acceptor": 100302924, "donor": 100302998, "exon_index": 14, "strand": "+", "coding": True}
                ],
                "exons_tx": [
                    {"index": 13, "start": 100302500, "end": 100302600, "seq_region_name": "13", "contains_coding": True},
                    {"index": 14, "start": 100302924, "end": 100302998, "seq_region_name": "13", "contains_coding": True},
                    {"index": 15, "start": 100306200, "end": 100306300, "seq_region_name": "13", "contains_coding": True},
                ],
            }
        },
    }
    final_payload: dict[str, Any] = {}

    context = run_goal_intake._ensure_explicit_splice_target_context(goal, final_payload)
    plan, _comparisons = run_goal_intake._compute_splice_site_plan(
        goal["splice_isoform_objectives"],
        goal["isoform_transcripts"],
        goal,
    )

    labels = sorted(f"{entry['site']}@{entry['position']}" for entry in plan)
    assert context["target_count"] == 1
    assert goal["explicit_splice_targets_source"] == "evaluation_targets"
    assert labels == ["acceptor@100305751", "donor@100305834"]
    assert all(entry["delta_direction"] == "decrease" for entry in plan)
    assert all(entry["source"] == "explicit_splice_target" for entry in plan)


def test_explicit_coordinate_splice_target_seeds_resolution_without_transcript_scope(monkeypatch) -> None:
    class Runtime:
        _compute_splice_site_plan = staticmethod(run_goal_intake._compute_splice_site_plan)
        _compute_steric_regions = staticmethod(run_goal_intake._compute_steric_regions)

    goal = {
        "gene_symbol": "PCCA",
        "species": "human",
        "intent": "splice_modulation",
        "target_context": "human hepatocyte-like cells",
        "explicit_splice_targets": [
            {
                "chrom": "chr13",
                "start": 100305751,
                "end": 100305834,
                "strand": "+",
                "desired_delta": "decrease",
                "label": "PCCA poison exon",
            }
        ],
        "explicit_splice_targets_source": "request_text",
        "splice_isoform_objectives": [
            {
                "objective": "skip_exon",
                "event_label": "PCCA poison exon skipping",
            }
        ],
    }
    final_payload: dict[str, Any] = {}

    result = pipeline_tools._seed_explicit_coordinate_splice_resolution(
        goal,
        final_payload,
        Runtime(),
    )

    labels = sorted(f"{entry['site']}@{entry['position']}" for entry in result["splice_site_plan"])
    assert labels == ["acceptor@100305751", "donor@100305834"]
    assert all(entry["delta_direction"] == "decrease" for entry in result["splice_site_plan"])
    assert result["resolved_transcript_strategy"]["requires_transcript_scope"] is False
    assert final_payload["splice_site_plan"] == result["splice_site_plan"]
    assert final_payload["steric_regions"]["global_region"]["start"] == 100305551
    assert final_payload["steric_regions"]["global_region"]["end"] == 100306034


def test_exploratory_windows_use_repaired_route_signal_plan(monkeypatch) -> None:
    class Runtime:
        @staticmethod
        def _compute_splice_site_plan(objectives, isoform_transcripts, structured_goal):
            return (
                [
                    {
                        "chrom": "chr1",
                        "position": 100,
                        "site": "donor",
                        "strand": "+",
                        "label": "donor@100",
                        "delta_direction": "decrease",
                        "present_in": ["pathogenic"],
                        "absent_in": ["beneficial"],
                        "covers_all_isoforms": False,
                    },
                    {
                        "chrom": "chr1",
                        "position": 240,
                        "site": "acceptor",
                        "strand": "+",
                        "label": "acceptor@240",
                        "delta_direction": "increase",
                        "present_in": ["beneficial"],
                        "absent_in": ["pathogenic"],
                        "covers_all_isoforms": False,
                    },
                    {
                        "chrom": "chr1",
                        "position": 360,
                        "site": "acceptor",
                        "strand": "+",
                        "label": "acceptor@360",
                        "delta_direction": "context",
                        "delta_per_isoform": {
                            "pathogenic": "decrease",
                            "beneficial": "increase",
                        },
                        "present_in": ["pathogenic", "beneficial"],
                        "absent_in": ["optional"],
                        "covers_all_isoforms": False,
                    },
                    {
                        "chrom": "chr1",
                        "position": 500,
                        "site": "donor",
                        "strand": "+",
                        "label": "donor@500",
                        "delta_direction": None,
                        "present_in": ["pathogenic", "beneficial"],
                        "absent_in": [],
                        "covers_all_isoforms": True,
                    },
                ],
                [],
            )

    monkeypatch.setattr(pipeline_tools, "_goal_runtime_module", lambda: Runtime())
    goal = {
        "gene_symbol": "TEST",
        "intent": "splice_modulation",
        "splice_site_plan": [],
        "splice_isoform_objectives": [
            {"isoform_label": "pathogenic", "desired_abundance": "decrease"},
            {"isoform_label": "beneficial", "desired_abundance": "increase"},
            {"isoform_label": "optional", "desired_abundance": "increase"},
        ],
        "isoform_transcripts": {
            "pathogenic": {
                "sites": [
                    {"chrom": "chr1", "acceptor": 40, "donor": 100, "strand": "+"},
                    {"chrom": "chr1", "acceptor": 360, "donor": 400, "strand": "+"},
                    {"chrom": "chr1", "acceptor": 450, "donor": 500, "strand": "+"},
                ]
            },
            "beneficial": {
                "sites": [
                    {"chrom": "chr1", "acceptor": 240, "donor": 300, "strand": "+"},
                    {"chrom": "chr1", "acceptor": 360, "donor": 400, "strand": "+"},
                    {"chrom": "chr1", "acceptor": 450, "donor": 500, "strand": "+"},
                ]
            },
            "optional": {
                "sites": [
                    {"chrom": "chr1", "acceptor": 450, "donor": 500, "strand": "+"},
                ]
            },
        },
    }

    programs, metadata = pipeline_tools._derive_exploratory_steric_programs(
        structured_goal=goal,
        base_program={
            "program_id": "steric_test",
            "modality": "steric",
            "targeting_policy": "default_exploratory",
            "run_splice_eval": True,
            "run_splice_router": True,
        },
    )

    assert metadata["source"] == "splice_site_plan"
    assert metadata["anchor_source"] == "objective_route_signal_sites"
    assert metadata["anchor_site_labels"] == ["donor@100", "acceptor@240", "acceptor@360"]
    assert programs[0]["routing_candidate_sites"] == ["donor@100", "acceptor@240", "acceptor@360"]


def test_exploratory_windows_union_steric_candidates_with_route_signal_sites() -> None:
    goal = {
        "gene_symbol": "TEST",
        "intent": "splice_modulation",
        "splice_isoform_objectives": [
            {"isoform_label": "pathogenic", "desired_abundance": "decrease"},
            {"isoform_label": "beneficial", "desired_abundance": "increase"},
            {"isoform_label": "optional", "desired_abundance": "increase"},
        ],
        "steric_site_candidates": [{"label": "donor@100", "site": "donor", "position": 100, "chrom": "chr1"}],
        "splice_site_plan": [
            {
                "chrom": "chr1",
                "position": 100,
                "site": "donor",
                "strand": "+",
                "label": "donor@100",
                "delta_direction": "decrease",
                "present_in": ["pathogenic"],
                "absent_in": ["beneficial", "optional"],
                "covers_all_isoforms": False,
            },
            {
                "chrom": "chr1",
                "position": 240,
                "site": "acceptor",
                "strand": "+",
                "label": "acceptor@240",
                "delta_direction": "increase",
                "present_in": ["beneficial"],
                "absent_in": ["pathogenic", "optional"],
                "covers_all_isoforms": False,
            },
            {
                "chrom": "chr1",
                "position": 360,
                "site": "acceptor",
                "strand": "+",
                "label": "acceptor@360",
                "delta_direction": "context",
                "delta_per_isoform": {
                    "pathogenic": "decrease",
                    "beneficial": "increase",
                },
                "present_in": ["pathogenic", "beneficial"],
                "absent_in": ["optional"],
                "covers_all_isoforms": False,
            },
        ],
    }

    programs, metadata = pipeline_tools._derive_exploratory_steric_programs(
        structured_goal=goal,
        base_program={
            "program_id": "steric_test",
            "modality": "steric",
            "targeting_policy": "default_exploratory",
            "run_splice_eval": True,
            "run_splice_router": True,
        },
    )

    assert metadata["anchor_source"] == "steric_site_candidates_plus_objective_route_signal_sites"
    assert metadata["anchor_site_labels"] == ["donor@100", "acceptor@240", "acceptor@360"]
    assert programs[0]["routing_candidate_sites"] == ["donor@100", "acceptor@240", "acceptor@360"]


def test_splice_switching_does_not_synthesize_gapmers_without_gapmer_request() -> None:
    goal = {
        "intent": "splice_modulation",
        "splice_isoform_objectives": [
            {"isoform_label": "pathogenic", "desired_abundance": "decrease", "role": "pathogenic"},
            {"isoform_label": "beneficial", "desired_abundance": "increase", "role": "beneficial"},
        ],
        "isoform_transcripts": {"pathogenic": {"sites": []}, "beneficial": {"sites": []}},
        "design_programs": [
            {
                "program_id": "steric_test",
                "modality": "steric",
                "intent": "splice_modulation",
            }
        ],
    }
    final_payload: dict[str, Any] = {}

    plan = pipeline_tools._ensure_splice_switch_gapmer_programs(goal, final_payload)

    assert plan["enabled"] is False
    assert plan["reason"] == "gapmer_not_requested"
    assert [program["modality"] for program in goal["design_programs"]] == ["steric"]
    assert "splice_switch_gapmer_plan" not in goal
    assert "splice_switch_gapmer_plan" not in final_payload


def test_conditional_gapmer_fallback_waits_for_failed_steric_routing(tmp_path: Path) -> None:
    goal = {
        "intent": "splice_modulation",
        "splice_isoform_objectives": [
            {
                "transcript_id": "ENST00000298910",
                "exon_label": "27",
                "objective_kind": "increase_skipping",
            }
        ],
        "design_programs": [
            {
                "program_id": "primary_steric_splice_switch",
                "modality": "steric",
                "intent": "splice_modulation",
            },
            {
                "program_id": "fallback_gapmer_knockdown",
                "modality": "gapmer",
                "priority": 2,
                "notes": "Secondary fallback objective: design gapmer ASOs if splice-switching is not viable.",
            },
        ],
    }
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    (run_dir / "goal_intake.json").write_text(
        json.dumps({"final_payload": {"structured_goal": goal}}),
        encoding="utf-8",
    )

    assert pipeline_tools._gapmer_programs_are_conditional_fallback_only(goal) is True
    assert pipeline_tools._resolve_intent_knockdown(str(run_dir)) is False

    (run_dir / "strategy_router_summary.json").write_text(
        json.dumps({"goals": [{"goal_label": "decrease acceptor@40303948", "count": 18}]}),
        encoding="utf-8",
    )
    assert pipeline_tools._resolve_intent_knockdown(str(run_dir)) is False

    (run_dir / "strategy_router_summary.json").write_text(
        json.dumps({"goals": [{"goal_label": "decrease acceptor@40303948", "count": 0}]}),
        encoding="utf-8",
    )
    assert pipeline_tools._resolve_intent_knockdown(str(run_dir)) is True


def test_routing_policy_and_route_artifact_preserve_declared_modalities(
    tmp_path: Path,
) -> None:
    goal = {
        "design_programs": [
            {
                "program_id": "primary_program",
                "modality": "steric",
                "execution_policy": {
                    "role": "primary",
                    "activation_condition": "always",
                },
            },
            {
                "program_id": "fallback_program",
                "modality": "gapmer",
                "execution_policy": {
                    "role": "fallback",
                    "activation_condition": "no_viable_primary_routes",
                },
            },
        ]
    }
    source = tmp_path / "source.route.json"
    source.write_text(
        json.dumps({"routes": [{"label": "route", "count": 2}]}),
        encoding="utf-8",
    )

    copied = pipeline_tools._copy_routing_artifact(
        str(source),
        tmp_path / "aggregate",
        "job",
        modality="steric",
    )
    policy = pipeline_tools._routing_policy_from_goal(
        goal,
        {"goals": [{"count": 2}]},
    )

    copied_payload = json.loads(Path(str(copied)).read_text(encoding="utf-8"))
    assert copied_payload["modality"] == "steric"
    assert copied_payload["routes"][0]["modality"] == "steric"
    assert policy == {
        "primary_modality": "steric",
        "conditional_fallback": {
            "modality": "gapmer",
            "activation_condition": "no_viable_primary_routes",
            "activated": False,
            "activation_reason": "viable_primary_routes_available",
        },
    }


def test_routing_job_does_not_repair_empty_splice_site_plan(monkeypatch) -> None:
    class Runtime:
        @staticmethod
        def _compute_splice_site_plan(objectives, isoform_transcripts, structured_goal):
            raise AssertionError("routing should not repair broad splice plans")

    monkeypatch.setattr(pipeline_tools, "_goal_runtime_module", lambda: Runtime())
    goal = {
        "splice_site_plan": [],
        "splice_isoform_objectives": [
            {"isoform_label": "progerin", "desired_abundance": "decrease"},
            {"isoform_label": "lamin_a", "desired_abundance": "increase"},
        ],
        "isoform_transcripts": {"progerin": {"sites": []}, "lamin_a": {"sites": []}},
    }

    scoped = pipeline_tools._goal_for_routing_job(
        structured_goal=goal,
        job={
            "job_id": "j1",
            "candidate_sites": ["donor@100"],
            "desired_delta": "decrease",
            "length": 17,
            "custom_constraint": {"kind": "preserve_me"},
        },
    )

    assert "routing_local_splice_site_plan" not in scoped
    assert scoped["routing_job"]["candidate_sites"] == ["donor@100"]
    assert scoped["routing_job"]["desired_delta"] == "decrease"
    assert scoped["routing_job"]["length"] == 17
    assert scoped["routing_job"]["custom_constraint"] == {"kind": "preserve_me"}


def test_splice_site_plan_direction_is_authoritative_for_site_routing() -> None:
    entry = {
        "site": "acceptor",
        "position": 100,
        "delta_direction": "increase",
    }

    assert (
        pipeline_tools._routing_context_expected_site_direction(
            entry,
            {"routing_job": {"desired_delta": "decrease"}},
        )
        == "increase"
    )


def test_unique_typed_route_target_scope_is_assigned_to_its_program() -> None:
    programs = [
        {
            "program_id": "productive_route",
            "desired_delta": "increase",
            "site_kind": "acceptor",
            "exon_label": 7,
        },
        {
            "program_id": "alternative_route",
            "desired_delta": "decrease",
            "site_kind": "acceptor",
            "exon_label": 7,
        },
    ]
    targets = [
        {
            "objective_type": "route_target",
            "desired_delta": "increase",
            "site_kind": "acceptor",
            "exon_label": 7,
            "isoform_label": "productive isoform",
        },
        {
            "objective_type": "route_target",
            "desired_delta": "decrease",
            "site_kind": "acceptor",
            "exon_label": 7,
            "isoform_label": "alternative isoform",
            "transcript_id": "ENST00000000001",
        },
    ]

    scoped = pipeline_tools._assign_unique_route_target_scopes(programs, targets)

    assert scoped[0]["isoform_label"] == "productive isoform"
    assert scoped[1]["isoform_label"] == "alternative isoform"
    assert scoped[1]["transcript_id"] == "ENST00000000001"
    assert "isoform_label" not in programs[0]


def test_ambiguous_route_target_scope_is_not_distributed_across_programs() -> None:
    programs = [
        {"program_id": "route_a", "desired_delta": "increase"},
        {"program_id": "route_b", "desired_delta": "increase"},
    ]
    targets = [
        {
            "objective_type": "route_target",
            "desired_delta": "increase",
            "isoform_label": "shared outcome",
        }
    ]

    scoped = pipeline_tools._assign_unique_route_target_scopes(programs, targets)

    assert all("isoform_label" not in program for program in scoped)


def test_routing_job_direction_is_used_when_site_direction_is_absent() -> None:
    assert (
        pipeline_tools._routing_context_expected_site_direction(
            {"site": "acceptor", "position": 100},
            {"routing_job": {"desired_delta": "decrease"}},
        )
        == "decrease"
    )


def test_scoped_decision_direction_survives_observation_reconciliation() -> None:
    assert (
        pipeline_tools._routing_context_expected_site_direction(
            {
                "site": "donor",
                "position": 100,
                "decision_direction": "decrease",
                "delta_direction": "increase",
            },
            {"routing_job": {"desired_delta": "increase"}},
        )
        == "decrease"
    )


def test_contextual_decision_direction_still_obeys_exon_effect_safety() -> None:
    assert (
        pipeline_tools._routing_context_expected_site_direction(
            {
                "site": "acceptor",
                "position": 100,
                "decision_direction": "decrease",
                "source_delta_direction": "context",
                "exon_indices": {"target": 7},
            },
            {"desired_effect": "inclusion", "exon_label": 7},
        )
        == "increase"
    )


def test_routing_job_uses_exact_isoform_direction_over_aggregate_site_direction() -> None:
    entry = {
        "site": "acceptor",
        "position": 100,
        "delta_direction": "increase",
        "delta_per_isoform": {
            "productive": "increase",
            "Pathogenic Isoform": "decrease",
        },
    }

    assert (
        pipeline_tools._routing_context_expected_site_direction(
            entry,
            {"routing_job": {"isoform_label": "pathogenic isoform"}},
        )
        == "decrease"
    )


def test_routing_job_materializes_scoped_decision_direction(monkeypatch) -> None:
    class Runtime:
        @staticmethod
        def _canonical_site_label(entry):
            return f"{entry.get('site')}@{entry.get('position')}"

    monkeypatch.setattr(pipeline_tools, "_goal_runtime_module", lambda: Runtime())
    goal = {
        "splice_site_plan": [
            {
                "site": "acceptor",
                "position": 100,
                "delta_direction": "increase",
                "delta_per_isoform": {
                    "productive": "increase",
                    "pathogenic": "decrease",
                },
            }
        ]
    }

    scoped = pipeline_tools._goal_for_routing_job(
        structured_goal=goal,
        job={
            "job_id": "pathogenic_route",
            "isoform_label": "pathogenic",
            "candidate_sites": ["acceptor@100"],
        },
    )

    assert scoped["routing_local_splice_site_plan"][0]["decision_direction"] == "decrease"
    assert "decision_direction" not in goal["splice_site_plan"][0]


def test_job_local_routing_direction_precedes_global_candidate_with_same_site(monkeypatch) -> None:
    class Runtime:
        @staticmethod
        def _canonical_site_label(entry):
            return f"{entry.get('site')}@{entry.get('position')}"

    monkeypatch.setattr(pipeline_tools, "_goal_runtime_module", lambda: Runtime())
    goal = {
        "routing_local_splice_site_plan": [
            {
                "site": "acceptor",
                "position": 100,
                "decision_direction": "decrease",
            }
        ],
        "steric_site_candidates": [
            {
                "site": "acceptor",
                "position": 100,
                "decision_direction": "increase",
            }
        ],
    }

    catalog = pipeline_tools._routing_catalog_entries(goal)
    anchors = pipeline_tools._routing_anchor_source_entries(goal)

    assert [entry["decision_direction"] for entry in catalog] == ["decrease"]
    assert [entry["decision_direction"] for entry in anchors] == ["decrease"]


def test_isoform_aware_local_plan_precedes_local_candidate_projection(monkeypatch) -> None:
    class Runtime:
        @staticmethod
        def _canonical_site_label(entry):
            return f"{entry.get('site')}@{entry.get('position')}"

    monkeypatch.setattr(pipeline_tools, "_goal_runtime_module", lambda: Runtime())
    goal = {
        "routing_local_splice_site_plan": [
            {
                "site": "acceptor",
                "position": 100,
                "decision_direction": "decrease",
                "delta_per_isoform": {"branch target": "decrease"},
            }
        ],
        "routing_local_steric_site_candidates": [
            {
                "site": "acceptor",
                "position": 100,
                "decision_direction": "increase",
            }
        ],
    }

    catalog = pipeline_tools._routing_catalog_entries(goal)
    anchors = pipeline_tools._routing_anchor_source_entries(goal)

    assert [entry["decision_direction"] for entry in catalog] == ["decrease"]
    assert [entry["decision_direction"] for entry in anchors] == ["decrease"]


def test_routing_job_scopes_existing_splice_site_plan_to_candidate_labels(monkeypatch) -> None:
    class Runtime:
        @staticmethod
        def _canonical_site_label(entry):
            return f"{entry.get('site')}@{entry.get('position')}"

    monkeypatch.setattr(pipeline_tools, "_goal_runtime_module", lambda: Runtime())
    goal = {
        "splice_site_plan": [
            {
                "chrom": "1",
                "position": 100,
                "site": "donor",
                "delta_direction": "decrease",
                "present_in": ["progerin"],
                "absent_in": ["lamin_a"],
            },
            {
                "chrom": "1",
                "position": 200,
                "site": "acceptor",
                "delta_direction": "increase",
                "present_in": ["lamin_a"],
                "absent_in": ["progerin"],
            },
        ],
    }

    scoped = pipeline_tools._goal_for_routing_job(
        structured_goal=goal,
        job={"job_id": "j1", "candidate_sites": ["donor@100"]},
    )

    assert [entry["position"] for entry in scoped["routing_local_splice_site_plan"]] == [100]
    assert scoped["routing_job"]["candidate_sites"] == ["donor@100"]


def test_routing_job_infers_site_kind_from_route_label_and_prefers_direct_targets(monkeypatch) -> None:
    class Runtime:
        @staticmethod
        def _canonical_site_label(entry):
            return f"{entry.get('site')}@{entry.get('position')}"

    monkeypatch.setattr(pipeline_tools, "_goal_runtime_module", lambda: Runtime())
    goal = {
        "splice_site_plan": [
            {
                "chrom": "17",
                "position": 100,
                "site": "acceptor",
                "delta_direction": "decrease",
                "unique_to_pathogenic": True,
            },
            {
                "chrom": "17",
                "position": 200,
                "site": "acceptor",
                "delta_direction": "context",
                "delta_per_isoform": {"pathogenic": "decrease", "productive": "increase"},
                "covers_all_isoforms": True,
            },
            {
                "chrom": "17",
                "position": 300,
                "site": "donor",
                "delta_direction": "decrease",
                "unique_to_pathogenic": True,
            },
        ],
    }

    acceptor_scoped = pipeline_tools._goal_for_routing_job(
        structured_goal=goal,
        job={"job_id": "j1", "program_id": "exon_acceptor_route", "label": "exon acceptor route"},
    )
    donor_scoped = pipeline_tools._goal_for_routing_job(
        structured_goal=goal,
        job={"job_id": "j2", "program_id": "exon_donor_route", "label": "exon donor route"},
    )

    assert acceptor_scoped["routing_job"]["site_kind"] == "acceptor"
    assert [
        pipeline_tools._routing_candidate_label(entry, Runtime())
        for entry in acceptor_scoped["routing_local_splice_site_plan"]
    ] == ["acceptor@100"]
    assert donor_scoped["routing_job"]["site_kind"] == "donor"
    assert [
        pipeline_tools._routing_candidate_label(entry, Runtime())
        for entry in donor_scoped["routing_local_splice_site_plan"]
    ] == ["donor@300"]


def test_objective_relevant_entries_do_not_promote_shared_context_over_requested_route(monkeypatch) -> None:
    class Runtime:
        @staticmethod
        def _canonical_site_label(entry):
            return f"{entry.get('site')}@{entry.get('position')}"

    monkeypatch.setattr(pipeline_tools, "_goal_runtime_module", lambda: Runtime())
    direct_acceptor = {
        "chrom": "17",
        "position": 100,
        "site": "acceptor",
        "delta_direction": "decrease",
        "unique_to_pathogenic": True,
        "observed_directional_max_delta": 0.2,
        "observed_support": 5,
        "decision_direction": "decrease",
        "source_delta_direction": "decrease",
    }
    shared_context_acceptor = {
        "chrom": "17",
        "position": 200,
        "site": "acceptor",
        "delta_direction": "context",
        "delta_per_isoform": {"pathogenic": "decrease", "productive": "increase"},
        "covers_all_isoforms": True,
        "pathogenic_isoforms": ["pathogenic"],
        "observed_directional_max_delta": 0.9,
        "observed_support": 20,
        "decision_direction": "decrease",
        "source_delta_direction": "context",
    }
    structured_goal = {
        "routing_job": {"job_label": "exon acceptor route", "site_kind": "acceptor"},
        "routing_local_splice_site_plan": [direct_acceptor],
        "splice_site_plan": [direct_acceptor, shared_context_acceptor],
    }

    prioritized = pipeline_tools._objective_relevant_routing_entries(
        structured_goal,
        [shared_context_acceptor, direct_acceptor],
        Runtime(),
    )

    assert [
        pipeline_tools._routing_candidate_label(entry, Runtime())
        for entry in prioritized
    ] == ["acceptor@100"]


def test_collapse_redundant_steric_jobs_keeps_distinct_acceptor_and_donor_routes() -> None:
    shared_rows = [["chr17", 1000, 1200, "+"]]
    jobs = [
        {
            "job_id": "steric_job_route_acceptor_window",
            "program_id": "route_acceptor_window",
            "label": "Exon splice acceptor window",
            "targeting_policy": "exploratory_derived",
            "region_rows": shared_rows,
            "run_splice_eval": True,
            "run_splice_router": True,
            "intent": "splice_modulation",
            "target_region": "splice",
        },
        {
            "job_id": "steric_job_route_donor_window",
            "program_id": "route_donor_window",
            "label": "Exon splice donor window",
            "targeting_policy": "exploratory_derived",
            "region_rows": shared_rows,
            "run_splice_eval": True,
            "run_splice_router": True,
            "intent": "splice_modulation",
            "target_region": "splice",
        },
    ]

    collapsed = pipeline_tools._collapse_redundant_steric_jobs(jobs)

    assert [job["job_id"] for job in collapsed] == [
        "steric_job_route_acceptor_window",
        "steric_job_route_donor_window",
    ]


def test_collapse_redundant_steric_jobs_keeps_distinct_requested_programs() -> None:
    shared_rows = [["chr17", 1000, 1200, "+"]]
    jobs = [
        {
            "job_id": "steric_job_first",
            "program_id": "first_derived",
            "parent_program_id": "first_requested_program",
            "label": "first route",
            "desired_delta": "increase",
            "targeting_policy": "exploratory_derived",
            "region_rows": shared_rows,
            "run_splice_eval": True,
            "run_splice_router": True,
            "intent": "splice_modulation",
            "site_kind": "acceptor",
        },
        {
            "job_id": "steric_job_second",
            "program_id": "second_derived",
            "parent_program_id": "second_requested_program",
            "label": "second route",
            "desired_delta": "decrease",
            "targeting_policy": "exploratory_derived",
            "region_rows": shared_rows,
            "run_splice_eval": True,
            "run_splice_router": True,
            "intent": "splice_modulation",
            "site_kind": "acceptor",
        },
    ]

    collapsed = pipeline_tools._collapse_redundant_steric_jobs(jobs)

    assert [job["parent_program_id"] for job in collapsed] == [
        "first_requested_program",
        "second_requested_program",
    ]
    assert [job["desired_delta"] for job in collapsed] == [
        "increase",
        "decrease",
    ]


def test_collapse_redundant_steric_jobs_keeps_distinct_length_variants() -> None:
    shared_rows = [["chr1", 100, 200, "+"]]
    jobs = [
        {
            "job_id": "variant_shorter",
            "program_id": "variant_shorter",
            "targeting_policy": "exploratory_derived",
            "region_rows": shared_rows,
            "intent": "binding",
            "length": 16,
        },
        {
            "job_id": "variant_longer",
            "program_id": "variant_longer",
            "targeting_policy": "exploratory_derived",
            "region_rows": shared_rows,
            "intent": "binding",
            "length": 17,
        },
    ]

    collapsed = pipeline_tools._collapse_redundant_steric_jobs(jobs)

    assert [job["length"] for job in collapsed] == [16, 17]


def test_job_candidate_site_labels_collects_all_window_labels() -> None:
    labels = pipeline_tools._job_candidate_site_labels(
        {
            "window_summaries": [
                {"site_labels": ["donor@100", "acceptor@200"]},
                {"site_labels": ["donor@300"]},
            ]
        }
    )

    assert labels == {"donor@100", "acceptor@200", "donor@300"}


def test_route_ranker_accepts_ranked_subset(monkeypatch, tmp_path: Path) -> None:
    class Runtime:
        @staticmethod
        def _canonical_site_label(entry):
            return f"{entry.get('site')}@{entry.get('position')}"

    monkeypatch.setattr(pipeline_tools, "_goal_runtime_module", lambda: Runtime())
    monkeypatch.setattr(
        pipeline_tools,
        "rank_route_candidates_with_llm",
        lambda payload: {
            "ranking_summary": "Kept the objective-specific donor only.",
            "ranked_routes": [
                {
                    "rank": 1,
                    "label": "donor@100",
                    "direction": "decrease",
                    "rationale": "Unique pathogenic donor with directional support.",
                    "confidence": "high",
                }
            ],
            "excluded_routes": [
                {
                    "label": "acceptor@200",
                    "direction": "decrease",
                    "reason": "Shared context site, not a focused route.",
                }
            ],
        },
    )

    result = pipeline_tools._rank_routing_entries_with_llm(
        structured_goal={},
        routing_entries=[
            {"site": "donor", "position": 100, "delta_direction": "decrease"},
            {"site": "acceptor", "position": 200, "decision_direction": "decrease"},
        ],
        debug_dir=str(tmp_path),
    )

    assert "error" not in result
    assert [entry["position"] for entry in result["ranked_entries"]] == [100]
    assert result["excluded_routes"][0]["label"] == "acceptor@200"


def test_route_ranker_appends_omitted_objective_aligned_routes(monkeypatch, tmp_path: Path) -> None:
    class Runtime:
        @staticmethod
        def _canonical_site_label(entry):
            return f"{entry.get('site')}@{entry.get('position')}"

    monkeypatch.setattr(pipeline_tools, "_goal_runtime_module", lambda: Runtime())
    monkeypatch.setattr(
        pipeline_tools,
        "rank_route_candidates_with_llm",
        lambda payload: {
            "ranking_summary": "Ranked only the unique pathogenic donor.",
            "ranked_routes": [
                {
                    "rank": 1,
                    "label": "donor@100",
                    "direction": "decrease",
                    "rationale": "Unique pathogenic donor with strong directional support.",
                    "confidence": "high",
                }
            ],
            "excluded_routes": [
                {
                    "label": "acceptor@200",
                    "direction": "decrease",
                    "reason": "Shared between beneficial and pathogenic isoforms.",
                }
            ],
        },
    )

    result = pipeline_tools._rank_routing_entries_with_llm(
        structured_goal={
            "splice_isoform_objectives": [
                {"isoform_label": "pathogenic", "desired_abundance": "decrease"},
                {"isoform_label": "beneficial", "desired_abundance": "increase"},
            ],
            "routing_job": {
                "job_id": "steric_job_isoform_balance",
                "job_label": "isoform_balance",
                "targeting_policy": "exploratory_derived",
                "candidate_sites": ["donor@100", "acceptor@200"],
            },
        },
        routing_entries=[
            {
                "site": "donor",
                "position": 100,
                "decision_direction": "decrease",
                "delta_per_isoform": {"pathogenic": "decrease"},
                "observed_directional_max_delta": 0.8,
                "observed_directional_support": 20,
                "observed_support": 40,
            },
            {
                "site": "acceptor",
                "position": 200,
                "decision_direction": "decrease",
                "delta_per_isoform": {
                    "pathogenic": "decrease",
                    "beneficial": "increase",
                },
                "observed_directional_max_delta": 0.3,
                "observed_directional_support": 12,
                "observed_support": 40,
            },
        ],
        debug_dir=str(tmp_path),
    )

    assert "error" not in result
    assert [entry["position"] for entry in result["ranked_entries"]] == [100, 200]
    assert result["ranked_entries"][1]["ranking_confidence"] == "low"
    assert "Retained omitted objective-relevant route" in result["ranked_entries"][1]["ranking_rationale"]
    output = json.loads((tmp_path / "llm_route_ranking_output.json").read_text(encoding="utf-8"))
    assert output["omitted_objective_route_rescue_applied"] is True


def test_route_ranker_treats_all_excluded_routes_as_skippable(monkeypatch, tmp_path: Path) -> None:
    class Runtime:
        @staticmethod
        def _canonical_site_label(entry):
            return f"{entry.get('site')}@{entry.get('position')}"

    monkeypatch.setattr(pipeline_tools, "_goal_runtime_module", lambda: Runtime())
    monkeypatch.setattr(
        pipeline_tools,
        "rank_route_candidates_with_llm",
        lambda payload: {
            "ranking_summary": "The only candidate is a shared off-objective acceptor.",
            "ranked_routes": [],
            "excluded_routes": [
                {
                    "label": "acceptor@200",
                    "direction": "decrease",
                    "reason": "Shared context site, not a focused route.",
                }
            ],
        },
    )

    result = pipeline_tools._rank_routing_entries_with_llm(
        structured_goal={},
        routing_entries=[
            {"site": "acceptor", "position": 200, "decision_direction": "decrease"},
        ],
        debug_dir=str(tmp_path),
    )

    assert result["error"] == "router_no_directional_targets"
    assert result["excluded_routes"][0]["label"] == "acceptor@200"


def test_route_ranker_rescues_explicit_user_window_when_llm_excludes_all(monkeypatch, tmp_path: Path) -> None:
    class Runtime:
        @staticmethod
        def _canonical_site_label(entry):
            return f"{entry.get('site')}@{entry.get('position')}"

    monkeypatch.setattr(pipeline_tools, "_goal_runtime_module", lambda: Runtime())
    monkeypatch.setattr(
        pipeline_tools,
        "rank_route_candidates_with_llm",
        lambda payload: {
            "ranking_summary": "The requested site has weak opposite-direction evidence.",
            "ranked_routes": [],
            "excluded_routes": [
                {
                    "label": "donor@100",
                    "direction": "increase",
                    "reason": "Requested increase is weak and opposite effects dominate.",
                }
            ],
        },
    )

    result = pipeline_tools._rank_routing_entries_with_llm(
        structured_goal={
            "per_site_window_program_preseed": {"status": "applied"},
            "routing_job": {
                "job_id": "steric_job_exon7_donor_window",
                "job_label": "exon7_donor_window",
                "targeting_policy": "explicit_user_region",
                "site_kind": "donor",
            },
        },
        routing_entries=[
            {
                "site": "donor",
                "position": 100,
                "decision_direction": "increase",
                "observed_directional_max_delta": 0.0075,
                "observed_opposite_max_delta": 0.9955,
                "observed_directional_support": 8,
                "observed_support": 42,
            },
        ],
        debug_dir=str(tmp_path),
    )

    assert "error" not in result
    assert result["ranked_entries"][0]["position"] == 100
    assert result["ranked_entries"][0]["ranking_confidence"] == "low"
    assert "User-requested explicit target window" in result["ranked_entries"][0]["ranking_rationale"]
    output = json.loads((tmp_path / "llm_route_ranking_output.json").read_text(encoding="utf-8"))
    assert output["explicit_user_requested_rescue_applied"] is True


def test_route_ranker_explicit_window_rescue_keeps_requested_site_kind(monkeypatch, tmp_path: Path) -> None:
    class Runtime:
        @staticmethod
        def _canonical_site_label(entry):
            return f"{entry.get('site')}@{entry.get('position')}"

    monkeypatch.setattr(pipeline_tools, "_goal_runtime_module", lambda: Runtime())
    monkeypatch.setattr(
        pipeline_tools,
        "rank_route_candidates_with_llm",
        lambda payload: {
            "ranking_summary": "The requested acceptor evidence is weak.",
            "ranked_routes": [],
            "excluded_routes": [
                {
                    "label": "acceptor@100",
                    "direction": "increase",
                    "reason": "Requested increase is weak.",
                },
                {
                    "label": "donor@200",
                    "direction": "increase",
                    "reason": "Off-objective for acceptor-window steric job.",
                },
            ],
        },
    )

    result = pipeline_tools._rank_routing_entries_with_llm(
        structured_goal={
            "per_site_window_program_preseed": {"status": "applied"},
            "routing_job": {
                "job_id": "steric_job_exon7_acceptor_window",
                "job_label": "exon7_acceptor_window",
                "targeting_policy": "explicit_user_region",
                "site_kind": "acceptor",
            },
        },
        routing_entries=[
            {
                "site": "acceptor",
                "position": 100,
                "decision_direction": "increase",
                "observed_directional_max_delta": 0.0011,
                "observed_opposite_max_delta": 0.9976,
                "observed_directional_support": 4,
                "observed_support": 27,
            },
            {
                "site": "donor",
                "position": 200,
                "decision_direction": "increase",
                "observed_directional_max_delta": 0.0075,
                "observed_opposite_max_delta": 0.9955,
                "observed_directional_support": 8,
                "observed_support": 42,
            },
        ],
        debug_dir=str(tmp_path),
    )

    assert "error" not in result
    assert [(entry["site"], entry["position"]) for entry in result["ranked_entries"]] == [("acceptor", 100)]


def test_route_ranker_rescues_aligned_contextual_routes(monkeypatch, tmp_path: Path) -> None:
    class Runtime:
        @staticmethod
        def _canonical_site_label(entry):
            return f"{entry.get('site')}@{entry.get('position')}"

    captured_payloads: list[dict[str, Any]] = []

    def fake_ranker(payload: dict[str, Any]) -> dict[str, Any]:
        captured_payloads.append(payload)
        return {
            "ranking_summary": "The only candidate is shared.",
            "ranked_routes": [],
            "excluded_routes": [
                {
                    "label": "acceptor@200",
                    "direction": "decrease",
                    "reason": "Shared context site.",
                }
            ],
        }

    monkeypatch.setattr(pipeline_tools, "_goal_runtime_module", lambda: Runtime())
    monkeypatch.setattr(pipeline_tools, "rank_route_candidates_with_llm", fake_ranker)

    result = pipeline_tools._rank_routing_entries_with_llm(
        structured_goal={
            "desired_effect": "inclusion",
            "exon_label": 7,
            "splice_isoform_objectives": [
                {"isoform_label": "pathogenic isoform", "desired_abundance": "decrease"},
                {"isoform_label": "beneficial isoform", "desired_abundance": "increase"},
            ]
        },
        routing_entries=[
            {
                "site": "acceptor",
                "position": 200,
                "decision_direction": "increase",
                "source_delta_direction": "context",
                "delta_per_isoform": {
                    "pathogenic isoform": "decrease",
                    "beneficial isoform": "increase",
                },
                "exon_indices": {"pathogenic isoform": 7, "beneficial isoform": 7},
                "covers_all_isoforms": True,
                "observed_directional_max_delta": 0.91,
                "observed_directional_support": 12,
            }
        ],
        debug_dir=str(tmp_path),
    )

    assert "error" not in result
    assert result["ranked_entries"][0]["position"] == 200
    assert result["ranked_entries"][0]["ranking_confidence"] == "medium"
    candidate_payload = captured_payloads[0]["candidates"][0]
    assert candidate_payload["delta_per_isoform"] == {
        "pathogenic isoform": "decrease",
        "beneficial isoform": "increase",
    }
    assert candidate_payload["isoform_objective_alignment"]["status"] == "aligned"


def test_route_ranker_rejects_contextual_route_opposite_exon_inclusion_direction(monkeypatch, tmp_path: Path) -> None:
    class Runtime:
        @staticmethod
        def _canonical_site_label(entry):
            return f"{entry.get('site')}@{entry.get('position')}"

    monkeypatch.setattr(pipeline_tools, "_goal_runtime_module", lambda: Runtime())
    monkeypatch.setattr(
        pipeline_tools,
        "rank_route_candidates_with_llm",
        lambda payload: {
            "ranking_summary": "No valid candidates.",
            "ranked_routes": [],
            "excluded_routes": [],
        },
    )

    result = pipeline_tools._rank_routing_entries_with_llm(
        structured_goal={
            "desired_effect": "inclusion",
            "exon_label": 7,
            "splice_isoform_objectives": [
                {"isoform_label": "pathogenic isoform", "desired_abundance": "decrease"},
                {"isoform_label": "beneficial isoform", "desired_abundance": "increase"},
            ],
        },
        routing_entries=[
            {
                "site": "acceptor",
                "position": 200,
                "decision_direction": "decrease",
                "source_delta_direction": "context",
                "delta_per_isoform": {
                    "pathogenic isoform": "decrease",
                    "beneficial isoform": "increase",
                },
                "exon_indices": {"pathogenic isoform": 7, "beneficial isoform": 7},
                "covers_all_isoforms": True,
                "observed_directional_max_delta": 0.91,
                "observed_directional_support": 12,
            }
        ],
        debug_dir=str(tmp_path),
    )

    assert result["error"] == "router_llm_ranking_no_serializable_candidates"


def test_observed_routes_do_not_flip_contextual_exon_inclusion_site_to_decrease(
    monkeypatch, tmp_path: Path
) -> None:
    class Runtime:
        @staticmethod
        def _canonical_site_label(entry):
            return f"{entry.get('site')}@{entry.get('position')}"

    monkeypatch.setattr(pipeline_tools, "_goal_runtime_module", lambda: Runtime())
    scored_path = tmp_path / "scored_tiles.json"
    scored_path.write_text(
        json.dumps(
            {
                "rows": [
                    {
                        "delta_names": "acceptor@200",
                        "delta_deletion": "-0.91",
                        "delta_replacement": "-0.87",
                    },
                    {
                        "delta_names": "acceptor@200",
                        "delta_deletion": "-0.74",
                        "delta_replacement": "-0.70",
                    },
                ]
            }
        ),
        encoding="utf-8",
    )

    entries = pipeline_tools._routing_target_entries(
        {
            "desired_effect": "inclusion",
            "exon_label": 7,
            "splice_site_plan": [
                {
                    "site": "acceptor",
                    "position": 200,
                    "delta_direction": "context",
                    "delta_per_isoform": {
                        "pathogenic isoform": "decrease",
                        "beneficial isoform": "increase",
                    },
                    "exon_indices": {"pathogenic isoform": 7, "beneficial isoform": 7},
                    "covers_all_isoforms": True,
                }
            ],
            "splice_isoform_objectives": [
                {"isoform_label": "pathogenic isoform", "desired_abundance": "decrease"},
                {"isoform_label": "beneficial isoform", "desired_abundance": "increase"},
            ],
        },
        scored_path=str(scored_path),
    )

    assert entries == []


def test_restart_from_stage_clears_consumed_clarification_restart(monkeypatch, tmp_path: Path) -> None:
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    state_path = tmp_path / "state.json"
    state_path.write_text(
        json.dumps(
            {
                "version": "2",
                "status": "initialized",
                "phase": "restart_intake_pending",
                "request": "Design SMN2 ASOs.",
                "flags": {},
                "run_dir": str(run_dir),
                "artifacts": {},
                "events": [],
                "last_error": None,
                "pending_questions": [],
                "clarification_queue": [
                    {
                        "task_id": "design_smn2",
                        "questions": ["Which sites should count as distinct sites?"],
                    }
                ],
                "blocked_task_ids": ["design_smn2"],
                "task_validation": {
                    "design_smn2": {
                        "status": "blocked",
                        "questions": ["Which sites should count as distinct sites?"],
                    }
                },
                "request_envelope": {
                    "request_text": "Design SMN2 ASOs.",
                    "tasks": [
                        {
                            "task_id": "design_smn2",
                            "task_type": "design_new_asos",
                            "label": "Design SMN2 ASOs",
                            "context": {
                                "gene_symbol": "SMN2",
                                "species": "human",
                                "target_context": "Human spinal cord motor neurons in an SMA model.",
                            },
                        }
                    ],
                    "blocked_tasks": [
                        {
                            "task_id": "design_smn2",
                            "questions": ["Which sites should count as distinct sites?"],
                            "missing_fields": ["context.structured_goal.region_constraints"],
                        }
                    ],
                },
                "task_graph": {
                    "tasks": [
                        {
                            "task_id": "design_smn2",
                            "task_type": "design_new_asos",
                            "label": "Design SMN2 ASOs",
                        }
                    ],
                    "dependencies": [],
                },
                "task_index": {
                    "tasks": [
                        {
                            "task_id": "design_smn2",
                            "task_type": "design_new_asos",
                            "label": "Design SMN2 ASOs",
                            "status": "blocked",
                        }
                    ]
                },
                "task_statuses": {
                    "design_smn2": {
                        "task_type": "design_new_asos",
                        "native": True,
                        "status": "blocked",
                    }
                },
                "clarification": {"pending_apply": False},
                "pending_clarification_restart_stage": "intake",
                "pending_clarification_issue_summary": "Clarification answers captured.",
                "stage_attempts": {},
                "force_rerun_stages": [],
            }
        ),
        encoding="utf-8",
    )
    monkeypatch.setattr(pipeline_tools, "_reject_disallowed_tool_call", lambda **_kwargs: None)

    result = pipeline_tools.restart_from_stage(
        str(state_path),
        stage="intake",
        issue_summary="Clarification answers captured.",
    )
    saved = json.loads(state_path.read_text(encoding="utf-8"))

    assert result["status"] == "ok"
    assert saved["pending_clarification_restart_stage"] == ""
    assert saved["pending_clarification_issue_summary"] == ""
    assert saved["status"] == "initialized"
    assert saved["pending_questions"] == []
    assert saved["clarification_queue"] == []
    assert saved["blocked_task_ids"] == []
    assert "design_smn2" not in saved["task_validation"]
    saved_envelope = saved["request_envelope"]
    assert not saved_envelope.get("blocked_tasks")
    assert "design_smn2" not in {
        str(task.get("task_id") or "").strip()
        for task in saved_envelope.get("tasks") or []
        if isinstance(task, dict)
    }
    assert all(
        str(payload.get("status") or "").strip().lower() != "blocked"
        for payload in (saved.get("task_statuses") or {}).values()
        if isinstance(payload, dict)
    )
    assert all(
        payload.get("status") != "blocked"
        for payload in saved["task_validation"].values()
        if isinstance(payload, dict)
    )
    planning = pipeline_tools._derive_next_actions(saved)
    assert planning["pending_questions"] == []
    assert planning["terminal_status"] == "in_progress"
    assert planning["next_actions"][0]["tool"] == "tool_parse_goal_request"


def test_explicit_restart_recovers_terminal_failure_and_resets_downstream_attempts(
    tmp_path: Path,
) -> None:
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    state_path = tmp_path / "state.json"
    state_path.write_text(
        json.dumps(
            {
                "version": "2",
                "status": "failed",
                "phase": "tool_gapmer_uniqueness_failed",
                "run_dir": str(run_dir),
                "events": [],
                "last_error": {"kind": "tool_gapmer_uniqueness_failed"},
                "last_tool": "evaluate_step_output",
                "stage_attempts": {
                    "resolution": 1,
                    "gapmer_tiling": 2,
                    "gapmer_scoring": 2,
                    "report": 1,
                },
            }
        ),
        encoding="utf-8",
    )

    result = pipeline_tools.restart_from_stage(
        str(state_path),
        stage="gapmer_tiling",
        issue_summary="A corrected implementation requires rebuilding the failed stage.",
    )
    saved = json.loads(state_path.read_text(encoding="utf-8"))

    assert result["status"] == "ok"
    assert saved["status"] == "routing_ready"
    assert saved["phase"] == "restart_gapmer_tiling_pending"
    assert saved["stage_attempts"] == {"resolution": 1}
    assert saved["last_error"] is None


def test_route_ranker_applies_persisted_strategy_order(monkeypatch, tmp_path: Path) -> None:
    class Runtime:
        @staticmethod
        def _canonical_site_label(entry):
            return f"{entry.get('site')}@{entry.get('position')}"

    monkeypatch.setattr(pipeline_tools, "_goal_runtime_module", lambda: Runtime())
    monkeypatch.setattr(
        pipeline_tools,
        "rank_route_candidates_with_llm",
        lambda payload: {
            "ranking_summary": "Sorted by raw observed effect.",
            "ranked_routes": [
                {
                    "rank": 1,
                    "label": "donor@200",
                    "direction": "decrease",
                    "rationale": "Slightly stronger raw delta.",
                    "confidence": "high",
                },
                {
                    "rank": 2,
                    "label": "acceptor@100",
                    "direction": "decrease",
                    "rationale": "Primary strategy target with strong support.",
                    "confidence": "high",
                },
            ],
        },
    )

    result = pipeline_tools._rank_routing_entries_with_llm(
        structured_goal={
            "splice_strategy_analysis": {
                "analysis": (
                    "## Ranked Steric Hindrance Targets\n"
                    "1) decrease target acceptor usage\n"
                    "- Primary targets: acceptor@100 (decrease)\n"
                    "2) decrease target donor usage\n"
                    "- Primary targets: donor@200 (decrease)\n"
                )
            }
        },
        routing_entries=[
            {"site": "acceptor", "position": 100, "delta_direction": "decrease"},
            {"site": "donor", "position": 200, "delta_direction": "decrease"},
        ],
        debug_dir=str(tmp_path),
    )

    assert "error" not in result
    assert result["strategy_priority_applied"] is False
    assert [entry["position"] for entry in result["ranked_entries"]] == [200, 100]
    assert "Slightly stronger raw delta" in result["ranked_entries"][0]["ranking_rationale"]


def test_route_ranker_keeps_metric_backed_llm_order_over_stale_strategy_order(monkeypatch, tmp_path: Path) -> None:
    class Runtime:
        @staticmethod
        def _canonical_site_label(entry):
            return f"{entry.get('site')}@{entry.get('position')}"

    monkeypatch.setattr(pipeline_tools, "_goal_runtime_module", lambda: Runtime())
    monkeypatch.setattr(
        pipeline_tools,
        "rank_route_candidates_with_llm",
        lambda payload: {
            "ranking_summary": "Ranked by observed directional leverage.",
            "ranked_routes": [
                {
                    "rank": 1,
                    "label": "donor@200",
                    "direction": "decrease",
                    "rationale": "Much stronger directional leverage.",
                    "confidence": "high",
                },
                {
                    "rank": 2,
                    "label": "acceptor@100",
                    "direction": "decrease",
                    "rationale": "On-objective but weaker.",
                    "confidence": "medium",
                },
            ],
        },
    )

    result = pipeline_tools._rank_routing_entries_with_llm(
        structured_goal={
            "splice_strategy_analysis": {
                "analysis": (
                    "## Ranked Steric Hindrance Targets\n"
                    "1) decrease target acceptor usage\n"
                    "- Primary targets: acceptor@100 (decrease)\n"
                    "2) decrease target donor usage\n"
                    "- Primary targets: donor@200 (decrease)\n"
                )
            }
        },
        routing_entries=[
            {
                "site": "acceptor",
                "position": 100,
                "delta_direction": "decrease",
                "observed_directional_max_delta": 0.0368,
                "observed_max_abs_delta": 0.0368,
            },
            {
                "site": "donor",
                "position": 200,
                "delta_direction": "decrease",
                "observed_directional_max_delta": 0.9953,
                "observed_max_abs_delta": 0.9953,
            },
        ],
        debug_dir=str(tmp_path),
    )

    assert "error" not in result
    assert result["strategy_priority_applied"] is False
    assert [entry["position"] for entry in result["ranked_entries"]] == [200, 100]


def test_design_objective_preseed_preserves_artifact_backed_isoform_objectives() -> None:
    preseed = pipeline_tools._prepare_design_objective_preseed(
        {
            "gene_symbol": "LMNA",
            "species": "human",
            "intent": "binding",
            "isoform_objectives": [
                {
                    "isoform_label": "progerin",
                    "direction": "decrease",
                    "transcript_id": "ENST00000368299",
                },
                {
                    "isoform_label": "snapgene_isoform",
                    "direction": "increase",
                    "definition_source": "artifact:/tmp/reference.gb",
                },
            ],
        },
        {"artifact_refs": [{"path": "/tmp/reference.gb"}]},
    )

    objectives = preseed["splice_isoform_objectives"]
    assert objectives[0]["desired_abundance"] == "decrease"
    assert objectives[0]["role"] == "context"
    assert objectives[1]["desired_abundance"] == "increase"
    assert objectives[1]["role"] == "context"
    assert objectives[1]["source"] == "snapgene"
    assert objectives[1]["snapgene_path"] == "/tmp/reference.gb"


def test_refseq_transcript_resolves_by_interval_context_when_xref_has_no_enst(monkeypatch) -> None:
    transcripts = [
        {
            "id": "ENST00000442489",
            "seq_region_name": "12",
            "start": 12717861,
            "end": 12721289,
            "strand": 1,
            "three_prime_terminal_exon_start": 12721036,
            "three_prime_terminal_exon_end": 12721289,
            "is_canonical": False,
            "is_mane_select": False,
            "biotype": "protein_coding",
        },
        {
            "id": "ENST00000228872",
            "seq_region_name": "12",
            "start": 12717368,
            "end": 12722369,
            "strand": 1,
            "three_prime_terminal_exon_start": 12721036,
            "three_prime_terminal_exon_end": 12722369,
            "is_canonical": True,
            "is_mane_select": False,
            "biotype": "protein_coding",
        },
    ]

    monkeypatch.setattr(run_goal_intake, "list_transcripts", lambda **kwargs: {"transcripts": transcripts})
    monkeypatch.setattr(run_goal_intake, "resolve_external_transcript_accession", lambda *args, **kwargs: None)

    def fake_fetch_transcript_structure(transcript_id: str) -> dict[str, Any]:
        assert transcript_id == "ENST00000228872"
        return {
            "seq_region_name": "12",
            "strand": 1,
            "exons_tx": [
                {"index": 1, "start": 12717368, "end": 12718314},
                {"index": 2, "start": 12718825, "end": 12719618},
                {"index": 3, "start": 12721036, "end": 12722369},
            ],
        }

    monkeypatch.setattr(run_goal_intake, "_fetch_transcript_structure", fake_fetch_transcript_structure)

    final_payload: dict[str, Any] = {
        "evaluation_targets": [
            {"chrom": "chr12", "start": 12721313, "end": 12721416, "strand": "+"},
            {"chrom": "chr12", "start": 12721400, "end": 12721460, "strand": "+"},
            {"chrom": "chr12", "start": 12721208, "end": 12721253, "strand": "+"},
            {"chrom": "chr12", "start": 12721281, "end": 12721326, "strand": "+"},
        ],
        "target_region": "utr3",
    }
    goal = {
        "gene_symbol": "CDKN1B",
        "species": "human",
        "target_region": "utr3",
        "splice_isoform_objectives": [
            {
                "isoform_label": "NM_004064.5",
                "transcript_id": "NM_004064.5",
                "desired_abundance": "increase",
            }
        ],
    }

    result = run_goal_intake._resolve_isoform_transcripts(
        goal,
        final_payload,
        allow_interactive=False,
    )

    resolved = result["isoform_transcripts"]["NM_004064.5"]
    assert resolved["transcript_id"] == "ENST00000228872"
    assert resolved["match_info"]["method"] == "external_accession_interval_context"
    assert resolved["match_info"]["confidence"] == "high"


def test_resolve_isoforms_leaves_other_isoform_bucket_unresolved_without_confirmation(monkeypatch) -> None:
    transcripts = [
        {
            "id": "ENST00000380743",
            "version": 9,
            "versioned_id": "ENST00000380743.9",
            "seq_region_name": "5",
            "start": 70000000,
            "end": 70009099,
            "strand": -1,
            "display_name": "SMN2-201",
            "description": "survival motor neuron 2 full-length transcript",
            "is_canonical": True,
            "is_mane_select": False,
            "biotype": "protein_coding",
        }
    ]

    monkeypatch.setattr(run_goal_intake, "list_transcripts", lambda **kwargs: {"transcripts": transcripts})

    aggregate_keys = {
        "excluded_exon_7_isoform_class",
        "aggregate:exon_presence:exon_7:excluded",
    }

    def fail_on_aggregate_label(*args: Any, **kwargs: Any) -> dict[str, Any]:
        isoform_label = args[1] if len(args) > 1 else kwargs.get("isoform_label")
        if isoform_label in aggregate_keys:
            raise AssertionError("aggregate isoform bucket should not be auto-matched")
        return {
            "transcript_id": "ENST00000380743",
            "confidence": "high",
            "method": "test",
            "rationale": "test fallback",
            "requires_confirmation": False,
        }

    monkeypatch.setattr(run_goal_intake, "auto_match_isoform_transcript", fail_on_aggregate_label)
    monkeypatch.setattr(run_goal_intake, "_generate_splice_strategy_analysis", lambda *args, **kwargs: None)

    def fake_fetch_transcript_structure(transcript_id: str) -> dict[str, Any]:
        assert transcript_id == "ENST00000380743.9"
        exons = [
            {"index": idx, "start": 70000000 + idx * 1000, "end": 70000099 + idx * 1000}
            for idx in range(1, 10)
        ]
        return {"id": "ENST00000380743", "versioned_id": "ENST00000380743.9", "seq_region_name": "5", "strand": -1, "exons_tx": exons}

    monkeypatch.setattr(run_goal_intake, "_fetch_transcript_structure", fake_fetch_transcript_structure)

    goal = {
        "gene_symbol": "SMN2",
        "species": "human",
        "intent": "splice_modulation",
        "exon_label": 7,
        "desired_effect": "inclusion",
        "focus_exons_only": True,
        "splice_isoform_objectives": [
            {
                "objective_kind": "increase_isoform",
                "objective_scope": "specific_transcript",
                "transcript_id": "ENST00000380743.9",
                "isoform_label": "ENST00000380743.9",
                "desired_abundance": "increase",
                "role": "beneficial",
            },
            {
                "objective_kind": "decrease_isoforms",
                "objective_scope": "aggregate_class",
                "isoform_label": "excluded_exon_7_isoform_class",
                "aggregate_class": {
                    "basis": "exon_presence",
                    "exon_label": 7,
                    "presence": "excluded",
                },
                "notes": "Decrease the structured aggregate class.",
                "desired_abundance": "decrease",
                "role": "pathogenic",
            },
            {
                "objective_kind": "decrease_isoforms",
                "objective_scope": "aggregate_class",
                "aggregate_class": {
                    "basis": "exon_presence",
                    "exon_label": 7,
                    "presence": "excluded",
                },
                "notes": "Decrease the same structured aggregate class without a display label.",
                "desired_abundance": "decrease",
                "role": "pathogenic",
            },
        ],
    }
    final_payload = deepcopy(goal)

    result = run_goal_intake._resolve_isoform_transcripts(
        goal,
        final_payload,
        allow_interactive=False,
    )

    assert "pending_isoform_transcript_confirmations" not in goal
    assert "unresolved_isoforms" not in goal
    assert result["isoform_transcripts"]["ENST00000380743.9"]["transcript_id"] == "ENST00000380743"
    assert not (aggregate_keys & set(result["isoform_transcripts"]))
    assert aggregate_keys <= set(goal["aggregate_isoform_classes"])


def test_user_locked_transcript_anchors_aggregate_exon_objective(monkeypatch) -> None:
    transcript_id = "ENST00000000001"
    locked_id = f"{transcript_id}.12"
    transcripts = [
        {
            "id": transcript_id,
            "version": 12,
            "versioned_id": locked_id,
            "seq_region_name": "7",
            "strand": 1,
            "display_name": "GENE_A-201",
            "biotype": "protein_coding",
        }
    ]
    monkeypatch.setattr(
        run_goal_intake,
        "list_transcripts",
        lambda **kwargs: {"transcripts": transcripts},
    )
    monkeypatch.setattr(
        run_goal_intake,
        "_fetch_transcript_structure",
        lambda observed_id: {
            "id": observed_id,
            "seq_region_name": "7",
            "strand": 1,
            "exons_tx": [
                {
                    "index": index,
                    "start": index * 100,
                    "end": index * 100 + 49,
                    "contains_coding": True,
                }
                for index in range(1, 7)
            ],
        },
    )
    monkeypatch.setattr(
        run_goal_intake,
        "_generate_splice_strategy_analysis",
        lambda *args, **kwargs: None,
    )
    goal = {
        "gene_symbol": "GENE_A",
        "species": "human",
        "intent": "splice_modulation",
        "exon_label": 4,
        "desired_effect": "skipping",
        "focus_exons_only": True,
        "transcript_policy": "user_locked",
        "transcript_lock": {
            "kind": "explicit_id",
            "value": locked_id,
            "version": 12,
            "source": "user_request",
        },
        "splice_isoform_objectives": [
            {
                "objective_scope": "aggregate_class",
                "aggregate_class": {
                    "basis": "exon_presence",
                    "exon_label": 4,
                    "presence": "included",
                },
                "exon_label": 4,
                "desired_effect": "skipping",
                "desired_abundance": "decrease",
            }
        ],
    }
    final_payload = deepcopy(goal)

    result = run_goal_intake._resolve_isoform_transcripts(
        goal,
        final_payload,
        allow_interactive=False,
    )

    assert result["isoform_transcripts"][locked_id]["transcript_id"] == transcript_id
    assert (
        result["isoform_transcripts"][locked_id]["match_info"]["method"]
        == "user_provided_transcript_id"
    )
    assert result["splice_site_plan"]
    assert {entry["exon_indices"][locked_id] for entry in result["splice_site_plan"]} == {
        4
    }
    assert goal["transcript_lock"] == {
        "kind": "explicit_id",
        "value": locked_id,
        "version": 12,
        "source": "user_request",
    }


def test_exon_only_splice_objective_defers_to_default_context() -> None:
    goal = {
        "gene_symbol": "Mapt",
        "species": "mouse",
        "gene_id": "ENSMUSG00000018411",
        "intent": "splice_modulation",
        "target_context": "mouse cortical neurons in a tauopathy model",
        "splice_isoform_objectives": [
            {
                "objective_scope": "default_exon_context",
                "exon_label": "10",
                "direction": "increase_inclusion",
            }
        ],
    }
    final_payload = deepcopy(goal)

    result = run_goal_intake._resolve_isoform_transcripts(
        goal,
        final_payload,
        allow_interactive=False,
    )

    assert result == {}
    assert goal["isoform_scope_source"] == "default_exon_splice_context"
    assert final_payload["isoform_scope_source"] == "default_exon_splice_context"


def test_co_skipping_exon_context_does_not_fabricate_single_exon_isoforms() -> None:
    request_text = (
        "Design splice-switching ASOs for human TTN NM_001267550.2 that induce either "
        "exons 325-326 or exons 326-327 co-skipping to bypass the exon-326 frameshift mutation."
    )
    intake = run_goal_intake._normalize_intake_output(
        {
            "status": "ready",
            "structured_goal": {
                "gene_symbol": "TTN",
                "species": "human",
                "intent": "splice_modulation",
                "transcript_policy": "user_locked",
                "transcript_lock": {
                    "kind": "reference_accession",
                    "value": "NM_001267550.2",
                    "note": "User specified the RefSeq transcript.",
                },
                "design_programs": [
                    {
                        "program_id": "ttn_325_326_coskipping",
                        "modality": "steric",
                        "intent": "splice_modulation",
                    }
                ],
                "splice_isoform_objectives": [
                    {
                        "objective_scope": "default_exon_context",
                        "exon_labels": [325, 326],
                        "desired_effect": "co_skip",
                        "direction": "decrease_inclusion",
                        "role": "therapeutic",
                    },
                    {
                        "objective_scope": "default_exon_context",
                        "exon_labels": [326, 327],
                        "desired_effect": "co_skip",
                        "direction": "decrease_inclusion",
                        "role": "therapeutic",
                    },
                ],
            },
        },
        request_text=request_text,
    )

    goal = intake["structured_goal"]
    labels = [
        str(objective.get("isoform_label") or "")
        for objective in goal["splice_isoform_objectives"]
        if isinstance(objective, dict)
    ]

    assert goal["isoform_scope_source"] == "default_exon_splice_context"
    assert len(goal["splice_isoform_objectives"]) == 2
    assert all(
        objective.get("objective_scope") == "default_exon_context"
        for objective in goal["splice_isoform_objectives"]
    )
    assert not any("full-length" in label or "Δ" in label for label in labels)


def test_co_skipping_request_overrides_compressed_single_exon_intake() -> None:
    request_text = (
        "Design splice-switching ASOs for human TTN NM_001267550.2 that induce either "
        "exons 325-326 or exons 326-327 co-skipping to bypass the exon-326 frameshift mutation."
    )
    intake = run_goal_intake._normalize_intake_output(
        {
            "status": "ready",
            "structured_goal": {
                "gene_symbol": "TTN",
                "species": "human",
                "intent": "splice_modulation",
                "target_context": "general human baseline context",
                "transcript_policy": "user_locked",
                "transcript_lock": {
                    "kind": "reference_accession",
                    "value": "NM_001267550.2",
                    "version": 2,
                    "note": "User explicitly specified using TTN transcript NM_001267550.2.",
                },
                "patient_variant": {
                    "gene_symbol": "TTN",
                    "kind": "cdna",
                    "notation": "c.82273C>T",
                },
                "design_programs": [
                    {
                        "program_id": "ttn_exon_326_skipping",
                        "modality": "steric",
                        "intent": "splice_modulation",
                        "desired_effect": "skipping",
                        "exon_label": 326,
                    }
                ],
                "splice_isoform_objectives": [
                    {
                        "objective_scope": "specific_transcript",
                        "reference_accession": "NM_001267550.2",
                        "isoform_label": "NM_001267550.2",
                        "exon_label": 326,
                        "desired_abundance": "decrease",
                        "desired_effect": "skipping",
                        "role": "pathogenic",
                    },
                    {
                        "objective_scope": "aggregate_class",
                        "aggregate_class": {
                            "basis": "exon_presence",
                            "exon_label": 326,
                            "presence": "excluded",
                        },
                        "desired_abundance": "increase",
                        "role": "beneficial",
                    },
                ],
            },
        },
        request_text=request_text,
    )

    goal = intake["structured_goal"]
    event_groups = [
        objective.get("exon_labels")
        for objective in goal["splice_isoform_objectives"]
        if isinstance(objective, dict)
    ]
    programs = goal["design_programs"]

    assert event_groups == [[325, 326], [326, 327]]
    assert all(
        objective.get("objective_scope") == "default_exon_context"
        and objective.get("desired_effect") == "skipping"
        and objective.get("desired_abundance") == "decrease"
        and "direction" not in objective
        for objective in goal["splice_isoform_objectives"]
    )
    assert all(objective.get("reference_accession") == "NM_001267550.2" for objective in goal["splice_isoform_objectives"])
    assert goal["transcript_lock"]["value"] == "NM_001267550.2"
    assert goal["patient_variant"]["notation"] == "c.82273C>T"
    assert [program.get("exon_labels") for program in programs] == [[325, 326], [326, 327]]
    assert all(
        program.get("desired_delta") == "decrease" and "direction" not in program
        for program in programs
    )


def test_co_skipping_splice_site_plan_keeps_outer_block_sites_only() -> None:
    isoform_transcripts = {
        "ENSTLOCKED": {
            "strand": "-",
            "sites": [
                {"chrom": "chr2", "acceptor": 1000, "donor": 1090, "exon_index": 324, "strand": "-", "coding": True},
                {"chrom": "chr2", "acceptor": 1200, "donor": 1290, "exon_index": 325, "strand": "-", "coding": True},
                {"chrom": "chr2", "acceptor": 1400, "donor": 1490, "exon_index": 326, "strand": "-", "coding": True},
                {"chrom": "chr2", "acceptor": 1600, "donor": 1690, "exon_index": 327, "strand": "-", "coding": True},
                {"chrom": "chr2", "acceptor": 1800, "donor": 1890, "exon_index": 328, "strand": "-", "coding": True},
            ],
            "exons_tx": [{"index": idx} for idx in range(324, 329)],
        }
    }
    structured_goal = {
        "splice_isoform_objectives": [
            {
                "objective_scope": "default_exon_context",
                "isoform_label": "ENSTLOCKED",
                "exon_labels": [325, 326],
                "desired_effect": "skipping",
                "desired_abundance": "decrease",
            }
        ],
        "focus_exons_only": True,
        "exon_labels": [325, 326],
    }

    plan, _comparisons = run_goal_intake._compute_splice_site_plan(
        structured_goal["splice_isoform_objectives"],
        isoform_transcripts,
        structured_goal,
    )

    assert [(entry["site"], entry["position"]) for entry in plan] == [
        ("acceptor", 1200),
        ("donor", 1490),
    ]
    assert all(entry["co_skipping_outer_site"] is True for entry in plan)


def test_default_splice_context_skips_unlocked_transcript_fetch_500(monkeypatch: Any) -> None:
    monkeypatch.setattr(
        run_goal_intake,
        "list_transcripts",
        lambda symbol, species="human", gene_id=None: {
            "transcripts": [
                {
                    "id": "ENST_FAIL",
                    "is_mane_select": True,
                    "is_canonical": False,
                    "length": 1000,
                    "exon_count": 4,
                },
                {
                    "id": "ENST_OK",
                    "is_mane_select": False,
                    "is_canonical": True,
                    "length": 900,
                    "exon_count": 4,
                },
            ]
        },
    )

    def fake_fetch(transcript_id: str) -> dict[str, Any]:
        if transcript_id == "ENST_FAIL":
            raise run_goal_intake.EnsemblServiceError("Ensembl REST lookup failed (500)")
        return {
            "id": transcript_id,
            "seq_region_name": "17",
            "strand": 1,
            "exons_tx": [
                {"index": 9, "start": 100, "end": 150, "contains_coding": True},
                {"index": 10, "start": 200, "end": 250, "contains_coding": True},
                {"index": 11, "start": 300, "end": 350, "contains_coding": True},
            ],
        }

    monkeypatch.setattr(run_goal_intake, "_fetch_transcript_structure", fake_fetch)
    monkeypatch.setattr(run_goal_intake, "_generate_splice_strategy_analysis", lambda *args, **kwargs: {})
    structured_goal = {
        "gene_symbol": "MAPT",
        "species": "human",
        "intent": "splice_modulation",
        "target_context": "human iPSC-derived cortical neurons",
        "focus_exons_only": True,
        "exon_label": 10,
        "splice_isoform_objectives": [
            {
                "objective_scope": "default_exon_context",
                "exon_label": 10,
                "desired_effect": "skipping",
                "desired_abundance": "decrease",
            }
        ],
    }
    final_payload = deepcopy(structured_goal)

    result = run_goal_intake._ensure_default_splice_context(structured_goal, final_payload)

    assert result["resolved_transcript_strategy"]["transcript_ids"] == ["ENST_OK"]
    assert "ENST_OK" in structured_goal["isoform_transcripts"]
    assert structured_goal["isoform_transcripts"]["ENST_OK"]["sites"]
    assert result["isoform_transcripts"]["ENST_OK"]["match_info"]["method"] == "default_splice_context_fallback"


def test_explicit_refseq_lock_from_request_is_preserved_for_coskipping() -> None:
    request = (
        "Design splice-switching ASOs for human TTN NM_001267550.2 that induce "
        "either exons 325-326 or exons 326-327 co-skipping."
    )
    goal = {
        "gene_symbol": "TTN",
        "species": "human",
        "intent": "splice_modulation",
        "target_context": "human heart / iPSC-derived cardiomyocytes",
    }
    final_payload = {"structured_goal": goal}

    changed = run_goal_intake._apply_explicit_transcript_lock_from_request_text(
        goal,
        request,
        final_payload,
    )
    run_goal_intake._apply_requested_co_skip_events_from_text(goal, request)

    assert changed is True
    assert goal["transcript_policy"] == "user_locked"
    assert goal["transcript_lock"]["kind"] == "reference_accession"
    assert goal["transcript_lock"]["value"] == "NM_001267550.2"
    assert {tuple(item["exon_labels"]) for item in goal["splice_isoform_objectives"]} == {
        (325, 326),
        (326, 327),
    }
    assert all(
        item.get("reference_accession") == "NM_001267550.2"
        for item in goal["splice_isoform_objectives"]
    )


def test_default_splice_context_candidate_order_ignores_gene_stable_ids() -> None:
    transcripts = [
        {
            "id": "ENSG00000155657",
            "object_type": "Gene",
            "is_mane_select": True,
            "is_canonical": True,
            "length": 1000000,
        },
        {
            "id": "ENST_OK",
            "object_type": "Transcript",
            "is_mane_select": False,
            "is_canonical": True,
            "length": 900,
        },
    ]

    selected = run_goal_intake._select_default_splice_context_transcript(transcripts)
    ordered = run_goal_intake._default_splice_context_candidate_order(transcripts, selected)

    assert selected["id"] == "ENST_OK"
    assert [entry["id"] for entry in ordered] == ["ENST_OK"]


def test_default_splice_context_retries_symbol_listing_after_gene_id_500(monkeypatch: Any) -> None:
    calls: list[Any] = []

    def fake_list_transcripts(symbol: str, species: str = "human", gene_id: Any = None) -> dict[str, Any]:
        calls.append(gene_id)
        if gene_id:
            raise run_goal_intake.EnsemblServiceError("Ensembl REST lookup failed (500)")
        return {
            "transcripts": [
                {
                    "id": "ENST_OK",
                    "is_mane_select": True,
                    "is_canonical": True,
                    "length": 900,
                    "exon_count": 3,
                }
            ]
        }

    def fake_fetch(transcript_id: str) -> dict[str, Any]:
        return {
            "id": transcript_id,
            "seq_region_name": "17",
            "strand": 1,
            "exons_tx": [
                {"index": 9, "start": 100, "end": 150, "contains_coding": True},
                {"index": 10, "start": 200, "end": 250, "contains_coding": True},
                {"index": 11, "start": 300, "end": 350, "contains_coding": True},
            ],
        }

    monkeypatch.setattr(run_goal_intake, "list_transcripts", fake_list_transcripts)
    monkeypatch.setattr(run_goal_intake, "_fetch_transcript_structure", fake_fetch)
    monkeypatch.setattr(run_goal_intake, "_generate_splice_strategy_analysis", lambda *args, **kwargs: {})

    goal = {
        "gene_symbol": "TTN",
        "gene_id": "ENSG00000155657",
        "species": "human",
        "intent": "splice_modulation",
        "target_context": "human heart / iPSC-derived cardiomyocytes",
        "focus_exons_only": True,
        "splice_isoform_objectives": [
            {
                "objective_scope": "default_exon_context",
                "exon_label": 10,
                "desired_effect": "skipping",
                "desired_abundance": "decrease",
            }
        ],
    }
    final_payload = deepcopy(goal)

    result = run_goal_intake._ensure_default_splice_context(goal, final_payload)

    assert calls == ["ENSG00000155657", None]
    assert result["resolved_transcript_strategy"]["transcript_ids"] == ["ENST_OK"]
    assert "ENST_OK" in goal["isoform_transcripts"]


def test_default_splice_context_uses_expanded_listing_structures(monkeypatch: Any) -> None:
    transcripts = [
        {
            "id": "ENST_OK",
            "is_mane_select": True,
            "is_canonical": True,
            "length": 900,
            "exon_count": 3,
        }
    ]
    structures = {
        "ENST_OK": {
            "id": "ENST_OK",
            "seq_region_name": "17",
            "strand": 1,
            "exons_tx": [
                {"index": 9, "start": 100, "end": 150, "contains_coding": True},
                {"index": 10, "start": 200, "end": 250, "contains_coding": True},
                {"index": 11, "start": 300, "end": 350, "contains_coding": True},
            ],
        }
    }

    monkeypatch.setattr(
        run_goal_intake,
        "get_gene_transcript_structures",
        lambda **kwargs: {"transcripts": transcripts, "structures": structures},
    )

    def fail_fetch(transcript_id: str) -> dict[str, Any]:
        raise AssertionError("expanded listing structure should avoid per-transcript fetch")

    monkeypatch.setattr(run_goal_intake, "_fetch_transcript_structure", fail_fetch)
    monkeypatch.setattr(run_goal_intake, "_generate_splice_strategy_analysis", lambda *args, **kwargs: {})

    goal = {
        "gene_symbol": "MAPT",
        "species": "human",
        "intent": "splice_modulation",
        "target_context": "human neurons",
        "focus_exons_only": True,
        "splice_isoform_objectives": [
            {
                "objective_scope": "default_exon_context",
                "exon_label": 10,
                "desired_effect": "skipping",
                "desired_abundance": "decrease",
            }
        ],
    }
    final_payload = deepcopy(goal)

    result = run_goal_intake._ensure_default_splice_context(goal, final_payload)

    assert result["resolved_transcript_strategy"]["transcript_ids"] == ["ENST_OK"]
    assert goal["isoform_transcripts"]["ENST_OK"]["exons_tx"] == structures["ENST_OK"]["exons_tx"]


def test_single_design_gene_preflight_skips_network_lookup() -> None:
    def fail_lookup(*args: Any, **kwargs: Any) -> dict[str, Any]:
        raise AssertionError("single-design preflight should not call Ensembl lookup")

    envelope = {
        "request_text": "Design splice-switching ASOs for human TTN.",
        "tasks": [
            {
                "task_id": "design_ttn",
                "task_type": "design_new_asos",
                "context": {
                    "structured_goal": {
                        "gene_symbol": "TTN",
                        "species": "human",
                        "intent": "splice_modulation",
                    }
                },
            }
        ],
    }

    result = pipeline_tools._preflight_request_design_genes(
        envelope,
        lookup_gene_id_fn=fail_lookup,
        lookup_gene_by_id_fn=fail_lookup,
    )

    assert result["status"] == "complete"
    assert result["request_envelope"] == envelope
    assert result["skipped"] == "single_design_task_resolved_downstream"


def test_goal_intake_step_autodefaults_missing_target_context(
    monkeypatch: Any,
) -> None:
    calls: list[str] = []

    async def fake_intake_once(prompt: str) -> dict[str, Any]:
        calls.append(prompt)
        if len(calls) == 1:
            return {
                "status": "need_clarification",
                "questions": [
                    "What is the target context (tissue/cell type/disease/model) for this TTN splice-switching ASO design (e.g., human skeletal muscle, cardiomyocytes/iPSC-CMs, DMD-like muscle model, or no specific context)?"
                ],
            }
        return {
            "status": "ready",
            "structured_goal": {
                "gene_symbol": "TTN",
                "species": "human",
                "intent": "splice_modulation",
                "target_context": "general human baseline context",
            },
        }

    monkeypatch.setattr(run_goal_intake, "_run_goal_intake_agent_once", fake_intake_once)

    intake = asyncio.run(
        run_goal_intake.run_goal_intake_step(
            "Design splice-switching ASOs for human TTN NM_001267550.2."
        )
    )

    assert intake["status"] == "ready"
    assert len(calls) == 2
    assert "No specific tissue/cell/model context; use general human baseline context." in calls[1]


def test_exon_only_splice_objective_seeds_default_splice_plan(monkeypatch) -> None:
    transcripts = [
        {
            "id": "ENSMUSTDEFAULT",
            "seq_region_name": "11",
            "start": 100000,
            "end": 120999,
            "strand": 1,
            "display_name": "Mapt-201",
            "description": "default Mapt transcript",
            "is_canonical": True,
            "length": 4000,
        }
    ]

    monkeypatch.setattr(run_goal_intake, "list_transcripts", lambda **kwargs: {"transcripts": transcripts})
    monkeypatch.setattr(run_goal_intake, "_generate_splice_strategy_analysis", lambda *args, **kwargs: {"status": "stubbed"})

    def fake_fetch_transcript_structure(transcript_id: str) -> dict[str, Any]:
        assert transcript_id == "ENSMUSTDEFAULT"
        exons = []
        for idx in range(1, 13):
            start = 100000 + idx * 1000
            end = start + 99
            exons.append(
                {
                    "index": idx,
                    "start": start,
                    "end": end,
                    "cdna_start": (idx - 1) * 100 + 1,
                    "cdna_end": idx * 100,
                    "coding_start": (idx - 1) * 100 + 1,
                    "coding_end": idx * 100,
                    "coding_start_genomic": start,
                    "coding_end_genomic": end,
                    "contains_coding": True,
                    "seq_region_name": "11",
                }
            )
        return {"seq_region_name": "11", "strand": 1, "exons_tx": exons}

    monkeypatch.setattr(run_goal_intake, "_fetch_transcript_structure", fake_fetch_transcript_structure)

    goal = {
        "gene_symbol": "Mapt",
        "species": "mouse",
        "gene_id": "ENSMUSG00000018411",
        "intent": "splice_modulation",
        "target_context": "mouse cortical neurons in a tauopathy model",
        "splice_isoform_objectives": [
            {"exon_label": "10", "direction": "increase_inclusion"}
        ],
    }
    final_payload = deepcopy(goal)

    assert run_goal_intake._should_seed_default_splice_context(goal) is True

    result = run_goal_intake._ensure_default_splice_context(goal, final_payload)

    assert result["resolved_transcript_strategy"]["mode"] == "default_splice_context"
    assert result["resolved_transcript_strategy"]["transcript_ids"] == ["ENSMUSTDEFAULT"]
    assert "ENSMUSTDEFAULT" in goal["isoform_transcripts"]
    seeded = [
        objective
        for objective in goal["splice_isoform_objectives"]
        if objective.get("isoform_label") == "ENSMUSTDEFAULT"
    ]
    assert seeded
    assert seeded[0]["desired_abundance"] == "increase"
    assert seeded[0]["exon_label"] == 10
    assert any(hint.get("isoform_label") == "ENSMUSTDEFAULT" for hint in goal["site_direction_hints"])
    assert result["splice_site_plan"]
    assert {10} == {
        int(value)
        for entry in result["splice_site_plan"]
        for value in (entry.get("exon_indices") or {}).values()
    }
    assert result["steric_regions"]["global_region"]["start"] is not None


def test_default_splice_context_honors_locked_refseq_transcript(monkeypatch: Any) -> None:
    transcripts = [
        {
            "id": "ENSTCANONICAL",
            "seq_region_name": "2",
            "start": 100000,
            "end": 450000,
            "strand": -1,
            "display_name": "TTN-canonical",
            "description": "canonical fallback transcript",
            "is_canonical": True,
            "length": 1000,
        },
        {
            "id": "ENSTLOCKED",
            "seq_region_name": "2",
            "start": 100000,
            "end": 450000,
            "strand": -1,
            "display_name": "TTN-locked",
            "description": "RefSeq locked transcript",
            "is_canonical": False,
            "length": 900,
        },
    ]

    monkeypatch.setattr(run_goal_intake, "list_transcripts", lambda **kwargs: {"transcripts": transcripts})
    monkeypatch.setattr(run_goal_intake, "_generate_splice_strategy_analysis", lambda *args, **kwargs: {"status": "stubbed"})

    def fake_resolve_external(accession: str, **kwargs: Any) -> dict[str, Any]:
        assert accession == "NM_001267550.2"
        return {
            "transcript_id": "ENSTLOCKED",
            "method": "test_refseq_xref",
            "confidence": "high",
            "rationale": "test resolved RefSeq accession",
        }

    def fake_fetch_transcript_structure(transcript_id: str) -> dict[str, Any]:
        assert transcript_id == "ENSTLOCKED"
        exons = []
        for idx in range(1, 331):
            start = 100000 + idx * 1000
            end = start + 99
            exons.append(
                {
                    "index": idx,
                    "start": start,
                    "end": end,
                    "cdna_start": (idx - 1) * 100 + 1,
                    "cdna_end": idx * 100,
                    "coding_start": (idx - 1) * 100 + 1,
                    "coding_end": idx * 100,
                    "coding_start_genomic": start,
                    "coding_end_genomic": end,
                    "contains_coding": True,
                    "seq_region_name": "2",
                }
            )
        return {"seq_region_name": "2", "strand": -1, "exons_tx": exons}

    monkeypatch.setattr(run_goal_intake, "resolve_external_transcript_accession", fake_resolve_external)
    monkeypatch.setattr(run_goal_intake, "_fetch_transcript_structure", fake_fetch_transcript_structure)

    goal = {
        "gene_symbol": "TTN",
        "species": "human",
        "intent": "splice_modulation",
        "target_context": "general human baseline context",
        "transcript_policy": "user_locked",
        "transcript_lock": {"kind": "reference_accession", "value": "NM_001267550.2"},
        "splice_isoform_objectives": [
            {
                "objective_scope": "default_exon_context",
                "exon_labels": [325, 326],
                "desired_effect": "skipping",
                "direction": "co_skipping",
                "desired_abundance": "decrease",
                "role": "beneficial",
            }
        ],
    }
    final_payload = deepcopy(goal)

    result = run_goal_intake._ensure_default_splice_context(goal, final_payload)

    assert result["resolved_transcript_strategy"]["transcript_ids"] == ["ENSTLOCKED"]
    assert "ENSTLOCKED" in goal["isoform_transcripts"]
    assert "ENSTCANONICAL" not in goal["isoform_transcripts"]
    assert result["isoform_transcripts"]["ENSTLOCKED"]["match_info"]["method"] == "test_refseq_xref"
    assert result["isoform_transcripts"]["ENSTLOCKED"]["match_info"]["reference_accession"] == "NM_001267550.2"
    assert result["splice_site_plan"]


def test_apply_transcript_confirmation_keeps_unasked_pending_items_clean(
    tmp_path: Path,
    monkeypatch: Any,
) -> None:
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    pending_items = [
        {
            "isoform_label": "TTN full-length exon 325 included",
            "suggested_transcript_id": "ENST00000589042",
            "method": "fallback",
            "confidence": "low",
            "rationale": "test pending item 1",
        },
        {
            "isoform_label": "TTN full-length exon 326 included",
            "suggested_transcript_id": "ENST00000589042",
            "method": "fallback",
            "confidence": "low",
            "rationale": "test pending item 2",
        },
    ]
    structured_goal = {
        "gene_symbol": "TTN",
        "species": "human",
        "intent": "splice_modulation",
        "pending_isoform_transcript_confirmations": deepcopy(pending_items),
    }
    pipeline_tools._save_goal_payload(
        str(run_dir),
        {"request": "Design TTN splice-switching ASOs."},
        deepcopy(structured_goal),
        deepcopy(structured_goal),
    )
    question = (
        "For isoform 'TTN full-length exon 325 included', should I use transcript ENST00000589042?"
    )
    state = {
        "status": "clarification_answers_ready",
        "phase": "clarification_answers_ready",
        "request": "Design TTN splice-switching ASOs.",
        "run_dir": str(run_dir),
        "clarification": {
            "kind": "isoform_transcript_confirmation",
            "pending_apply": True,
            "restart_stage": "resolution",
            "last_answers": [{"question": question, "answer": "accept"}],
        },
        "pending_questions": [],
        "clarification_qa_history": [],
        "force_rerun_stages": [],
    }
    state_path = tmp_path / "run_state.json"
    state_path.write_text(json.dumps(state), encoding="utf-8")

    def fake_interpreter(answer: Any, **kwargs: Any) -> dict[str, Any]:
        assert str(answer).strip() == "accept"
        return {
            "decision": "accept",
            "transcript_id": kwargs["suggested_transcript_id"],
            "rationale": "accepted",
        }

    monkeypatch.setattr(
        pipeline_tools,
        "interpret_transcript_confirmation_answer",
        fake_interpreter,
    )
    monkeypatch.setattr(
        pipeline_tools,
        "list_transcripts",
        lambda **_kwargs: {"transcripts": []},
    )

    result = pipeline_tools.apply_clarification_to_goal(str(state_path))
    saved_goal = json.loads(resolve_run_path(run_dir, "structured_goal.json").read_text(encoding="utf-8"))

    assert result["status"] == "ok"
    assert result["applied_transcript_confirmations"] == 1
    assert saved_goal["isoform_transcript_confirmations"][0]["isoform_label"] == pending_items[0]["isoform_label"]
    assert saved_goal["pending_isoform_transcript_confirmations"] == [pending_items[1]]
    assert "error" not in saved_goal["pending_isoform_transcript_confirmations"][0]


def test_binding_utr_transcript_context_does_not_require_splice_router(tmp_path: Path) -> None:
    run_dir = tmp_path / "binding_utr_run"
    run_dir.mkdir()
    goal = {
        "gene_symbol": "CDKN1B",
        "species": "human",
        "intent": "binding",
        "target_region": "utr3",
        "splice_isoform_objectives": [
            {
                "isoform_label": "NM_004064.5",
                "transcript_id": "NM_004064.5",
                "desired_abundance": "increase",
            }
        ],
        "design_programs": [
            {
                "program_id": "steric_20nt_target_site_blockers",
                "modality": "steric",
                "intent": "binding",
                "target_region": "utr3",
            }
        ],
    }
    payload = {
        "final_payload": {
            **deepcopy(goal),
            "structured_goal": deepcopy(goal),
        },
    }
    (run_dir / "goal_intake.json").write_text(json.dumps(payload), encoding="utf-8")

    assert pipeline_tools._resolve_run_splice_router(str(run_dir)) is False
    assert pipeline_tools._resolve_run_splice_eval(str(run_dir)) is False


def test_binding_interval_context_does_not_become_explicit_splice_preseed() -> None:
    goal = {
        "gene_symbol": "CDKN1B",
        "species": "human",
        "intent": "binding",
        "target_region": "utr3",
        "transcript_id": "NM_004064.5",
        "design_programs": [
            {
                "program_id": "steric_cdkn1b_01",
                "modality": "steric",
                "intent": "binding",
            }
        ],
        "region_constraints": [
            {
                "constraint_id": "pum1_eclip_sites_user_bed",
                "source": "user_provided_bed_intervals",
                "aso_overlap_requirement": "any_overlap",
                "bed_intervals": [
                    {"chrom": "chr12", "start": 12721313, "end": 12721416, "name": "PUM1_eCLIP_site_1"},
                ],
            }
        ],
    }
    request_text = (
        "Design target-site blocker ASOs for CDKN1B 3' UTR of transcript NM_004064.5. "
        "Use chr12:12721313-12721416 exactly; any overlap counts."
    )
    state = {
        "request": request_text,
        "design_objective_preseed": None,
        "clarification_qa_history": [],
    }

    preseed = pipeline_tools._prepare_design_objective_preseed(goal, {"transcript_id": "NM_004064.5"})

    assert preseed is not None
    assert "splice_isoform_objectives" not in preseed
    assert "evaluation_targets" not in preseed
    assert pipeline_tools._request_coordinate_preseed_if_available(state, goal) == {}


def test_binding_evaluation_targets_do_not_become_explicit_splice_targets() -> None:
    goal = {
        "intent": "binding",
        "run_splice_router": False,
        "run_splice_eval": False,
        "design_programs": [
            {
                "program_id": "primary",
                "modality": "steric",
                "intent": "binding",
                "run_splice_router": False,
                "run_splice_eval": False,
            }
        ],
        "evaluation_targets": [
            {
                "label": "supplied binding interval",
                "transcript_id": "NM_000001.1",
                "chrom": "chr1",
                "start": 100,
                "end": 140,
                "strand": "+",
                "desired_delta": "increase",
            }
        ],
    }

    assert run_goal_intake._evaluation_targets_to_explicit_splice_targets(goal) == []
    assert run_goal_intake._ensure_explicit_splice_target_context(goal, {}) == {}
    assert "explicit_splice_targets" not in goal


def test_knockdown_direction_does_not_seed_splice_isoform_objective() -> None:
    goal = {
        "intent": "knockdown",
        "design_programs": [{
            "program_id": "exploratory_gapmers", "intent": "knockdown",
            "modality": "gapmer", "desired_delta": "decrease",
        }],
    }

    preseed = pipeline_tools._prepare_design_objective_preseed(
        goal, {"transcript_id": "ENST00000252486.5"})

    assert not preseed or "splice_isoform_objectives" not in preseed


def test_scoped_design_attests_planner_transcript_without_user_lock(monkeypatch) -> None:
    transcript_chooser_agent = importlib.import_module("aso_agents.transcript_chooser_agent")
    monkeypatch.setattr(transcript_chooser_agent, "local_reference_gene_transcripts",
                        lambda *args, **kwargs: [{
                            "id": "ENST00000252486", "versioned_id": "ENST00000252486.9",
                        }])
    task = {"task_id": "design", "task_type": "design_new_asos", "context": {
        "gene_symbol": "APOE", "species": "human", "transcript_id": "ENST00000252486.5",
        "structured_goal": {"intent": "knockdown"},
    }}
    request = pipeline_tools._build_scoped_design_request_text(
        "Retrieve the versioned primary transcript yourself and design gapmers.",
        task, {"tasks": [task]})

    assert '"transcript_id": "ENST00000252486.9"' in request
    assert "ENST00000252486.5" not in request
    assert '"source": "installed_annotation"' in request
    assert '"user_locked": false' in request


def test_scoped_design_preserves_explicit_user_transcript_version(monkeypatch) -> None:
    transcript_chooser_agent = importlib.import_module("aso_agents.transcript_chooser_agent")
    monkeypatch.setattr(transcript_chooser_agent, "local_reference_gene_transcripts",
                        lambda *args, **kwargs: (_ for _ in ()).throw(AssertionError("not called")))
    task = {"task_id": "design", "task_type": "design_new_asos", "context": {
        "gene_symbol": "APOE", "species": "human", "transcript_id": "ENST00000252486.5",
    }}
    request = pipeline_tools._build_scoped_design_request_text(
        "Design against ENST00000252486.5 exactly.", task, {"tasks": [task]})

    assert '"transcript_id": "ENST00000252486.5"' in request


def test_submitted_set_identity_survives_evaluation_that_drops_csv_columns(tmp_path) -> None:
    path = tmp_path / "submitted.csv"
    path.write_text(
        "set_id,candidate_id,molecule_5to3,intended_chemistry\n"
        "A,A1,CGTAAAGACCCCTACAGCGC,full-PS gapmer\n"
        "B,B1,GCGCTGTAGGGGTCTTTACG,full-PS gapmer\n",
        encoding="utf-8",
    )
    evaluated = {"provided_asos": [
        {"name": "A1", "sequence": "CGTAAAGACCCCTACAGCGC"},
        {"name": "B1", "sequence": "GCGCTGTAGGGGTCTTTACG"},
    ]}
    context = {"artifact_refs": [{"path": str(path)}]}

    sets = pipeline_tools._collect_aso_compare_sets(
        state={}, task={}, task_context=context,
        request_envelope={}, evaluated_payload=evaluated,
    )

    assert [(item["set_id"], item["source_type"]) for item in sets] == [
        ("A", "provided"), ("B", "provided")]
    assert sets[0]["sequence_set"] == {"CGTAAAGACCCCTACAGCGC"}
    assert sets[1]["sequence_set"] == {"GCGCTGTAGGGGTCTTTACG"}


def test_design_objective_preseed_derives_isoform_objectives_from_evaluation_targets() -> None:
    preseed = pipeline_tools._prepare_design_objective_preseed(
        {
            "gene_symbol": "LMNA",
            "species": "human",
            "intent": "splice_modulation",
            "evaluation_targets": [
                {
                    "isoform_label": "progerin",
                    "desired_delta": "decrease",
                    "transcript_id": "ENST00000368299",
                },
                {
                    "isoform_label": "Lamin A (isoform defined by SnapGene/GenBank artifact)",
                    "desired_delta": "increase",
                },
            ],
        },
        {"artifact_refs": [{"path": "/tmp/reference.gb"}]},
    )

    objectives = preseed["splice_isoform_objectives"]
    assert objectives[0]["transcript_id"] == "ENST00000368299"
    assert objectives[0]["desired_abundance"] == "decrease"
    assert objectives[1]["desired_abundance"] == "increase"
    assert objectives[1]["snapgene_path"] == "/tmp/reference.gb"
    assert objectives[1]["source"] == "snapgene"


def test_design_objective_preseed_normalizes_if_possible_evaluation_targets() -> None:
    preseed = pipeline_tools._prepare_design_objective_preseed(
        {
            "gene_symbol": "LMNA",
            "species": "human",
            "intent": "splice_modulation",
            "evaluation_targets": [
                {
                    "isoform_label": "progerin",
                    "desired_delta": "decrease",
                    "transcript_id": "ENST00000368299",
                },
                {
                    "isoform_label": "laminA_from_snapgene_file",
                    "desired_delta": "increase",
                },
                {
                    "isoform_label": "Lamin C",
                    "desired_delta": "increase_if_possible",
                    "transcript_id": "ENST00000677389",
                },
            ],
        },
        {"artifact_refs": [{"path": "/tmp/reference.gb"}]},
    )

    objectives = preseed["splice_isoform_objectives"]
    assert [objective["desired_abundance"] for objective in objectives] == [
        "decrease",
        "increase",
    ]
    assert objectives[1]["source"] == "snapgene"
    assert objectives[1]["snapgene_path"] == "/tmp/reference.gb"
    assert all(
        objective.get("transcript_id") != "ENST00000677389"
        for objective in objectives
    )


def test_design_objective_preseed_derives_isoform_objectives_from_design_programs() -> None:
    preseed = pipeline_tools._prepare_design_objective_preseed(
        {
            "gene_symbol": "LMNA",
            "species": "human",
            "intent": "splice_modulation",
            "target_context": "splice modulation / isoform rebalancing",
            "design_programs": [
                {
                    "program_id": "p1_decrease",
                    "objective": "decrease_isoform",
                    "target_transcript_id": "ENST00000368299",
                    "modality": "steric",
                },
                {
                    "program_id": "p2_increase_artifact",
                    "objective": "increase_isoform",
                    "target_isoform_source": {"type": "artifact", "artifact_ref": "/tmp/reference.gb"},
                    "modality": "steric",
                },
                {
                    "program_id": "p3_increase_laminc",
                    "objective": "increase_isoform",
                    "isoform_label": "Lamin C",
                    "modality": "steric",
                },
            ],
        },
        {},
    )

    objectives = preseed["splice_isoform_objectives"]
    assert [objective["desired_abundance"] for objective in objectives] == ["decrease", "increase", "increase"]
    assert objectives[0]["transcript_id"] == "ENST00000368299"
    assert objectives[1]["isoform_label"] == "SnapGene/GenBank artifact isoform"
    assert objectives[1]["snapgene_path"] == "/tmp/reference.gb"
    assert objectives[1]["source"] == "snapgene"
    assert preseed["structured_goal_overlay"]["intent"] == "splice_modulation"


def test_design_objective_preseed_does_not_turn_parameter_variants_into_splice_outcomes() -> None:
    goal = {
        "gene_symbol": "GENE1",
        "species": "human",
        "intent": "knockdown",
        "design_programs": [
            {
                "program_id": "length_variant_1",
                "modality": "gapmer",
                "intent": "knockdown",
                "objective": "Decrease the requested target.",
                "desired_delta": "decrease",
                "isoform_label": "selected transcript",
                "length": 16,
            },
            {
                "program_id": "length_variant_2",
                "modality": "gapmer",
                "intent": "knockdown",
                "objective": "Decrease the requested target.",
                "desired_delta": "decrease",
                "isoform_label": "selected transcript",
                "length": 17,
            },
        ],
    }

    normalized = pipeline_tools._normalized_design_objective_preseed(goal)
    finalized = pipeline_tools._finalize_design_task_contract(
        {
            "task_id": "design",
            "task_type": "design_new_asos",
            "context": {"structured_goal": goal},
        },
        {},
    )

    assert "splice_isoform_objectives" not in normalized
    assert (
        "splice_isoform_objectives"
        not in finalized["context"]["structured_goal"]
    )


def test_design_objective_preseed_normalizes_isoform_modulation_objectives() -> None:
    preseed = pipeline_tools._prepare_design_objective_preseed(
        {
            "gene_symbol": "LMNA",
            "species": "human",
            "intent": "binding",
            "target_context": "splice modulation and isoform rebalancing",
            "isoform_modulation_objectives": [
                {
                    "isoform_label": "progerin",
                    "desired_delta": "decrease",
                    "transcript_id": "ENST00000368299",
                },
                {
                    "isoform_label": "SnapGene artifact isoform",
                    "desired_delta": "increase",
                },
                {
                    "isoform_label": "Lamin C",
                    "desired_delta": "increase",
                    "transcript_id": "ENST00000677389",
                },
            ],
        },
        {"artifact_refs": [{"path": "/tmp/reference.gb"}]},
    )

    objectives = preseed["splice_isoform_objectives"]
    assert [objective["desired_abundance"] for objective in objectives] == ["decrease", "increase", "increase"]
    assert objectives[0]["role"] == "context"
    assert objectives[1]["source"] == "snapgene"
    assert objectives[1]["snapgene_path"] == "/tmp/reference.gb"
    assert objectives[2]["transcript_id"] == "ENST00000677389"
    assert preseed["structured_goal_overlay"]["intent"] == "splice_modulation"


def test_design_objective_preseed_preserves_artifact_only_isoform_objective() -> None:
    preseed = pipeline_tools._prepare_design_objective_preseed(
        {
            "gene_symbol": "LMNA",
            "species": "human",
            "target_context": "splice-modulating isoform switching",
            "isoform_objectives": [
                {
                    "transcript_id": "ENST00000368299",
                    "direction": "decrease",
                    "priority": "top",
                },
                {
                    "source": "snapgene_artifact",
                    "artifact_ref_task_id": "t1_analyze_snapgene_isoform",
                    "direction": "increase",
                    "priority": "high",
                },
                {
                    "transcript_id": "ENST00000677389",
                    "direction": "increase",
                    "priority": "if_possible",
                },
            ],
        },
        {"artifact_refs": [{"path": "/tmp/reference.gb"}]},
    )

    objectives = preseed["splice_isoform_objectives"]
    assert [objective["desired_abundance"] for objective in objectives] == ["decrease", "increase", "increase"]
    assert objectives[0]["isoform_label"] == "ENST00000368299"
    assert objectives[1]["isoform_label"] == "SnapGene/GenBank artifact isoform"
    assert objectives[1]["source"] == "snapgene"
    assert objectives[1]["snapgene_path"] == "/tmp/reference.gb"
    assert objectives[2]["isoform_label"] == "ENST00000677389"


def test_design_objective_preseed_uses_prioritized_objectives_with_artifact_ref() -> None:
    preseed = pipeline_tools._prepare_design_objective_preseed(
        {
            "gene_symbol": "LMNA",
            "species": "human",
            "target_context": "splice-modulating isoform switching",
            "prioritized_objectives": [
                {
                    "priority": 1,
                    "objective": "decrease_isoform",
                    "transcript_id": "ENST00000368299",
                    "isoform_label": "progerin",
                },
                {
                    "priority": 2,
                    "objective": "increase_isoform",
                    "artifact_ref": "/tmp/reference.gb",
                    "isoform_label": "Lamin A (from file)",
                },
                {
                    "priority": 3,
                    "objective": "increase_isoform",
                    "transcript_id": "ENST00000677389",
                    "isoform_label": "Lamin C",
                },
            ],
        },
        {"artifact_refs": [{"path": "/tmp/reference.gb"}]},
    )

    objectives = preseed["splice_isoform_objectives"]
    assert [objective["desired_abundance"] for objective in objectives] == ["decrease", "increase", "increase"]
    assert objectives[0]["isoform_label"] == "progerin"
    assert objectives[1]["isoform_label"] == "Lamin A (from file)"
    assert objectives[1]["source"] == "snapgene"
    assert objectives[1]["snapgene_path"] == "/tmp/reference.gb"
    assert objectives[2]["transcript_id"] == "ENST00000677389"


def test_design_objective_preseed_uses_isoform_targets_with_source_artifact_task() -> None:
    preseed = pipeline_tools._prepare_design_objective_preseed(
        {
            "gene_symbol": "LMNA",
            "species": "human",
            "target_context": "splice-modulating isoform switching",
            "isoform_targets": [
                {
                    "isoform_label": "progerin",
                    "transcript_id": "ENST00000368299",
                    "direction": "decrease",
                    "priority": "highest",
                },
                {
                    "isoform_label": "laminA_from_snapgene",
                    "direction": "increase",
                    "priority": "high",
                    "source_artifact_task": "t1_analyze_snapgene_isoform",
                },
                {
                    "isoform_label": "laminC",
                    "transcript_id": "ENST00000677389",
                    "direction": "increase",
                    "priority": "medium",
                },
            ],
        },
        {"artifact_refs": [{"path": "/tmp/reference.gb"}]},
    )

    objectives = preseed["splice_isoform_objectives"]
    assert [objective["desired_abundance"] for objective in objectives] == ["decrease", "increase", "increase"]
    assert objectives[1]["isoform_label"] == "laminA_from_snapgene"
    assert objectives[1]["source"] == "snapgene"
    assert objectives[1]["snapgene_path"] == "/tmp/reference.gb"


def test_lmna_isoform_objectives_surface_shared_and_unique_routes() -> None:
    isoform_transcripts = {
        "progerin": {
            "chrom": "chr1",
            "strand": "+",
            "sites": [
                {
                    "exon_index": 10,
                    "chrom": "chr1",
                    "strand": "+",
                    "acceptor": 156137654,
                    "donor": 156137743,
                    "acceptor_plan_ok": True,
                    "donor_plan_ok": True,
                    "coding": True,
                },
                {
                    "exon_index": 11,
                    "chrom": "chr1",
                    "strand": "+",
                    "acceptor": 156138488,
                    "donor": 156138607,
                    "acceptor_plan_ok": True,
                    "donor_plan_ok": True,
                    "coding": True,
                },
                {
                    "exon_index": 12,
                    "chrom": "chr1",
                    "strand": "+",
                    "acceptor": 156139080,
                    "donor": 156139317,
                    "acceptor_plan_ok": True,
                    "donor_plan_ok": True,
                    "coding": True,
                },
            ],
        },
        "Lamin A": {
            "chrom": "chr1",
            "strand": "+",
            "sites": [
                {
                    "exon_index": 10,
                    "chrom": "chr1",
                    "strand": "+",
                    "acceptor": 156137654,
                    "donor": 156137743,
                    "acceptor_plan_ok": True,
                    "donor_plan_ok": True,
                    "coding": True,
                },
                {
                    "exon_index": 11,
                    "chrom": "chr1",
                    "strand": "+",
                    "acceptor": 156138488,
                    "donor": 156138757,
                    "acceptor_plan_ok": True,
                    "donor_plan_ok": True,
                    "coding": True,
                },
                {
                    "exon_index": 12,
                    "chrom": "chr1",
                    "strand": "+",
                    "acceptor": 156139080,
                    "donor": 156140081,
                    "acceptor_plan_ok": True,
                    "donor_plan_ok": True,
                    "coding": True,
                },
            ],
        },
        "Lamin C": {
            "chrom": "chr1",
            "strand": "+",
            "sites": [
                {
                    "exon_index": 10,
                    "chrom": "chr1",
                    "strand": "+",
                    "acceptor": 156137654,
                    "donor": 156137865,
                    "acceptor_plan_ok": True,
                    "donor_plan_ok": True,
                    "coding": True,
                }
            ],
        },
    }
    objectives = [
        {"isoform_label": "progerin", "desired_abundance": "decrease", "role": "pathogenic"},
        {"isoform_label": "Lamin A", "desired_abundance": "increase", "role": "beneficial"},
        {"isoform_label": "Lamin C", "desired_abundance": "increase", "role": "beneficial"},
    ]
    structured_goal = {
        "gene_symbol": "LMNA",
        "species": "human",
        "target_context": "human fibroblasts",
        "intent": "splice_modulation",
        "splice_isoform_objectives": objectives,
        "isoform_transcripts": isoform_transcripts,
    }

    splice_plan, _comparisons = run_goal_intake._compute_splice_site_plan(
        objectives,
        isoform_transcripts,
        structured_goal,
    )
    structured_goal["splice_site_plan"] = splice_plan
    programs, _meta = pipeline_tools._derive_exploratory_steric_programs(
        structured_goal=structured_goal,
        base_program={
            "program_id": "steric_lmna_routes",
            "label": "LMNA isoform route scan",
            "modality": "steric",
            "run_splice_eval": True,
            "run_splice_router": True,
        },
    )

    assert programs
    route_labels = set(programs[0]["routing_candidate_sites"])
    assert {
        "donor@156137743",
        "acceptor@156138488",
        "donor@156138607",
        "donor@156138757",
        "acceptor@156139080",
    }.issubset(route_labels)


def test_design_objective_preseed_uses_nested_isoform_modulation_dict() -> None:
    preseed = pipeline_tools._prepare_design_objective_preseed(
        {
            "gene_symbol": "LMNA",
            "species": "human",
            "target_context": "splice-modulating isoform switching",
            "isoform_modulation": {
                "primary_decrease": {
                    "transcript_id": "ENST00000368299",
                    "isoform_label": "progerin",
                },
                "secondary_increase": {
                    "source": "artifact",
                    "artifact_ref": "/tmp/reference.gb",
                    "isoform_label": "laminA_from_snapgene_file",
                },
                "tertiary_increase_if_possible": {
                    "transcript_id": "ENST00000677389",
                    "isoform_label": "Lamin C",
                },
            },
        },
        {"artifact_refs": [{"path": "/tmp/reference.gb"}]},
    )

    objectives = preseed["splice_isoform_objectives"]
    assert [objective["desired_abundance"] for objective in objectives] == ["decrease", "increase", "increase"]
    assert objectives[1]["source"] == "snapgene"
    assert objectives[1]["snapgene_path"] == "/tmp/reference.gb"


def test_task_context_merges_default_and_task_specific_context() -> None:
    context = pipeline_tools._task_context_from_request_envelope(
        {
            "default_context": {
                "artifact_refs": [{"path": "/tmp/reference.gb"}],
                "species": "human",
            },
            "gene_symbol": "LMNA",
        },
        {
            "task_type": "design_new_asos",
            "context": {
                "target_context": "fibroblast",
            },
        },
    )

    assert context["artifact_refs"] == [{"path": "/tmp/reference.gb"}]
    assert context["species"] == "human"
    assert context["gene_symbol"] == "LMNA"
    assert context["target_context"] == "fibroblast"


def test_design_preseed_uses_nested_program_objectives_and_dependency_artifacts() -> None:
    request_envelope = {
        "default_context": {
            "gene_symbol": "LMNA",
            "species": "human",
            "target_context": "fibroblast splice modulation",
        },
        "tasks": [
            {
                "task_id": "t1_analyze_reference",
                "task_type": "analyze_artifact",
                "context": {
                    "artifact_refs": [
                        {
                            "path": "/tmp/reference.gb",
                            "artifact_type": "snapgene_genbank",
                        }
                    ]
                },
            },
            {
                "task_id": "t2_design",
                "task_type": "design_new_asos",
                "depends_on": ["t1_analyze_reference"],
                "context": {
                    "structured_goal": {
                        "intent": "binding",
                        "design_programs": [
                            {
                                "program_id": "p1",
                                "modality": "steric",
                                "objectives": [
                                    {
                                        "type": "decrease_isoform",
                                        "isoform_label": "progerin",
                                        "transcript_id": "ENST00000368299",
                                    },
                                    {
                                        "type": "increase_isoform",
                                        "isoform_label": "Lamin A (from SnapGene file)",
                                        "source": {"artifact_task_id": "t1_analyze_reference"},
                                    },
                                    {
                                        "type": "increase_isoform_if_possible",
                                        "isoform_label": "Lamin C",
                                        "transcript_id": "ENST00000677389",
                                    },
                                ],
                            }
                        ],
                    }
                },
            },
        ],
    }
    task = request_envelope["tasks"][1]

    context = pipeline_tools._task_context_from_request_envelope(request_envelope, task)
    preseed = pipeline_tools._prepare_design_objective_preseed(
        task["context"]["structured_goal"],
        context,
    )

    objectives = preseed["splice_isoform_objectives"]
    assert [objective["desired_abundance"] for objective in objectives] == ["decrease", "increase", "increase"]
    assert objectives[0]["transcript_id"] == "ENST00000368299"
    assert objectives[1]["source"] == "snapgene"
    assert objectives[1]["snapgene_path"] == "/tmp/reference.gb"
    assert objectives[2]["transcript_id"] == "ENST00000677389"
    assert preseed["structured_goal_overlay"]["intent"] == "splice_modulation"


def test_design_preseed_infers_transcript_objective_from_task_context() -> None:
    preseed = pipeline_tools._prepare_design_objective_preseed(
        {
            "gene_symbol": "LMNA",
            "species": "human",
            "target_context": "splice/isoform modulation to reduce target isoform abundance",
            "design_programs": [
                {
                    "modality": "steric",
                    "objective": "decrease_target_isoform",
                }
            ],
        },
        {
            "transcript_id": "ENST00000368299",
            "target_context": "decrease progerin isoform",
        },
    )

    objectives = preseed["splice_isoform_objectives"]
    assert objectives == [
        {
            "isoform_label": "ENST00000368299",
            "transcript_id": "ENST00000368299",
            "desired_abundance": "decrease",
            "role": "context",
        }
    ]


def test_design_preseed_infers_artifact_objective_from_task_context() -> None:
    preseed = pipeline_tools._prepare_design_objective_preseed(
        {
            "gene_symbol": "LMNA",
            "species": "human",
            "target_context": "increase isoform represented by provided SnapGene file",
            "design_programs": [
                {
                    "modality": "steric",
                    "objective": "increase_target_isoform",
                }
            ],
        },
        {
            "artifact_refs": [{"path": "/tmp/reference.gb"}],
            "target_context": "increase the isoform/construct represented by the provided file",
        },
    )

    objectives = preseed["splice_isoform_objectives"]
    assert objectives == [
        {
            "isoform_label": "SnapGene/GenBank artifact isoform",
            "desired_abundance": "increase",
            "role": "context",
            "snapgene_path": "/tmp/reference.gb",
            "source": "snapgene",
        }
    ]


def test_large_composite_graph_schedules_designs_then_downstream_steps(tmp_path: Path) -> None:
    state = _state_from_raw_envelope(
        {
            "schema_version": "v3",
            "request_text": (
                "Design human and mouse MAPT gapmers, evaluate my provided ASOs, compare the "
                "generated sets, trace where my ASOs drop out, and write one synthesis."
            ),
            "request_mode": "composite",
            "shared_inputs": {
                "provided_asos": [
                    {"name": "lab_aso_1", "sequence": "ACGTACGTACGTACGTACGT"},
                    {"name": "lab_aso_2", "sequence": "TGCATGCATGCATGCATGCA"},
                ]
            },
            "tasks": [
                {
                    "task_id": "design_human_mapt",
                    "task_type": "design_new_asos",
                    "label": "Design human MAPT gapmers",
                    "can_run_in_parallel": True,
                        "context": {
                            "gene_symbol": "MAPT",
                            "species": "human",
                            "target_context": "human cortical neurons in tauopathy model",
                            "structured_goal": {
                                "gene_symbol": "MAPT",
                                "species": "human",
                                "target_context": "human cortical neurons in tauopathy model",
                                "intent": "knockdown",
                                "design_programs": [{"modality": "gapmer"}],
                            },
                    },
                },
                {
                    "task_id": "design_mouse_mapt",
                    "task_type": "design_new_asos",
                    "label": "Design mouse Mapt gapmers",
                    "can_run_in_parallel": True,
                        "context": {
                            "gene_symbol": "Mapt",
                            "species": "mouse",
                            "target_context": "mouse cortical neurons in tauopathy model",
                            "structured_goal": {
                                "gene_symbol": "Mapt",
                                "species": "mouse",
                                "target_context": "mouse cortical neurons in tauopathy model",
                                "intent": "knockdown",
                                "design_programs": [{"modality": "gapmer"}],
                            },
                    },
                },
                {
                    "task_id": "evaluate_lab_asos",
                    "task_type": "evaluate_existing_aso_set",
                    "label": "Evaluate provided ASOs",
                    "context": {},
                },
                {
                    "task_id": "compare_generated_sets",
                    "task_type": "compare_aso_sets",
                    "label": "Compare generated human and mouse sets",
                    "context": {},
                },
                {
                    "task_id": "trace_lab_asos",
                    "task_type": "trace_filter_decisions",
                    "label": "Trace provided ASO filtering",
                    "context": {},
                },
            ],
        },
        tmp_path,
    )

    envelope = state["request_envelope"]
    task_by_id = {task["task_id"]: task for task in envelope["tasks"]}
    assert "task_build_composite_report" in task_by_id
    assert task_by_id["evaluate_lab_asos"]["context"]["provided_asos"][0]["name"] == "lab_aso_1"
    assert task_by_id["trace_lab_asos"]["context"]["provided_asos"][1]["sequence"] == "TGCATGCATGCATGCATGCA"
    assert set(task_by_id["task_build_composite_report"]["depends_on"]) == {
        "design_human_mapt",
        "design_mouse_mapt",
        "evaluate_lab_asos",
        "compare_generated_sets",
        "trace_lab_asos",
    }

    validation = pipeline_tools._request_validation_summary(envelope)
    assert validation["blocked_task_ids"] == []
    assert validation["orphan_blocked_task_ids"] == []

    assert set(pipeline_tools._runnable_task_ids_for_state(state)) == {
        "design_human_mapt",
        "design_mouse_mapt",
    }
    planning = pipeline_tools._derive_next_actions(state)
    assert planning["terminal_status"] == "in_progress"
    assert planning["runnable_task_ids"] == ["design_human_mapt", "design_mouse_mapt"]
    assert _next_tool(planning) == "tool_execute_design_task"

    _mark_completed(state, "design_human_mapt", "design_mouse_mapt")
    assert pipeline_tools._runnable_task_ids_for_state(state) == ["evaluate_lab_asos"]

    _mark_completed(state, "evaluate_lab_asos")
    assert pipeline_tools._runnable_task_ids_for_state(state) == ["compare_generated_sets"]

    _mark_completed(state, "compare_generated_sets")
    assert pipeline_tools._runnable_task_ids_for_state(state) == ["trace_lab_asos"]

    _mark_completed(state, "trace_lab_asos")
    assert pipeline_tools._runnable_task_ids_for_state(state) == ["task_build_composite_report"]


def test_generated_set_comparison_without_provided_asos_is_runnable(tmp_path: Path) -> None:
    state = _state_from_raw_envelope(
        {
            "schema_version": "v3",
            "request_text": "Design human and mouse MAPT ASOs and compare the generated sets.",
            "request_mode": "composite",
            "tasks": [
                {
                    "task_id": "design_human",
                    "task_type": "design_new_asos",
                    "context": {
                        "gene_symbol": "MAPT",
                        "species": "human",
                        "structured_goal": {
                            "gene_symbol": "MAPT",
                            "species": "human",
                            "intent": "knockdown",
                            "design_programs": [{"modality": "gapmer"}],
                        },
                    },
                },
                {
                    "task_id": "design_mouse",
                    "task_type": "design_new_asos",
                    "context": {
                        "gene_symbol": "Mapt",
                        "species": "mouse",
                        "structured_goal": {
                            "gene_symbol": "Mapt",
                            "species": "mouse",
                            "intent": "knockdown",
                            "design_programs": [{"modality": "gapmer"}],
                        },
                    },
                },
                {"task_id": "compare_generated", "task_type": "compare_aso_sets", "context": {}},
            ],
        },
        tmp_path,
    )

    validation = pipeline_tools._request_validation_summary(state["request_envelope"])
    compare_validation = validation["tasks"]["compare_generated"]
    assert compare_validation["status"] == "runnable"
    assert "provided_asos_missing" not in compare_validation["issues"]

    _mark_completed(state, "design_human", "design_mouse")
    assert pipeline_tools._runnable_task_ids_for_state(state) == ["compare_generated"]


def test_same_gene_isoform_design_tasks_merge_before_execution() -> None:
    raw_envelope = {
        "schema_version": "v3",
        "request_text": "Decrease progerin, increase a SnapGene LMNA isoform, and increase Lamin C.",
        "request_mode": "composite",
        "tasks": [
            {
                "task_id": "parse_snapgene",
                "task_type": "analyze_artifact",
                "context": {
                    "artifact_refs": [
                        {"path": "/tmp/lmna_isoform.gb", "type": "snapgene_genbank"}
                    ]
                },
            },
            {
                "task_id": "design_progerin",
                "task_type": "design_new_asos",
                "context": {
                    "gene_symbol": "LMNA",
                    "species": "human",
                    "target_context": "fibroblasts",
                        "structured_goal": {
                            "gene_symbol": "LMNA",
                            "species": "human",
                            "intent": "splice_modulation",
                            "design_programs": [
                                {
                                    "program_id": "decrease_progerin",
                                    "label": "decrease progerin",
                                    "modality": "steric",
                                    "intent": "splice_modulation",
                                    "objective": "decrease_target_isoform_expression_or_splicing",
                                    "desired_delta": "decrease",
                                    "transcript_id": "ENST00000368299",
                                    "isoform_label": "progerin",
                                }
                            ],
                            "splice_isoform_objectives": [
                                {
                                    "objective_scope": "specific_transcript",
                                    "transcript_id": "ENST00000368299",
                                    "isoform_label": "progerin",
                                    "desired_abundance": "decrease",
                                    "role": "pathogenic",
                                }
                            ],
                    },
                },
            },
            {
                "task_id": "design_snapgene_isoform",
                "task_type": "design_new_asos",
                "depends_on": ["parse_snapgene"],
                "context": {
                    "gene_symbol": "LMNA",
                    "species": "human",
                    "target_context": "fibroblasts",
                    "artifact_refs": [
                        {"path": "/tmp/lmna_isoform.gb", "type": "snapgene_genbank"}
                    ],
                        "structured_goal": {
                            "gene_symbol": "LMNA",
                            "species": "human",
                            "intent": "splice_modulation",
                            "design_programs": [
                                {
                                    "program_id": "increase_supplied_isoform",
                                    "label": "increase supplied isoform",
                                    "modality": "steric",
                                    "intent": "splice_modulation",
                                    "objective": "increase_target_isoform_expression_or_splicing",
                                    "desired_delta": "increase",
                                    "isoform_label": "snapgene_isoform",
                                }
                            ],
                            "splice_isoform_objectives": [
                                {
                                    "objective_scope": "specific_isoform",
                                    "isoform_label": "snapgene_isoform",
                                    "desired_abundance": "increase",
                                    "role": "beneficial",
                                    "snapgene_path": "/tmp/lmna_isoform.gb",
                                }
                            ],
                    },
                },
            },
            {
                "task_id": "design_lamin_c",
                "task_type": "design_new_asos",
                "context": {
                    "gene_symbol": "LMNA",
                    "species": "human",
                    "target_context": "fibroblasts",
                        "structured_goal": {
                            "gene_symbol": "LMNA",
                            "species": "human",
                            "intent": "splice_modulation",
                            "design_programs": [
                                {
                                    "program_id": "increase_lamin_c",
                                    "label": "increase Lamin C",
                                    "modality": "steric",
                                    "intent": "splice_modulation",
                                    "objective": "increase_target_isoform_expression_or_splicing",
                                    "desired_delta": "increase",
                                    "transcript_id": "ENST00000677389",
                                    "isoform_label": "Lamin C",
                                }
                            ],
                            "splice_isoform_objectives": [
                                {
                                    "objective_scope": "specific_transcript",
                                    "transcript_id": "ENST00000677389",
                                    "isoform_label": "Lamin C",
                                    "desired_abundance": "increase",
                                    "role": "beneficial",
                                }
                            ],
                    },
                },
            },
            {"task_id": "report", "task_type": "build_composite_report"},
        ],
    }

    envelope = pipeline_tools._build_request_envelope(
        request_text=raw_envelope["request_text"],
        request_id="lmna_isoform_merge",
        structured_goal={},
        raw_envelope=raw_envelope,
    )

    design_tasks = [task for task in envelope["tasks"] if task["task_type"] == "design_new_asos"]
    assert len(design_tasks) == 1
    design_task = design_tasks[0]
    assert design_task["depends_on"] == ["parse_snapgene"]
    assert design_task["merged_from_task_ids"] == [
        "design_progerin",
        "design_snapgene_isoform",
        "design_lamin_c",
    ]
    goal = design_task["context"]["structured_goal"]
    assert goal["intent"] == "splice_modulation"
    assert "transcript_id" not in goal
    assert [program["program_id"] for program in goal["design_programs"]] == [
        "decrease_progerin",
        "increase_supplied_isoform",
        "increase_lamin_c",
    ]
    objectives = goal["splice_isoform_objectives"]
    assert [(item["isoform_label"], item["desired_abundance"]) for item in objectives] == [
        ("progerin", "decrease"),
        ("snapgene_isoform", "increase"),
        ("Lamin C", "increase"),
    ]
    assert objectives[1]["snapgene_path"] == "/tmp/lmna_isoform.gb"
    report_task = next(task for task in envelope["tasks"] if task["task_type"] == "build_composite_report")
    assert report_task["depends_on"] == ["parse_snapgene", "design_progerin"]


def test_snapgene_isoform_variants_do_not_generate_patient_reference_by_default(
    tmp_path: Path,
    monkeypatch,
) -> None:
    reference = tmp_path / "reference.fa"
    reference.write_text(">chr1\nACGT\n", encoding="utf-8")
    structured_goal = {"gene_symbol": "LMNA", "species": "human"}
    final_payload: dict[str, Any] = {}
    annotations = [
        {
            "reference_fasta_path": str(reference),
            "patient_variants": [{"chrom": "chr1", "pos": 2, "ref": "C", "alt": "T"}],
        }
    ]

    def fail_patient_reference_write(*_args: Any, **_kwargs: Any) -> None:
        raise AssertionError("whole-genome patient FASTA should not be generated for isoform-only SnapGene use")

    monkeypatch.setattr(run_goal_intake, "_write_patient_reference", fail_patient_reference_write)

    result = run_goal_intake._apply_patient_reference_from_snapgene_annotations(
        structured_goal,
        final_payload,
        str(tmp_path),
        annotations,
    )

    evidence = structured_goal["snapgene_variant_evidence"]
    assert result is None
    assert "patient_reference" not in structured_goal
    assert evidence["patient_reference_fasta_generated"] is False
    assert Path(evidence["vcf_path"]).exists()
    assert final_payload["snapgene_variant_evidence"] == evidence
    assert annotations[0]["snapgene_variant_evidence"] == evidence


def test_snapgene_patient_variant_request_can_generate_patient_reference(
    tmp_path: Path,
    monkeypatch,
) -> None:
    reference = tmp_path / "reference.fa"
    reference.write_text(">chr1\nACGT\n", encoding="utf-8")
    structured_goal = {
        "gene_symbol": "LMNA",
        "species": "human",
        "patient_variant": {"chr": "chr1", "pos": 2, "ref": "C", "alt": "T"},
    }
    final_payload: dict[str, Any] = {}
    annotations = [
        {
            "reference_fasta_path": str(reference),
            "patient_variants": [{"chrom": "chr1", "pos": 2, "ref": "C", "alt": "T"}],
        }
    ]

    def write_stub(_reference_path: Path, _variants: list[dict[str, Any]], output_path: Path, **_kwargs: Any) -> None:
        output_path.write_text(">patient\nATGT\n", encoding="utf-8")

    monkeypatch.setattr(run_goal_intake, "_write_patient_reference", write_stub)

    result = run_goal_intake._apply_patient_reference_from_snapgene_annotations(
        structured_goal,
        final_payload,
        str(tmp_path),
        annotations,
    )

    assert result is not None
    assert result["fasta_path"].endswith("patient_reference.fa")
    assert Path(result["fasta_path"]).exists()
    assert structured_goal["patient_reference"] == result
    assert final_payload["patient_reference"] == result
    assert annotations[0]["patient_reference"] == result


def test_prepare_snapgene_objectives_skips_patient_variant_derivation_for_isoform_reference(
    tmp_path: Path,
    monkeypatch,
) -> None:
    snapgene_path = tmp_path / "reference.gb"
    snapgene_path.write_text("LOCUS       TEST                     12 bp    DNA\nORIGIN\n        1 acgtacgtacgt\n//\n", encoding="utf-8")
    structured_goal = {
        "gene_symbol": "LMNA",
        "species": "human",
        "splice_isoform_objectives": [
            {
                "isoform_label": f"snapgene:{snapgene_path}",
                "desired_abundance": "increase",
                "role": "beneficial",
            }
        ],
    }
    final_payload: dict[str, Any] = {}
    derive_values: list[bool] = []

    def fake_ingest_snapgene_file(_path: str, **kwargs: Any) -> dict[str, Any]:
        derive_values.append(bool(kwargs.get("derive_patient_variants")))
        return {
            "snapgene_path": str(snapgene_path),
            "exons": [
                {
                    "order": 1,
                    "chrom": "chr1",
                    "genomic_start": 1,
                    "genomic_end": 12,
                    "strand": "+",
                    "sequence": "ACGTACGTACGT",
                }
            ],
            "patient_variants": [{"chrom": "chr1", "pos": 2, "ref": "C", "alt": "T"}],
        }

    monkeypatch.setattr(run_goal_intake, "ingest_snapgene_file", fake_ingest_snapgene_file)
    monkeypatch.setattr(
        run_goal_intake,
        "_build_snapgene_transcript_entry",
        lambda annotation, label: {
            "transcript_id": "SNAPGENE::reference",
            "source": "snapgene",
            "snapgene_path": annotation["snapgene_path"],
            "chrom": "chr1",
            "strand": "+",
            "exons_tx": annotation["exons"],
        },
    )

    result = run_goal_intake._prepare_snapgene_objectives(
        structured_goal,
        final_payload,
        str(tmp_path),
        user_request=f"Design ASOs using {snapgene_path} as an isoform reference.",
    )

    assert derive_values == [False]
    assert result["preseeded_isoforms"] == ["SnapGene isoform (reference.gb)"]
    assert "patient_reference" not in structured_goal


def test_prepare_snapgene_objectives_derives_patient_variants_for_patient_reference_request(
    tmp_path: Path,
    monkeypatch,
) -> None:
    snapgene_path = tmp_path / "reference.gb"
    snapgene_path.write_text("LOCUS       TEST                     12 bp    DNA\nORIGIN\n        1 acgtacgtacgt\n//\n", encoding="utf-8")
    structured_goal = {
        "gene_symbol": "LMNA",
        "species": "human",
        "patient_variant": {"chr": "chr1", "pos": 2, "ref": "C", "alt": "T"},
        "splice_isoform_objectives": [
            {
                "isoform_label": f"snapgene:{snapgene_path}",
                "desired_abundance": "increase",
                "role": "beneficial",
            }
        ],
    }
    final_payload: dict[str, Any] = {}
    derive_values: list[bool] = []

    def fake_ingest_snapgene_file(_path: str, **kwargs: Any) -> dict[str, Any]:
        derive_values.append(bool(kwargs.get("derive_patient_variants")))
        return {
            "snapgene_path": str(snapgene_path),
            "exons": [
                {
                    "order": 1,
                    "chrom": "chr1",
                    "genomic_start": 1,
                    "genomic_end": 12,
                    "strand": "+",
                    "sequence": "ACGTACGTACGT",
                }
            ],
            "patient_variants": [],
        }

    monkeypatch.setattr(run_goal_intake, "ingest_snapgene_file", fake_ingest_snapgene_file)
    monkeypatch.setattr(
        run_goal_intake,
        "_build_snapgene_transcript_entry",
        lambda annotation, label: {
            "transcript_id": "SNAPGENE::reference",
            "source": "snapgene",
            "snapgene_path": annotation["snapgene_path"],
            "chrom": "chr1",
            "strand": "+",
            "exons_tx": annotation["exons"],
        },
    )

    run_goal_intake._prepare_snapgene_objectives(
        structured_goal,
        final_payload,
        str(tmp_path),
        user_request=f"Design against patient reference {snapgene_path}.",
    )

    assert derive_values == [True]


def test_same_gene_splice_merge_preserves_scoped_site_programs() -> None:
    acceptor_program = pipeline_tools._splice_relative_program_spec(
        exon_label=7,
        transcript_id="ENST00000380743",
        site="acceptor",
        rel_start=-25,
        rel_end=24,
        label="exon7_acceptor_window",
        candidate_limit=30,
        target_context="spinal cord motor neurons",
    )
    donor_program = pipeline_tools._splice_relative_program_spec(
        exon_label=7,
        transcript_id="ENST00000380743",
        site="donor",
        rel_start=-9,
        rel_end=60,
        label="exon7_donor_window",
        candidate_limit=30,
        target_context="spinal cord motor neurons",
    )
    raw_envelope = {
        "schema_version": "v3",
        "request_text": "Design SMN2 exon 7 acceptor and donor splice ASOs.",
        "request_mode": "composite",
        "tasks": [
            {
                "task_id": "design_acceptor",
                "task_type": "design_new_asos",
                "context": {
                    "gene_symbol": "SMN2",
                    "species": "human",
                    "target_context": "spinal cord motor neurons",
                    "structured_goal": {
                        "gene_symbol": "SMN2",
                        "species": "human",
                        "target_context": "spinal cord motor neurons",
                        "intent": "splice_modulation",
                        "splice_isoform_objectives": [
                            {"transcript_id": "ENST00000626847", "desired_abundance": "decrease"}
                        ],
                        "design_programs": [acceptor_program],
                    },
                },
            },
            {
                "task_id": "design_donor",
                "task_type": "design_new_asos",
                "context": {
                    "gene_symbol": "SMN2",
                    "species": "human",
                    "target_context": "spinal cord motor neurons",
                    "structured_goal": {
                        "gene_symbol": "SMN2",
                        "species": "human",
                        "target_context": "spinal cord motor neurons",
                        "intent": "splice_modulation",
                        "splice_isoform_objectives": [
                            {"transcript_id": "ENST00000380743", "desired_abundance": "increase"}
                        ],
                        "design_programs": [donor_program],
                    },
                },
            },
        ],
    }

    envelope = pipeline_tools._build_request_envelope(
        request_text=raw_envelope["request_text"],
        request_id="smn2_scoped_merge",
        structured_goal={},
        raw_envelope=raw_envelope,
    )

    design_tasks = [task for task in envelope["tasks"] if task["task_type"] == "design_new_asos"]
    assert len(design_tasks) == 1
    programs = design_tasks[0]["context"]["structured_goal"]["design_programs"]
    assert [program["program_id"] for program in programs] == [
        "steric_exon7_acceptor_window",
        "steric_exon7_donor_window",
    ]


def test_isoform_conversion_keeps_explicit_site_windows_without_exploration(
    monkeypatch,
    tmp_path: Path,
) -> None:
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    goal = {
        "gene_symbol": "SMN2",
        "species": "human",
        "intent": "splice_modulation",
        "target_context": "spinal cord motor neurons",
        "splice_isoform_objectives": [
            {"transcript_id": "ENST00000626847", "desired_abundance": "decrease", "conversion_role": "from"},
            {"transcript_id": "ENST00000380743", "desired_abundance": "increase", "conversion_role": "to"},
        ],
        "isoform_transcripts": {
            "ENST00000626847": {"sites": [{"site": "donor", "position": 100}]},
            "ENST00000380743": {"sites": [{"site": "donor", "position": 100}]},
        },
        "design_programs": [
            pipeline_tools._splice_relative_program_spec(
                exon_label=7,
                transcript_id="ENST00000380743",
                site="acceptor",
                rel_start=-25,
                rel_end=24,
                label="exon7_acceptor_window",
                candidate_limit=30,
                target_context="spinal cord motor neurons",
            ),
            pipeline_tools._splice_relative_program_spec(
                exon_label=7,
                transcript_id="ENST00000380743",
                site="donor",
                rel_start=-9,
                rel_end=60,
                label="exon7_donor_window",
                candidate_limit=30,
                target_context="spinal cord motor neurons",
            ),
        ],
    }
    (run_dir / "goal_intake.json").write_text(
        json.dumps({"final_payload": {**deepcopy(goal), "structured_goal": deepcopy(goal)}}),
        encoding="utf-8",
    )
    state_path = tmp_path / "run_state.json"
    state_path.write_text(
        json.dumps(
            {
                "run_dir": str(run_dir),
                "status": "region_constraints_ready",
                "phase": "region_constraints_ready",
                "request": "Design SMN2 exon 7 splice-switching ASOs.",
                "pending_questions": [],
                "last_error": None,
            }
        ),
        encoding="utf-8",
    )

    def fail_derivation(**_kwargs):
        raise AssertionError("explicit user site windows should not trigger exploratory derivation")

    monkeypatch.setattr(pipeline_tools, "_reject_disallowed_tool_call", lambda **_kwargs: None)
    monkeypatch.setattr(pipeline_tools, "_resolve_intent_steric", lambda _run_dir: True)
    monkeypatch.setattr(pipeline_tools, "_resolve_run_splice_eval", lambda _run_dir: True)
    monkeypatch.setattr(
        pipeline_tools,
        "_collect_exploratory_splice_site_entries",
        lambda _goal: ([{"site": "donor", "position": 100}], "unit"),
    )
    monkeypatch.setattr(
        pipeline_tools,
        "_splice_modulation_isoform_conversion_policy",
        lambda *_args, **_kwargs: {"enabled": True, "priority": "primary"},
    )
    monkeypatch.setattr(pipeline_tools, "_derive_exploratory_steric_programs", fail_derivation)

    result = pipeline_tools.tool_derive_exploratory_steric_regions(str(state_path))

    assert result["status"] == "ok"
    updated = json.loads(resolve_run_path(run_dir, "structured_goal.json").read_text())
    assert updated["steric_exploration_plan"]["status"] == "skipped"
    assert [program["program_id"] for program in updated["design_programs"]] == [
        "steric_exon7_acceptor_window",
        "steric_exon7_donor_window",
    ]


def test_same_gene_design_merge_does_not_merge_cross_species_tasks() -> None:
    raw_envelope = {
        "schema_version": "v3",
        "request_text": "Design human and mouse MAPT splice ASOs.",
        "request_mode": "composite",
        "tasks": [
            {
                "task_id": "design_human_mapt",
                "task_type": "design_new_asos",
                "context": {
                    "gene_symbol": "MAPT",
                    "species": "human",
                    "structured_goal": {
                        "gene_symbol": "MAPT",
                        "species": "human",
                        "intent": "splice_modulation",
                        "splice_isoform_objectives": [
                            {"isoform_label": "human_a", "desired_abundance": "decrease"}
                        ],
                        "design_programs": [{"modality": "steric"}],
                    },
                },
            },
            {
                "task_id": "design_mouse_mapt",
                "task_type": "design_new_asos",
                "context": {
                    "gene_symbol": "Mapt",
                    "species": "mouse",
                    "structured_goal": {
                        "gene_symbol": "Mapt",
                        "species": "mouse",
                        "intent": "splice_modulation",
                        "splice_isoform_objectives": [
                            {"isoform_label": "mouse_a", "desired_abundance": "decrease"}
                        ],
                        "design_programs": [{"modality": "steric"}],
                    },
                },
            },
        ],
    }

    envelope = pipeline_tools._build_request_envelope(
        request_text=raw_envelope["request_text"],
        request_id="cross_species_no_merge",
        structured_goal={},
        raw_envelope=raw_envelope,
    )

    design_tasks = [task for task in envelope["tasks"] if task["task_type"] == "design_new_asos"]
    assert [task["task_id"] for task in design_tasks] == ["design_human_mapt", "design_mouse_mapt"]


def test_orphan_blocked_task_pauses_for_clarification_without_api_call() -> None:
    envelope = {
        "schema_version": "v3",
        "request_text": "Evaluate these ASOs and also design rat MAPT ASOs.",
        "request_mode": "composite",
        "tasks": [
            {
                "task_id": "evaluate_lab_asos",
                "task_type": "evaluate_existing_aso_set",
                "context": {
                    "gene_symbol": "MAPT",
                    "species": "human",
                    "provided_asos": [{"name": "lab_aso_1", "sequence": "ACGTACGTACGTACGTACGT"}],
                },
            }
        ],
        "blocked_tasks": [
            {
                "task_id": "design_rat_mapt",
                "missing_fields": ["species_reference_data"],
                "questions": ["Should I skip rat MAPT or provide/download the missing rat references?"],
                "reason": "The design task was omitted from tasks but listed as blocked.",
            }
        ],
    }

    validation = pipeline_tools._request_validation_summary(envelope)
    assert validation["orphan_blocked_task_ids"] == ["design_rat_mapt"]

    repair = pipeline_tools._repair_request_envelope_if_needed(
        request_text=envelope["request_text"],
        run_dir="",
        structured_goal={},
        request_envelope=envelope,
        qa_history=[],
    )

    assert repair["status"] == "need_clarification"
    assert repair["reason"] == "blocking_request_fields_missing"
    assert repair["questions"] == ["Should I skip rat MAPT or provide/download the missing rat references?"]


def test_request_text_intervals_and_union_answer_clear_binding_coverage_block() -> None:
    request_text = (
        "Design 20mer target-site blocker ASOs for human CDKN1B 3' UTR in Human HEK293T cells in an "
        "in vitro CDKN1B/p27 mRNA and protein upregulation assay. Target these hg38 intervals and "
        "treat them as a union; any overlap counts and 1 nt overlap is sufficient. "
        "PUM1 site: hg38 chr12:12721313-12721416 (+). "
        "miRNA site: hg38 chr12:12721400-12721460 (+). "
        "eCLIP peak A: hg38 chr12:12721208-12721253 (+). "
        "eCLIP peak B: hg38 chr12:12721281-12721326 (+)."
    )
    envelope = pipeline_tools._build_request_envelope(
        request_text=request_text,
        request_id="cdkn1b_binding_intervals",
        structured_goal={},
        raw_envelope={
            "schema_version": "v3",
            "request_text": request_text,
            "request_mode": "composite",
            "tasks": [
                {
                    "task_id": "design_cdkn1b_3utr_blockers",
                    "task_type": "design_new_asos",
                    "context": {
                        "gene_symbol": "CDKN1B",
                        "species": "human",
                        "target_context": (
                            "Human HEK293T cells in an in vitro CDKN1B/p27 mRNA and protein "
                            "upregulation assay after target-site blocker ASO treatment."
                        ),
                        "structured_goal": {
                            "gene_symbol": "CDKN1B",
                            "species": "human",
                            "target_context": (
                                "Human HEK293T cells in an in vitro CDKN1B/p27 mRNA and protein "
                                "upregulation assay after target-site blocker ASO treatment."
                            ),
                            "intent": "binding",
                            "target_region": "utr3",
                            "aso_length": 20,
                            "design_programs": [
                                {"program_id": "steric_cdkn1b", "modality": "steric", "intent": "binding"}
                            ],
                        },
                    },
                }
            ],
            "blocked_tasks": [
                {
                    "task_id": "design_cdkn1b_3utr_blockers",
                    "missing_fields": ["interval_coverage_policy"],
                    "questions": [
                        "Should candidates be returned to cover any site (union) or separately ensure coverage per interval?"
                    ],
                    "reason": "Coverage policy is ambiguous.",
                }
            ],
        },
    )
    qa_history = [
        {
            "question": "Should candidates be returned to cover any site (union) or separately ensure coverage per interval?",
            "answer": "Use union coverage: candidates may overlap any one of the listed intervals. Keep a combined ranked output.",
        }
    ]

    changed = pipeline_tools._apply_request_text_binding_interval_constraints_to_envelope(
        envelope,
        request_text,
        qa_history,
    )
    validation = pipeline_tools._request_validation_summary(envelope)
    task = envelope["tasks"][0]
    goal = task["context"]["structured_goal"]
    constraint = goal["region_constraints"][0]

    assert changed is True
    assert "blocked_tasks" not in envelope
    assert validation["blocked_task_ids"] == []
    assert goal["aso_overlap_requirement"] == "any_overlap"
    assert goal["minimum_overlap_bp"] == 1
    assert constraint["operator"] == "union"
    assert constraint["aso_overlap_requirement"] == "any_overlap"
    assert constraint["minimum_overlap_bp"] == 1
    assert len(constraint["intervals_bed0_half_open"]) == 4
    assert {
        (row["chrom"], row["start"], row["end"])
        for row in constraint["intervals_bed0_half_open"]
    } == {
        ("chr12", 12721313, 12721416),
        ("chr12", 12721400, 12721460),
        ("chr12", 12721208, 12721253),
        ("chr12", 12721281, 12721326),
    }


def test_request_text_interval_constraints_accept_any_overlap_phrase() -> None:
    request_text = (
        "Design 20-nt steric ASOs for human CDKN1B in HEK293T cells to block "
        "PUM1/eCLIP site 1 chr12:12721313-12721416 (+) and miR-221 seed window "
        "chr12:12721208-12721253 (+) with any overlap. Use hg38."
    )

    constraints = pipeline_tools._explicit_bed_interval_constraints_from_request_text(request_text)

    assert len(constraints) == 1
    assert constraints[0]["aso_overlap_requirement"] == "any_overlap"
    assert {
        (row["chrom"], row["start"], row["end"])
        for row in constraints[0]["intervals_bed0_half_open"]
    } == {
        ("chr12", 12721313, 12721416),
        ("chr12", 12721208, 12721253),
    }


def test_request_text_intervals_override_inferred_binding_sources() -> None:
    request_text = (
        "Design 20-nt steric ASOs for human CDKN1B in HEK293T cells to block "
        "PUM1/eCLIP site 1 chr12:12721313-12721416 (+) and miR-221 seed window "
        "chr12:12721208-12721253 (+) with any overlap. Use hg38."
    )
    inferred_constraints = [
        {
            "constraint_id": "pum1_live_lookup",
            "source_type": "rbp_binding",
            "rbp_binding": {"rbp": "PUM1"},
        }
    ]
    envelope = pipeline_tools._build_request_envelope(
        request_text=request_text,
        request_id="cdkn1b_explicit_override",
        structured_goal={},
        raw_envelope={
            "request_text": request_text,
            "structured_goal": {
                "gene_symbol": "CDKN1B",
                "intent": "binding",
                "region_constraints": deepcopy(inferred_constraints),
            },
            "tasks": [
                {
                    "task_id": "design_cdkn1b",
                    "task_type": "design_new_asos",
                    "context": {
                        "structured_goal": {
                            "gene_symbol": "CDKN1B",
                            "intent": "binding",
                            "region_constraints": deepcopy(inferred_constraints),
                        }
                    },
                }
            ],
        },
    )

    changed = pipeline_tools._apply_request_text_binding_interval_constraints_to_envelope(
        envelope,
        request_text,
    )

    assert changed is True
    root_constraint = envelope["structured_goal"]["region_constraints"][0]
    task = envelope["tasks"][0]
    task_constraint = task["context"]["structured_goal"]["region_constraints"][0]
    assert root_constraint["source_type"] == "explicit_rows"
    assert task_constraint["source_type"] == "explicit_rows"
    assert len(task_constraint["intervals_bed0_half_open"]) == 2
    assert task["context"]["region_constraints"] == task["context"]["structured_goal"]["region_constraints"]


def test_mixed_genome_build_interval_request_blocks_before_design() -> None:
    request_text = (
        "Design steric ASOs for human CDKN1B in HEK293T cells using hg19 "
        "chr12:12721208-12721253 for one regulatory site and hg38 "
        "chr12:12721313-12721416 for a second regulatory site. Use any overlap with each site."
    )
    envelope = pipeline_tools._build_request_envelope(
        request_text=request_text,
        request_id="cdkn1b_mixed_builds",
        structured_goal={},
        raw_envelope={
            "schema_version": "v3",
            "request_text": request_text,
            "tasks": [
                {
                    "task_id": "design_cdkn1b_sites",
                    "task_type": "design_new_asos",
                    "context": {
                        "gene_symbol": "CDKN1B",
                        "species": "human",
                        "target_context": "HEK293T cells",
                        "structured_goal": {
                            "gene_symbol": "CDKN1B",
                            "species": "human",
                            "target_context": "HEK293T cells",
                            "intent": "binding",
                        },
                    },
                }
            ],
        },
    )

    changed = pipeline_tools._apply_request_text_binding_interval_constraints_to_envelope(
        envelope,
        request_text,
        [],
    )
    validation = pipeline_tools._request_validation_summary(envelope)
    task_validation = validation["tasks"]["design_cdkn1b_sites"]
    constraint = envelope["tasks"][0]["context"]["structured_goal"]["region_constraints"][0]

    assert changed is True
    assert "genome_build" not in constraint
    assert {row["genome_build"] for row in constraint["intervals_bed0_half_open"]} == {"GRCh37", "GRCh38"}
    assert task_validation["status"] == "blocked"
    assert "mixed_genome_build_coordinates" in task_validation["issues"]
    assert validation["runnable_task_ids"] == []
    assert validation["blocked_task_ids"] == ["design_cdkn1b_sites"]
    assert any("genome build" in question.lower() for question in task_validation["questions"])


def test_mixed_genome_build_request_blocks_split_design_tasks() -> None:
    request_text = (
        "Design steric ASOs for human CDKN1B in HEK293T cells using hg19 "
        "chr12:12721208-12721253 for one regulatory site and hg38 "
        "chr12:12721313-12721416 for a second regulatory site. Use any overlap with each site."
    )
    envelope = {
        "schema_version": "v3",
        "request_text": request_text,
        "tasks": [
            {
                "task_id": "design_cdkn1b_site1_hg19",
                "task_type": "design_new_asos",
                "context": {
                    "gene_symbol": "CDKN1B",
                    "species": "human",
                    "target_context": "HEK293T cells",
                    "structured_goal": {
                        "gene_symbol": "CDKN1B",
                        "species": "human",
                        "target_context": "HEK293T cells",
                        "region_constraints": [
                            {
                                "coordinate_space": "genomic",
                                "intervals_bed0_half_open": [
                                    {
                                        "chrom": "chr12",
                                        "start": 12721208,
                                        "end": 12721253,
                                        "strand": "+",
                                        "genome_build": "GRCh37",
                                    }
                                ],
                            }
                        ],
                    },
                },
            },
            {
                "task_id": "design_cdkn1b_site2_hg38",
                "task_type": "design_new_asos",
                "context": {
                    "gene_symbol": "CDKN1B",
                    "species": "human",
                    "target_context": "HEK293T cells",
                    "structured_goal": {
                        "gene_symbol": "CDKN1B",
                        "species": "human",
                        "target_context": "HEK293T cells",
                        "region_constraints": [
                            {
                                "coordinate_space": "genomic",
                                "intervals_bed0_half_open": [
                                    {
                                        "chrom": "chr12",
                                        "start": 12721313,
                                        "end": 12721416,
                                        "strand": "+",
                                        "genome_build": "GRCh38",
                                    }
                                ],
                            }
                        ],
                    },
                },
            },
            {
                "task_id": "compile_report",
                "task_type": "build_composite_report",
                "depends_on": ["design_cdkn1b_site1_hg19", "design_cdkn1b_site2_hg38"],
                "context": {},
            },
        ],
    }

    validation = pipeline_tools._request_validation_summary(envelope)

    assert validation["blocked_task_ids"] == ["design_cdkn1b_site1_hg19", "design_cdkn1b_site2_hg38"]
    assert validation["tasks"]["compile_report"]["status"] == "runnable"
    for task_id in validation["blocked_task_ids"]:
        entry = validation["tasks"][task_id]
        assert "mixed_genome_build_coordinates" in entry["issues"]
        assert any("genome build" in question.lower() for question in entry["questions"])


def test_generated_filter_trace_uses_dynamic_steric_ranked_csv(tmp_path: Path) -> None:
    run_dir = tmp_path / "run"
    tiling_dir = run_dir / "tiling"
    tiling_dir.mkdir(parents=True)
    row = {
        "Chrom": "chr12",
        "Start": "12721214",
        "End": "12721234",
        "Sequence": "CTAAAAGCGTTGGATGTAGC",
        "Off_target_expression_tissue": "CellXGene Census (kidney)",
        "Off_target_expression_details": [],
    }
    for filename in ("tiles.json", "initial_filtering_tiles.json", "steric_scored_uniqueness.json"):
        (tiling_dir / filename).write_text(json.dumps({"rows": [row]}), encoding="utf-8")
    ranked_csv = tiling_dir / "utr3'_steric_uniqueness_ranked.csv"
    ranked_csv.write_text(
        "Sequence,Chrom,Start,End,Off_target_expression_tissue\n"
        "CTAAAAGCGTTGGATGTAGC,chr12,12721214,12721234,CellXGene Census (kidney)\n",
        encoding="utf-8",
    )

    pipeline_tools._emit_v2_generated_design_artifacts(str(run_dir))
    trace = json.loads(resolve_run_path(run_dir, "generated_filter_trace.json").read_text(encoding="utf-8"))
    generated = json.loads(resolve_run_path(run_dir, "generated_aso_set.json").read_text(encoding="utf-8"))
    trace_row = trace["rows"][0]

    assert generated["row_count"] == 1
    assert trace_row["stage_presence"]["final_output"] is True
    assert trace_row["terminal_status"] == "final_output"
    assert trace_row["eligible"] is True
    assert trace_row["stage_results"]["expression_annotation"]["status"] == "pass"


def test_generated_artifacts_use_strategy_routes_as_final_outputs(tmp_path: Path) -> None:
    run_dir = tmp_path / "run"
    tiling_dir = run_dir / "tiling"
    routes_dir = run_dir / "strategy_routes"
    tiling_dir.mkdir(parents=True)
    routes_dir.mkdir(parents=True)
    routed = {
        "Chrom": "chr5",
        "Start": "70070760",
        "End": "70070775",
        "Sequence": "AACAAATGCTGGCAT",
    }
    intermediate = {
        "Chrom": "chr5",
        "Start": "70070761",
        "End": "70070776",
        "Sequence": "ACAAATGCTGGCATT",
    }
    for filename in ("tiles.json", "initial_filtering_tiles.json", "steric_scored_uniqueness.json"):
        (tiling_dir / filename).write_text(json.dumps({"rows": [routed, intermediate]}), encoding="utf-8")
    (tiling_dir / "steric_uniqueness_ranked.csv").write_text(
        "Sequence,Chrom,Start,End\n"
        "ACAAATGCTGGCATT,chr5,70070761,70070776\n",
        encoding="utf-8",
    )
    route_csv = routes_dir / "steric_job_exon7_iss_n1__rank01_donor_70070751.route.csv"
    route_csv.write_text(
        "sequence,chrom,start_bp,end_bp\n"
        "AACAAATGCTGGCAT,chr5,70070760,70070775\n",
        encoding="utf-8",
    )

    pipeline_tools._emit_v2_generated_design_artifacts(str(run_dir))
    generated = json.loads(resolve_run_path(run_dir, "generated_aso_set.json").read_text(encoding="utf-8"))
    trace = json.loads(resolve_run_path(run_dir, "generated_filter_trace.json").read_text(encoding="utf-8"))
    terminal_by_sequence = {row["sequence"]: row["terminal_status"] for row in trace["rows"]}

    assert generated["row_count"] == 1
    assert generated["rows"][0]["sequence"] == "AACAAATGCTGGCAT"
    assert generated["rows"][0]["job_id"] == "steric_job_exon7_iss_n1"
    assert generated["rows"][0]["job_label"] == "exon7_iss_n1"
    assert generated["rows"][0]["route_label"] == "donor@70070751"
    assert generated["rows"][0]["site_label"] == "donor@70070751"
    assert generated["sources"]["final_output"] == [str(route_csv)]
    assert terminal_by_sequence["AACAAATGCTGGCAT"] == "final_output"
    assert terminal_by_sequence["ACAAATGCTGGCATT"] != "final_output"


def test_generated_artifacts_apply_candidate_limit_per_site_for_split_programs(tmp_path: Path) -> None:
    run_dir = tmp_path / "run"
    routes_dir = run_dir / "strategy_routes"
    routes_dir.mkdir(parents=True)
    (run_dir / "goal_intake.json").write_text(
        json.dumps(
            {
                "request": "Return at least 2 candidates per site.",
                "final_payload": {
                    "status": "ready",
                    "structured_goal": {
                        "gene_symbol": "SMN2",
                        "candidate_limit_per_site": 2,
                        "min_asos_per_site": 2,
                        "per_site_window_program_preseed": {"status": "applied"},
                        "design_programs": [
                            {"program_id": "steric_site_a", "label": "site_a", "candidate_limit": 2},
                            {"program_id": "steric_site_b", "label": "site_b", "candidate_limit": 2},
                            {"program_id": "steric_site_c", "label": "site_c", "candidate_limit": 2},
                        ],
                    },
                },
            }
        ),
        encoding="utf-8",
    )

    for job_label, base_start in (("site_a", 100), ("site_b", 200)):
        rows = []
        sequence_pool = [
            "ACGTACGTACGTACGTACGT",
            "CGTACGTACGTACGTACGTA",
            "GTACGTACGTACGTACGTAC",
        ]
        for idx in range(3):
            rows.append(
                {
                    "sequence": sequence_pool[idx],
                    "chrom": "chr1",
                    "start_bp": str(base_start + idx),
                    "end_bp": str(base_start + idx + 20),
                    "selected_label": f"acceptor@{base_start}",
                }
            )
        route_csv = routes_dir / f"steric_job_steric_{job_label}__rank01_acceptor_{base_start}.route.csv"
        route_csv.write_text(
            "sequence,chrom,start_bp,end_bp,selected_label\n"
            + "\n".join(
                f"{row['sequence']},{row['chrom']},{row['start_bp']},{row['end_bp']},{row['selected_label']}"
                for row in rows
            )
            + "\n",
            encoding="utf-8",
        )

    pipeline_tools._emit_v2_generated_design_artifacts(str(run_dir))
    generated = json.loads(resolve_run_path(run_dir, "generated_aso_set.json").read_text(encoding="utf-8"))
    counts = {}
    for row in generated["rows"]:
        counts[row["job_label"]] = counts.get(row["job_label"], 0) + 1

    assert generated["row_count"] == 4
    assert counts == {"site_a": 2, "site_b": 2}
    assert generated["candidate_selection"]["selection_scope"] == "per_site"
    assert {
        (group["group_label"], group["eligible_count"], group["selected_count"], group["shortfall_count"])
        for group in generated["candidate_selection"]["groups"]
    } == {("site_a", 3, 2, 0), ("site_b", 3, 2, 0), ("site_c", 0, 0, 2)}


def test_request_text_per_site_candidate_count_populates_exact_keys() -> None:
    state = {"request": "Design SMN2 exon 7 ASOs and return at least 20 candidates per site."}
    structured_goal = {
        "gene_symbol": "SMN2",
        "design_programs": [
            {
                "program_id": "steric_acceptor",
                "label": "acceptor",
                "modality": "steric",
                "region_constraints": [{"splice_relative_window": {"anchor_site": "acceptor"}}],
            },
            {
                "program_id": "steric_donor",
                "label": "donor",
                "modality": "steric",
                "region_constraints": [{"splice_relative_window": {"anchor_site": "donor"}}],
            },
        ],
    }
    final_payload = {"structured_goal": structured_goal}

    changed = pipeline_tools._apply_candidate_count_from_request_text_to_goal(
        state,
        structured_goal,
        final_payload,
    )

    assert changed is True
    assert structured_goal["candidate_limit_per_site"] == 20
    assert structured_goal["min_asos_per_site"] == 20
    assert "candidate_limit" not in structured_goal
    assert all(program["candidate_limit"] == 20 for program in structured_goal["design_programs"])


def test_generated_artifacts_apply_candidate_limit_per_route_for_single_exploratory_job(tmp_path: Path) -> None:
    run_dir = tmp_path / "run"
    routes_dir = run_dir / "strategy_routes"
    routes_dir.mkdir(parents=True)
    (run_dir / "goal_intake.json").write_text(
        json.dumps(
            {
                "request": "Return 2 candidates for each route.",
                "final_payload": {
                    "status": "ready",
                    "structured_goal": {
                        "gene_symbol": "MAPT",
                        "candidate_limit_per_site": 2,
                        "design_programs": [
                            {
                                "program_id": "steric_mapt_01_explore",
                                "label": "steric_mapt_01",
                                "candidate_limit": 2,
                                "candidate_sites": ["acceptor@46010310", "donor@46010402"],
                                "routing_candidate_sites": ["acceptor@46010310", "donor@46010402"],
                            }
                        ],
                    },
                },
            }
        ),
        encoding="utf-8",
    )

    route_specs = [
        ("rank01_acceptor_46010310", "acceptor@46010310", "ACGTACGTACGTACGTACGT", 100),
        ("rank02_donor_46010402", "donor@46010402", "TTGCATTGCATTGCATTGCA", 200),
    ]
    for route_name, route_label, sequence_seed, base_start in route_specs:
        rows = []
        for idx in range(3):
            rows.append(
                {
                    "sequence": sequence_seed[idx:] + sequence_seed[:idx],
                    "chrom": "chr17",
                    "start_bp": str(base_start + idx),
                    "end_bp": str(base_start + idx + 20),
                    "selected_label": route_label,
                }
            )
        route_csv = routes_dir / f"steric_job_steric_mapt_01_explore__{route_name}.route.csv"
        route_csv.write_text(
            "sequence,chrom,start_bp,end_bp,selected_label\n"
            + "\n".join(
                f"{row['sequence']},{row['chrom']},{row['start_bp']},{row['end_bp']},{row['selected_label']}"
                for row in rows
            )
            + "\n",
            encoding="utf-8",
        )

    pipeline_tools._emit_v2_generated_design_artifacts(str(run_dir))
    generated = json.loads(resolve_run_path(run_dir, "generated_aso_set.json").read_text(encoding="utf-8"))
    counts: dict[str, int] = {}
    for row in generated["rows"]:
        counts[row["route_label"]] = counts.get(row["route_label"], 0) + 1

    assert generated["row_count"] == 4
    assert counts == {"acceptor@46010310": 2, "donor@46010402": 2}
    assert generated["candidate_selection"]["selection_scope"] == "per_site"
    assert generated["candidate_selection"]["group_scope"] == "route"
    assert {
        (group["group_label"], group["eligible_count"], group["selected_count"], group["shortfall_count"])
        for group in generated["candidate_selection"]["groups"]
    } == {("acceptor@46010310", 3, 2, 0), ("donor@46010402", 3, 2, 0)}


def test_offtarget_expression_annotations_sync_to_per_job_tiles(tmp_path: Path) -> None:
    source = tmp_path / "tiling" / "initial_filtering_tiles.json"
    target = tmp_path / "steric_jobs" / "job_acceptor" / "initial_filtering_tiles.json"
    source.parent.mkdir(parents=True)
    target.parent.mkdir(parents=True)
    source.write_text(
        json.dumps(
            {
                "schema": {
                    "columns": {
                        "Sequence": "string",
                        "Off_target_expression_sum": "number",
                        "Off_target_expression_sum_0mm": "number",
                        "Off_target_expression_weighted_risk": "number",
                        "Off_target_expression_tissue": "string",
                        "Target_gene_expression": "number",
                    }
                },
                "rows": [
                    {
                        "Sequence": "AACAAATGCTGGCAT",
                        "Chrom": "chr5",
                        "Start": 70070760,
                        "End": 70070775,
                        "Off_target_expression_sum": 42.5,
                        "Off_target_expression_sum_0mm": 2.5,
                        "Off_target_expression_weighted_risk": 50.0,
                        "Off_target_expression_tissue": "CellXGene Census (brain)",
                        "Target_gene_expression": 1.2,
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    target.write_text(
        json.dumps(
            {
                "schema": {"columns": {"Sequence": "string"}},
                "rows": [
                    {
                        "Sequence": "AACAAATGCTGGCAT",
                        "Chrom": "chr5",
                        "Start": 70070760,
                        "End": 70070775,
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    final_payload = {
        "steric_tiling_results_by_job": {
            "job_acceptor": {
                "filtered_tiles_path": str(target),
            }
        }
    }

    result = pipeline_tools._sync_offtarget_expression_to_design_result_tiles(
        source_json_path=str(source),
        structured_goal={},
        final_payload=final_payload,
        modality="steric",
        result_key="filtered_tiles_path",
    )

    updated = json.loads(target.read_text(encoding="utf-8"))
    row = updated["rows"][0]
    assert result["status"] == "updated"
    assert row["Off_target_expression_sum"] == 42.5
    assert row["Off_target_expression_sum_0mm"] == 2.5
    assert row["Off_target_expression_weighted_risk"] == 50.0
    assert row["Off_target_expression_tissue"] == "CellXGene Census (brain)"
    assert row["Target_gene_expression"] == 1.2
    assert updated["schema"]["columns"]["Off_target_expression_weighted_risk"] == "number"


def test_generated_filter_trace_includes_gapmer_job_initial_filter_source(tmp_path: Path) -> None:
    run_dir = tmp_path / "run"
    gapmer_dir = run_dir / "gapmer_tiling"
    job_dir = run_dir / "gapmer_jobs" / "gapmer_job_pcsk9"
    gapmer_dir.mkdir(parents=True)
    job_dir.mkdir(parents=True)
    kept = {
        "Name": "ENST00000302118_cdna826-845",
        "Sequence": "TGGTGGAGGTGTATCTCCTA",
        "Genomic_span_chrom": "chr1",
        "Genomic_span_start": 55052290,
        "Genomic_span_end": 55052309,
        "Off_target_expression_tissue": "CellXGene Census (liver)",
    }
    removed = {
        "Name": "ENST00000302118_cdna827-846",
        "Sequence": "GGTGGAGGTGTATCTCCTAA",
        "Genomic_span_chrom": "chr1",
        "Genomic_span_start": 55052291,
        "Genomic_span_end": 55052310,
    }
    (gapmer_dir / "tiles.json").write_text(json.dumps({"rows": [kept, removed]}), encoding="utf-8")
    (job_dir / "tiles.json").write_text(json.dumps({"rows": [kept, removed]}), encoding="utf-8")
    (job_dir / "initial_filtering_tiles.json").write_text(json.dumps({"rows": [kept]}), encoding="utf-8")
    for path in (gapmer_dir / "gapmer_scored_uniqueness.json", job_dir / "gapmer_scored_uniqueness.json"):
        path.write_text(json.dumps({"rows": [kept]}), encoding="utf-8")
    (gapmer_dir / "gapmer_uniqueness_ranked.csv").write_text(
        "Sequence,Name,Genomic_span_chrom,Genomic_span_start,Genomic_span_end\n"
        "TGGTGGAGGTGTATCTCCTA,ENST00000302118_cdna826-845,chr1,55052290,55052309\n",
        encoding="utf-8",
    )

    pipeline_tools._emit_v2_generated_design_artifacts(str(run_dir))
    trace = json.loads(resolve_run_path(run_dir, "generated_filter_trace.json").read_text(encoding="utf-8"))
    terminal_by_sequence = {row["sequence"]: row["terminal_status"] for row in trace["rows"]}

    assert str(job_dir / "initial_filtering_tiles.json") in trace["sources"]["initial_filter"]
    assert terminal_by_sequence["TGGTGGAGGTGTATCTCCTA"] == "final_output"
    assert terminal_by_sequence["GGTGGAGGTGTATCTCCTAA"] == "filtered_during_initial_filter"


def test_blocked_non_design_task_does_not_prevent_partial_composite_report(tmp_path: Path) -> None:
    state = _state_from_raw_envelope(
        {
            "schema_version": "v3",
            "request_text": "Answer a question and inspect a run I have not provided yet.",
            "request_mode": "composite",
            "tasks": [
                {
                    "task_id": "answer_phylo",
                    "task_type": "answer_question",
                    "label": "Explain phyloP",
                    "context": {"question": "What does phyloP mean?"},
                },
                {
                    "task_id": "inspect_missing_run",
                    "task_type": "analyze_artifact",
                    "label": "Inspect prior run",
                    "context": {},
                },
            ],
        },
        tmp_path,
    )

    validation = pipeline_tools._request_validation_summary(state["request_envelope"])
    assert validation["blocked_task_ids"] == ["inspect_missing_run"]
    assert validation["runnable_task_ids"] == ["answer_phylo", "task_build_composite_report"]

    _mark_completed(state, "answer_phylo")
    assert pipeline_tools._runnable_task_ids_for_state(state) == ["task_build_composite_report"]

    planning = pipeline_tools._derive_next_actions(state)
    assert planning["terminal_status"] == "in_progress"
    assert _next_tool(planning) == "tool_build_composite_report"


def test_missing_artifact_ref_question_includes_structured_path(tmp_path: Path) -> None:
    missing_path = tmp_path / "missing_prior_run"
    state = _state_from_raw_envelope(
        {
            "schema_version": "v3",
            "request_text": f"Inspect {missing_path} and write a short synthesis.",
            "request_mode": "composite",
            "tasks": [
                {
                    "task_id": "inspect_missing_run",
                    "task_type": "analyze_artifact",
                    "label": "Inspect prior run",
                    "context": {"artifact_refs": [{"path": str(missing_path)}]},
                },
            ],
        },
        tmp_path,
    )

    validation = pipeline_tools._request_validation_summary(state["request_envelope"])
    questions = pipeline_tools._clarification_questions_for_request_issues([], state["request_envelope"])

    assert validation["blocked_task_ids"] == ["inspect_missing_run"]
    assert questions
    assert str(missing_path) in questions[0]


def test_existing_local_artifact_path_clears_model_blocked_tasks(tmp_path: Path) -> None:
    snapgene_path = tmp_path / "isoform.gb"
    snapgene_path.write_text("LOCUS       TEST                     12 bp    DNA\nORIGIN\n        1 acgtacgtacgt\n//\n")
    envelope = {
        "schema_version": "v3",
        "request_text": f"Design LMNA ASOs using {snapgene_path}.",
        "request_mode": "composite",
        "tasks": [
            {
                "task_id": "parse_snapgene",
                "task_type": "analyze_artifact",
                "context": {"artifact_refs": [{"path": str(snapgene_path), "artifact_type": "genbank"}]},
            },
            {
                "task_id": "design_lmna",
                "task_type": "design_new_asos",
                "depends_on": ["parse_snapgene"],
                    "context": {
                        "gene_symbol": "LMNA",
                        "species": "human",
                        "target_context": "human dermal fibroblasts in progeria",
                        "structured_goal": {
                            "gene_symbol": "LMNA",
                            "species": "human",
                            "target_context": "human dermal fibroblasts in progeria",
                            "intent": "splice_modulation",
                            "design_programs": [
                                {
                                    "program_id": "steric",
                                    "label": "steric",
                                    "modality": "steric",
                                    "intent": "splice_modulation",
                                    "objective": "decrease target isoform abundance",
                                    "desired_delta": "decrease",
                                }
                            ],
                        },
                },
            },
        ],
        "task_dependencies": [
            {"upstream_task_id": "parse_snapgene", "downstream_task_id": "design_lmna"}
        ],
        "blocked_tasks": [
            {
                "task_id": "parse_snapgene",
                "missing_fields": ["artifact_file_content_or_accessible_uri"],
                "questions": ["Please upload/provide the file contents."],
                "reason": "The referenced artifact path is on the user's local filesystem and is not accessible.",
            },
            {
                "task_id": "design_lmna",
                "missing_fields": ["snapgene_isoform_definition"],
                "questions": ["Once the SnapGene isoform is identified, confirm it."],
                "reason": "The design depends on the artifact parser output.",
            },
        ],
    }

    validation = pipeline_tools._request_validation_summary(envelope)
    assert validation["blocked_task_ids"] == []

    repaired = pipeline_tools._repair_request_envelope_if_needed(
        request_text=envelope["request_text"],
        run_dir=str(tmp_path),
        structured_goal={},
        request_envelope=envelope,
        qa_history=[],
    )
    assert repaired["status"] == "ready"
    assert "blocked_tasks" not in repaired["request_envelope"]


def test_artifact_excerpt_reads_genbank_text(tmp_path: Path) -> None:
    snapgene_path = tmp_path / "reference.gb"
    _write_test_genbank(
        snapgene_path,
        [
            SeqFeature(
                FeatureLocation(10, 40, strand=1),
                type="misc_feature",
                qualifiers={
                    "label": ["SnapGene isoform (construct.dna) exon 1"],
                    "note": ["Isoform: SnapGene isoform (construct.dna); exon 1"],
                },
            )
        ],
    )

    excerpt = pipeline_tools._artifact_excerpt(snapgene_path)

    assert "LOCUS" in excerpt
    assert "SnapGene isoform (construct.dna) exon 1" in excerpt
    assert "binary_or_unsupported" not in excerpt


def test_snapgene_genbank_artifact_analysis_is_deterministic(tmp_path: Path, monkeypatch) -> None:
    snapgene_path = tmp_path / "reference.gb"
    _write_test_genbank(
        snapgene_path,
        [
            SeqFeature(
                FeatureLocation(10, 40, strand=1),
                type="misc_feature",
                qualifiers={"label": ["ENST000001 exon 1"], "note": ["Isoform: ENST000001; exon 1"]},
            ),
            SeqFeature(
                FeatureLocation(100, 130, strand=1),
                type="misc_feature",
                qualifiers={
                    "label": ["SnapGene isoform (construct.dna) exon 1"],
                    "note": ["Isoform: SnapGene isoform (construct.dna); exon 1"],
                },
            ),
            SeqFeature(
                FeatureLocation(200, 236, strand=1),
                type="misc_feature",
                qualifiers={
                    "label": ["SnapGene isoform (construct.dna) exon 2"],
                    "note": ["Isoform: SnapGene isoform (construct.dna); exon 2"],
                },
            ),
            SeqFeature(
                FeatureLocation(250, 270, strand=1),
                type="misc_feature",
                qualifiers={
                    "label": ["increase donor@123"],
                    "note": ["Strategy: increase donor@123", "Sequence: ACGTACGTACGTACGTACGT"],
                },
            ),
        ],
    )
    state = _state_from_raw_envelope(
        {
            "schema_version": "v3",
            "request_text": f"Inspect {snapgene_path}.",
            "request_mode": "composite",
            "tasks": [
                {
                    "task_id": "parse_snapgene",
                    "task_type": "analyze_artifact",
                    "label": "Parse SnapGene isoform",
                    "context": {"artifact_refs": [{"path": str(snapgene_path), "artifact_type": "genbank"}]},
                }
            ],
        },
        tmp_path,
    )
    state_path = tmp_path / "run_state.json"
    state_path.write_text(json.dumps(state), encoding="utf-8")

    def _fail_workbench(*_args: Any, **_kwargs: Any) -> dict[str, Any]:
        raise AssertionError("workbench response should not be called for parseable GenBank artifacts")

    monkeypatch.setattr(pipeline_tools, "_run_workbench_response_sync", _fail_workbench)

    result = pipeline_tools.tool_analyze_artifact(str(state_path))

    assert result["status"] == "ok"
    assert Path(result["artifacts"]["artifact_index_json"]).exists()
    output_path = Path(result["artifacts"].get("json") or result["artifacts"]["response_json"])
    payload = json.loads(output_path.read_text())
    assert payload["analysis_method"] == "deterministic_snapgene_genbank_parser"
    assert payload["artifact_index"]["aggregate"]["file_count_indexed"] >= 1
    assert payload["artifact_summaries"][0]["parseable"] is True
    assert "SnapGene isoform (construct.dna)" in payload["markdown"]
    assert "increase donor@123" in payload["markdown"]


def test_task_scoped_response_metadata_excludes_sibling_artifact_path(tmp_path: Path) -> None:
    missing_path = tmp_path / "missing_prior_run"
    state = _state_from_raw_envelope(
        {
            "schema_version": "v3",
            "request_text": f"Explain phyloP, then inspect {missing_path} and synthesize.",
            "request_mode": "composite",
            "tasks": [
                {
                    "task_id": "answer_phylo",
                    "task_type": "answer_question",
                    "label": "Explain phyloP",
                    "context": {"question": "What does phyloP mean for ASO prioritization?"},
                },
                {
                    "task_id": "inspect_missing_run",
                    "task_type": "analyze_artifact",
                    "label": "Inspect prior run",
                    "context": {"artifact_refs": [{"path": str(missing_path)}]},
                },
            ],
        },
        tmp_path,
    )
    answer_task = state["task_graph"]["tasks"][0]
    scoped = pipeline_tools._task_scoped_request_envelope(state["request_envelope"], answer_task)
    metadata = pipeline_tools._task_scoped_response_metadata(
        parent_request_text=state["request"],
        request_envelope=scoped,
        task_context=scoped["task_context"],
    )

    assert "phyloP" in metadata["request_text"]
    assert str(missing_path) not in metadata["request_text"]
    assert str(missing_path) not in metadata["request_envelope"]["request_text"]
    assert "parent_request_text" not in metadata["request_envelope"]


def test_single_task_scoped_response_keeps_immutable_parent_request() -> None:
    parent_request = (
        "Design steric ASOs for the seventh exon in reviewed motor neurons."
    )
    envelope = {
        "request_mode": "single_task",
        "request_text": parent_request,
        "request_summary": "Design steric ASOs for the fifth exon.",
        "tasks": [
            {
                "task_id": "design",
                "task_type": "design_new_asos",
                "label": "Design the requested ASOs",
                "context": {},
            }
        ],
    }

    metadata = pipeline_tools._task_scoped_response_metadata(
        parent_request_text=parent_request,
        request_envelope=envelope,
        task_context={},
    )

    assert parent_request in metadata["request_text"]
    assert "fifth exon" not in metadata["request_text"]
    assert metadata["request_envelope"]["request_text"] == metadata["request_text"]


def test_dedupe_clarification_questions_prefers_file_content_prompt(tmp_path: Path) -> None:
    missing_file = tmp_path / "uploaded_isoform.dna"
    questions = pipeline_tools._dedupe_clarification_questions(
        [
            (
                f"The referenced artifact path does not exist locally: {missing_file}. "
                "Please provide a valid run directory, artifact file, or ASO set source."
            ),
            (
                "Please upload/provide the uploaded_isoform.dna file contents or paste the "
                "DNA/CDS/mRNA sequence, or provide the corresponding transcript ID (ENST...)."
            ),
        ]
    )

    assert len(questions) == 1
    assert "file contents" in questions[0]
    assert "transcript ID" in questions[0]


def test_cyclic_dependencies_are_pruned_so_large_graph_can_start(tmp_path: Path) -> None:
    state = _state_from_raw_envelope(
        {
            "schema_version": "v3",
            "request_text": "Answer two ASO questions and synthesize them.",
            "request_mode": "composite",
            "tasks": [
                {
                    "task_id": "answer_a",
                    "task_type": "answer_question",
                    "depends_on": ["answer_b"],
                    "context": {"question": "What is an ASO gapmer?"},
                },
                {
                    "task_id": "answer_b",
                    "task_type": "answer_question",
                    "depends_on": ["answer_a"],
                    "context": {"question": "What is a steric-blocking ASO?"},
                },
            ],
        },
        tmp_path,
    )

    task_by_id = {task["task_id"]: task for task in state["request_envelope"]["tasks"]}
    assert task_by_id["answer_a"]["depends_on"] == []
    assert task_by_id["answer_b"]["depends_on"] == ["answer_a"]
    assert pipeline_tools._runnable_task_ids_for_state(state) == ["answer_a"]

    planning = pipeline_tools._derive_next_actions(state)
    assert planning["terminal_status"] == "in_progress"
    assert planning["stop_reason"] == ""
    assert _next_tool(planning) == "tool_answer_question"


def test_self_dependencies_are_removed_from_task_graph(tmp_path: Path) -> None:
    state = _state_from_raw_envelope(
        {
            "schema_version": "v3",
            "request_text": "Explain why ASO GC content matters.",
            "request_mode": "single_task",
            "tasks": [
                {
                    "task_id": "answer_gc",
                    "task_type": "answer_question",
                    "depends_on": ["answer_gc"],
                    "context": {"question": "Why does ASO GC content matter?"},
                }
            ],
        },
        tmp_path,
    )

    task = state["request_envelope"]["tasks"][0]
    assert task["task_id"] == "answer_gc"
    assert task["depends_on"] == []
    assert pipeline_tools._runnable_task_ids_for_state(state) == ["answer_gc"]


def test_duplicate_task_ids_are_preserved_with_unique_ids(tmp_path: Path) -> None:
    state = _state_from_raw_envelope(
        {
            "schema_version": "v3",
            "request_text": "Explain ASO gapmers and explain steric ASOs.",
            "request_mode": "single_task",
            "response_mode": "report",
            "tasks": [
                {
                    "task_id": "answer_step",
                    "task_type": "answer-question",
                    "context": {"question": "What is an ASO gapmer?"},
                },
                {
                    "task_id": "answer_step",
                    "task_type": "answer question",
                    "context": {"question": "What is a steric-blocking ASO?"},
                },
            ],
        },
        tmp_path,
    )

    envelope = state["request_envelope"]
    task_ids = [task["task_id"] for task in envelope["tasks"]]
    task_types = [task["task_type"] for task in envelope["tasks"]]

    assert task_ids[:2] == ["answer_step", "answer_step_02"]
    assert task_types[:2] == ["answer_question", "answer_question"]
    assert envelope["request_mode"] == "composite"
    assert envelope["response_mode"] == "composite"
    assert pipeline_tools._runnable_task_ids_for_state(state) == ["answer_step", "answer_step_02"]


def test_single_design_task_ignores_stale_composite_response_mode(tmp_path: Path) -> None:
    state = _state_from_raw_envelope(
        {
            "schema_version": "v3",
            "request_text": "Design PCCA splice-switching ASOs.",
            "request_mode": "composite",
            "response_mode": "composite",
            "tasks": [
                {
                    "task_id": "design_pcca_acceptor_sso",
                    "task_type": "design_new_asos",
                    "context": {
                        "gene_symbol": "PCCA",
                        "species": "human",
                        "target_context": "human hepatocyte-like cells",
                    },
                }
            ],
        },
        tmp_path,
    )

    envelope = state["request_envelope"]
    assert envelope["request_mode"] == "single_task"
    assert envelope["response_mode"] == "report"


def test_single_design_integrity_ignores_stale_composite_response_mode(tmp_path: Path) -> None:
    request_envelope = {
        "request_id": "unit_pcca",
        "request_text": "Design PCCA splice-switching ASOs.",
        "request_mode": "single_task",
        "response_mode": "composite",
        "workflow_name": "design_new_asos",
        "tasks": [
            {
                "task_id": "design_pcca_acceptor_sso",
                "task_type": "design_new_asos",
                "context": {"gene_symbol": "PCCA", "target_context": "human hepatocyte-like cells"},
            }
        ],
    }
    (tmp_path / "report.md").write_text("# Report\n", encoding="utf-8")
    (tmp_path / "goal_intake_report.html").write_text("<html></html>\n", encoding="utf-8")
    (tmp_path / "goal_intake.json").write_text(
        json.dumps(
            {
                "structured_goal": {
                    "gene_symbol": "PCCA",
                    "intent": "binding",
                    "target_region": "utr3",
                },
                "final_payload": {"request_envelope": request_envelope},
            }
        ),
        encoding="utf-8",
    )

    integrity = pipeline_tools._extract_pipeline_integrity(
        str(tmp_path),
        {
            "exists": {
                "report_md": True,
                "report_html": True,
                "composite_report_md": False,
                "composite_report_json": False,
                "task_index_json": False,
            },
            "markers": {},
        },
        skip_router=True,
    )

    # Stale composite mode is ignored, but report-only output is not a delivered design.
    assert integrity["ok"] is False
    assert "requested_candidate_count_not_delivered" in integrity["blocking_errors"]
    assert "composite_report_md_missing" not in integrity["blocking_errors"]
    assert "design_outcome_missing:design_pcca_acceptor_sso" not in integrity["blocking_errors"]
    assert "stale_single_task_composite_response_mode_ignored" in integrity["warnings"]


def test_unsupported_only_task_requests_clarification_without_fatal_planner_state(tmp_path: Path) -> None:
    state = _state_from_raw_envelope(
        {
            "schema_version": "v3",
            "request_text": "Prioritize these ASOs with a custom ranker.",
            "request_mode": "single_task",
            "tasks": [
                {
                    "task_id": "custom_rank",
                    "task_type": "rank_existing_asos",
                    "context": {
                        "provided_asos": [{"name": "lab_aso_1", "sequence": "ACGTACGTACGTACGTACGT"}]
                    },
                }
            ],
        },
        tmp_path,
    )

    validation = pipeline_tools._request_validation_summary(state["request_envelope"])
    assert validation["blocked_task_ids"] == ["custom_rank"]
    assert validation["tasks"]["custom_rank"]["issues"] == ["unsupported_task_type"]

    repair = pipeline_tools._repair_request_envelope_if_needed(
        request_text=state["request"],
        run_dir=str(tmp_path),
        structured_goal={},
        request_envelope=state["request_envelope"],
        qa_history=[],
    )
    assert repair["status"] == "need_clarification"
    assert "unsupported task type" in repair["questions"][0].lower()

    planning = pipeline_tools._derive_next_actions(state)
    assert planning["terminal_status"] == "needs_user_input"
    assert planning["stop_reason"] == "pending_clarification_questions"
    assert planning["pending_questions"]


def test_unsupported_non_design_task_allows_partial_composite_after_runnable_work(tmp_path: Path) -> None:
    state = _state_from_raw_envelope(
        {
            "schema_version": "v3",
            "request_text": "Explain ASO GC content and do a custom prioritization step.",
            "request_mode": "composite",
            "tasks": [
                {
                    "task_id": "answer_gc",
                    "task_type": "answer_question",
                    "context": {"question": "Why does ASO GC content matter?"},
                },
                {
                    "task_id": "custom_rank",
                    "task_type": "custom_prioritization",
                    "context": {"provided_asos": [{"sequence": "ACGTACGTACGTACGTACGT"}]},
                },
            ],
        },
        tmp_path,
    )

    validation = pipeline_tools._request_validation_summary(state["request_envelope"])
    assert validation["blocked_task_ids"] == ["custom_rank"]
    assert pipeline_tools._runnable_task_ids_for_state(state) == ["answer_gc"]

    _mark_completed(state, "answer_gc")
    assert pipeline_tools._runnable_task_ids_for_state(state) == ["task_build_composite_report"]

    planning = pipeline_tools._derive_next_actions(state)
    assert planning["terminal_status"] == "in_progress"
    assert _next_tool(planning) == "tool_build_composite_report"


def test_near_duplicate_artifact_clarification_questions_are_collapsed() -> None:
    questions = pipeline_tools._normalize_questions_input(
        [
            "The referenced artifact path does not exist locally. Please provide a valid run directory, artifact file, or ASO set source.",
            "The referenced path '/tmp/definitely_missing_agentaso_run' is not accessible. Please provide an accessible run directory path, or upload/attach the run artifacts.",
            "To produce the final synthesis, I need the artifact inspection results. Please provide an accessible run directory path or upload/attach the run artifacts for '/tmp/definitely_missing_agentaso_run'.",
            "Provide an accessible artifact/run output for t2_inspect_run so the final synthesis can incorporate the run inspection.",
        ]
    )

    assert len(questions) == 1
    assert "/tmp/definitely_missing_agentaso_run" in questions[0]


def test_distinct_missing_artifact_paths_keep_distinct_questions() -> None:
    questions = pipeline_tools._normalize_questions_input(
        [
            "The referenced path '/tmp/missing_run_a' is not accessible. Please provide an accessible run directory path.",
            "The referenced path '/tmp/missing_run_b' is not accessible. Please provide an accessible run directory path.",
        ]
    )

    assert len(questions) == 2
    assert any("/tmp/missing_run_a" in question for question in questions)
    assert any("/tmp/missing_run_b" in question for question in questions)


def test_distinct_task_specific_artifact_questions_keep_distinct_questions() -> None:
    questions = pipeline_tools._normalize_questions_input(
        [
            "Provide an accessible artifact/run output for t2_inspect_run so the final synthesis can incorporate the run inspection.",
            "Provide an accessible artifact/run output for t3_compare_run so the final synthesis can incorporate the run inspection.",
        ]
    )

    assert len(questions) == 2
    assert any("t2_inspect_run" in question for question in questions)
    assert any("t3_compare_run" in question for question in questions)


def test_question_normalization_rejects_punctuation_only_fragments() -> None:
    meaningful = "Which user-controlled input should be used?"

    assert pipeline_tools._normalize_questions_input(
        [",", "...", "!?", meaningful]
    ) == [meaningful]


def test_structured_splice_objective_overrides_generic_binding_intent() -> None:
    goal = {
        "gene_symbol": "SMN2",
        "species": "human",
        "intent": "binding",
        "splice_objective": {
            "exon_label": "exon 7",
            "desired_delta": "increase_inclusion",
            "transcript_policy": "canonical_transcript",
        },
        "design_programs": [
            {
                "modality": "steric",
                "intent": "binding",
                "program_id": "steric_smn2_01",
            }
        ],
    }
    program = goal["design_programs"][0]

    assert pipeline_tools._request_contract_infer_intent(goal, {}, goal["design_programs"]) == "splice_modulation"
    assert pipeline_tools._design_request_requires_splice_eval(program, fallback_goal=goal) is True
    assert pipeline_tools._design_request_requires_splice_router(program, fallback_goal=goal) is True


def test_gapmer_exon_region_constraint_preserved_in_request_contract() -> None:
    task = {
        "task_id": "design_pcsk9_gapmers_ex4",
        "task_type": "design_new_asos",
        "context": {
            "gene_symbol": "PCSK9",
            "species": "human",
            "transcript_id": "ENST00000302118",
            "structured_goal": {
                "gene_symbol": "PCSK9",
                "species": "human",
                "intent": "knockdown",
                "target_region": "cds",
                "region_constraints": [
                    {
                        "constraint_kind": "transcript_region",
                        "operator": "include",
                        "transcript_id": "ENST00000302118",
                        "region": {"exon_label": "exon 4"},
                    }
                ],
                "design_programs": [
                    {
                        "modality": "gapmer",
                        "program_id": "gapmer_pcsk9_01",
                        "target_region": "cds",
                    }
                ],
            },
        },
    }

    finalized = pipeline_tools._finalize_design_task_contract(task, {})
    goal = finalized["context"]["structured_goal"]
    program = goal["design_programs"][0]

    assert goal["exon_label"] == 4
    assert goal["focus_exons_only"] is True
    assert program["exon_label"] == 4
    assert program["focus_exons_only"] is True


def test_sequence_filters_are_normalized_parameter_constraints() -> None:
    constraints = parse_parameter_constraints(
        "",
        {
            "sequence_filters": {
                "gc_content_min": 0.0,
                "gc_content_max": 1.0,
                "homopolymer_max_run": 4,
            }
        },
    )

    assert get_gc_bounds(constraints) == (0.0, 1.0)
    assert constraints["gc_content"]["min_percent"] == 0.0
    assert constraints["gc_content"]["max_percent"] == 100.0
    assert constraints["gc_content"]["min_inclusive"] is True
    assert constraints["gc_content"]["max_inclusive"] is True
    assert constraints["homopolymers"]["per_base_max"] == {"A": 4, "C": 4, "G": 4, "T": 4}
    assert "gc_content_min" not in constraints
    assert "gc_content_max" not in constraints


def test_open_open_gc_bounds_override_initial_filter_defaults(tmp_path: Path) -> None:
    constraints = parse_parameter_constraints("Accept GC content greater than 0 and less than 1.", {})
    constraints_from_llm_schema = {"gc_fraction": {"min_exclusive": 0.0, "max_exclusive": 1.0}}
    tiles_path = tmp_path / "tiles.json"
    tiles_path.write_text(
        json.dumps(
            {
                "schema": {"table": "aso_tiles"},
                "rows": [
                    {"Chrom": "chr1", "Start": "1", "End": "21", "No_homopolymer": "True", "GC": "0.25", "Sequence": "ACATATATATATATATATAT"},
                    {"Chrom": "chr1", "Start": "2", "End": "22", "No_homopolymer": "True", "GC": "0.70", "Sequence": "GCGCGCATATGCGCATATGC"},
                    {"Chrom": "chr1", "Start": "3", "End": "23", "No_homopolymer": "True", "GC": "1.00", "Sequence": "GCGCGCGCGCGCGCGCGCGC"},
                ],
            }
        ),
        encoding="utf-8",
    )

    summary = run_goal_intake._apply_initial_tiling_filters(str(tiles_path), constraints)

    assert get_gc_bounds(constraints) == (0.0, 1.0)
    assert get_gc_bounds(constraints_from_llm_schema) == (0.0, 1.0)
    assert get_gc_inclusive_flags(constraints) == (False, False)
    assert get_gc_inclusive_flags(constraints_from_llm_schema) == (False, False)
    assert constraints["gc_content"]["min_inclusive"] is False
    assert constraints["gc_content"]["max_inclusive"] is False
    assert "gc_fraction_min_exclusive" not in constraints
    assert "gc_fraction_max_exclusive" not in constraints
    assert "min_exclusive" not in constraints["gc_content"]
    assert "max_exclusive" not in constraints["gc_content"]
    assert summary["constraint_summary"]["gc_min_percent"] == 0.0
    assert summary["constraint_summary"]["gc_max_percent"] == 100.0
    assert summary["constraint_summary"]["gc_min_exclusive"] is True
    assert summary["constraint_summary"]["gc_max_exclusive"] is True
    assert summary["passed"]["count"] == 2
    assert summary["filtered_gc_out_of_range"]["count"] == 1


def test_apply_parameter_constraints_rewrites_model_aliases_to_exact_keys() -> None:
    structured_goal = {
        "design_programs": [
            {
                "program_id": "steric_cdkn1b_01",
                "modality": "steric",
                "gc_content_range": {"min_exclusive": 0.0, "max_exclusive": 1.0},
            }
        ],
        "sequence_constraints": {
            "gc_content_range": {"min_exclusive": 0.0, "max_exclusive": 1.0},
            "disallowed_homopolymer_runs": [
                {"base": "G", "max_allowed_run_length": 3},
            ],
        },
        "candidate_filters": {
            "gc_content_fraction": {"min_exclusive": 0.0, "max_exclusive": 1.0},
        },
        "offtarget_constraints": {
            "max_offtarget_mappings": 2,
        },
    }

    normalized_goal = apply_parameter_constraints_to_goal(
        structured_goal,
        "Design ASOs with GC content greater than 0 and less than 1 and fewer than 3 off-target mappings.",
    )
    constraints = normalized_goal["parameter_constraints"]

    assert normalized_goal is structured_goal
    assert "gc_content_range" not in normalized_goal["design_programs"][0]
    assert "sequence_constraints" not in normalized_goal
    assert "candidate_filters" not in normalized_goal
    assert "offtarget_constraints" not in normalized_goal
    assert constraints["gc_content"] == {
        "sources": [
            "structured_goal.sequence_constraints",
            "structured_goal.candidate_filters",
            "structured_goal.design_programs",
            "user_text",
        ],
        "min": 0.0,
        "min_percent": 0.0,
        "min_inclusive": False,
        "max": 1.0,
        "max_percent": 100.0,
        "max_inclusive": False,
    }
    assert constraints["max_binding_sites"] == 2
    assert constraints["binding_sites"]["max"] == 2
    assert constraints["homopolymers"]["per_base_max"]["G"] == 3
    assert "gc_content_range" not in constraints
    assert "max_offtarget_mappings" not in constraints


def test_parameter_constraints_strip_gc_content_fraction_bounds_from_sequence_block() -> None:
    goal = {
        "parameter_constraints": {
            "gc_content": {
                "min": 0.35,
                "max": 0.65,
                "min_inclusive": False,
                "max_inclusive": False,
            }
        },
        "sequence_constraints": {
            "gc_content_min_fraction": 0.35,
            "gc_content_max_fraction": 0.65,
        },
    }

    normalized = apply_parameter_constraints_to_goal(
        goal,
        "Use GC content greater than 35% and less than 65%.",
    )

    assert "sequence_constraints" not in normalized
    assert normalized["parameter_constraints"]["gc_content"]["min"] == 0.35
    assert normalized["parameter_constraints"]["gc_content"]["max"] == 0.65
    assert normalized["parameter_constraints"]["gc_content"]["min_inclusive"] is False
    assert normalized["parameter_constraints"]["gc_content"]["max_inclusive"] is False


def test_parameter_constraints_fold_common_envelope_aliases_after_default_fill() -> None:
    structured_goal = {
        "parameter_constraints": {
            "homopolymers": {
                "per_base_max": {"A": 4, "C": 4, "G": 4, "T": 4},
            }
        },
        "homopolymers": {
            "per_base_max": {"A": 4, "C": 4, "G": 4, "T": 4},
        },
        "sequence_constraints": {
            "gc_percent_range": [35, 65],
            "max_homopolymer_runs": {"C": 8},
        },
        "off_target_policy": {
            "no_exact_off_target_gene_hits": True,
        },
    }

    normalized = apply_parameter_constraints_to_goal(
        structured_goal,
        "Apply GC 35-65%, no C runs greater than 8, and no exact off-target gene hits.",
    )
    constraints = normalized["parameter_constraints"]

    assert constraints["gc_content"]["min"] == 0.35
    assert constraints["gc_content"]["max"] == 0.65
    assert constraints["max_binding_sites"] == 0
    assert constraints["homopolymers"]["per_base_max"]["C"] == 8
    assert normalized["homopolymer_limits"]["C"] == 8
    assert normalized["homopolymer_limits"]["A"] == 4


def test_request_envelope_embedded_goals_use_exact_parameter_constraint_keys() -> None:
    request_text = "Accept GC content greater than 0 and less than 1."
    raw_goal = {
        "gene_symbol": "CDKN1B",
        "species": "human",
        "gc_content_constraint": {
            "min_fraction_exclusive": 0.0,
            "max_fraction_exclusive": 1.0,
        },
        "design_programs": [
            {
                "program_id": "steric_cdkn1b_01",
                "modality": "steric",
                "gc_content_range": {"min_exclusive": 0.0, "max_exclusive": 1.0},
            }
        ],
    }
    envelope = pipeline_tools._build_request_envelope(
        request_text=request_text,
        request_id="goal_test_CDKN1B",
        structured_goal=raw_goal,
        raw_envelope={
            "tasks": [
                {
                    "task_id": "design_cdkn1b",
                    "task_type": "design_new_asos",
                    "context": {"structured_goal": deepcopy(raw_goal)},
                }
            ]
        },
    )

    serialized = json.dumps(envelope)
    assert "gc_content_constraint" not in serialized
    assert "gc_content_range" not in serialized
    assert "min_exclusive" not in serialized
    assert "min_fraction_exclusive" not in serialized
    top_constraints = envelope["structured_goal"]["parameter_constraints"]["gc_content"]
    task_constraints = envelope["tasks"][0]["context"]["structured_goal"]["parameter_constraints"]["gc_content"]
    assert top_constraints["min_inclusive"] is False
    assert top_constraints["max_inclusive"] is False
    assert task_constraints["min_inclusive"] is False
    assert task_constraints["max_inclusive"] is False


def test_binding_intervals_do_not_become_splice_targets_during_envelope_normalization() -> None:
    request_text = (
        "Design binding ASOs against the supplied genomic interval "
        "chr2:100-140 on the plus strand."
    )
    raw_goal = {
        "gene_symbol": "GENE1",
        "species": "human",
        "target_context": "reviewed cells",
        "intent": "binding",
        "evaluation_targets": [
            {
                "label": "supplied interval",
                "chrom": "chr2",
                "start": 100,
                "end": 140,
                "strand": "+",
                "desired_delta": "increase",
            }
        ],
        "region_constraints": [
            {
                "constraint_id": "supplied_intervals",
                "source_type": "explicit_rows",
                "coordinate_space": "genomic",
                "operator": "union",
                "aso_overlap_requirement": "any_overlap",
                "intervals_bed0_half_open": [
                    {
                        "label": "supplied interval",
                        "chrom": "chr2",
                        "start": 100,
                        "end": 140,
                        "strand": "+",
                    }
                ],
            }
        ],
        "design_programs": [
            {
                "program_id": "primary",
                "label": "primary",
                "modality": "steric",
                "intent": "binding",
                "objective": "bind the supplied interval",
                "desired_delta": "increase",
            }
        ],
    }
    envelope = pipeline_tools._build_request_envelope(
        request_text=request_text,
        request_id="binding_intervals",
        structured_goal={},
        raw_envelope={
            "tasks": [
                {
                    "task_id": "design",
                    "task_type": "design_new_asos",
                    "context": {"structured_goal": raw_goal},
                }
            ]
        },
    )

    goal = envelope["tasks"][0]["context"]["structured_goal"]
    assert goal["intent"] == "binding"
    assert goal["region_constraints"]
    assert "explicit_splice_targets" not in goal
    assert "splice_isoform_objectives" not in goal
    assert goal.get("focus_exons_only") is not True


def test_request_envelope_embedded_goals_surface_splice_targets_and_per_route_counts() -> None:
    request_text = (
        "For PCCA in human hepatocytes, skip the patient poison exon at "
        "chr13:100861220-100861336 on the plus strand. Return 8 candidates per route."
    )
    raw_goal = {
        "gene_symbol": "PCCA",
        "species": "human",
        "target_context": "human hepatocytes",
        "intent": "splice_modulation",
        "evaluation_targets": {
            "gene_symbol": "PCCA",
            "chrom": "chr13",
            "strand": "+",
            "start": 100861220,
            "end": 100861336,
        },
        "candidate_limit": 8,
        "design_programs": [
            {"program_id": "pcca_steric", "modality": "steric", "intent": "splice_modulation"},
            {"program_id": "pcca_gapmer", "modality": "gapmer", "intent": "knockdown"},
        ],
    }
    envelope = pipeline_tools._build_request_envelope(
        request_text=request_text,
        request_id="goal_test_PCCA",
        structured_goal={},
        raw_envelope={
            "tasks": [
                {
                    "task_id": "design_pcca",
                    "task_type": "design_new_asos",
                    "context": {"structured_goal": deepcopy(raw_goal)},
                }
            ]
        },
    )

    goal = envelope["tasks"][0]["context"]["structured_goal"]
    target = goal["explicit_splice_targets"][0]
    assert goal["candidate_limit_per_site"] == 8
    assert goal["explicit_splice_targets_source"] == "evaluation_targets"
    assert target["chrom"] == "chr13"
    assert target["start"] == 100861220
    assert target["end"] == 100861336
    assert target["strand"] == "+"


def test_request_envelope_embedded_goals_extract_request_text_splice_coordinates() -> None:
    request_text = (
        "For PCCA in human hepatocytes, design steric ASOs to increase productive PCCA by skipping "
        "the patient poison exon 14x at chr13:100861220-100861336 on the plus strand. "
        "Return 8 candidates per route."
    )
    raw_goal = {
        "gene_symbol": "PCCA",
        "species": "human",
        "target_context": (
            "Human hepatocytes; skip poison exon 14x at chr13:100861220-100861336 (+)."
        ),
        "intent": "splice_modulation",
        "splice_isoform_objectives": [
            {
                "objective_scope": "default_exon_context",
                "exon_label": "14x",
                "desired_abundance": "decrease",
                "desired_effect": "skipping",
                "role": "pathogenic",
            }
        ],
        "candidate_limit": 8,
        "design_programs": [
            {"program_id": "pcca_steric", "modality": "steric", "intent": "splice_modulation"},
        ],
    }
    envelope = pipeline_tools._build_request_envelope(
        request_text=request_text,
        request_id="goal_test_PCCA_text_coords",
        structured_goal={},
        raw_envelope={
            "tasks": [
                {
                    "task_id": "design_pcca",
                    "task_type": "design_new_asos",
                    "context": {"structured_goal": deepcopy(raw_goal)},
                }
            ]
        },
    )

    goal = envelope["tasks"][0]["context"]["structured_goal"]
    target = goal["explicit_splice_targets"][0]
    assert goal["explicit_splice_targets_source"] == "request_text"
    assert goal["candidate_limit_per_site"] == 8
    assert target["chrom"] == "chr13"
    assert target["start"] == 100861220
    assert target["end"] == 100861336
    assert target["strand"] == "+"
    assert target["source"] == "request_text"


def test_same_gene_multimodal_comparison_retains_independent_task_caps() -> None:
    request_text = (
        "For PCCA in human hepatocytes, design steric ASOs to skip poison exon 14x using "
        "custom GenBank pcca_14x.dna. Also include a fallback gapmer strategy that selectively "
        "depletes exon-14x-containing transcripts. Return 8 candidates per route and report which "
        "branch performed better."
    )
    envelope = pipeline_tools._build_request_envelope(
        request_text=request_text,
        request_id="goal_test_PCCA_branch_merge",
        structured_goal={},
        raw_envelope={
            "default_context": {"artifact_refs": [{"path": "pcca_14x.dna"}]},
            "tasks": [
                {
                    "task_id": "design_steric_pcca_ex14x_skip",
                    "task_type": "design_new_asos",
                    "context": {
                        "artifact_refs": [{"path": "pcca_14x.dna"}],
                        "gene_symbol": "PCCA",
                        "species": "human",
                        "target_context": "human hepatocytes",
                        "structured_goal": {
                            "gene_symbol": "PCCA",
                            "species": "human",
                            "target_context": "human hepatocytes",
                            "intent": "splice_modulation",
                            "candidate_limit": 8,
                            "design_programs": [
                                {
                                    "program_id": "steric_pcca",
                                    "modality": "steric",
                                    "intent": "splice_modulation",
                                }
                            ],
                            "splice_isoform_objectives": [
                                {
                                    "objective_scope": "default_exon_context",
                                    "exon_label": "14x",
                                    "desired_abundance": "decrease",
                                    "role": "pathogenic",
                                }
                            ],
                        },
                    },
                },
                {
                    "task_id": "design_gapmer_pcca_ex14x_deplete",
                    "task_type": "design_new_asos",
                    "context": {
                        "artifact_refs": [{"path": "pcca_14x.dna"}],
                        "gene_symbol": "PCCA",
                        "species": "human",
                        "target_context": "human hepatocytes",
                        "structured_goal": {
                            "gene_symbol": "PCCA",
                            "species": "human",
                            "target_context": "human hepatocytes",
                            "intent": "splice_modulation",
                            "candidate_limit": 8,
                            "design_programs": [
                                {
                                    "program_id": "gapmer_pcca",
                                    "modality": "gapmer",
                                    "intent": "knockdown",
                                }
                            ],
                            "splice_isoform_objectives": [
                                {
                                    "objective_scope": "aggregate_class",
                                    "aggregate_class": {
                                        "basis": "exon_presence",
                                        "exon_label": 14,
                                        "presence": "included",
                                    },
                                    "desired_abundance": "decrease",
                                    "role": "pathogenic",
                                }
                            ],
                        },
                    },
                },
                {
                    "task_id": "compare_steric_vs_gapmer_routes",
                    "task_type": "compare_aso_sets",
                    "depends_on": [
                        "design_steric_pcca_ex14x_skip",
                        "design_gapmer_pcca_ex14x_deplete",
                    ],
                    "inputs": ["steric_candidate_table_top8", "gapmer_candidate_table_top8"],
                    "context": {"artifact_refs": [{"path": "pcca_14x.dna"}]},
                },
            ],
        },
    )
    design_tasks = [
        task for task in envelope["tasks"] if task["task_type"] == "design_new_asos"
    ]
    assert len(design_tasks) == 2
    goals = [task["context"]["structured_goal"] for task in design_tasks]
    modalities = {program["modality"] for goal in goals for program in goal["design_programs"]}
    validation = pipeline_tools._request_validation_summary(envelope)

    assert modalities == {"steric", "gapmer"}
    assert all(goal["candidate_limit"] == 8 for goal in goals)
    comparison = next(task for task in envelope["tasks"] if task["task_type"] == "compare_aso_sets")
    assert set(comparison["depends_on"]) == {task["task_id"] for task in design_tasks}
    assert validation["blocked_task_ids"] == []


def test_runtime_projection_does_not_invent_a_missing_design_branch_from_request_prose() -> None:
    request_text = (
        "For PCCA in human hepatocytes, design steric ASOs to skip poison exon 14x. "
        "Also include a fallback gapmer strategy that selectively depletes exon-14x-containing transcripts. "
        "Return 8 candidates per route."
    )
    envelope = pipeline_tools._build_request_envelope(
        request_text=request_text,
        request_id="goal_test_PCCA_gapmer_branch_repair",
        structured_goal={},
        raw_envelope={
            "tasks": [
                {
                    "task_id": "design_steric_pcca_ex14x_skip",
                    "task_type": "design_new_asos",
                    "merged_from_task_ids": [
                        "design_steric_pcca_ex14x_skip",
                        "design_gapmer_pcca_ex14x_deplete",
                    ],
                    "context": {
                        "gene_symbol": "PCCA",
                        "species": "human",
                        "target_context": "human hepatocytes",
                        "structured_goal": {
                            "gene_symbol": "PCCA",
                            "species": "human",
                            "target_context": "human hepatocytes",
                            "intent": "splice_modulation",
                            "candidate_limit": 8,
                            "design_programs": [
                                {
                                    "program_id": "steric_pcca",
                                    "modality": "steric",
                                    "intent": "splice_modulation",
                                }
                            ],
                        },
                    },
                }
            ],
        },
    )

    goal = envelope["tasks"][0]["context"]["structured_goal"]
    modalities = {program["modality"] for program in goal["design_programs"]}

    assert modalities == {"steric"}
    assert goal["candidate_limit_per_site"] == 8


def test_goal_schema_key_violation_detector_finds_executable_aliases() -> None:
    raw = {
        "status": "ready",
        "structured_goal": {
            "constraints": {
                "gc_content_fraction": {"min_exclusive": 0.0, "max_exclusive": 1.0},
                "max_offtarget_mappings": 2,
            },
            "binding_site_constraints": {
                "tile_overlap_requirement": "any_overlap",
                "min_overlap_bp": 1,
            },
        },
    }

    violations = run_goal_intake._find_goal_schema_key_violations(raw)
    by_key = {item["key"]: item["replacement"] for item in violations}

    assert by_key["gc_content_fraction"] == "constraints.gc_content"
    assert by_key["min_exclusive"] == "constraints.gc_content.min with min_inclusive=false"
    assert by_key["max_exclusive"] == "constraints.gc_content.max with max_inclusive=false"
    assert by_key["max_offtarget_mappings"] == "constraints.max_binding_sites"
    assert by_key["binding_site_constraints"] == "aso_overlap_requirement and minimum_overlap_bp"
    assert by_key["tile_overlap_requirement"] == "aso_overlap_requirement"
    assert by_key["min_overlap_bp"] == "minimum_overlap_bp"


def test_goal_schema_key_violation_detector_flags_splice_objective_scope_repairs() -> None:
    raw = {
        "status": "ready",
        "structured_goal": {
            "gene_symbol": "GENE",
            "species": "human",
            "intent": "splice_modulation",
            "splice_isoform_objectives": [
                {
                    "isoform_label": "named isoform without transcript",
                    "desired_abundance": "increase",
                },
                {
                    "target_class": "exon_presence",
                    "exon_label": 7,
                    "aggregate_presence": "excluded",
                    "desired_abundance": "decrease",
                },
                {
                    "objective_scope": "aggregate_class",
                    "aggregate_class": {"basis": "exon_presence", "exon_label": 7},
                    "desired_abundance": "decrease",
                },
            ],
        },
    }

    violations = run_goal_intake._find_goal_schema_key_violations(raw)

    assert any(item["key"] == "objective_scope" for item in violations)
    assert any(item["key"] == "target_class" for item in violations)
    assert any(item["key"] == "aggregate_presence" for item in violations)
    assert any(item["key"] == "aggregate_class" for item in violations)
    assert all(
        item.get("category") == "splice_isoform_objective"
        for item in violations
        if "splice_isoform_objectives" in item.get("path", "")
    )


def test_goal_intake_reprompts_when_model_uses_constraint_aliases(monkeypatch: Any) -> None:
    class DummyResult:
        def __init__(self, payload: dict[str, Any]) -> None:
            self.final_output = json.dumps(payload)

    calls: list[str] = []

    async def fake_run(agent: Any, *, input: str, max_steps: int) -> DummyResult:
        calls.append(str(getattr(agent, "name", "")))
        if calls[-1] == "ASO Goal Intake":
            return DummyResult(
                {
                    "status": "ready",
                    "structured_goal": {
                        "gene_symbol": "LRRK2",
                        "species": "human",
                        "intent": "knockdown",
                        "constraints": {
                            "gc_content_fraction": {"min_exclusive": 0.0, "max_exclusive": 1.0},
                            "max_offtarget_mappings": 2,
                        },
                        "design_programs": [
                            {
                                "program_id": "gapmer_lrrk2",
                                "modality": "gapmer",
                                "intent": "knockdown",
                                "gc_content_range": {"min_exclusive": 0.0, "max_exclusive": 1.0},
                            }
                        ],
                    },
                }
            )
        if calls[-1] == "ASO Physical Site Scope Extractor":
            return DummyResult({"feature_region": "none", "require_complete_site_in_one_feature": False,
                                "evidence_quote": ""})
        if calls[-1] == "ASO Goal Intake Clarification Necessity Reviewer":
            assert '"status":"ready"' in input
            return DummyResult(
                {
                    "blocking_question_indexes": [],
                    "additional_blocking_questions": [],
                    "reason": None,
                }
            )
        assert calls[-1] == "ASO Goal Schema Repair"
        assert "gc_content_fraction" in input
        assert "constraints.gc_content" in input
        assert "constraints.max_binding_sites" in input
        assert "objective_scope" in input
        assert "aggregate_class" in input
        assert "Do not invent transcript IDs" in input
        return DummyResult(
            {
                "status": "ready",
                "structured_goal": {
                    "gene_symbol": "LRRK2",
                    "species": "human",
                    "intent": "knockdown",
                    "constraints": {
                        "gc_content": {
                            "min": 0.0,
                            "max": 1.0,
                            "min_inclusive": False,
                            "max_inclusive": False,
                            "min_percent": 0.0,
                            "max_percent": 100.0,
                        },
                        "max_binding_sites": 2,
                    },
                    "design_programs": [
                        {
                            "program_id": "gapmer_lrrk2",
                            "modality": "gapmer",
                            "intent": "knockdown",
                        }
                    ],
                },
            }
        )

    monkeypatch.setattr(run_goal_intake.Runner, "run", fake_run)

    intake = asyncio.run(
        run_goal_intake._run_goal_intake_agent_once(
            "Design LRRK2 gapmers with GC content greater than 0 and less than 1 and at most 2 binding sites."
        )
    )

    serialized = json.dumps(intake)
    assert calls == [
        "ASO Goal Intake",
        "ASO Goal Schema Repair",
        "ASO Physical Site Scope Extractor",
        "ASO Goal Intake Clarification Necessity Reviewer",
    ]
    assert "gc_content_fraction" not in serialized
    assert "gc_content_range" not in serialized
    assert "max_offtarget_mappings" not in serialized
    assert "min_exclusive" not in serialized
    constraints = intake["structured_goal"]["constraints"]
    assert constraints["gc_content"]["min_inclusive"] is False
    assert constraints["gc_content"]["max_inclusive"] is False
    assert constraints["max_binding_sites"] == 2


def test_goal_intake_reprompts_when_splice_objective_scope_is_missing(monkeypatch: Any) -> None:
    class DummyResult:
        def __init__(self, payload: dict[str, Any]) -> None:
            self.final_output = json.dumps(payload)

    calls: list[str] = []

    async def fake_run(agent: Any, *, input: str, max_steps: int) -> DummyResult:
        calls.append(str(getattr(agent, "name", "")))
        if calls[-1] == "ASO Goal Intake":
            return DummyResult(
                {
                    "status": "ready",
                    "structured_goal": {
                        "gene_symbol": "GENE",
                        "species": "human",
                        "target_context": "disease-relevant cells",
                        "intent": "splice_modulation",
                        "exon_label": 7,
                        "splice_isoform_objectives": [
                            {
                                "isoform_label": "qualitative group",
                                "desired_abundance": "decrease",
                                "role": "pathogenic",
                            }
                        ],
                    },
                }
            )
        if calls[-1] == "ASO Goal Intake Clarification Necessity Reviewer":
            assert '"status":"ready"' in input
            return DummyResult(
                {
                    "blocking_question_indexes": [],
                    "additional_blocking_questions": [],
                    "reason": None,
                }
            )
        assert calls[-1] == "ASO Goal Schema Repair"
        assert "objective_scope" in input
        assert "aggregate_class" in input
        return DummyResult(
            {
                "status": "ready",
                "structured_goal": {
                    "gene_symbol": "GENE",
                    "species": "human",
                    "target_context": "disease-relevant cells",
                    "intent": "splice_modulation",
                    "exon_label": 7,
                    "splice_isoform_objectives": [
                        {
                            "objective_scope": "aggregate_class",
                            "isoform_label": "qualitative group",
                            "aggregate_class": {
                                "basis": "exon_presence",
                                "exon_label": 7,
                                "presence": "excluded",
                            },
                            "desired_abundance": "decrease",
                            "role": "pathogenic",
                        }
                    ],
                },
            }
        )

    monkeypatch.setattr(run_goal_intake.Runner, "run", fake_run)

    intake = asyncio.run(
        run_goal_intake._run_goal_intake_agent_once(
            "Design splice-switching ASOs for GENE in disease-relevant cells."
        )
    )

    objectives = intake["structured_goal"]["splice_isoform_objectives"]
    assert calls == [
        "ASO Goal Intake",
        "ASO Goal Schema Repair",
        "ASO Goal Intake Clarification Necessity Reviewer",
    ]
    assert objectives[0]["objective_scope"] == "aggregate_class"
    assert objectives[0]["aggregate_class"]["presence"] == "excluded"


def test_pipeline_parameter_constraints_include_sequence_filters() -> None:
    merged = pipeline_tools._merge_parameter_constraints(
        {
            "sequence_filters": {
                "gc_content_min": 0.0,
                "gc_content_max": 1.0,
                "homopolymer_max_run": 4,
            },
            "parameter_constraints": {
                "max_binding_sites": 1,
            },
        }
    )

    assert get_gc_bounds(merged) == (0.0, 1.0)
    assert merged["max_binding_sites"] == 1
    assert merged["homopolymers"]["per_base_max"]["A"] == 4
    assert "gc_content_min" not in merged
    assert "gc_content_max" not in merged


def test_offtarget_mapping_aliases_are_normalized_to_binding_constraints() -> None:
    constraints = parse_parameter_constraints(
        "",
        {
            "offtarget_constraints": {
                "max_offtarget_mappings": 2,
            }
        },
    )

    assert constraints["max_binding_sites"] == 2
    assert constraints["binding_sites"]["max"] == 2
    assert "max_offtarget_mappings" not in constraints


def test_offtarget_max_total_mappings_alias_is_normalized() -> None:
    constraints = parse_parameter_constraints(
        "",
        {
            "off_target_constraints": {
                "max_total_mappings": 2,
            }
        },
    )

    assert constraints["max_binding_sites"] == 2
    assert constraints["binding_sites"]["max"] == 2


def test_offtarget_mapping_language_sets_strict_less_than_limit() -> None:
    constraints = parse_parameter_constraints("Require fewer than 3 off-target mappings.", {})

    assert constraints["max_binding_sites"] == 2
    assert constraints["binding_sites"]["max"] == 2


def test_sequence_constraints_base_specific_homopolymer_runs_are_normalized() -> None:
    constraints = parse_parameter_constraints(
        "",
        {
            "sequence_constraints": {
                "disallowed_homopolymer_runs": [
                    {
                        "base": "G",
                        "max_allowed_run_length": 3,
                    }
                ]
            }
        },
    )

    per_base = get_homopolymer_per_base_limits(constraints)
    assert per_base["G"] == 3
    assert sequence_passes_homopolymer_limits("AAAGGG", per_base) is True
    assert sequence_passes_homopolymer_limits("AAAGGGG", per_base) is False


def test_filter_block_homopolymer_exclusions_are_canonicalized() -> None:
    goal = {
        "filters": {
            "exclude_homopolymers": [
                {
                    "base": "G",
                    "min_run_length": 4,
                }
            ],
            "off_target_mappings": {
                "max_mappings": 2,
            },
        }
    }

    constraints = parse_parameter_constraints("", goal)

    assert constraints["max_binding_sites"] == 2
    assert constraints["binding_sites"]["max"] == 2
    assert constraints["homopolymers"]["per_base_max"]["G"] == 3
    assert "GGGG" in constraints["homopolymers"]["avoid_sequences"]


def test_gapmer_binding_constraint_prefers_explicit_off_target_count() -> None:
    passing_row = {
        "Sequence": "ACGTACGTACGTACGTACGT",
        "Times_mapped": 15,
        "Times_mapped_off_target_genes": 1,
    }
    failing_row = {
        "Sequence": "ACGTACGTACGTACGTACGT",
        "Times_mapped": 1,
        "off_target_gene_count": 3,
    }
    constraints = {"max_binding_sites": 2}

    assert get_binding_count_from_row(passing_row) == 1
    assert evaluate_row_against_constraints(passing_row, constraints) == (True, [])
    assert get_binding_count_from_row(failing_row) == 3
    assert evaluate_row_against_constraints(failing_row, constraints) == (
        False,
        ["binding_sites_exceed"],
    )


def test_gapmer_binding_constraint_preserves_explicit_zero() -> None:
    row = {
        "Times_mapped": 15,
        "Times_mapped_off_target_genes": 0,
        "off_target_gene_count": 4,
    }

    assert get_binding_count_from_row(row) == 0


def test_raw_gapmer_binding_constraint_is_deferred_until_off_target_mapping(tmp_path: Path) -> None:
    row = {
        "Sequence": "ACGTACGTACGTACGTACGT",
        "GC": 0.5,
        "No_homopolymer": True,
        "Times_mapped": 15,
        "source": "gapmer_cdna",
    }
    tiles_path = tmp_path / "raw_gapmer_tiles.json"
    tiles_path.write_text(json.dumps({"rows": [row]}), encoding="utf-8")

    assert get_binding_count_from_row(row) is None
    summary = run_goal_intake._apply_initial_tiling_filters(
        str(tiles_path),
        {"max_binding_sites": 2},
    )

    assert summary["passed"]["count"] == 1
    assert summary["filtered_binding_sites_exceed"]["count"] == 0


def test_overmapped_gapmer_lower_bound_fails_hard_binding_limit() -> None:
    row = {
        "Sequence": "ACGTACGTACGTACGTACGT",
        "source": "gapmer_cdna",
        "Times_mapped": 15,
        "Times_mapped_off_target_genes": 0,
        "off_target_mapping_count_lower_bound": 2001,
    }

    assert get_binding_count_from_row(row) == 2001
    assert evaluate_row_against_constraints(row, {"max_binding_sites": 2}) == (
        False,
        ["binding_sites_exceed"],
    )


def test_user_text_homopolymer_exclusions_are_canonicalized() -> None:
    constraints = parse_parameter_constraints(
        "Exclude sequences with G homopolymer runs of length 4 or greater.",
        {},
    )

    assert constraints["homopolymers"]["per_base_max"]["G"] == 3
    assert "GGGG" in constraints["homopolymers"]["avoid_sequences"]


def test_user_text_exact_gc_no_g_and_total_binding_site_are_canonicalized() -> None:
    request_text = (
        "Require GC content exactly 50%, exclude G homopolymer runs of length 1 or greater, "
        "and require at most one total binding site."
    )
    goal = apply_parameter_constraints_to_goal({}, request_text)
    constraints = goal["parameter_constraints"]

    assert constraints["gc_content"]["min"] == 0.5
    assert constraints["gc_content"]["max"] == 0.5
    assert constraints["max_binding_sites"] == 1
    assert constraints["binding_sites"]["max"] == 1
    assert constraints["homopolymers"]["per_base_max"]["G"] == 0
    assert goal["homopolymer_limits"]["G"] == 0
    assert sequence_passes_homopolymer_limits("ACACACACACACACACACAC", {"G": 0}) is True
    assert sequence_passes_homopolymer_limits("ACACACACACACACACACAG", {"G": 0}) is False


def test_pipeline_parameter_constraints_include_offtarget_constraint_aliases() -> None:
    merged = pipeline_tools._merge_parameter_constraints(
        {
            "off_target_constraints": {
                "max_total_mappings": 2,
            },
            "sequence_constraints": {
                "disallowed_homopolymer_runs": [
                    {
                        "base": "G",
                        "max_allowed_run_length": 3,
                    }
                ]
            },
        }
    )

    assert merged["max_binding_sites"] == 2
    assert merged["binding_sites"]["max"] == 2
    assert merged["homopolymers"]["per_base_max"]["G"] == 3
    assert "max_total_mappings" not in merged
    assert "disallowed_homopolymer_runs" not in merged


def test_design_objective_preseed_preserves_direct_sequence_filters() -> None:
    normalized = pipeline_tools._normalized_design_objective_preseed(
        {
            "splice_isoform_objectives": [
                {
                    "isoform_label": "target isoform",
                    "transcript_id": "ENST000001",
                    "desired_abundance": "decrease",
                }
            ],
            "sequence_filters": {
                "gc_content_min": 0.0,
                "gc_content_max": 1.0,
            },
        }
    )

    assert normalized["structured_goal_overlay"]["sequence_filters"] == {
        "gc_content_min": 0.0,
        "gc_content_max": 1.0,
    }


def test_design_objective_preseed_preserves_offtarget_constraints() -> None:
    normalized = pipeline_tools._normalized_design_objective_preseed(
        {
            "splice_isoform_objectives": [
                {
                    "isoform_label": "ENST00000298910",
                    "transcript_id": "ENST00000298910",
                    "desired_abundance": "decrease",
                }
            ],
            "offtarget_constraints": {
                "max_offtarget_mappings": 2,
            },
            "sequence_constraints": {
                "disallowed_homopolymer_runs": [
                    {
                        "base": "G",
                        "max_allowed_run_length": 3,
                    }
                ]
            },
        }
    )

    assert normalized["structured_goal_overlay"]["offtarget_constraints"] == {
        "max_offtarget_mappings": 2,
    }
    assert normalized["structured_goal_overlay"]["sequence_constraints"] == {
        "disallowed_homopolymer_runs": [
            {
                "base": "G",
                "max_allowed_run_length": 3,
            }
        ]
    }


def test_report_manifest_uses_effective_parameter_constraints() -> None:
    report = pipeline_tools._build_report_payload_manifest(
        structured_goal={
            "gene_symbol": "LRRK2",
            "off_target_constraints": {
                "max_total_mappings": 2,
            },
            "sequence_constraints": {
                "disallowed_homopolymer_runs": [
                    {
                        "base": "G",
                        "max_allowed_run_length": 3,
                    }
                ]
            },
        },
        final_payload={
            "parameter_constraints": {
                "homopolymers": {
                    "per_base_max": {
                        "A": 4,
                        "C": 4,
                        "G": 4,
                        "T": 4,
                    }
                }
            }
        },
    )

    assert report["parameter_constraints"]["max_binding_sites"] == 2
    assert report["parameter_constraints"]["homopolymers"]["per_base_max"]["G"] == 3


def test_homopolymer_exclusions_tighten_default_homopolymer_map() -> None:
    merged = pipeline_tools._merge_parameter_constraints(
        {
            "parameter_constraints": {
                "homopolymers": {
                    "per_base_max": {
                        "A": 4,
                        "C": 4,
                        "G": 4,
                        "T": 4,
                    },
                    "avoid_sequences": ["AAAAA", "CCCCC", "GGGGG", "TTTTT"],
                }
            },
            "sequence_constraints": {
                "homopolymer_exclusions": [
                    {
                        "base": "G",
                        "max_run_length": 3,
                    }
                ]
            },
        }
    )

    assert merged["homopolymers"]["per_base_max"]["G"] == 3
    assert "GGGG" in merged["homopolymers"]["avoid_sequences"]


def test_empty_gapmer_results_preserve_parameter_constraint_reason(tmp_path: Path) -> None:
    scored = tmp_path / "gapmer_scored_uniqueness.json"
    scored.write_text(json.dumps({"rows": []}), encoding="utf-8")

    result = pipeline_tools._apply_gapmer_constraints_to_result(
        gapmer_results={
            "scored_tiles_path": str(scored),
            "initial_filtering": {
                "passed": {"count": 0},
                "total_rows": 170,
                "constraint_summary": {"max_binding_sites": 2},
                "filtered_binding_sites_exceed": {"count": 170},
            },
        },
        structured_goal={},
        runtime=object(),
    )

    filtering = result["constraint_filtering"]
    assert filtering["applied"] is True
    assert filtering["reason"] == "no_gapmer_candidates_after_parameter_constraints"
    assert filtering["kept_count"] == 0
    assert filtering["removed_count"] == 170
    assert filtering["constraint_summary"]["max_binding_sites"] == 2


def test_shared_transcript_gapmer_filter_precedes_ranking_and_panel_limit(tmp_path: Path) -> None:
    scored = tmp_path / "scored.json"
    rows = [
        {"Name": "wrong_high_rank", "Sequence": "AAAA", "isoform_uniqueness_labels": ["ENST00000000002"], "priority": 0},
        {"Name": "shared_one", "Sequence": "CCCC", "isoform_uniqueness_labels": ["ENST00000000001", "ENST00000000002"], "priority": 1},
        {"Name": "shared_two", "Sequence": "GGGG", "isoform_uniqueness_labels": ["ENST00000000001", "ENST00000000002"], "priority": 2},
    ]
    scored.write_text(json.dumps({"rows": rows}), encoding="utf-8")

    def rank(path: str) -> dict[str, Any]:
        payload = json.loads(Path(path).read_text(encoding="utf-8"))
        payload["rows"].sort(key=lambda row: row["priority"])
        for index, row in enumerate(payload["rows"], 1):
            row["selection_rank"] = index
        Path(path).write_text(json.dumps(payload), encoding="utf-8")
        return {"row_count": len(payload["rows"]), "status": "complete"}

    goal = {
        "candidate_limit": 1,
        "isoform_scope": ["ENST00000000001.5", "ENST00000000002.7", "ENST00000000003.1"],
        "design_programs": [{"modality": "gapmer", "target_type": "transcript_cohort"}],
        "evaluation_targets": [
            {"target_type": "transcript", "objective_type": "primary_outcome", "desired_delta": "decrease", "transcript_id": transcript}
            for transcript in ("ENST00000000001.5", "ENST00000000002.7")
        ],
    }
    result = pipeline_tools._apply_gapmer_constraints_to_result(
        gapmer_results={"scored_tiles_path": str(scored)},
        structured_goal=goal,
        runtime=SimpleNamespace(
            _annotate_gapmer_selection_ranks=rank,
            filter_tiles_json=lambda path, constraints: {"applied": False, "output_path": path},
        ),
        apply_candidate_limit=True,
    )
    assert "error" not in result
    selected = json.loads(Path(result["gapmer_results"]["scored_tiles_path"]).read_text())["rows"]
    assert [row["Name"] for row in selected] == ["shared_one"]
    assert result["gapmer_results"]["transcript_cohort_filtering"]["failed_count"] == 1


def test_shared_transcript_gapmer_requires_measured_binding_labels(tmp_path: Path) -> None:
    scored = tmp_path / "scored.json"
    scored.write_text(json.dumps({"rows": [{"Name": "unchecked"}]}), encoding="utf-8")
    goal = {
        "design_programs": [{"modality": "gapmer", "target_type": "transcript_cohort"}],
        "evaluation_targets": [
            {"target_type": "transcript", "objective_type": "primary_outcome", "desired_delta": "decrease", "transcript_id": transcript}
            for transcript in ("ENST00000000001", "ENST00000000002")
        ],
    }
    result = pipeline_tools._apply_gapmer_constraints_to_result(
        gapmer_results={"scored_tiles_path": str(scored)}, structured_goal=goal,
        runtime=SimpleNamespace(
            filter_tiles_json=lambda path, constraints: {"applied": False, "output_path": path},
        ),
    )
    assert result["error"] == "gapmer_transcript_cohort_evidence_incomplete"


def test_ordinary_gapmer_is_not_subject_to_transcript_cohort_filter(tmp_path: Path) -> None:
    scored = tmp_path / "scored.json"
    scored.write_text(json.dumps({"rows": [{"Name": "single", "isoform_uniqueness_labels": ["ENST00000000002"]}]}), encoding="utf-8")
    result = pipeline_tools._apply_gapmer_constraints_to_result(
        gapmer_results={"scored_tiles_path": str(scored)},
        structured_goal={"design_programs": [{"modality": "gapmer", "target_type": "transcript"}]},
        runtime=SimpleNamespace(
            filter_tiles_json=lambda path, constraints: {"applied": False, "output_path": path},
        ),
    )
    assert "error" not in result
    assert "transcript_cohort_filtering" not in result["gapmer_results"]


def test_shared_transcript_filter_does_not_remove_other_gapmer_programs(tmp_path: Path) -> None:
    scored = tmp_path / "scored.json"
    scored.write_text(json.dumps({"rows": [
        {"Name": "cohort_miss", "program_id": "shared", "isoform_uniqueness_labels": ["ENST00000000001"]},
        {"Name": "single_target", "program_id": "ordinary", "isoform_uniqueness_labels": ["ENST00000000001"]},
        {"Name": "cohort_hit", "program_id": "shared", "isoform_uniqueness_labels": ["ENST00000000001", "ENST00000000002"]},
    ]}), encoding="utf-8")
    goal = {"design_programs": [
        {"program_id": "shared", "modality": "gapmer", "target_type": "transcript_cohort",
         "transcript_ids": ["ENST00000000001.1", "ENST00000000002.2"]},
        {"program_id": "ordinary", "modality": "gapmer", "target_type": "transcript"},
    ]}
    result = pipeline_tools._apply_gapmer_constraints_to_result(
        gapmer_results={"scored_tiles_path": str(scored)}, structured_goal=goal,
        runtime=SimpleNamespace(
            filter_tiles_json=lambda path, constraints: {"applied": False, "output_path": path},
        ),
    )
    assert "error" not in result
    rows = json.loads(Path(result["gapmer_results"]["scored_tiles_path"]).read_text())["rows"]
    assert {row["Name"] for row in rows} == {"single_target", "cohort_hit"}
    assert result["gapmer_results"]["transcript_cohort_filtering"]["checked_count"] == 2


def test_empty_gapmer_job_reports_roll_up_to_aggregate_parameter_reason() -> None:
    filtering = pipeline_tools._aggregate_empty_gapmer_job_constraint_reports(
        [
            {
                "applied": True,
                "reason": "no_gapmer_candidates_after_parameter_constraints",
                "removed_count": 100,
                "constraint_summary": {"max_binding_sites": 2},
                "filtered_binding_sites_exceed": {"count": 100},
            },
            {
                "applied": True,
                "reason": "no_gapmer_candidates_after_parameter_constraints",
                "removed_count": 70,
                "filtered_binding_sites_exceed": {"count": 70},
            },
        ],
        output_path="/tmp/gapmer_scored_uniqueness.json",
    )

    assert filtering is not None
    assert filtering["applied"] is True
    assert filtering["reason"] == "no_gapmer_candidates_after_parameter_constraints"
    assert filtering["removed_count"] == 170
    assert filtering["filtered_binding_sites_exceed"]["count"] == 170
    assert filtering["constraint_summary"]["max_binding_sites"] == 2


def test_filter_bucket_export_clears_stale_empty_job_counts(tmp_path: Path) -> None:
    source = tmp_path / "filtered_gc_out_of_range.json"
    source.write_text(
        json.dumps(
            {
                "schema": {"table": "tiles", "primary_key": "id", "columns": {}},
                "rows": [
                    {"id": "job1_row", "Chrom": "chr1", "Start": 101, "End": 110, "Sequence": "G" * 18}
                ],
            }
        ),
        encoding="utf-8",
    )
    structured_goal: dict[str, Any] = {}
    final_payload: dict[str, Any] = {
        "steric_design_jobs": [
            {
                "job_id": "job1",
                "label": "job1",
                "output_dir": str(tmp_path / "job1"),
                "region_rows": [["chr1", 100, 120, "job1", 0, "+"]],
            },
            {
                "job_id": "job2",
                "label": "job2",
                "output_dir": str(tmp_path / "job2"),
                "region_rows": [["chr1", 200, 220, "job2", 0, "+"]],
            },
        ],
        "steric_tiling_results_by_job": {
            "job2": {
                "filtered_gc_out_of_range_path": "/tmp/stale.json",
                "filtered_gc_out_of_range_path_row_count": 159,
            }
        },
    }

    pipeline_tools._export_job_filter_bucket_views(
        run_dir=str(tmp_path),
        structured_goal=structured_goal,
        final_payload=final_payload,
        modality="steric",
        filter_summary={"filtered_gc_out_of_range": {"count": 1, "path": str(source)}},
        bucket_specs=[("filtered_gc_out_of_range", "filtered_gc_out_of_range_path", "filtered_gc_out_of_range.json")],
    )

    results = final_payload["steric_tiling_results_by_job"]
    assert results["job1"]["filtered_gc_out_of_range_path_row_count"] == 1
    assert results["job2"]["filtered_gc_out_of_range_path_row_count"] == 0
    assert Path(results["job2"]["filtered_gc_out_of_range_path"]).exists()


def test_filter_bucket_export_writes_empty_job_counts_for_empty_source(tmp_path: Path) -> None:
    source = tmp_path / "filtered_gc_out_of_range.json"
    source.write_text(
        json.dumps({"schema": {"table": "tiles", "primary_key": "id", "columns": {}}, "rows": []}),
        encoding="utf-8",
    )
    structured_goal: dict[str, Any] = {}
    final_payload: dict[str, Any] = {
        "steric_design_jobs": [
            {
                "job_id": "job1",
                "label": "job1",
                "output_dir": str(tmp_path / "job1"),
                "region_rows": [["chr1", 100, 120, "job1", 0, "+"]],
            }
        ],
        "steric_tiling_results_by_job": {
            "job1": {
                "filtered_gc_out_of_range_path": "/tmp/stale.json",
                "filtered_gc_out_of_range_path_row_count": 142,
            }
        },
    }

    pipeline_tools._export_job_filter_bucket_views(
        run_dir=str(tmp_path),
        structured_goal=structured_goal,
        final_payload=final_payload,
        modality="steric",
        filter_summary={"filtered_gc_out_of_range": {"count": 0, "path": str(source)}},
        bucket_specs=[("filtered_gc_out_of_range", "filtered_gc_out_of_range_path", "filtered_gc_out_of_range.json")],
    )

    result = final_payload["steric_tiling_results_by_job"]["job1"]
    assert result["filtered_gc_out_of_range_path_row_count"] == 0
    assert Path(result["filtered_gc_out_of_range_path"]).exists()


def test_steric_scoring_restart_invalidates_root_filter_buckets(tmp_path: Path) -> None:
    paths = {Path(path).relative_to(tmp_path) for path in pipeline_tools._known_artifact_paths_by_stage(str(tmp_path), "steric_scoring")}

    assert Path("tiling/filtered_gc_out_of_range.json") in paths
    assert Path("tiling/filtered_no_homopolymer.json") in paths
    assert Path("tiling/filtered_binding_sites_exceed.json") in paths


def test_gapmer_scoring_uses_projected_mutant_cdna_for_uniqueness(
    tmp_path: Path,
    monkeypatch,
) -> None:
    tiles_path = tmp_path / "gapmer_tiles.json"
    tiles_path.write_text(
        json.dumps({"rows": [{"Sequence": "A" * 20}]}),
        encoding="utf-8",
    )
    observed: dict[str, Any] = {}

    def project_variant(goal, transcript, label, wildtype_cdna):
        observed["projection_input"] = wildtype_cdna
        return {
            "status": "resolved",
            "transcript_id": transcript["transcript_id"],
            "cdna_position": 3,
            "wildtype_allele": "A",
            "mutant_allele": "C",
            "mutant_cdna": "AACA",
            "source": {"provider": "test"},
        }

    def build_kmer_index(cdna_map, lengths):
        observed["cdna_map"] = dict(cdna_map)
        observed["lengths"] = list(lengths)
        return {}

    runtime = SimpleNamespace(
        _apply_initial_tiling_filters=lambda path, constraints: {
            "passed": {"path": path, "count": 1},
            "total_rows": 1,
        },
        _resolve_patient_reference_fasta=lambda goal: None,
        _build_isoform_cdna_sequences=lambda transcripts, species, reference_fasta=None: {
            "primary": "AAAA"
        },
        _patient_variant_cdna_projection=project_variant,
        _build_gapmer_kmer_index_by_len=build_kmer_index,
        _annotate_gapmer_uniqueness=lambda path, index, label, **kwargs: path,
    )
    monkeypatch.setattr(
        pipeline_tools,
        "_annotate_gapmer_offtargets_by_species",
        lambda **kwargs: {"complete": True},
    )
    monkeypatch.setattr(
        pipeline_tools,
        "_annotate_repeat_risk_json",
        lambda **kwargs: {"status": "complete"},
    )

    result = pipeline_tools._score_gapmer_tiles_json(
        gapmer_results={
            "json_path": str(tiles_path),
            "json_path_unfiltered": str(tmp_path / "removed_tiling_checkpoint.json"),
            "tile_len": 20,
        },
        structured_goal={
            "species": "human",
            "patient_variant": {"notation": "reference:c.3A>C"},
        },
        runtime=runtime,
        isoform_transcripts={"primary": {"transcript_id": "transcript-1"}},
        all_transcripts_full={"primary": {"transcript_id": "transcript-1"}},
        iso_label_primary="primary",
    )

    assert "error" not in result
    assert observed["projection_input"] == "AAAA"
    assert observed["cdna_map"]["primary"] == "AACA"
    assert observed["lengths"] == [20]
    assert result["gapmer_results"]["variant_projection"]["mutant_allele"] == "C"


def test_gapmer_scoring_restart_preserves_per_job_tiling_checkpoint(tmp_path: Path) -> None:
    run_dir = tmp_path / "run"
    job_dir = run_dir / "artifacts" / "design" / "gapmer_jobs" / "job-1"
    checkpoint = job_dir / "artifacts" / "design" / "gapmer_tiling" / "tiles.json"
    checkpoint.parent.mkdir(parents=True)
    checkpoint.write_text(json.dumps({"rows": [{"Sequence": "A" * 20}]}), encoding="utf-8")
    working_tiles = job_dir / "tiles.json"
    working_tiles.write_text(json.dumps({"rows": [{"Sequence": "A" * 20}]}), encoding="utf-8")
    scored = job_dir / "gapmer_scored_uniqueness.json"
    scored.write_text(json.dumps({"rows": []}), encoding="utf-8")

    invalidation_paths = {
        Path(path).relative_to(run_dir)
        for path in pipeline_tools._known_artifact_paths_by_stage(str(run_dir), "gapmer_scoring")
    }

    assert Path("artifacts/design/gapmer_jobs/job-1/tiles.json") in invalidation_paths
    assert Path("artifacts/design/gapmer_jobs/job-1/gapmer_scored_uniqueness.json") in invalidation_paths
    assert Path("artifacts/design/gapmer_jobs/job-1") not in invalidation_paths
    assert Path("artifacts/design/gapmer_jobs/job-1/artifacts/design/gapmer_tiling/tiles.json") not in invalidation_paths
