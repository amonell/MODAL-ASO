import hashlib
import json
from pathlib import Path

import openpyxl
import pytest

from aso_agents.design_workbooks import DesignWorkbook, validate_design_workbooks
from test_design_workbooks import context, make_book


def change_observations(declared, update):
    meta = declared["completed_run"]["design_evidence"][0]["observations"]
    path = Path(meta["path"])
    data = json.loads(path.read_text())
    update(data["rows"])
    path.write_text(json.dumps(data))
    meta["sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()


def test_invalid_recorded_aliases_fail_together_with_canonical_names_available(tmp_path):
    declared, path, _ = context(tmp_path)
    writer = DesignWorkbook(declared)
    with pytest.raises(ValueError) as error:
        writer.add_table("Library", writer.rows(), task_id="design", source_columns={
            "Target": "recorded:/target_5to3", "Length": "recorded:/length_nt"})
    message = str(error.value)
    assert "Target <- recorded:/target_5to3" in message
    assert "Length <- recorded:/length_nt" in message
    assert "without recorded:/" in message and "Canonical fields:" in message
    assert not path.exists()
    writer.add_table("Library", writer.rows(), task_id="design", source_columns={
        "Target": "target_5to3", "Length": "length_nt"})
    writer.save(path)
    values = list(openpyxl.load_workbook(path, data_only=True)["Library"].values)
    row = dict(zip(values[0], values[1]))
    assert row["Target"] == "AACCGTAACCGTAACCG" and row["Length"] == 17


@pytest.mark.parametrize("known_value", [None, 0, -0.3])
def test_real_field_null_zero_and_row_specific_missing_values_are_preserved(tmp_path, known_value):
    declared, path, payload = context(tmp_path, second=True)
    change_observations(declared, lambda rows: rows[0]["recorded_fields"].update({"/optional_measurement": known_value}))
    writer = DesignWorkbook(declared)
    writer.add_table("Library", writer.rows(), task_id="design",
                     source_columns={"Measurement": "recorded:/optional_measurement"})
    result = writer.save(path)
    validate_design_workbooks(payload, result, [path])
    values = list(openpyxl.load_workbook(path, data_only=True)["Library"].values)
    col = values[0].index("Measurement")
    assert [row[col] for row in values[1:]] == [known_value, None]


def test_other_producer_fields_cannot_validate_an_invented_field_in_this_task(tmp_path):
    declared, _, _ = context(tmp_path, second=True)
    def change(rows):
        rows[1]["binding_identity"]["task_id"] = "other-task"
        rows[1]["binding_match_key"] = "separate-other-task-binding"
        rows[1]["recorded_fields"]["/other_only"] = 0
    change_observations(declared, change)
    writer = DesignWorkbook(declared)
    with pytest.raises(ValueError, match="absent from every source row in this task"):
        writer.add_table("Library", writer.rows(task_id="design"), task_id="design",
                         source_columns={"Other": "recorded:/other_only"})


def test_actual_workbook_validation_rejects_a_forged_blank_source_alias(tmp_path):
    declared, path, payload = context(tmp_path)
    _, result = make_book(declared, path)
    result["design_tables"][0]["columns"]["Unknown measurement"] = "recorded:/not_in_sources"
    with pytest.raises(ValueError, match="Unknown recorded design columns"):
        validate_design_workbooks(payload, result, [path])
